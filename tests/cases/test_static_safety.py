"""Isolated control-flow and numeric regressions; no Office rendering."""
import copy,json,math,subprocess,sys,tempfile,unittest
from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from check_signoff_scope import check,sha,impacted_pages
from connector_integrity import boundary_distance
from verify_coverage_review import validate_region_contract
from font_evidence import review_fonts
from PIL import Image,ImageDraw
from types import SimpleNamespace
from build_pptx import run as build_run
from compute_capacity import _master_txstyle_sz

class StaticSafety(unittest.TestCase):
    def test_builder_strict_empty_expected_and_new_output_directory(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);template=root/'template.pptx';edits=root/'edits.json';out=root/'new'/'result.pptx'
            p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);sh=s.shapes.add_textbox(Pt(50),Pt(60),Pt(100),Pt(30));sh.text='old';p.save(template)
            spec={'selected_slides':[1],'edits':[{'slide':1,'address':{'shape_id':sh.shape_id,'paragraph':0,'run':0},'expected_text':'','new_text':'new'}]};edits.write_text(json.dumps(spec))
            args=SimpleNamespace(template=template,edits=edits,output=out,detail=None,strict=True,no_lint=False,dry_run=False)
            self.assertNotEqual(build_run(args),0);self.assertFalse(out.exists())
            spec['edits'][0]['expected_text']='old';edits.write_text(json.dumps(spec))
            self.assertEqual(build_run(args),0);self.assertEqual(Presentation(out).slides[0].shapes[0].text,'new')

    def test_master_font_namespace_resolved(self):
        p=Presentation();master=p.slide_master
        e=master._element.xpath('./p:txStyles/p:titleStyle/a:lvl1pPr/a:defRPr')[0];e.set('sz','3800')
        self.assertEqual(_master_txstyle_sz(master,'TITLE',0),38)

    def test_complete_8pt_chain_and_copied_or_late_attempt_rejected(self):
        # Synthetic render-shaped evidence only; no claim of visual acceptance.
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            def make(name,size,x,end):
                folder=root/name;folder.mkdir();p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540)
                slide=p.slides.add_slide(p.slide_layouts[6]);sh=slide.shapes.add_textbox(Pt(x),Pt(100),Pt(100),Pt(30));sh.name='chart-note:dense:label';sh.text='same necessary label';sh.text_frame.paragraphs[0].runs[0].font.size=Pt(size)
                path=folder/'snapshot.pptx';p.save(path)
                image=Image.new('RGB',(960,540),'white');ImageDraw.Draw(image).rectangle((x,100,x+70,120),fill='black');image.save(folder/'slide-01.png')
                rec={'slide':1,'shape_id':sh.shape_id,'shape':sh.name,'audit_path':str(sh.shape_id),'kind':'text','text':sh.text,'run_font_sizes_pt':[size], 'left_pt':x,'top_pt':100,'width_pt':100,'height_pt':30,'bound_left_pt':x,'bound_top_pt':100,'bound_width_pt':70,'bound_height_pt':20}
                report={'schema':'render-evidence/v2','sha256':sha(path),'snapshot':str(path),'slide_count':1,'status':'failed_preflight' if size==10 else 'rendered_pending_visual_review','issues':[{'reason':'synthetic conflict'}] if size==10 else [],'raster_hashes':{'1':sha(folder/'slide-01.png')},'text_records':[rec],'chart_geometry_records':[],'manual_review':[{'slide':1,'audit_path':'master-layout','review_id':'m1','reason':'master_layout_visual_review'}],'render_finished_utc':end,'render_started_utc':'2026-10-18T12:00:00+00:00'}
                (folder/'audit.json').write_text(json.dumps(report))
                return path,folder,report,rec
            a,ad,ar,at=make('a',10,20,'2026-10-18T10:00:00+00:00')
            b,bd,br,bt=make('b',10,25,'2026-10-18T11:00:00+00:00')
            p,d,render,r=make('final',8,25,'2026-10-18T12:01:00+00:00')
            item={'exception_id':'local-1','slide':1,'audit_path':r['audit_path'],'content_role':'complex_graph_auxiliary','justification':'Synthetic contract only','final_visual_observation':'Synthetic observation','observed_10pt_conflict':'Synthetic conflict','content_must_remain':'Necessary label','layout_and_split_analysis':'Synthetic required-field test','attempts':[]}
            for path,folder,report,text in ((a,ad,ar,at),(b,bd,br,bt)):
                item['attempts'].append({'pptx':str(path),'render_dir':str(folder),'slide':1,'audit_path':text['audit_path'],'observed_conflict':'Synthetic region conflict','layout_action':'Move target box','crop_box_px':[0,0,960,540],'crop_path':str(folder/'slide-01.png')})
            self.assertEqual(review_fonts(p,render,{'font_exceptions':[item]}),[])
            bad=copy.deepcopy(item);bad['attempts'][1]=copy.deepcopy(bad['attempts'][0])
            self.assertTrue(review_fonts(p,render,{'font_exceptions':[bad]}))
            br['render_finished_utc']='2026-10-18T13:00:00+00:00';(bd/'audit.json').write_text(json.dumps(br))
            self.assertTrue(review_fonts(p,render,{'font_exceptions':[item]}))

    def test_eccentric_ellipse_uses_shortest_distance(self):
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);e=s.shapes.add_shape(MSO_SHAPE.OVAL,Pt(0),Pt(0),Pt(200),Pt(20))
        e.fill.solid();e.fill.fore_color.rgb=RGBColor(20,30,40)
        x=100+100/math.sqrt(2);y=10+10/math.sqrt(2)
        normal=[(x-100)/10000,(y-10)/100];length=math.hypot(*normal)
        for d in (.4,-.4,.8):
            pt=[x+d*normal[0]/length,y+d*normal[1]/length]
            self.assertAlmostEqual(boundary_distance(pt,e),abs(d),places=7)

    def test_external_transparent_containers_rejected(self):
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);ids=[];containers={}
        for y in (100,200):
            text=s.shapes.add_textbox(Pt(50),Pt(y),Pt(100),Pt(20));text.text='label';ids.append(text.shape_id)
            c=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(40),Pt(y-10),Pt(130),Pt(40));c.fill.solid();c.fill.fore_color.rgb=RGBColor(0,0,0);c.line.fill.background()
            alpha=OxmlElement('a:alpha');alpha.set('val','0');c._element.xpath('./p:spPr/a:solidFill/*')[0].append(alpha);containers[str(text.shape_id)]=c.shape_id
        f={'shape_ids':ids};r={'relation_contract':{'kind':'separate_visible_containers','shape_ids':ids,'containers':containers,'visible_relation_explanation':'Transparent proxy must fail'}}
        self.assertIsNotNone(validate_region_contract(f,r,s,1))

    def test_importing_fixture_does_not_write_files(self):
        with tempfile.TemporaryDirectory() as td:
            cmd=[sys.executable,'-c',f"import sys;sys.path.insert(0,{str(Path(__file__).parent)!r});import test_complex_visuals"]
            result=subprocess.run(cmd,cwd=td,capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr);self.assertEqual(list(Path(td).iterdir()),[])

    def test_full_and_post_signoff_scope(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);before=root/'before.pptx';after=root/'after.pptx'
            p=Presentation()
            for _ in range(2):s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_textbox(Pt(50),Pt(60),Pt(100),Pt(30)).text='original'
            p.save(before);p.slides[0].shapes[0].text='corrected';p.save(after)
            self.assertEqual(impacted_pages(before,after)[0],[1])
            self.assertRaises(ValueError,check,after,{'mode':'full','sha256':sha(after),'checked_slides':[1]})
            self.assertEqual(check(after,{'mode':'full','sha256':sha(after),'checked_slides':[1,2]})['mode'],'full')
            baseline=root/'baseline.json';record={'status':'signed_off','mode':'full','unresolved':[],'sha256':sha(before),'checked_slides':[1,2],'signed_at':'2026-10-18T10:00:00+00:00'}
            baseline.write_text(json.dumps(record))
            spec={'mode':'amendment','sha256':sha(after),'checked_slides':[1],'baseline_signoff':str(baseline),'baseline_signoff_sha256':sha(baseline),'baseline_pptx':str(before),'user_reported_issue':{'source':'user','reference':'conversation/user-issue','quote':'Fix page 1','reported_at':'2026-10-18T11:00:00+00:00'}}
            self.assertEqual(check(after,spec)['required_slides'],[1])
            for field,value in [('source','agent'),('reported_at','2026-10-18T09:00:00+00:00')]:
                bad=copy.deepcopy(spec);bad['user_reported_issue'][field]=value;self.assertRaises(ValueError,check,after,bad)
            for status in ('pending','evidence_complete_pending_visual_signoff'):
                baseline.write_text(json.dumps({**record,'status':status}));bad={**spec,'baseline_signoff_sha256':sha(baseline)};self.assertRaises(ValueError,check,after,bad)

    def test_shared_theme_change_affects_every_page(self):
        with tempfile.TemporaryDirectory() as td:
            a=Path(td)/'a.pptx';b=Path(td)/'b.pptx';p=Presentation()
            for _ in range(2):p.slides.add_slide(p.slide_layouts[6])
            p.save(a)
            with ZipFile(a) as z:parts={n:z.read(n) for n in z.namelist()}
            theme=next(n for n in parts if n.startswith('ppt/theme/') and n.endswith('.xml'));parts[theme]+=b'\n'
            with ZipFile(b,'w') as z:
                for n,data in parts.items():z.writestr(n,data)
            self.assertEqual(impacted_pages(a,b)[0],[1,2])

    def test_missing_8pt_evidence_and_8pt_table_fail(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'f.pptx';p=Presentation();p.slide_height=Pt(540);p.slides.add_slide(p.slide_layouts[6]);p.save(path)
            row={'slide':1,'audit_path':'2','text':'aux label','shape':'chart-note:dense:label','top_pt':100,'run_font_sizes_pt':[16,8]}
            self.assertTrue(review_fonts(path,{'text_records':[row]},{}))
            e={'exception_id':'e1','slide':1,'audit_path':'2','content_role':'complex_graph_auxiliary','justification':'Complex label','final_visual_observation':'Local view checked'}
            self.assertTrue(review_fonts(path,{'text_records':[row]},{'font_exceptions':[e]}))
            table={**row,'audit_path':'2/table/1/1','shape':None};e={**e,'audit_path':table['audit_path'],'content_role':'ordinary_data_table'}
            self.assertTrue(review_fonts(path,{'text_records':[table]},{'font_exceptions':[e]}))
            table['run_font_sizes_pt']=[10]
            self.assertEqual(review_fonts(path,{'text_records':[table]},{'font_exceptions':[e]}),[])

if __name__=='__main__':unittest.main()
