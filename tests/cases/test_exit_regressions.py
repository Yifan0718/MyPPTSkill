"""Isolated fault-injection fixtures, not rendering or signoff evidence."""
import sys,unittest,tempfile,json,subprocess,shutil,importlib.util
from pathlib import Path
from unittest.mock import patch
B=Path(__file__).resolve().parent
sys.dont_write_bytecode=True

import pptx_delivery_shell as shell
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
class Regression(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.d=Path(self.tmp.name)
  self.p=self.d/'input.pptx';p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);t=s.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1));t.text='Test';self.sid=t.shape_id;p.save(self.p)
  self.png=self.d/'first.png';Image.new('RGB',(960,540),'blue').save(self.png);self.proof=self.d/'proof.json';self.before=self.p.read_bytes()
 def finalize(self):return shell.finalize(self.p,shell.sha(self.before),self.png,self.proof)
 def test_proof_success(self):
  self.finalize();self.assertEqual([],shell.verify_shell(self.p,self.png));self.assertTrue(self.proof.exists())
 def test_proof_failure_restores_source(self):
  replace=shell.os.replace
  def fail(src,dst):
   if Path(dst)==self.proof:raise OSError('proof blocked')
   return replace(src,dst)
  with patch.object(shell.os,'replace',side_effect=fail),self.assertRaisesRegex(RuntimeError,'original PPTX restored'):self.finalize()
  self.assertEqual(self.before,self.p.read_bytes());self.assertFalse(self.proof.exists())
 def test_proof_failure_preserves_external_change(self):
  replace=shell.os.replace
  def fail(src,dst):
   if Path(dst)==self.proof:self.p.write_bytes(b'external edit');raise OSError('proof blocked')
   return replace(src,dst)
  with patch.object(shell.os,'replace',side_effect=fail),self.assertRaisesRegex(RuntimeError,'rollback refused or failed'):self.finalize()
  self.assertEqual(b'external edit',self.p.read_bytes())
 def test_proof_rollback_failure_propagates(self):
  replace=shell.os.replace;calls=[]
  def fail(src,dst):
   calls.append(dst)
   if len(calls)>1:raise OSError('commit/rollback blocked')
   return replace(src,dst)
  with patch.object(shell.os,'replace',side_effect=fail),self.assertRaisesRegex(RuntimeError,'rollback refused or failed'):self.finalize()
  self.assertNotEqual(self.before,self.p.read_bytes());self.assertFalse(self.proof.exists())
 def capacity(self,sid,font=18):
  inp=self.d/'detail.json';inp.write_text(json.dumps({'pages':[{'slide_number':1,'text_slots':[{'address':{'shape_id':sid},'editable':True,'font_size_pt':font,'max_chars':9999}]}]}));out=self.d/'out.json'
  r=subprocess.run([sys.executable,'-X','utf8','-B',str(B/'compute_capacity.py'),str(inp),str(self.p),'-o',str(out)],capture_output=True)
  return r,out
 def test_capacity_missing_slot_fails_without_output(self):
  r,out=self.capacity(99999);self.assertNotEqual(0,r.returncode);self.assertFalse(out.exists())
 def test_capacity_valid_slot_succeeds(self):
  r,out=self.capacity(self.sid);self.assertEqual(0,r.returncode,r.stderr);self.assertLess(json.loads(out.read_text())['pages'][0]['text_slots'][0]['max_chars'],9999)
 def test_capacity_absent_font_allows_heuristic(self):
  r,out=self.capacity(self.sid,None);self.assertEqual(0,r.returncode,r.stderr)
if __name__=='__main__':unittest.main(verbosity=2)
