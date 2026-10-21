"""Reproduce and guard the remaining paths from upstream issues 9 and 10."""
import contextlib,io,json,subprocess,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch,Mock
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import local_update_guard as guard
import render_slides as renderer
import compute_capacity

class UpstreamIssueRegressions(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.w=Path(self.temp.name)
 def root(self,name='live'):
  p=self.w/name;p.mkdir();(p/'VERSION').write_text('1.0');(p/'SKILL.md').write_text('fixture');return p
 def test_update_source_shapes_fail_cleanly_before_any_staging(self):
  for i,value in enumerate([None,{},[],{'update_source':None},{'update_source':''},{'update_source':42}]):
   root=self.root(str(i));(root/'updates.json').write_text(json.dumps(value));stage=self.w/('stage-'+str(i));err=io.StringIO()
   with patch.object(sys,'argv',['guard','prepare','--root',str(root),'--stage',str(stage)]),patch.object(guard.subprocess,'run') as launch,contextlib.redirect_stderr(err):code=guard.main()
   self.assertEqual(code,4);self.assertNotIn('Traceback',err.getvalue());self.assertFalse(stage.exists());self.assertFalse((root/'local-overrides').exists());launch.assert_not_called()
 def test_explicit_local_candidate_works_with_unset_source(self):
  root=self.root();(root/'updates.json').write_text(json.dumps({'update_source':None}));candidate=self.root('candidate');(candidate/'VERSION').write_text('1.1');stage=self.w/'stage'
  with patch.object(sys,'argv',['guard','prepare','--root',str(root),'--stage',str(stage),'--candidate',str(candidate)]),patch.object(guard.subprocess,'run') as launch,contextlib.redirect_stdout(io.StringIO()):code=guard.main()
  self.assertEqual(code,0);self.assertTrue((stage/'plan.json').is_file());launch.assert_not_called()
 def render_fixture(self,expected):
  p=Presentation();p.slides.add_slide(p.slide_layouts[6]);source=self.w/'deck.pptx';p.save(source);selected=[]
  def convert(cmd,**kw):
   if cmd[0]=='soffice':
    d=Path(cmd[cmd.index('--outdir')+1]);(d/'a-foreign.pdf').write_bytes(b'foreign')
    if expected:(d/'input.pdf').write_bytes(b'expected')
   else:
    selected.append(Path(cmd[4]).name);Image.new('RGB',(80,45),'white').save(Path(cmd[-1]).parent/'slide-1.png')
   return Mock(returncode=0)
  return source,selected,convert
 def test_extra_pdf_does_not_change_snapshot_pdf_selection(self):
  p,selected,convert=self.render_fixture(True)
  with patch.object(renderer,'_run_converter',side_effect=convert):images=renderer.render(p,self.w/'out')
  self.assertEqual(selected,['input.pdf']);self.assertEqual(len(images),1)
 def test_foreign_pdf_alone_never_reaches_rasterization(self):
  p,selected,convert=self.render_fixture(False)
  with patch.object(renderer,'_run_converter',side_effect=convert):
   with self.assertRaisesRegex(RuntimeError,'expected input.pdf'):renderer.render(p,self.w/'out')
  self.assertEqual(selected,[])
 def capacity_fixture(self,missing_paragraph=False,bad_json=False):
  p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);sh=s.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1));sh.text_frame.clear()
  if missing_paragraph:
   for pa in list(sh.text_frame._txBody.p_lst):sh.text_frame._txBody.remove(pa)
  pptx=self.w/'input.pptx';p.save(pptx);detail=self.w/'detail.json';out=self.w/'result.json';out.write_text('preserved')
  detail.write_text('[' if bad_json else json.dumps({'pages':[{'slide_number':1,'text_slots':[{'address':{'shape_id':sh.shape_id,'paragraph':0},'role':'body','editable':True}]}]}))
  run=subprocess.run([sys.executable,'-B','-X','utf8',str(Path(compute_capacity.__file__)),str(detail),str(pptx),'-o',str(out)],capture_output=True,encoding='utf8',timeout=30)
  return run,out
 def test_normal_empty_text_frame_still_computes_capacity(self):
  run,out=self.capacity_fixture();self.assertEqual(run.returncode,0,run.stderr);self.assertIn('pages',json.loads(out.read_text()))
 def test_missing_paragraph_cli_returns_structured_failure(self):
  run,out=self.capacity_fixture(missing_paragraph=True);self.assertEqual(run.returncode,3);self.assertNotIn('Traceback',run.stderr);self.assertEqual(json.loads(run.stderr)['status'],'failed');self.assertEqual(out.read_text(),'preserved')
 def test_malformed_capacity_json_preserves_previous_output(self):
  run,out=self.capacity_fixture(bad_json=True);self.assertEqual(run.returncode,3);self.assertNotIn('Traceback',run.stderr);self.assertIn('JSONDecodeError',json.loads(run.stderr)['reason']);self.assertEqual(out.read_text(),'preserved')

if __name__=='__main__':unittest.main()
