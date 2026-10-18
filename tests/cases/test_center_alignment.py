"""Synthetic pixel/native inventory regressions, never real rendering/signoff."""
import sys
sys.dont_write_bytecode = True
import unittest
from unittest.mock import patch
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement
import center_alignment as ca


def fixture(offset=0, multiple=False, multiline=False, texture=False, shadow=False, extra=False, grouped=False):
    p=Presentation();p.slide_width=Pt(400);p.slide_height=Pt(400)
    s=p.slides.add_slide(p.slide_layouts[6]);data=ChartData();data.categories=['A','B'];data.add_series('S',[1,2])
    chart=s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT,Pt(25),Pt(25),Pt(350),Pt(350),data)
    im=Image.new('RGB',(400,400),'white');d=ImageDraw.Draw(im)
    if texture:
        for y in range(0,400,4):d.line((0,y,399,y),fill=(195,205,215))
    if shadow:d.ellipse((60,60,361,361),fill=(180,180,180))
    d.ellipse((49,49,350,350),fill='#304774');d.ellipse((99,99,300,300),fill='white')
    font=ImageFont.load_default(size=26)
    texts=['12.52','kW'] if multiple else ['12.52\nkW' if multiline else '12.52 kW']
    raw=[d.multiline_textbbox((0,0),text,font=font,spacing=8,align='center') for text in texts]
    heights=[b[3]-b[1] for b in raw];total=sum(heights)+12*(len(texts)-1)
    y=200-total/2+offset;regions=[];members=[]
    group=s.shapes.add_group_shape() if grouped else None
    for text,b,h in zip(texts,raw,heights):
        w=b[2]-b[0];x=200-w/2
        d.multiline_text((x-b[0],y-b[1]),text,font=font,spacing=8,align='center',fill='#304774')
        region=[int(x)-2,int(y)-2,int(x+w)+3,int(y+h)+3]
        container=group.shapes if group is not None else s.shapes
        member=container.add_textbox(Pt(region[0]),Pt(region[1]),Pt(region[2]-region[0]),Pt(region[3]-region[1]));member.text=text
        member.name='background:renaming-does-not-exempt-center-text'
        members.append(member);regions.append(region);y+=h+12
    if extra:d.rectangle((130,130,140,140),fill='red')
    render={'sha256':'SYNTHETIC','raster_hashes':{'1':'SYNTHETIC'},'text_records':[],
            'chart_geometry_records':[{'slide':1,'shape_id':chart.shape_id,'plot_bounds_pt':[70,70,250,250]}]}
    return p,s,chart,members,render,im,regions


class CenterAlignmentTests(unittest.TestCase):
    def scan(self,**kwargs):
        p,s,c,m,r,im,regions=fixture(**kwargs)
        found=ca.scan_doughnut_centers(p,1,r,im)
        self.assertEqual(len(found),1)
        return found[0],(p,s,c,m,r,im,regions)

    def test_centered_still_requires_local_and_full_page(self):
        f,_=self.scan()
        self.assertEqual(f['reason'],ca.REVIEW)
        self.assertEqual(f['measurement_status'],'measured_aligned')
        self.assertTrue(f['local_and_full_page_review_required'])

    def test_actual_ink_high_is_misaligned(self):
        f,_=self.scan(offset=-14)
        self.assertEqual(f['reason'],ca.MISALIGNED)
        self.assertLess(f['delta_px'][1],-12)

    def test_exact_tolerance_is_flagged_consistently(self):
        _,_,_,_,_,im,_=fixture()
        measured={'host_bbox_px':[50,50,350,350],'hole_bbox_px':[100,100,300,300],
                  'text_group_bbox_px':[170,192,230,212],'member_ink_bounds_px':[[170,192,230,212]],
                  'line_ink_bounds_px':[[170,192,230,212]]}
        with patch.object(ca,'_segment',return_value=(measured,None)):
            f=ca.measure_visible_hole(im,[25,25,375,375],[[160,180,240,220]],1)
        self.assertEqual(f['delta_px'],[0,2]);self.assertEqual(f['measurement_status'],'measured_misalignment')

    def test_plot_center_and_frame_are_not_used_as_visible_center(self):
        f,_=self.scan()
        self.assertEqual(f['hole_center_px'],[200,200])
        self.assertNotEqual(f['hole_center_px'],[160,160])

    def test_multiple_members_are_jointly_centered(self):
        f,_=self.scan(multiple=True)
        self.assertEqual(f['measurement_status'],'measured_aligned')
        self.assertEqual(len(f['text_ids']),2)
        self.assertEqual(len(f['member_ink_bounds_px']),2)
        self.assertGreater(len(f['line_ink_bounds_px']),1)

    def test_multiline_is_jointly_centered(self):
        f,_=self.scan(multiline=True)
        self.assertEqual(f['measurement_status'],'measured_aligned')
        self.assertGreater(len(f['line_ink_bounds_px']),1)

    def test_multiple_members_high_rejected(self):
        f,_=self.scan(multiple=True,offset=-14)
        self.assertEqual(f['reason'],ca.MISALIGNED)

    def test_unknown_texture_is_explicit_manual(self):
        f,_=self.scan(texture=True)
        self.assertEqual(f['measurement_status'],'manual_review_required')

    def test_shadow_is_not_declared_correct(self):
        f,_=self.scan(shadow=True)
        self.assertEqual(f['measurement_status'],'manual_review_required')

    def test_unassigned_ink_preserves_known_host(self):
        f,_=self.scan(extra=True)
        self.assertEqual(f['measurement_status'],'manual_review_required')
        self.assertIn('host_bbox_px',f['measurement_known_fields'])
        self.assertIn('hole_bbox_px',f['measurement_known_fields'])
        self.assertNotIn('text_group_bbox_px',f['measurement_known_fields'])

    def test_nested_members_are_manual_with_top_level_owner_ids(self):
        f,(_,s,_,m,_,_,_)=self.scan(grouped=True,multiple=True)
        self.assertEqual(f['measurement_status'],'manual_review_required')
        self.assertTrue(set(f['shape_ids'])<={sh.shape_id for sh in s.shapes})
        self.assertTrue(all('/group/' in p for p in f['text_audit_paths']))
        self.assertEqual(set(f['text_leaf_ids']),{x.shape_id for x in m})

    def test_optional_dependency_absent_is_manual_not_import_failure(self):
        with patch.object(ca,'cv2',None):
            f,_=self.scan()
        self.assertEqual(f['measurement_status'],'manual_review_required')

    def test_native_shadow_effect_requires_review(self):
        p,s,c,m,r,im,regions=fixture()
        effect=OxmlElement('a:effectLst');effect.append(OxmlElement('a:outerShdw'))
        properties=OxmlElement('c:spPr');properties.append(effect);c.chart._element.append(properties)
        found=ca.scan_doughnut_centers(p,1,r,im)
        self.assertEqual(found[0]['measurement_status'],'manual_review_required')

    def test_no_overlay_text_has_no_center_relationship(self):
        p,s,c,m,r,im,regions=fixture()
        for member in m:s.shapes._spTree.remove(member._element)
        self.assertEqual(ca.scan_doughnut_centers(p,1,r,im),[])

    def test_text_ink_outside_hole_is_not_silently_truncated(self):
        p,s,c,m,r,im,regions=fixture()
        ImageDraw.Draw(im).rectangle((195,363,205,370),fill='#304774')
        regions[0][3]=374
        result=ca.measure_visible_hole(im,[25,25,375,375],regions,1)
        self.assertEqual(result['measurement_status'],'manual_review_required')
        self.assertIn('truncated',result['measurement_limitation'])


if __name__=='__main__':unittest.main(verbosity=2)
