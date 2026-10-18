"""Numeric axes must not absorb a full-width conclusion or exempt real offsets."""
import hashlib,tempfile,unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_CONNECTOR
from audit_layout_coverage import audit

class AxisRegionScope(unittest.TestCase):
    def scan(self,*,grid=True,values=('0','100','200'),summary_width=850,narrow_shift=0):
        with tempfile.TemporaryDirectory() as td:
            p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540)
            s=p.slides.add_slide(p.slide_layouts[6]);records=[];tick_ids=[]
            for i,value in enumerate(values):
                y=100+i*60;sh=s.shapes.add_textbox(Pt(35),Pt(y),Pt(36),Pt(25));sh.text=value;tick_ids.append(sh.shape_id)
                ink=35+(narrow_shift if i==2 else 0)
                records.append({'slide':1,'shape_id':sh.shape_id,'shape':sh.name,'kind':'text','text':value,'font_pt':12,'paragraph_alignment':1,'left_pt':35,'top_pt':y,'width_pt':36,'height_pt':25,'bound_left_pt':ink,'bound_top_pt':y,'bound_width_pt':20,'bound_height_pt':15})
                if grid:s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(100),Pt(y+7.5),Pt(900),Pt(y+7.5))
            sh=s.shapes.add_textbox(Pt(36),Pt(300),Pt(summary_width),Pt(30));sh.text='A whole-page conclusion'
            records.append({'slide':1,'shape_id':sh.shape_id,'shape':sh.name,'kind':'text','text':sh.text,'font_pt':16,'paragraph_alignment':1,'left_pt':36,'top_pt':300,'width_pt':summary_width,'height_pt':30,'bound_left_pt':39,'bound_top_pt':300,'bound_width_pt':300,'bound_height_pt':20})
            path=Path(td)/'fixture.pptx';p.save(path)
            result=audit(path,{'sha256':hashlib.sha256(path.read_bytes()).hexdigest().upper(),'slide_count':1,'status':'rendered_pending_visual_review','issues':[],'text_records':records})
            return result,tick_ids,sh.shape_id

    def test_axis_and_conclusion_remain_reviewed_separately(self):
        result,ticks,summary=self.scan()
        region=[f for f in result['findings'] if f['reason']=='axis_stack_spanning_region_review']
        self.assertEqual(len(region),1)
        self.assertEqual(region[0]['axis_tick_ids'],ticks)
        self.assertEqual(region[0]['spanning_text_id'],summary)
        self.assertGreaterEqual(len(region[0]['matched_grid_ids']),3)

    def test_numeric_kpis_without_grid_do_not_gain_axis_exemption(self):
        for kwargs in ({'grid':False},{'values':('0','100','300')},{'summary_width':500}):
            result,_,_=self.scan(**kwargs)
            self.assertFalse(any(f['reason']=='axis_stack_spanning_region_review' for f in result['findings']))

    def test_tick_stack_itself_still_detects_offset(self):
        result,ticks,summary=self.scan(narrow_shift=12)
        self.assertTrue(any(f['reason']=='rendered_column_text_anchor_mismatch' and set(f['shape_ids'])==set(ticks) for f in result['findings']))

if __name__=='__main__':unittest.main()
