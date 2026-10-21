"""Independent whole-page whitespace sweep and rendered-text alignment candidates.

python audit_layout_coverage.py candidate.pptx render/audit.json coverage.json
Exit 2: review findings; exit 3: invalid/missing evidence; exit 0: no findings.
This is a diagnostic sweep, NOT visual signoff. Never silently ignore findings.
"""
from pathlib import Path
import argparse
import hashlib
import json
import math
import statistics
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from connector_integrity import inventory as connection_inventory
from center_alignment import scan_doughnut_centers
from render_evidence import native_inventory
from audit_visual_components import scan as scan_components


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def ink_bounds(record, image, pixels_per_pt):
    """Conservative raster measurement for an isolated text-range crop.
    Return None for uncertain backgrounds/thresholds, never fabricate a pass.
    DrawingML bounds are not glyph ink bounds (side bearings/fallback fonts).
    """
    from PIL import ImageChops
    pad=2*pixels_per_pt
    x0=max(0,math.floor(record['bound_left_pt']*pixels_per_pt-pad))
    y0=max(0,math.floor(record['bound_top_pt']*pixels_per_pt))
    x1=min(image.width,math.ceil((record['bound_left_pt']+record['bound_width_pt'])*pixels_per_pt+pad))
    y1=min(image.height,math.ceil((record['bound_top_pt']+record['bound_height_pt'])*pixels_per_pt))
    if x1<=x0 or y1<=y0:return None
    crop=image.crop((x0,y0,x1,y1)).convert('RGB')
    corners=[crop.getpixel((x,y)) for x,y in ((0,0),(crop.width-1,0),(0,crop.height-1),(crop.width-1,crop.height-1))]
    bg=tuple(int(statistics.median(c[k] for c in corners)) for k in range(3))
    if any(max(abs(c[k]-bg[k]) for k in range(3))>20 for c in corners):return None
    from PIL import Image
    diff=ImageChops.difference(crop,Image.new('RGB',crop.size,bg))
    channels=diff.split();contrast=ImageChops.lighter(ImageChops.lighter(channels[0],channels[1]),channels[2])
    boxes=[]
    for threshold in (30,70,110):
        bounds=contrast.point(lambda v:255 if v>threshold else 0).getbbox()
        if bounds is None:return None
        boxes.append(bounds)
    if any(max(b[k] for b in boxes)-min(b[k] for b in boxes)>1 for k in range(4)):return None
    b=boxes[0]
    return {'bounds_px':[x0+b[0],y0+b[1],x0+b[2],y0+b[3]],
            'bounds_pt':[(x0+b[0])/pixels_per_pt,(y0+b[1])/pixels_per_pt,(x0+b[2])/pixels_per_pt,(y0+b[3])/pixels_per_pt],
            'crop_px':[x0,y0,x1,y1],'thresholds':[30,70,110]}


def segment_hits_box(x1, y1, x2, y2, box):
    """Liang-Barsky clipping; also handles horizontal/vertical connectors."""
    left,top,right,bottom=box
    dx,dy=x2-x1,y2-y1;low,high=0.0,1.0
    for p,q in ((-dx,x1-left),(dx,right-x1),(-dy,y1-top),(dy,bottom-y1)):
        if abs(p)<1e-12:
            if q<0:return False
        else:
            t=q/p
            if p<0:low=max(low,t)
            else:high=min(high,t)
            if low>high:return False
    return True


def sweep(objects, axis):
    """Merge overlapping projections, exposing EVERY gap between occupied bands."""
    lo, hi = (0, 2) if axis == 'horizontal' else (1, 3)
    bands = []
    for obj in sorted(objects, key=lambda x: x['bounds'][lo]):
        start, end = obj['bounds'][lo], obj['bounds'][hi]
        if bands and start <= bands[-1]['end'] + .25:
            bands[-1]['end'] = max(end, bands[-1]['end'])
            bands[-1]['shape_ids'].append(obj['id'])
        else:
            bands.append({'start': start, 'end': end, 'shape_ids': [obj['id']]})
    return bands, [bands[i+1]['start'] - bands[i]['end'] for i in range(len(bands)-1)]


def audit(pptx, render):
    # COM can return JSON null for table-cell names. Preserve every record;
    # normalize only the optional display name, never omit unnamed cells.
    render={**render,'text_records':[{**r,'shape':r.get('shape') or ''} for r in render.get('text_records',[])]}
    digest = sha(pptx)
    prs = Presentation(pptx)
    if render.get('sha256', '').upper() != digest or render.get('slide_count') != len(prs.slides):
        raise ValueError('Render evidence hash/page count differs from current PPTX')
    # A completed renderer can report layout preflight failures. Keep those raw
    # failures intact and continue diagnostic coverage; final verification still
    # validates render evidence and every preflight disposition before signoff.
    if render.get('status') not in ('rendered_pending_visual_review','failed_preflight'):
        raise ValueError('Rendering must complete before diagnostic coverage can run')
    center_inventory = native_inventory(prs)
    scale = prs.slide_height / 12700 / 540
    width = prs.slide_width / 12700
    findings, pages = [], []
    def flag(n, kind, detail):
        findings.append({'id': f'p{n}-{kind}-{len(findings)+1}', 'slide': n,
                         'reason': kind, **detail})
    for n, slide in enumerate(prs.slides, 1):
        for candidate in scan_components(slide,n,render,scale):
            kind=candidate.pop("reason");flag(n,kind,candidate)
        for connection in connection_inventory(slide):
            flag(n, 'connector_endpoint_review_required', connection)
        raster=None;pixel_bounds={};raster_error=None;raster_path=None
        try:
            from PIL import Image
            if not render.get('snapshot'):raise ValueError('No render snapshot path')
            raster_path=Path(render['snapshot']).parent/f'slide-{n:02}.png'
            with Image.open(raster_path) as loaded:raster=loaded.convert('RGB')
        except (ImportError,OSError,ValueError) as exc:raster_error=str(exc)
        for centered in scan_doughnut_centers(prs,n,render,raster,center_inventory):
            kind=centered.pop('reason')
            flag(n,kind,centered)
        optical_inventory=[]
        for record in render.get('text_records',[]):
            if record.get('slide')!=n:continue
            needed=('bound_left_pt','bound_top_pt','bound_width_pt','bound_height_pt')
            measured=None
            if raster is not None and all(k in record for k in needed):
                measured=ink_bounds(record,raster,raster.width/width)
            optical_inventory.append({'shape_id':record.get('shape_id'),'shape':record.get('shape'),
                'kind':record.get('kind'),'text':record.get('text'),'ink':measured,
                'status':'measured_pending_visual_review' if measured else 'local_visual_review_required'})
        # Mandatory optical review for mixed typography, even when geometric
        # tolerances pass. Include headings/subheadings excluded from body scans.
        all_texts=[r for r in render.get('text_records',[]) if r.get('slide')==n and r.get('kind')=='text'
                   and not r.get('shape','').startswith('page-number:')]
        for i,a in enumerate(all_texts):
            for b in all_texts[i+1:]:
                styles_differ=any(a.get(k)!=b.get(k) for k in ('font_pt','font_name','font_bold','font_italic'))
                heading_pair=(a.get('shape','').startswith('title:') and b.get('shape','').startswith('subtitle:')) or (b.get('shape','').startswith('title:') and a.get('shape','').startswith('subtitle:'))
                if not styles_differ and not heading_pair:continue
                yw=min(a['top_pt']+a['height_pt'],b['top_pt']+b['height_pt'])-max(a['top_pt'],b['top_pt'])
                xgap=max(a['left_pt'],b['left_pt'])-min(a['left_pt']+a['width_pt'],b['left_pt']+b['width_pt'])
                xw=min(a['left_pt']+a['width_pt'],b['left_pt']+b['width_pt'])-max(a['left_pt'],b['left_pt'])
                ygap=max(a['top_pt'],b['top_pt'])-min(a['top_pt']+a['height_pt'],b['top_pt']+b['height_pt'])
                row=yw>=.5*min(a['height_pt'],b['height_pt']) and 0<=xgap<=width*.3
                column=xw>=.8*min(a['width_pt'],b['width_pt']) and 0<=ygap<=max(a['font_pt'],b['font_pt'],12*scale)*2
                if row or column:
                    flag(n,'mixed_typography_optical_review',{'shape_ids':[a.get('shape_id'),b.get('shape_id')],
                         'texts':[a.get('text'),b.get('text')],'font_sizes_pt':[a['font_pt'],b['font_pt']],
                         'relationship':'row' if row else 'stack','reason_detail':'Inspect actual glyph edges/center or baseline locally, then return to full page even if coordinates match.'})
        objects, excluded = [], []
        for s in slide.shapes:
            # Exclusion is narrow and visible in the report, never by user-picked IDs.
            role = s.name.split(':', 1)[0]
            if role in ('background', 'page-number', 'title', 'subtitle') or s.shape_type == MSO_SHAPE_TYPE.LINE:
                excluded.append({'id': s.shape_id, 'name': s.name, 'reason': 'title/footer/background/connector'})
                if role in ('background', 'title', 'subtitle'):
                    flag(n, 'excluded_role_requires_visual_verification', {'shape_id':s.shape_id,'shape':s.name})
                continue
            if s.width <= 0 or s.height <= 0:
                excluded.append({'id': s.shape_id, 'name': s.name, 'reason': 'zero_area'})
                continue
            objects.append({'id': s.shape_id, 'name': s.name,
                            'bounds': [s.left/12700, s.top/12700, (s.left+s.width)/12700, (s.top+s.height)/12700]})
            if s.shape_type == MSO_SHAPE_TYPE.GROUP or s.has_chart or s.has_table or s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                flag(n, 'internal_layout_requires_review', {'shape_id': s.shape_id, 'shape': s.name})
        scans = []
        # Global vertical bands plus each occupied row's horizontal scan: a full-width
        # conclusion cannot hide column gaps higher up the page.
        bands, gaps = sweep(objects, 'vertical')
        scans.append({'axis': 'vertical', 'bands': bands, 'gaps_pt': gaps})
        for band in bands:
            row = [x for x in objects if x['id'] in band['shape_ids']]
            hb, hg = sweep(row, 'horizontal')
            scans.append({'axis': 'horizontal', 'row_shape_ids': band['shape_ids'], 'bands': hb, 'gaps_pt': hg})
        for scan in scans:
            gs = scan['gaps_pt']
            if any(g < 12*scale-.25*scale for g in gs):
                flag(n, 'narrow_occupied_band_gap', {'axis':scan['axis'],'gaps_pt':gs,'bands':scan['bands']})
            if len(gs) >= 2:
                target = sum(gs)/len(gs)
                if any(abs(g-target) > max(3*scale, target*.2) for g in gs):
                    flag(n, 'uneven_occupied_band_gaps', {'axis': scan['axis'], 'gaps_pt': gs, 'bands': scan['bands']})
        if len(bands) < 2:
            flag(n, 'body_band_scan_inconclusive', {'reason_detail': 'Overlapping containers or a single module; inspect internal whitespace visually'})
        texts = [r for r in render.get('text_records', []) if r.get('slide') == n and r.get('kind') == 'text'
                 and not r.get('shape', '').startswith(('title:', 'subtitle:', 'page-number:'))]
        expected={s.shape_id for s in slide.shapes if s.has_text_frame and s.text.strip()
                  and not s.name.startswith(('title:', 'subtitle:', 'page-number:'))}
        available={r.get('shape_id') for r in texts}
        if not expected.issubset(available):
            raise ValueError(f'Slide {n}: rendered text records missing for shapes {sorted(expected-available)}')
        for r in texts:
            for key in ('bound_left_pt', 'bound_top_pt', 'bound_width_pt', 'bound_height_pt'):
                if key not in r or not math.isfinite(float(r[key])):
                    raise ValueError(f'Slide {n}: missing actual rendered text bounds; rerender with current renderer')
            if raster is not None:
                pixel_bounds[r.get('shape_id')]=ink_bounds(r,raster,raster.width/width)
        # Geometry candidates deliberately err toward review. Bounding rectangles
        # are not glyph masks; every candidate still needs inspection of the ink.
        for i,a in enumerate(texts):
            for b in texts[i+1:]:
                iw=min(a['bound_left_pt']+a['bound_width_pt'],b['bound_left_pt']+b['bound_width_pt'])-max(a['bound_left_pt'],b['bound_left_pt'])
                ih=min(a['bound_top_pt']+a['bound_height_pt'],b['bound_top_pt']+b['bound_height_pt'])-max(a['bound_top_pt'],b['bound_top_pt'])
                if iw>.5*scale and ih>.5*scale:
                    flag(n,'rendered_text_rectangles_overlap',{'shape_ids':[a.get('shape_id'),b.get('shape_id')],
                        'texts':[a['text'],b['text']],'intersection_width_pt':iw,'intersection_height_pt':ih})
        for connector in slide.shapes:
            if connector.shape_type!=MSO_SHAPE_TYPE.LINE:continue
            try:
                coords=[getattr(connector,key)/12700 for key in ('begin_x','begin_y','end_x','end_y')]
                # Elbow/curved paths cannot be approximated as a diagonal for clearance.
                geometry=connector._element.xpath('./p:spPr/a:prstGeom')
                if not geometry or geometry[0].get('prst') not in ('line','straightConnector1'):
                    flag(n,'nonstraight_connector_path_requires_review',{'shape_id':connector.shape_id})
                    continue
            except (AttributeError,ValueError,TypeError):
                flag(n,'connector_geometry_unavailable',{'shape_id':connector.shape_id});continue
            for r in texts:
                pad=1*scale
                box=[r['bound_left_pt']-pad,r['bound_top_pt']-pad,
                     r['bound_left_pt']+r['bound_width_pt']+pad,r['bound_top_pt']+r['bound_height_pt']+pad]
                if segment_hits_box(*coords,box):
                    flag(n,'connector_crosses_or_touches_text',{'connector_id':connector.shape_id,
                         'text_id':r.get('shape_id'),'text':r['text'],'segment_pt':coords,'text_clearance_box_pt':box})
        pairs = []
        for i, a in enumerate(texts):
            for b in texts[i+1:]:
                # Similar one-line neighbours whose containing boxes share a row.
                # Different font sizes still need visual optical alignment review.
                if min(a['height_pt'], b['height_pt']) <= 0:
                    continue
                overlap = min(a['top_pt']+a['height_pt'], b['top_pt']+b['height_pt'])-max(a['top_pt'], b['top_pt'])
                horizontal_gap = max(a['left_pt'], b['left_pt'])-min(a['left_pt']+a['width_pt'], b['left_pt']+b['width_pt'])
                if overlap < .5*min(a['height_pt'], b['height_pt']) or not 0 <= horizontal_gap <= width*.3:
                    continue
                if any(r['font_pt'] <= 0 or r['bound_height_pt'] > r['font_pt']*1.6 for r in (a,b)):
                    flag(n,'multiline_or_mixed_font_row_alignment_requires_review',{
                        'shape_ids':[a.get('shape_id'),b.get('shape_id')],
                        'reason_detail':'Choose and measure actual first-line or whole-block alignment; unknown/multiline is not a pass.'})
                    continue
                top_delta = abs(a['bound_top_pt']-b['bound_top_pt'])
                center_delta = abs(a['bound_top_pt']+a['bound_height_pt']/2-b['bound_top_pt']-b['bound_height_pt']/2)
                pair = {'shape_ids': [a.get('shape_id'), b.get('shape_id')], 'shapes': [a['shape'], b['shape']],
                        'top_delta_pt': top_delta, 'center_delta_pt': center_delta}
                pairs.append(pair)
                # Do not opportunistically accept top alignment when the visible
                # centers disagree (e.g. a large numeral beside a smaller heading).
                if center_delta > 2*scale:
                    flag(n, 'rendered_row_text_misaligned', pair)
                ink_a=pixel_bounds.get(a.get('shape_id'));ink_b=pixel_bounds.get(b.get('shape_id'))
                if ink_a and ink_b:
                    ab,bb=ink_a['bounds_px'],ink_b['bounds_px']
                    ink_delta=abs((ab[1]+ab[3])/2-(bb[1]+bb[3])/2)
                    threshold=max(2,.75*scale*raster.width/width)
                    if ink_delta>=threshold:
                        flag(n,'rendered_ink_row_misaligned',{'shape_ids':pair['shape_ids'],
                             'ink_bounds_px':[ab,bb],'center_delta_px':ink_delta,'threshold_px':threshold})
                else:
                    flag(n,'glyph_ink_alignment_requires_review',{'shape_ids':pair['shape_ids'],
                         'reason_detail':'Ink comparison unavailable; coordinate agreement cannot authorize optical alignment'})
        # Vertical stacks: compare actual visible left edges, not textbox edges.
        # These are candidates because chart data labels can legitimately move.
        vertical_pairs=[]
        text_shapes={sh.shape_id:sh for sh in slide.shapes if sh.has_text_frame}
        def anchor_kind(r):
            if 'paragraph_alignment' in r:
                return {1:'left',2:'center',3:'right'}.get(r['paragraph_alignment'],'mixed')
            sh=text_shapes.get(r.get('shape_id'))
            if sh is None:return 'unknown'
            if sh.is_placeholder and any(p.alignment is None for p in sh.text_frame.paragraphs if p.text.strip()):return 'unknown'
            kinds={p.alignment or PP_ALIGN.LEFT for p in sh.text_frame.paragraphs if p.text.strip()}
            if len(kinds)!=1:return 'mixed'
            return {PP_ALIGN.LEFT:'left',PP_ALIGN.CENTER:'center',PP_ALIGN.RIGHT:'right'}.get(next(iter(kinds)),'mixed')
        def anchor_position(r,kind):
            return r['left_pt']+({'left':0,'center':.5,'right':1}.get(kind,0))*r['width_pt']
        relation_pairs=[]
        ordered=sorted(texts,key=lambda r:r['bound_top_pt'])
        for i,a in enumerate(ordered):
            compatible=[]
            uncertain=[]
            for b in ordered[i+1:]:
                gap=b['bound_top_pt']-a['bound_top_pt']-a['bound_height_pt']
                xover=min(a['left_pt']+a['width_pt'],b['left_pt']+b['width_pt'])-max(a['left_pt'],b['left_pt'])
                comparable_widths=max(a['width_pt'],b['width_pt'])<=3*min(a['width_pt'],b['width_pt'])
                ak,bk=anchor_kind(a),anchor_kind(b)
                delta=abs(anchor_position(a,ak)-anchor_position(b,bk))
                shared_anchor=ak==bk and ak in ('left','center','right') and delta<=12*scale
                if (comparable_widths or shared_anchor) and 0<=gap<=100*scale and xover>=.8*min(a['width_pt'],b['width_pt']):
                    if ak==bk and ak in ('left','center','right') and delta<=12*scale:
                        compatible.append((gap,b))
                    else:
                        uncertain.append({'shape_ids':[a.get('shape_id'),b.get('shape_id')],
                            'texts':[a['text'],b['text']],'native_anchors':[ak,bk],
                            'frame_bounds_pt':[[r['left_pt'],r['top_pt'],r['width_pt'],r['height_pt']] for r in (a,b)],
                            'anchor_delta_pt':delta,'vertical_gap_pt':gap})
            if compatible:
                gap,b=min(compatible,key=lambda x:x[0])
                left_delta=abs(a['bound_left_pt']-b['bound_left_pt'])
                pair={'shape_ids':[a.get('shape_id'),b.get('shape_id')], 'shapes':[a['shape'],b['shape']],
                      'left_delta_pt':left_delta,'vertical_gap_pt':gap,'default_alignment':'left'}
                pair['anchor']=anchor_kind(a)
                vertical_pairs.append(pair)
            if uncertain:
                nearest=min(x['vertical_gap_pt'] for x in uncertain)
                relation_pairs.extend(x for x in uncertain if x['vertical_gap_pt']<=nearest+.25*scale)
        # Check the whole connected stack so two matching items cannot hide a
        # third item using a different anchor.
        components=[]
        for pair in vertical_pairs:
            merged=set(pair['shape_ids']);keep=[]
            for component in components:
                if merged & component:merged |= component
                else:keep.append(component)
            components=keep+[merged]
        by_id={r.get('shape_id'):r for r in texts}
        # Do not let an aligned title/value pair hide the explanation below it.
        # Build compact mixed-anchor stacks independently of the author's groups.
        mixed_components=[]
        for pair in vertical_pairs+relation_pairs:
            merged=set(pair['shape_ids']);keep=[]
            for component in mixed_components:
                if merged & component:merged|=component
                else:keep.append(component)
            mixed_components=keep+[merged]
        for component in mixed_components:
            group=[by_id[sid] for sid in sorted(component)]
            if len(group)<3 or len({anchor_kind(r) for r in group})<2:continue
            if max(r['width_pt'] for r in group)>.35*width:continue
            overlap=min(r['left_pt']+r['width_pt'] for r in group)-max(r['left_pt'] for r in group)
            if overlap<.8*min(r['width_pt'] for r in group):continue
            relation_pairs.append({'shape_ids':sorted(component),'native_anchors':[anchor_kind(r) for r in group],
                'reason_detail':'Compact title/value/explanation stack needs one actual visual anchor; different paragraph modes are not an exemption'})
        for relation in relation_pairs:
            relation['rendered_bounds_pt']=[[by_id[sid]['bound_left_pt'],by_id[sid]['bound_top_pt'],
                by_id[sid]['bound_width_pt'],by_id[sid]['bound_height_pt']] for sid in relation['shape_ids']]
            inks=[pixel_bounds.get(sid) for sid in relation['shape_ids']]
            relation['ink_bounds_px']=[v['bounds_px'] for v in inks] if all(inks) else None
            relation['threshold_px']=max(2,.75*scale*raster.width/width) if raster is not None else None
            flag(n,'text_region_relation_requires_review',relation)
        for component in components:
            # A full-width conclusion below a real linear numeric tick stack is
            # another visual region, not one more axis label. Preserve both the
            # region-review finding and the tick stack's own alignment checks.
            import re
            wide=[sid for sid in component if by_id[sid]['width_pt']>=width*.65]
            rest=set(component)-set(wide)
            if len(wide)==1 and len(rest)>=3:
                broad=by_id[wide[0]];ticks=[by_id[sid] for sid in sorted(rest,key=lambda sid:by_id[sid]['bound_top_pt'])]
                numeric=all(re.fullmatch(r'[+−-]?\d+(?:\.\d+)?%?',r['text'].strip()) for r in ticks)
                if numeric:
                    values=[float(r['text'].strip().replace('−','-').rstrip('%')) for r in ticks]
                    centers=[r['bound_top_pt']+r['bound_height_pt']/2 for r in ticks]
                    dv=values[-1]-values[0]
                    slope=(centers[-1]-centers[0])/dv if dv else 0
                    linear=slope!=0 and len(set(values))==len(values) and all(abs(y-(centers[0]+slope*(v-values[0])))<=2*scale for v,y in zip(values,centers))
                    contained=all(r['width_pt']<=width*.15 and broad['left_pt']<=r['left_pt']+12*scale and broad['left_pt']+broad['width_pt']>=r['left_pt']+r['width_pt']-.5*scale for r in ticks)
                    below=broad['bound_top_pt']>=max(r['bound_top_pt']+r['bound_height_pt'] for r in ticks)
                    grid_ids=set()
                    for tick,cy in zip(ticks,centers):
                        for line in slide.shapes:
                            if line.shape_type!=MSO_SHAPE_TYPE.LINE:continue
                            geom=line._element.xpath('./p:spPr/a:prstGeom')
                            if not geom or geom[0].get('prst') not in ('line','straightConnector1'):continue
                            if any(e.get('type','none')!='none' for e in line._element.xpath('./p:spPr/a:ln/a:headEnd | ./p:spPr/a:ln/a:tailEnd')):continue
                            x1,y1,x2,y2=[getattr(line,k)/12700 for k in ('begin_x','begin_y','end_x','end_y')]
                            if abs(y1-y2)<=.25*scale and abs(x2-x1)>=width*.3 and 0<=min(x1,x2)-(tick['bound_left_pt']+tick['bound_width_pt'])<=width*.15 and abs(cy-y1)<=max(2*scale,tick.get('font_pt',12)*.35):grid_ids.add(line.shape_id)
                    if linear and contained and below and len(grid_ids)>=3:
                        all_rows=[by_id[sid] for sid in sorted(component)]
                        flag(n,'axis_stack_spanning_region_review',{'shape_ids':sorted(component),'axis_tick_ids':sorted(rest),'spanning_text_id':wide[0],'matched_grid_ids':sorted(grid_ids),
                            'native_anchors':[anchor_kind(r) for r in all_rows],
                            'frame_bounds_pt':[[r['left_pt'],r['top_pt'],r['width_pt'],r['height_pt']] for r in all_rows],
                            'reason_detail':'Full-width region below a linear numeric tick stack; verify actual spanning relationship, then review tick alignment separately'})
                        component=rest
            group=[by_id[sid] for sid in component]
            kind=anchor_kind(group[0]);positions=[anchor_position(r,kind) for r in group]
            # Prevent transitive anchor drift from manufacturing one giant column.
            if max(positions)-min(positions)>12*scale:
                flag(n,'text_region_relation_requires_review',{'shape_ids':sorted(component),
                     'native_anchors':[anchor_kind(r) for r in group],
                     'frame_bounds_pt':[[r['left_pt'],r['top_pt'],r['width_pt'],r['height_pt']] for r in group],
                     'anchor_delta_pt':max(positions)-min(positions),'reason_detail':'transitive projection chain is not proof of one column'})
                continue
            lefts=[r['bound_left_pt'] for r in group]
            centers=[r['bound_left_pt']+r['bound_width_pt']/2 for r in group]
            rights=[r['bound_left_pt']+r['bound_width_pt'] for r in group]
            selected={'left':lefts,'center':centers,'right':rights}[kind]
            measured=[pixel_bounds.get(sid) for sid in sorted(component)]
            if max(selected)-min(selected)>2*scale:
                flag(n,'rendered_column_text_anchor_mismatch',{'shape_ids':sorted(component),
                     'left_delta_pt':max(lefts)-min(lefts),'center_delta_pt':max(centers)-min(centers),
                     'right_delta_pt':max(rights)-min(rights),'default_alignment':kind,
                     'ink_bounds_px':[m['bounds_px'] for m in measured] if all(measured) else None,
                     'threshold_px':max(2,.75*scale*raster.width/width) if raster is not None else None})
            if any(x is None for x in measured):
                flag(n,'glyph_ink_alignment_requires_review',{'shape_ids':sorted(component),
                     'reason_detail':raster_error or 'Ambiguous background or antialias bounds; inspect actual glyph edges locally'})
            else:
                bs=[m['bounds_px'] for m in measured]
                ls=[b[0] for b in bs];cs=[(b[0]+b[2])/2 for b in bs];rs=[b[2] for b in bs]
                tolerance=max(2,.75*scale*raster.width/width)
                selected={'left':ls,'center':cs,'right':rs}[kind]
                if max(selected)-min(selected)>=tolerance:
                    flag(n,'rendered_ink_column_misaligned',{'shape_ids':sorted(component),
                         'ink_bounds_px':bs,'left_delta_px':max(ls)-min(ls),
                         'center_delta_px':max(cs)-min(cs),'right_delta_px':max(rs)-min(rs),
                         'threshold_px':tolerance,'pixels_per_pt':raster.width/width,'default_alignment':kind})
        # Whole chart + sidebar, not just each column's own internals.
        for shape in slide.shapes:
            if not shape.has_chart:continue
            l,t,w,h=[v/12700 for v in (shape.left,shape.top,shape.width,shape.height)]
            for side in ('left','right'):
                neighbours=[r for r in texts if
                    ((r['left_pt']>=l+w and r['left_pt']-(l+w)<width*.25) if side=='right' else
                     (r['left_pt']+r['width_pt']<=l and l-r['left_pt']-r['width_pt']<width*.25))
                    and r['bound_top_pt']>=t-10*scale and r['bound_top_pt']+r['bound_height_pt']<=t+h+30*scale]
                if len(neighbours)<2:continue
                # A whole text group below this chart belongs to a later panel
                # row, not this chart's sidebar. Keep genuine sidebar members
                # together (including a final explanation extending below it).
                if not any(min(r['bound_top_pt']+r['bound_height_pt'],t+h)>max(r['bound_top_pt'],t) for r in neighbours):
                    continue
                nt=min(r['bound_top_pt'] for r in neighbours);nb=max(r['bound_top_pt']+r['bound_height_pt'] for r in neighbours)
                delta=(nt+nb)/2-(t+h/2)
                plot=next((c.get('plot_bounds_pt') for c in render.get('chart_geometry_records',[]) if c.get('slide')==n and c.get('shape_id')==shape.shape_id),None)
                plot_delta=(nt+nb)/2-(plot[1]+plot[3])/2 if plot else None
                if abs(delta)>6*scale:
                    flag(n,'chart_sidebar_vertical_balance',{'chart_id':shape.shape_id,'side':side,
                        'text_ids':[r.get('shape_id') for r in neighbours], 'chart_bounds_pt':[l,t,l+w,t+h],
                        'sidebar_top_bottom_pt':[nt,nb],'center_delta_pt':delta,
                        'actual_plot_bounds_pt':plot,'plot_center_delta_pt':plot_delta,
                        'requires_actual_plot_bounds_review':True})
        # Header-to-body and body-to-conclusion must remain explicitly visible;
        for chart_record in render.get('chart_geometry_records',[]):
            if chart_record.get('slide')!=n:continue
            chart=next((sh for sh in slide.shapes if sh.shape_id==chart_record['shape_id']),None)
            if chart is None:continue
            for table in slide.shapes:
                if not table.has_table:continue
                overlap=min(chart.left+chart.width,table.left+table.width)-max(chart.left,table.left)
                gap=(table.top-chart.top-chart.height)/12700
                if overlap>=.8*min(chart.width,table.width) and -5*scale<=gap<=120*scale:
                    right=chart_record['plot_bounds_pt'][2];table_right=(table.left+table.width)/12700
                    if abs(right-table_right)>2*scale:
                        flag(n,'chart_table_visible_edge_requires_review',{'shape_ids':[chart.shape_id,table.shape_id],
                             'plot_right_pt':right,'table_right_pt':table_right,'delta_pt':abs(right-table_right),
                             'reason_detail':'Review visible plot/axis extent versus table edge; equal outer chart/table frames do not prove visual alignment'})
        # excluding title roles from the body sweep does not exempt these gaps.
        headings=[r for r in render.get('text_records',[]) if r.get('slide')==n and r.get('shape','').startswith(('title:','subtitle:'))]
        if headings and objects:
            header_bottom=max(r.get('bound_top_pt',r['top_pt'])+r['bound_height_pt'] for r in headings)
            body_top=min(o['bounds'][1] for o in objects)
            flag(n,'header_body_boundary_review',{'header_bottom_pt':header_bottom,'body_top_pt':body_top,
                 'gap_pt':body_top-header_bottom,'required':'Compare actual ink/plot bounds above AND below the main visual; record both gaps, not title coordinates alone.'})
        pages.append({'slide': n, 'objects': objects, 'excluded': excluded, 'scans': scans, 'text_alignment_pairs': pairs,'vertical_text_pairs':vertical_pairs,'region_relations':relation_pairs,
                      'raster_path':str(raster_path) if raster_path else None,'raster_sha256':sha(raster_path) if raster is not None else None,
                      'ink_bounds':{str(k):v for k,v in pixel_bounds.items()},'optical_text_inventory':optical_inventory,
                      'optical_review_targets':[{'shape_id':sh.shape_id,'name':sh.name} for sh in slide.shapes]})
    title_groups={}
    for page in pages:
        for row in page['optical_text_inventory']:
            role=(row.get('shape') or '').split(':',1)[0]
            if role not in ('title','subtitle'):continue
            original=next((r for r in render['text_records'] if r.get('slide')==page['slide'] and r.get('shape_id')==row.get('shape_id') and r.get('kind')=='text'),None)
            if original is None:continue
            key=(role,original.get('font_pt'),original.get('font_bold'),original.get('paragraph_alignment',1))
            title_groups.setdefault(key,[]).append((page['slide'],row,original))
    for key,group in title_groups.items():
        if len({n for n,_,_ in group})<2:continue
        measured=[(n,row,original) for n,row,original in group if row['ink'] is not None]
        if len(measured)!=len(group):
            for n,row,_ in group:
                if row['ink'] is None:flag(n,'cross_page_title_ink_requires_review',{'shape_id':row['shape_id'],'role':key[0]})
        if not measured:continue
        def xanchor(row):
            b=row['ink']['bounds_pt'];return (b[0]+b[2])/2 if key[3]==2 else b[2] if key[3]==3 else b[0]
        xs=[xanchor(row) for _,row,_ in measured];ys=[row['ink']['bounds_pt'][1] for _,row,_ in measured]
        tx,ty=statistics.median(xs),statistics.median(ys)
        for n,row,_ in measured:
            dx,dy=abs(xanchor(row)-tx),abs(row['ink']['bounds_pt'][1]-ty)
            if dx>1.2*scale or dy>1.2*scale:
                flag(n,'cross_page_title_ink_misaligned',{'shape_id':row['shape_id'],'role':key[0],
                    'peer_slides':[a for a,_,_ in measured],'target_anchor_pt':[tx,ty],'delta_pt':[dx,dy],
                    'reason_detail':'Compare actual ink among same-role/font/alignment peers; frame offsets alone are not failures'})
    if sha(pptx) != digest:
        raise ValueError('Input changed during coverage audit')
    return {'sha256': digest, 'slide_count': len(prs.slides), 'pages': pages,
            'findings': findings, 'status': 'review_required' if findings else 'coverage_pass_pending_visual_review'}


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('pptx',type=Path);ap.add_argument('render',type=Path);ap.add_argument('report',type=Path)
    a=ap.parse_args();render_data={}
    try:
        render_data=json.loads(a.render.read_text(encoding='utf-8-sig'))
        result=audit(a.pptx,render_data)
        code=2 if result['findings'] else 0
    except Exception as exc:
        result={'status':'failed','error':str(exc)};code=3
    from report_io import emit_report
    report_error=emit_report(a.report,result,inputs=(a.pptx,a.render),frozen_dirs=((Path(render_data['snapshot']).parent,) if isinstance(render_data,dict) and isinstance(render_data.get('snapshot'),str) and render_data['snapshot'] else ()))
    if report_error:return report_error
    print(result['status']);return code

if __name__=='__main__':raise SystemExit(main())
