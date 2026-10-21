"""Verify current audience/feedback evidence; human reading establishes semantics."""
from pathlib import Path
import argparse,hashlib,json,math
from PIL import Image,ImageChops
from pptx import Presentation
from render_evidence import native_inventory
from layout_visual_review import _covered

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest().upper()
def load(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def text_ok(v):return isinstance(v,str) and len(v.strip())>=8 and v.strip().lower() not in ('verified','no issue','通过通过通过通过')
def verify(pptx,render_dir,review):
    pptx=Path(pptx);render_dir=Path(render_dir);before=sha(pptx);prs=Presentation(pptx);count=len(prs.slides);issues=[]
    def need(ok,msg):
        if not ok:issues.append(msg)
    if not isinstance(review,dict):return {'sha256':before,'status':'failed','issues':['Audience review must be an object']}
    render=load(render_dir/'audit.json');native=native_inventory(prs)
    if not isinstance(render,dict):return {'sha256':before,'status':'failed','issues':['Render audit must be an object']}
    need(isinstance(render.get('sha256'),str) and render['sha256'].upper()==before and render.get('slide_count')==count,'Render does not match candidate')
    need(review.get('schema')=='audience-review/v1','Missing audience-review/v1 schema')
    need(review.get('sha256')==before,'Audience review candidate hash mismatch')
    need(text_ok(review.get('audience')),'Describe the target audience and prior knowledge')
    need(review.get('expression_mode') in ('advantages','advantage_led','academic','objective'),'Unknown expression mode')
    nav=review.get('chapter_navigation',{})
    if not isinstance(nav,dict):need(False,'chapter_navigation must be an object');nav={}
    need(nav.get('mode') in ('none','section_first','divider','persistent'),'Missing chapter navigation decision')
    need(nav.get('basis') in ('user','inferred','delegated') and text_ok(nav.get('description')),'Chapter decision needs its actual basis')
    need(review.get('unresolved')==[],'Unresolved audience findings remain')
    records={};rows=render.get('text_records',[])
    if not isinstance(rows,list):need(False,'Render text_records must be an array');rows=[]
    for index,record in enumerate(rows):
        if not isinstance(record,dict) or type(record.get('slide'))!=int or not isinstance(record.get('audit_path'),str) or not isinstance(record.get('text'),str):
            need(False,f'Render text_records[{index}] needs integer slide, string audit_path and text');continue
        key=(record['slide'],record['audit_path'])
        if key in records:need(False,'Duplicate render text record: '+str(key));continue
        records[key]=record
    def evidence(ev,current):
        if not isinstance(ev,dict):need(False,'Invalid evidence object');return
        n=ev.get('slide');need(type(n)==int and 1<=n<=current,'Definition/evidence must be on this page or a preceding page')
        if type(n)!=int or not 1<=n<=count:return
        image_fields=any(k in ev for k in ('crop_path','crop_box_px','crop_sha256','image_sha256'))
        if image_fields and 'audit_path' in ev:need(False,'Evidence must be either text or actual image, not mixed fields')
        if 'audit_path' in ev and not image_fields:
            if not isinstance(ev['audit_path'],str):need(False,'Evidence audit_path must be a string');return
            r=records.get((n,ev['audit_path']));q=ev.get('quote')
            need(bool(r) and isinstance(q,str) and bool(q.strip()) and q in r.get('text',''),'Quoted explanation is absent from actual rendered text')
        else:
            try:
                p=render_dir/f'slide-{n:02}.png';cp=Path(ev['crop_path']);b=ev['crop_box_px']
                with Image.open(p) as im,Image.open(cp) as crop:
                    need(isinstance(b,list) and len(b)==4 and all(type(x)==int for x in b),'Invalid visual evidence crop')
                    if not isinstance(b,list) or len(b)!=4 or not all(type(x)==int for x in b):return
                    valid_box=0<=b[0]<b[2]<=im.width and 0<=b[1]<b[3]<=im.height
                    need(valid_box,'Visual evidence crop outside page')
                    if not valid_box:return
                    need(ev.get('image_sha256')==sha(p) and ev.get('crop_sha256')==sha(cp),'Stale visual evidence')
                    actual=im.convert('RGB').crop(b);crop=crop.convert('RGB')
                    need(actual.size==crop.size and ImageChops.difference(actual,crop).getbbox() is None,'Visual evidence differs from actual render')
                    need(text_ok(ev.get('observation')),'Visual evidence needs actual reading observation')
            except (OSError,KeyError,TypeError,ValueError):need(False,'Missing or invalid visual evidence')
    pages=review.get('pages',[])
    if not isinstance(pages,list) or any(not isinstance(p,dict) for p in pages):need(False,'Invalid audience page list');pages=[]
    nums=[p.get('slide') for p in pages]
    need(all(type(n)==int for n in nums) and sorted(n for n in nums if type(n)==int)==list(range(1,count+1)),'Audience review must cover all pages exactly once')
    example_pages=[]
    for p in pages:
        n=p.get('slide')
        if type(n)!=int or not 1<=n<=count:continue
        image=render_dir/f'slide-{n:02}.png'
        need(image.is_file() and p.get('image_sha256')==sha(image),'Missing/stale audience page image')
        page_size=None
        try:
            with Image.open(image) as full:page_size=full.size
        except OSError as error:need(False,f'Page {n} image unavailable: '+str(error))
        need(p.get('result')=='pass' and p.get('unresolved')==[],f'Page {n} comprehension unresolved')
        for key in ('purpose','plain_language','operation_or_reading','metrics_or_symbols','evidence_to_conclusion'):
            need(text_ok(p.get(key)),f'Page {n} needs specific reader reconstruction: {key}')
        need(p.get('kind') in ('general','result','comparison','example'),f'Page {n} unknown content role')
        if p.get('kind')=='example':example_pages.append(n)
        visual=p.get('visual_checks',{})
        if not isinstance(visual,dict):need(False,f'Page {n} visual_checks must be an object');visual={}
        for key in ('content_reading','layout','local_detail','full_page_return','cross_page','palette'):
            need(text_ok(visual.get(key)),f'Page {n} lacks actual visual-stage observation: {key}')
        targets=[]
        def collect(shapes,parent=''):
            for sh in shapes:
                path=(parent+'/group/' if parent else '')+str(sh.shape_id)
                if sh.has_table or sh.has_chart or sh.shape_type==13 or hasattr(sh,'shapes'):targets.append(path)
                if hasattr(sh,'shapes'):collect(sh.shapes,path)
        collect(prs.slides[n-1].shapes)
        need(isinstance(p.get('visual_targets_reviewed'),list) and all(isinstance(x,str) for x in p['visual_targets_reviewed']) and sorted(p['visual_targets_reviewed'])==sorted(targets),f'Page {n} visual component inventory incomplete')
        local=p.get('local_views',[])
        need(isinstance(local,list) and (bool(local) or not targets),f'Page {n} complex components need actual local views')
        if isinstance(local,list):
            crop_boxes=[]
            for ev in local:
                evidence(ev,n)
                need(isinstance(ev,dict) and ev.get('slide')==n and 'crop_path' in ev,f'Page {n} local view must be current actual pixels')
                if isinstance(ev,dict) and isinstance(ev.get('crop_box_px'),list):
                    if page_size:need(ev['crop_box_px']!=[0,0,*page_size],f'Page {n} full image cannot replace local inspection')
                    b=ev['crop_box_px']
                    if len(b)==4 and all(type(v)==int for v in b):crop_boxes.append(b)
            if page_size:
                width,height=page_size;pp=width/(prs.slide_width/12700)
                for target in targets:
                    b=native['bounds'].get((n,target))
                    region=[max(0,math.floor(b[0]*pp)),max(0,math.floor(b[1]*pp)),min(width,math.ceil(b[2]*pp)),min(height,math.ceil(b[3]*pp))] if b else [0,0,width,height]
                    need(_covered(region,crop_boxes),f'Page {n} local views omit component {target}; include its full visible context')

        evs=p.get('evidence',[])
        need(isinstance(evs,list) and bool(evs),f'Page {n} missing visible explanation evidence')
        if isinstance(evs,list):
            need(any(isinstance(e,dict) and e.get('slide')==n for e in evs),f'Page {n} only cites elsewhere; inspect current page')
            for ev in evs:evidence(ev,n)
    examples=review.get('examples',[])
    if not isinstance(examples,list) or any(not isinstance(x,dict) for x in examples):examples=[];need(False,'Invalid example records')
    need(sorted(x.get('slide',0) for x in examples if type(x.get('slide'))==int)==sorted(example_pages),'Example selection records must match example pages')
    for x in examples:
        need(type(x.get('slide'))==int and 1<=x['slide']<=count,'Example slide must be a valid integer page number')
        for key in ('purpose','selection_scope','source_locator','why_representative','what_visible','claim_check'):
            need(text_ok(x.get(key)),'Example selection needs '+key)
        source=x.get('source_path');need(isinstance(source,str) and bool(source),'Example source_path must be a nonempty string')
        path=Path(source) if isinstance(source,str) else Path('.')
        need(path.is_file() and x.get('source_sha256')==sha(path),'Example source missing or stale')
        need(isinstance(x.get('alternatives_considered'),list) and bool(x['alternatives_considered']) and all(text_ok(a) for a in x['alternatives_considered']),'Record candidates considered or the concrete single-source situation')
    definitions=review.get('definitions',[])
    if not isinstance(definitions,list):need(False,'Invalid definitions');definitions=[]
    for d in definitions:
        if not isinstance(d,dict):need(False,'Invalid definition');continue
        need(all(isinstance(d.get(k),str) and bool(d[k].strip()) for k in ('term','display_name')),'Definition needs term and display_name')
        first=d.get('introduced_on');uses=d.get('used_on',[])
        need(isinstance(d.get('evidence'),dict) and d['evidence'].get('slide')==first,'Definition evidence must match introduced_on')
        need(type(first)==int and isinstance(uses,list) and bool(uses) and all(type(n)==int and 1<=first<=n<=count for n in uses),'Definition appears after use or outside deck')
        if type(first)==int:evidence(d.get('evidence'),first)
    feedback=review.get('feedback',[])
    if not isinstance(feedback,list):feedback=[];need(False,'Invalid feedback list')
    seen=set()
    for f in feedback:
        if not isinstance(f,dict):need(False,'Invalid feedback item');continue
        fid=f.get('id');need(isinstance(fid,str) and bool(fid) and fid not in seen,'Feedback IDs missing/duplicate');seen.add(str(fid))
        need(f.get('status')=='resolved' and text_ok(f.get('issue')) and text_ok(f.get('observation')),'Feedback needs actual resolution')
        need(isinstance(f.get('category'),str) and bool(f['category'].strip()),'Feedback category is required')
        affected=f.get('affected_slides',[]);similar=f.get('similar_checked_slides',[])
        need(isinstance(affected,list) and bool(affected) and all(type(n)==int and 1<=n<=count for n in affected),'Invalid affected pages')
        need(isinstance(similar,list) and bool(similar) and all(type(n)==int and 1<=n<=count for n in similar),'Missing similar-item inspection')
        if not isinstance(affected,list) or not all(type(n)==int for n in affected):affected=[]
        changes=f.get('changes',[]);need(isinstance(changes,list) and bool(changes),'Feedback only answered in chat; missing artifact evidence')
        if isinstance(changes,list):
            need(set(affected).issubset({e.get('slide') for e in changes if isinstance(e,dict) and type(e.get('slide'))==int}),'Feedback closure omits affected pages')
            for ev in changes:evidence(ev,count)
    need(sha(pptx)==before,'Candidate changed during audience verification')
    return {'sha256':before,'status':'failed' if issues else 'audience_evidence_complete_pending_semantic_signoff','issues':issues}

def main():
    ap=argparse.ArgumentParser();[ap.add_argument(n,type=Path) for n in ('pptx','render_dir','review','report')];a=ap.parse_args()
    try:result=verify(a.pptx,a.render_dir,load(a.review))
    except Exception as e:result={'status':'failed','issues':[str(e)]}
    from report_io import emit_report
    report_error=emit_report(a.report,result,inputs=(a.pptx,a.review),frozen_dirs=(a.render_dir,))
    if report_error:return report_error
    print(result['status']);return 2 if result['issues'] else 0
if __name__=='__main__':raise SystemExit(main())
