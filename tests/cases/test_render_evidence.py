"""Synthetic fixtures only: contract/propagation tests, never visual signoff."""
import sys
sys.dont_write_bytecode = True
import copy
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Pt
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from render_evidence import validate_render_evidence, validate_manual_dispositions, digest


class GateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='schema-test-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        p = Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
        title = s.shapes.add_textbox(Pt(20), Pt(20), Pt(100), Pt(30)); title.text = 'Title'; title.name = 'title:test'
        group = s.shapes.add_group_shape()
        child = group.shapes.add_textbox(Pt(20), Pt(80), Pt(100), Pt(30)); child.text = 'Group text'
        nested = group.shapes.add_group_shape()
        nested_text = nested.shapes.add_textbox(Pt(20), Pt(120), Pt(100), Pt(30)); nested_text.text = 'Nested text'
        table = s.shapes.add_table(2, 2, Pt(200), Pt(50), Pt(100), Pt(50)); table.table.cell(0, 1).text = 'Cell'
        data = ChartData(); data.categories = ['A']; data.add_series('Series', [1])
        chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(200), Pt(200), Pt(200), Pt(100), data)
        pic_file = self.root / 'fixture.png'; Image.new('RGB', (20, 20), '#123456').save(pic_file)
        pic = s.shapes.add_picture(str(pic_file), Pt(400), Pt(30))
        self.pptx = self.root / 'input.pptx'; p.save(self.pptx)
        snapshot = self.root / 'snapshot.pptx'; shutil.copyfile(self.pptx, snapshot)
        self.image = self.root / 'slide-01.png'
        synthetic = Image.new('RGB', (960, 540), '#e0e0e0'); draw = ImageDraw.Draw(synthetic)
        for x in range(0, 960, 5): draw.line((x, 0, x, 539), fill='#123456')
        synthetic.save(self.image)
        self.title_path = str(title.shape_id)
        self.group_path = str(group.shape_id)
        self.nested_path = f'{group.shape_id}/group/{nested.shape_id}'
        self.chart_path = str(chart.shape_id)
        self.paths = [self.title_path, f'{group.shape_id}/group/{child.shape_id}', f'{self.nested_path}/group/{nested_text.shape_id}', f'{table.shape_id}/table/1/2']
        def text(path):
            values=dict(zip(self.paths,['Title','Group text','Nested text','Cell']))
            return {'slide': 1, 'audit_path': path, 'text': values[path], 'run_font_sizes_pt':[12], 'kind': 'deliberately arbitrary', 'left_pt': 20, 'top_pt': 20, 'width_pt': 100, 'height_pt': 30, 'bound_left_pt': 25, 'bound_top_pt': 25, 'bound_width_pt': 90, 'bound_height_pt': 20}
        manual = [
            {'slide': 1, 'audit_path': self.group_path, 'reason': 'group_geometry_requires_visual_review'},
            {'slide': 1, 'audit_path': self.nested_path, 'reason': 'group_geometry_requires_visual_review'},
            {'slide': 1, 'audit_path': self.chart_path, 'reason': 'chart_axes_curves_annotations_require_visual_review'},
            {'slide': 1, 'audit_path': str(pic.shape_id), 'reason': 'picture_internal_visual_review'},
            {'slide': 1, 'audit_path': 'master-layout', 'reason': 'master_layout_visual_review'},
        ]
        for i, item in enumerate(manual, 1): item['review_id'] = f'm{i}'
        self.render = {'schema': 'render-evidence/v2', 'sha256': digest(self.pptx), 'snapshot': str(snapshot), 'slide_count': 1, 'status': 'rendered_pending_visual_review', 'issues': [], 'raster_hashes': {'1': digest(self.image)}, 'text_records': [text(path) for path in self.paths], 'chart_geometry_records': [{'slide': 1, 'audit_path': self.chart_path, 'plot_bounds_pt': [220, 210, 380, 280]}], 'manual_review': manual}

    def valid(self, data=None):
        return validate_render_evidence(self.pptx, self.render if data is None else data, self.root)

    def test_positive_inventory_and_kind_independence(self):
        self.assertEqual(self.valid()['texts'], {(1, path) for path in self.paths})

    def test_missing_each_text_category_rejected(self):
        for path in self.paths:
            with self.subTest(path=path):
                data = copy.deepcopy(self.render)
                data['text_records'] = [r for r in data['text_records'] if r['audit_path'] != path]
                with self.assertRaisesRegex(ValueError, 'Text inventory mismatch'): self.valid(data)

    def test_schema_snapshot_raster_and_chart_negative_mutations(self):
        mutations = [
            ('old schema', lambda d: d.update(schema='render-evidence/v1')),
            ('missing snapshot', lambda d: d.update(snapshot=str(self.root/'missing.pptx'))),
            ('fake snapshot contents', lambda d: d.update(snapshot=str(self.image))),
            ('missing raster', lambda d: d.update(raster_hashes={})),
            ('stale raster hash', lambda d: d.update(raster_hashes={'1': '0'*64})),
            ('extra raster', lambda d: d['raster_hashes'].update({'2': '0'*64})),
            ('missing chart', lambda d: d.update(chart_geometry_records=[])),
            ('duplicate text', lambda d: d['text_records'].append(copy.deepcopy(d['text_records'][0]))),
            ('unknown text', lambda d: d['text_records'][0].update(audit_path='9999')),
            ('nonfinite geometry', lambda d: d['text_records'][0].update(bound_width_pt=float('nan'))),
            ('duplicate chart', lambda d: d['chart_geometry_records'].append(copy.deepcopy(d['chart_geometry_records'][0]))),
            ('unknown chart', lambda d: d['chart_geometry_records'][0].update(audit_path='9999')),
            ('missing image manual', lambda d: d['manual_review'].pop(3)),
            ('missing master manual', lambda d: d['manual_review'].pop()),
            ('duplicate review id', lambda d: d['manual_review'][1].update(review_id='m1')),
            ('unknown manual path', lambda d: d['manual_review'][0].update(audit_path='9999')),
            ('renderer failed', lambda d: d.update(status='failed_preflight')),
            ('unresolved issue', lambda d: d.update(issues=[{'reason': 'font_too_small'}])),
        ]
        for name, change in mutations:
            with self.subTest(name=name):
                data = copy.deepcopy(self.render); change(data)
                with self.assertRaises((ValueError, OSError)): self.valid(data)

    def chart_error(self):
        self.render['chart_geometry_records'] = []
        self.render['manual_review'].append({'review_id': 'm6', 'slide': 1, 'audit_path': self.chart_path, 'reason': 'native_plot_bounds_unavailable'})

    def dispositions(self):
        crop = self.root/'crop.png'
        with Image.open(self.image) as image: image.crop((0, 0, 960, 540)).save(crop)
        return {'render_manual_reviews': [dict(review_id=r['review_id'], slide=r['slide'], audit_path=r['audit_path'], decision='verified_in_context', returned_to_full_page=True, image_sha256=digest(self.image), observation='Synthetic schema test only; not a visual signoff', crop_box_px=[0, 0, 960, 540], crop_path=str(crop)) for r in self.render['manual_review']]}

    def test_chart_error_is_pending_not_closed_by_reason(self):
        self.chart_error(); self.valid()
        self.assertTrue(validate_manual_dispositions(self.render, {}, self.root))
        review = self.dispositions()
        self.assertEqual(validate_manual_dispositions(self.render, review, self.root), [])
        review['render_manual_reviews'].pop()
        self.assertTrue(validate_manual_dispositions(self.render, review, self.root))

    def test_conflicting_chart_error_and_geometry_rejected(self):
        self.render['manual_review'].append({'review_id': 'm6', 'slide': 1, 'audit_path': self.chart_path, 'reason': 'native_plot_bounds_unavailable'})
        with self.assertRaisesRegex(ValueError, 'simultaneously'): self.valid()

    def test_manual_evidence_negative_mutations(self):
        mutations = [
            ('no crop', lambda d: d['render_manual_reviews'][0].pop('crop_path')),
            ('wrong target', lambda d: d['render_manual_reviews'][0].update(audit_path='9999')),
            ('duplicate disposition', lambda d: d['render_manual_reviews'].append(copy.deepcopy(d['render_manual_reviews'][0]))),
            ('unresolved decision', lambda d: d['render_manual_reviews'][0].update(decision='ignored')),
            ('placeholder observation', lambda d: d['render_manual_reviews'][0].update(observation='OK')),
            ('no full page', lambda d: d['render_manual_reviews'][0].update(returned_to_full_page=False)),
            ('stale source hash', lambda d: d['render_manual_reviews'][0].update(image_sha256='0'*64)),
            ('unrelated pixels', lambda d: d['render_manual_reviews'][0].update(crop_path=str(self.root/'fixture.png'))),
        ]
        for name, change in mutations:
            with self.subTest(name=name):
                review = self.dispositions(); change(review)
                self.assertTrue(validate_manual_dispositions(self.render, review, self.root))


    def test_unrelated_current_pixels_cannot_close_manual(self):
        review = self.dispositions()
        unrelated = self.root/'unrelated.png'
        with Image.open(self.image) as image: image.crop((800, 400, 900, 500)).save(unrelated)
        review['render_manual_reviews'][0].update(crop_box_px=[800, 400, 900, 500], crop_path=str(unrelated))
        errors = validate_manual_dispositions(self.render, review, self.root)
        self.assertTrue(any('native object bounds' in r['reason'] for r in errors))

    def test_constant_pixels_cannot_close_manual(self):
        Image.new('RGB', (960, 540), 'white').save(self.image)
        self.render['raster_hashes']['1'] = digest(self.image)
        self.valid()  # Structure/hash is valid; this is explicitly not visual acceptance.
        errors = validate_manual_dispositions(self.render, self.dispositions(), self.root)
        self.assertTrue(any('Blank/constant pixels' in r['reason'] for r in errors))

    def test_changed_png_rejected_without_updating_recorded_hash(self):
        Image.new('RGB', (960, 540), 'red').save(self.image)
        with self.assertRaisesRegex(ValueError, 'Raster missing/stale'): self.valid()


if __name__ == '__main__': unittest.main(verbosity=2)
