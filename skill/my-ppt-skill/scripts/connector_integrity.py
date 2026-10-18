"""Connection inventory and fail-closed endpoint evidence validation.

No role/name exclusions. Native geometry is necessary, never visual signoff.
"""
import math
import hashlib
import csv
import json
import copy
from pathlib import Path
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageChops


class CheckFailure(str):
    """An original diagnostic with its category set at the actual check site."""
    def __new__(cls, message, category='integrity'):
        value = super().__new__(cls, message)
        value.category = category
        return value


def layout_failure(message):
    return CheckFailure(message, 'layout')


def inventory(slide):
    result = []
    def visit(shapes, nested=False):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(s.shapes, True)
                continue
            ends = s._element.xpath('./p:spPr/a:ln/a:headEnd | ./p:spPr/a:ln/a:tailEnd')
            arrows = {e.tag.rsplit('}', 1)[-1]: e.get('type', 'none') for e in ends}
            presets=s._element.xpath('./p:spPr/a:prstGeom')
            block_arrow=bool(presets and 'arrow' in presets[0].get('prst','').lower())
            if s.shape_type != MSO_SHAPE_TYPE.LINE and not block_arrow and not any(v != 'none' for v in arrows.values()):
                continue
            row = {'shape_id': s.shape_id, 'name': s.name, 'arrows': arrows,
                   'endpoints_pt': None, 'geometry_supported': False, 'preset':None,
                   'line_width_pt': float(s.line.width/12700) if s.line.width else 1.0,
                   'block_arrow':block_arrow}
            try:
                x = s._element.xpath('./p:spPr/a:xfrm')[0]
                geom = s._element.xpath('./p:spPr/a:prstGeom')[0].get('prst')
                row['preset']=geom
                if not nested and int(x.get('rot', '0')) == 0 and geom in ('line', 'straightConnector1', 'bentConnector2', 'bentConnector3', 'bentConnector4', 'bentConnector5', 'curvedConnector2', 'curvedConnector3', 'curvedConnector4', 'curvedConnector5'):
                    row['endpoints_pt'] = [[s.begin_x/12700, s.begin_y/12700], [s.end_x/12700, s.end_y/12700]]
                    row['geometry_supported'] = True
            except (AttributeError, IndexError, ValueError, TypeError):
                pass
            result.append(row)
    visit(slide.shapes)
    return result


def boundary_distance(point, shape):
    """Measure supported visible geometry, never an invisible bounding proxy."""
    if shape.rotation != 0 or shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return None
    geom = shape._element.xpath('./p:spPr/a:prstGeom')
    if not geom or geom[0].get('prst') not in ('rect','ellipse','diamond'):
        return None
    # Textbox bounds and invisible dummy rectangles must not become targets.
    fill = shape._element.xpath('./p:spPr/a:solidFill | ./p:spPr/a:gradFill | ./p:spPr/a:pattFill')
    line = shape._element.xpath('./p:spPr/a:ln/a:solidFill')
    if not fill and not line:
        return None
    for paint in fill+line:
        alpha=paint.xpath('.//a:alpha')
        if alpha and any(int(a.get('val','0'))<50000 for a in alpha):return None
    x,y = point
    l,t,r,b = shape.left/12700,shape.top/12700,(shape.left+shape.width)/12700,(shape.top+shape.height)/12700
    if geom[0].get('prst')=='ellipse':
        rx,ry=(r-l)/2,(b-t)/2;cx,cy=(r+l)/2,(b+t)/2
        if rx<=0 or ry<=0:return None
        # Closest-point Lagrange multiplier, not radial projection (which
        # overestimates distance for eccentric ellipses).
        a,b,u,v=rx,ry,abs(x-cx),abs(y-cy)
        if a<b:a,b,u,v=b,a,v,u
        if v==0:
            cutoff=(a*a-b*b)/a
            if u>=cutoff:return abs(u-a)
            q=a*a*u/(a*a-b*b)
            return math.hypot(u-q,b*math.sqrt(max(0,1-(q/a)**2)))
        if u==0:return abs(v-b)
        lo=-b*b;hi=max(a*u+b*v,a*a,b*b)
        for _ in range(100):
            m=(lo+hi)/2
            if m==lo or m==hi:break
            value=(a*u/(m+a*a))**2+(b*v/(m+b*b))**2
            if value>1:lo=m
            else:hi=m
        qx=a*a*u/(hi+a*a);qy=b*b*v/(hi+b*b)
        return math.hypot(u-qx,v-qy)
    if geom[0].get('prst')=='diamond':
        vertices=[((l+r)/2,t),(r,(t+b)/2),((l+r)/2,b),(l,(t+b)/2)]
        distances=[]
        for a,c in zip(vertices,vertices[1:]+vertices[:1]):
            dx,dy=c[0]-a[0],c[1]-a[1];den=dx*dx+dy*dy
            if not den:return None
            u=max(0,min(1,((x-a[0])*dx+(y-a[1])*dy)/den))
            distances.append(math.hypot(x-a[0]-u*dx,y-a[1]-u*dy))
        return min(distances)
    return min(math.hypot(x-l,y-min(max(y,t),b)), math.hypot(x-r,y-min(max(y,t),b)),
               math.hypot(x-min(max(x,l),r),y-t), math.hypot(x-min(max(x,l),r),y-b))


def endpoint_error(row, endpoint, proof, slide, rows, scale):
    """Legacy strict API: the first original diagnostic, with no visual bypass."""
    failures = endpoint_findings(row, endpoint, proof, slide, rows, scale)
    return failures[0] if failures else None


def endpoint_findings(row, endpoint, proof, slide, rows, scale):
    """Keep layout judgments separate and continue all applicable contracts."""
    failures = []
    if not isinstance(proof, dict):
        return [CheckFailure('Every endpoint needs a structured evidence record')]
    supported = row['geometry_supported']
    if not supported:
        failures.append(layout_failure('Unsupported transformed/group/custom connection: normalize to native slide-space connectors and rerender'))
    point = row['endpoints_pt'][endpoint] if supported else None
    kind = proof.get('kind')
    has_arrow = row['arrows'].get('headEnd' if endpoint == 0 else 'tailEnd', 'none') != 'none'
    # A 3.6pt terminal in the real regression has correct endpoint coordinates
    # but renders as a detached tiny triangle. Never let that pass numerically.
    if supported and has_arrow and row.get('preset') in ('line','straightConnector1'):
        length=math.dist(*row['endpoints_pt'])
        minimum=max(6*scale,4*row.get('line_width_pt',1))
        if length < minimum:
            failures.append(layout_failure(f'Arrow terminal segment too short ({length:.4f} pt < {minimum:.4f} pt); lengthen the approach and inspect actual tip'))
    shapes = {}
    def visit(items):
        for shape in items:
            shapes[shape.shape_id] = shape
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes)
    visit(slide.shapes)
    if kind in ('data_marker_center','data_reference_line'):
        data_role_error(row, endpoint, proof, slide, scale, _failures=failures)
        return failures
    if kind in ('icon_stroke','arrow_axis','axis_tick'):
        structural_role_error(row, proof, slide, rows, scale, _failures=failures)
        return failures
    if kind == 'shape_boundary':
        target = shapes.get(proof.get('target_shape_id'))
        if target is None or target.shape_id == row['shape_id']:
            failures.append(CheckFailure('Missing real target shape'))
            return failures
        distance = boundary_distance(point, target) if supported else None
        if distance is None:
            paint = target._element.xpath('./p:spPr/a:solidFill | ./p:spPr/a:gradFill | ./p:spPr/a:pattFill | ./p:spPr/a:ln/a:solidFill')
            visible = bool(paint) and not any(int(a.get('val','0')) < 50000 for p in paint for a in p.xpath('.//a:alpha'))
            message = 'Target boundary is not supported visible geometry; add verified geometry support, do not use an invisible proxy'
            failures.append(layout_failure(message) if visible else CheckFailure(message))
            return failures
    elif kind == 'line_junction':
        target = next((r for r in rows if r['shape_id']==proof.get('target_shape_id')), None)
        if has_arrow:
            failures.append(CheckFailure('Arrow tip must reach its target node, not an unrelated line junction'))
        if target is None or target['shape_id']==row['shape_id']:
            failures.append(CheckFailure('Invalid line junction'))
            return failures
        k=proof.get('target_endpoint')
        if type(k) is not int or k not in (0,1):
            failures.append(CheckFailure('A junction must identify the exact other endpoint'))
            return failures
        if not target['geometry_supported']:
            failures.append(layout_failure('Invalid line junction'))
        if not supported or not target['geometry_supported']:
            return failures
        distance=math.dist(point,target['endpoints_pt'][k])
    elif kind == 'open_nonrelational':
        if any(v!='none' for v in row['arrows'].values()) or row.get('block_arrow'):
            failures.append(CheckFailure('Arrow-bearing lines cannot be self-exempted as open decoration/axis'))
        if proof.get('purpose') not in ('axis','separator','dimension','legend_sample'):
            failures.append(CheckFailure('Unconnected flow lines are not nonrelational decoration'))
        return failures
    else:
        failures.append(CheckFailure('Every endpoint needs a real target/junction or a verified nonrelational purpose'))
        return failures
    if not math.isfinite(distance) or distance > .5*scale:
        failures.append(layout_failure(f'Endpoint detached from target by {distance:.4f} pt; repair and rerender (limit {0.5*scale:.4f} pt)'))
    return failures


def data_role_error(row, endpoint, proof, slide, scale, _failures=None):
    failures = []
    error = _data_role_error(row, endpoint, proof, slide, scale, failures)
    if error:
        failures.append(CheckFailure(error))
    if _failures is not None:
        _failures.extend(failures)
    return failures[0] if failures else None


def _data_role_error(row, endpoint, proof, slide, scale, failures):
    """Verify actual data-marker center or equation-defined reference geometry."""
    if any(v!='none' for v in row['arrows'].values()) or row.get('block_arrow'):
        return 'Data-line roles cannot exempt an arrow-bearing relationship'
    if not row['geometry_supported'] or row.get('preset') not in ('line','straightConnector1'):
        failures.append(layout_failure('Data role requires actual straight geometry'))
    mapping=proof.get('data_mapping',{})
    if not isinstance(mapping,dict):return 'Invalid data role evidence: data_mapping must be an object'
    source=Path(mapping.get('source_path',''))
    if not source.is_file() or hashlib.sha256(source.read_bytes()).hexdigest().upper()!=mapping.get('source_sha256'):
        return 'Data role needs the exact source file and hash'
    def number(k):
        v=mapping.get(k)
        if type(v) not in (int,float) or not math.isfinite(v):raise ValueError('Missing finite mapping '+k)
        return v
    try:
        point=row['endpoints_pt'][endpoint] if row['endpoints_pt'] is not None else None
        if proof['kind']=='data_marker_center':
            shapes={s.shape_id:s for s in slide.shapes};target=shapes.get(proof.get('target_shape_id'))
            if target is None:return 'Missing real data marker'
            geom=target._element.xpath('./p:spPr/a:prstGeom')
            if not geom or geom[0].get('prst')!='ellipse' or target.rotation!=0 or not target._element.xpath('./p:spPr/a:solidFill'):
                return 'Data marker must be an actual solid ellipse'
            if any(int(a.get('val','0'))<50000 for a in target._element.xpath('./p:spPr/a:solidFill//a:alpha')):
                return 'Transparent proxy cannot serve as a data marker'
            if max(target.width,target.height)/12700>20*scale:
                failures.append(layout_failure('Marker too large for this data-point contract'))
            center=[(target.left+target.width/2)/12700,(target.top+target.height/2)/12700]
            if point is not None and math.dist(point,center)>.5*scale:
                failures.append(layout_failure('Data segment endpoint does not match actual marker center'))
            axis=mapping.get('axis')
            if axis not in ('x','y'):return 'Declare the numeric data axis'
            with source.open(encoding='utf-8-sig',newline='') as f:records=list(csv.DictReader(f))
            index=mapping.get('row_index');column=mapping.get('column')
            if type(index) is not int or not 0<=index<len(records) or column not in records[index]:return 'Source CSV row/column missing'
            value=float(records[index][column])
            if not math.isfinite(value):return 'Source value is nonfinite'
            mapped=number('offset_pt')+number('scale_pt')*value
            if number('scale_pt')==0:return 'Data-axis scale cannot be zero'
            if abs(mapped-center[0 if axis=='x' else 1])>.5*scale:return 'Marker position disagrees with source data mapping'
            return None
        # Reference lines are mathematical comparisons, not nodes/flow arrows.
        # Equation and axis transform are bound to an inspectable JSON sidecar.
        if source.suffix.lower()!='.json':return 'Reference line requires a JSON equation/axis mapping source'
        spec=json.loads(source.read_text(encoding='utf-8-sig'))
        if not isinstance(spec,dict) or not {'equation','x_axis','y_axis','endpoints_data'}.issubset(spec):return 'Incomplete reference equation specification'
        if spec.get('equation')!='y=x':return 'Only equality reference is currently supported; add tested mapping for another equation'
        values=spec.get('endpoints_data')
        if not isinstance(values,list) or len(values)!=2:return 'Reference needs two explicit data endpoints'
        expected=[]
        for pair in values:
            if not isinstance(pair,list) or len(pair)!=2 or any(type(v) not in (int,float) or not math.isfinite(v) for v in pair):return 'Invalid reference values'
            if abs(pair[0]-pair[1])>1e-9:return 'Reference values contradict y=x'
            q=[]
            for value,key in zip(pair,('x_axis','y_axis')):
                a=spec[key];offset=a.get('offset_pt');factor=a.get('scale_pt')
                if any(type(v) not in (int,float) or not math.isfinite(v) for v in (offset,factor)) or factor==0:return 'Invalid reference axis mapping'
                q.append(offset+factor*value)
            expected.append(q)
        if row['endpoints_pt'] is None:
            return 'Data mapping cannot be validated without actual reference endpoint coordinates'
        if any(math.dist(a,b)>.5*scale for a,b in zip(row['endpoints_pt'],expected)):
            return 'Reference line disagrees with actual equation coordinates'
        if not spec.get('source_data_path') or not spec.get('source_data_sha256'):return 'Reference axis mapping needs underlying chart data provenance'
        data=Path(spec['source_data_path'])
        if not data.is_file() or hashlib.sha256(data.read_bytes()).hexdigest().upper()!=spec['source_data_sha256']:return 'Reference chart data missing/stale'
        return None
    except (OSError,ValueError,TypeError,KeyError,IndexError) as exc:
        return 'Invalid data role evidence: '+str(exc)


def structural_role_error(row, proof, slide, rows, scale, _failures=None):
    """Check complete real-object declarations even if geometry is ambiguous."""
    failures=[]
    index={r['shape_id']:r for r in rows};shapes={s.shape_id:s for s in slide.shapes}
    kind=proof.get('kind')
    fields = ('cluster_shape_ids',) if kind=='icon_stroke' else ('tick_shape_ids','label_shape_ids')
    for field in fields:
        ids=proof.get(field)
        minimum=3 if field=='cluster_shape_ids' else 2
        valid=isinstance(ids,list) and all(type(i) is int for i in ids)
        if not valid or len(ids)<minimum or len(ids)!=len(set(ids)):
            failures.append(CheckFailure('Structural role requires complete distinct native objects: '+field))
            continue
        known=shapes if field=='label_shape_ids' else index
        if any(i not in known for i in ids):
            failures.append(CheckFailure('Structural role references a missing real object: '+field))
        if field=='label_shape_ids':
            if any(i not in shapes or not shapes[i].has_text_frame or not shapes[i].text.strip() for i in ids):
                failures.append(CheckFailure('Axis label is missing/empty'))
        elif any(i in index and (index[i].get('block_arrow') or any(v!='none' for v in index[i]['arrows'].values())) for i in ids):
            failures.append(CheckFailure('Structural stroke declarations cannot contain arrow-bearing relationships'))
    if kind=='icon_stroke' and row['shape_id'] not in (proof.get('cluster_shape_ids') or []):
        failures.append(CheckFailure('Icon needs complete cluster of at least three actual strokes'))
    if kind!='icon_stroke':
        axis=proof.get('axis_shape_id')
        if type(axis) is not int or axis not in index:
            failures.append(CheckFailure('Axis declaration is missing its real native connector'))
        if kind=='arrow_axis' and row['shape_id']!=axis:
            failures.append(CheckFailure('Axis role must refer to the reviewed line itself'))
        if kind=='axis_tick' and row['shape_id'] not in (proof.get('tick_shape_ids') or []):
            failures.append(CheckFailure('Tick role must include the reviewed stroke'))
    if not failures and kind=='icon_stroke':
        # Inventory omissions remain integrity failures even when an earlier
        # spacing/extent judgment also fails. Do not let that judgment hide an
        # adjacent native stroke omitted from the declared cluster.
        cluster=[index[i] for i in proof['cluster_shape_ids']];tol=.5*scale
        def direction(r):
            if not r['geometry_supported'] or r.get('preset') not in ('line','straightConnector1'):
                return None
            (x,y),(a,b)=r['endpoints_pt']
            if abs(y-b)<=tol and abs(x-a)>tol:return 0
            if abs(x-a)<=tol and abs(y-b)>tol:return 1
            return None
        d=direction(row)
        if d is not None and all(direction(r)==d for r in cluster):
            spans=[sorted(p[d] for p in r['endpoints_pt']) for r in cluster]
            positions=sorted(r['endpoints_pt'][0][1-d] for r in cluster)
            step=max(b-a for a,b in zip(positions,positions[1:]))
            for other in rows:
                if other['shape_id'] in proof['cluster_shape_ids'] or direction(other)!=d or any(v!='none' for v in other['arrows'].values()):continue
                span=sorted(p[d] for p in other['endpoints_pt']);pos=other['endpoints_pt'][0][1-d]
                if all(abs(span[k]-spans[0][k])<=tol for k in (0,1)) and positions[0]-step-tol<=pos<=positions[-1]+step+tol:
                    failures.append(CheckFailure('Icon cluster omits a matching adjacent stroke'))
    if not failures:
        error=_structural_role_error(row,proof,slide,rows,scale)
        if error:failures.append(error if isinstance(error,CheckFailure) else CheckFailure(error))
    if _failures is not None:_failures.extend(failures)
    return failures[0] if failures else None


def _structural_role_error(row, proof, slide, rows, scale):
    """Validate real geometry for nonrelational marks; a role string is not proof."""
    index={r['shape_id']:r for r in rows};shapes={s.shape_id:s for s in slide.shapes}
    kind=proof.get('kind');tol=.5*scale
    def straight(r):return r and r['geometry_supported'] and r.get('preset') in ('line','straightConnector1')
    def arrow(r):return any(v!='none' for v in r['arrows'].values())
    def axis_direction(r):
        (x,y),(a,b)=r['endpoints_pt']
        if abs(y-b)<=tol and abs(x-a)>tol:return 0
        if abs(x-a)<=tol and abs(y-b)>tol:return 1
        return None
    if kind=='icon_stroke':
        ids=proof.get('cluster_shape_ids',[])
        if len(ids)!=len(set(ids)) or len(ids)<3 or row['shape_id'] not in ids:return 'Icon needs complete cluster of at least three actual strokes'
        cluster=[index.get(i) for i in ids]
        if any(not straight(r) or arrow(r) for r in cluster):return layout_failure('Icon strokes must be actual unarrowed straight strokes')
        direction=axis_direction(row)
        if direction is None or any(axis_direction(r)!=direction for r in cluster):return layout_failure('Icon stroke cluster must be parallel')
        points=[p for r in cluster for p in r['endpoints_pt']]
        if any(max(p[k] for p in points)-min(p[k] for p in points)>48*scale for k in (0,1)):
            return layout_failure('Icon cluster exceeds compact 48pt footprint')
        spans=[sorted(p[direction] for p in r['endpoints_pt']) for r in cluster]
        if any(max(s[k] for s in spans)-min(s[k] for s in spans)>tol for k in (0,1)):
            return layout_failure('Icon strokes do not share a common extent')
        positions=sorted(r['endpoints_pt'][0][1-direction] for r in cluster)
        gaps=[b-a for a,b in zip(positions,positions[1:])]
        if min(gaps)<2*scale or max(gaps)-min(gaps)>tol:return layout_failure('Icon strokes must be distinct and regularly spaced')
        # Require the full local pattern, not a selected subset of matching strokes.
        for other in rows:
            if other['shape_id'] in ids or not straight(other) or arrow(other) or axis_direction(other)!=direction:continue
            span=sorted(p[direction] for p in other['endpoints_pt'])
            pos=other['endpoints_pt'][0][1-direction]
            if all(abs(span[k]-spans[0][k])<=tol for k in (0,1)) and positions[0]-max(gaps)-tol<=pos<=positions[-1]+max(gaps)+tol:
                return 'Icon cluster omits a matching adjacent stroke'
        return None
    axis=index.get(proof.get('axis_shape_id'))
    if not straight(axis):return layout_failure('Axis needs actual straight slide-space geometry')
    direction=axis_direction(axis)
    if direction is None or math.dist(*axis['endpoints_pt'])<48*scale:return layout_failure('Axis must be horizontal/vertical with measurable extent')
    if kind=='arrow_axis' and row['shape_id']!=axis['shape_id']:return 'Axis role must refer to the reviewed line itself'
    ids=proof.get('tick_shape_ids',[])
    if len(ids)<2 or len(ids)!=len(set(ids)):return 'Axis needs at least two distinct actual ticks'
    if kind=='axis_tick' and row['shape_id'] not in ids:return 'Tick role must include the reviewed stroke'
    ticks=[index.get(i) for i in ids];lo,hi=sorted(p[direction] for p in axis['endpoints_pt']);cross=axis['endpoints_pt'][0][1-direction]
    positions=[]
    for tick in ticks:
        if not straight(tick) or arrow(tick) or axis_direction(tick)!=1-direction:return layout_failure('Ticks must be perpendicular unarrowed strokes')
        length=math.dist(*tick['endpoints_pt'])
        if not 2*scale<=length<=30*scale:return layout_failure('Tick size inconsistent with an axis')
        pos=tick['endpoints_pt'][0][direction];a,b=sorted(p[1-direction] for p in tick['endpoints_pt'])
        if not lo-tol<=pos<=hi+tol or not a-tol<=cross<=b+tol:return layout_failure('Tick does not meet the actual axis')
        positions.append(pos)
    if len(set(round(v,3) for v in positions))!=len(positions):return layout_failure('Axis ticks overlap')
    # Actual text labels distinguish a timeline/coordinate axis from a free flow arrow.
    labels=proof.get('label_shape_ids',[])
    if len(labels)<2 or len(labels)!=len(set(labels)):return 'Axis needs actual labels, not a free arrow relabeled as an axis'
    matched=set()
    for sid in labels:
        text=shapes.get(sid)
        if text is None or not text.has_text_frame or not text.text.strip():return 'Axis label is missing/empty'
        center=[(text.left+text.width/2)/12700,(text.top+text.height/2)/12700]
        near=min(range(len(positions)),key=lambda i:abs(center[direction]-positions[i]))
        if abs(center[direction]-positions[near])>36*scale or abs(center[1-direction]-cross)>48*scale:return layout_failure('Axis labels are not located by their ticks')
        matched.add(near)
    if len(matched)<2:return layout_failure('Labels must identify at least two separate ticks')
    return None


def connector_finding(n, row, endpoint, proof, record):
    """Intact native rule input and endpoint relation; no role-name whitelist."""
    return {'id':f'connector:{n}:{row["shape_id"]}:{endpoint}', 'slide':n,
            'type':'connector_endpoint_layout', 'shape_id':row['shape_id'],
            'endpoint':endpoint, 'native_connection':copy.deepcopy(row),
            'endpoint_proof':copy.deepcopy({k:v for k,v in proof.items() if k!='layout_visual_review'}),
            'semantic_relation':record.get('semantic_relation'),
            'path_observation':record.get('path_observation')}


def verify_connections(prs, review, render_dir, *, render=None, resolutions=None):
    """Strict by default; explicit current independent review may close layout.

    All native entries, both endpoints, exact crops, data and semantic contracts
    are still checked. The caller must validate the complete render evidence.
    Only accepted resolutions are returned separately; original errors stay in
    each resolution and the unmodified renderer/review records.
    """
    issues=[]
    expected={(n,r['shape_id']):r for n,s in enumerate(prs.slides,1) for r in inventory(s)}
    records=review.get('connector_reviews',[])
    if not isinstance(records,list) or any(not isinstance(r,dict) for r in records):
        return [{'reason':'Connector reviews must be structured records'}]
    keys=[(r.get('slide'),r.get('shape_id')) for r in records]
    if any(type(n) is not int or type(sid) is not int for n,sid in keys) or len(keys)!=len(set(keys)) or set(keys)!=set(expected):
        return [{'reason':'Connector reviews must exactly cover the independent native inventory; no skip/name exclusions'}]
    scale=prs.slide_height/12700/540
    for record in records:
        n,sid=record['slide'],record['shape_id'];row=expected[n,sid]
        def fail(msg):issues.append({'slide':n,'shape_id':sid,'reason':msg})
        image_path=Path(render_dir)/f'slide-{n:02}.png'
        if not image_path.is_file():fail('Missing actual connection render');continue
        source_hash=hashlib.sha256(image_path.read_bytes()).hexdigest().upper()
        if record.get('image_sha256')!=source_hash or record.get('returned_to_full_page') is not True:
            fail('Connection review must bind current image and return to full page')
        if not record.get('path_observation') or not record.get('semantic_relation'):
            fail('Missing actual path clearance, direction and source/target semantic observations')
        proofs=record.get('endpoints',[])
        if not isinstance(proofs,list) or len(proofs)!=2 or any(not isinstance(p,dict) for p in proofs):
            fail('Both endpoints require independent evidence');continue
        try:
            with Image.open(image_path) as src:im=src.convert('RGB')
        except (ValueError,OSError) as exc:
            fail('Cannot read current connection render: '+str(exc));continue
        pp=im.width/(prs.slide_width/12700)
        for k,proof in enumerate(proofs):
            try:
                errors=endpoint_findings(row,k,proof,prs.slides[n-1],inventory(prs.slides[n-1]),scale)
            except (KeyError,ValueError,TypeError,OSError,AttributeError) as exc:
                fail('Endpoint rule execution/evidence failed: '+str(exc));continue
            layout_errors=[str(e) for e in errors if getattr(e,'category','integrity')=='layout']
            for error in errors:
                if getattr(error,'category','integrity')!='layout':fail(str(error))
            if not proof.get('visible_tip_and_boundary_observation'):
                fail('Missing endpoint-specific actual tip/contact/clearance observation')
            # Verify crop pixels, not just a filename or a claimed source hash.
            try:
                box=proof['crop_box_px'];path=Path(proof['crop_path'])
                if len(box)!=4 or any(type(v) is not int for v in box):raise ValueError('Invalid crop box')
                l,t,r,b=box
                if not (0<=l<r<=im.width and 0<=t<b<=im.height):raise ValueError('Crop outside image')
                if row['endpoints_pt'] is None:
                    layout_errors.append('Endpoint local evidence invalid: Endpoint geometry unavailable')
                else:
                    x,y=[v*pp for v in row['endpoints_pt'][k]]
                    pad=8*scale*pp
                    if not (l<=x-pad and r>=x+pad and t<=y-pad and b>=y+pad):raise ValueError('Crop must include endpoint and surrounding boundary/clearance')
                with Image.open(path) as crop:actual=crop.convert('RGB')
                expected_crop=im.crop(box)
                if actual.size!=expected_crop.size or ImageChops.difference(actual,expected_crop).getbbox():raise ValueError('Crop is not the exact current rendered region')
            except (KeyError,ValueError,TypeError,OSError) as exc:
                fail('Endpoint local evidence invalid: '+str(exc))
            if layout_errors:
                disposition=proof.get('layout_visual_review')
                if disposition is None or render is None:
                    for error in layout_errors:fail(error)
                else:
                    from layout_visual_review import validate_layout_visual_review
                    finding=connector_finding(n,row,k,proof,record)
                    error,resolution=validate_layout_visual_review(
                        finding,disposition,image_path,prs.slides[n-1],render,
                        review.get('author_id'),layout_errors)
                    if error:
                        for original in layout_errors:fail(original)
                        fail('Connector visual reconciliation invalid: '+error)
                    elif resolutions is not None:
                        resolutions.append(resolution)
    return issues
