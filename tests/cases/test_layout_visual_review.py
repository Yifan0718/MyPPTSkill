"""Synthetic review-protocol tests only; no real-deck signoff is fabricated."""
import copy
import json
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from audit_layout_coverage import audit
from layout_visual_review import CODE,layout_context,validate_layout_visual_review
from measurement_fallback import digest

NOTE='Synthetic unittest evidence; not a real observation, actual reviewer provenance, or deliverable signoff.'
ROW_ERROR='Actual row-center mismatch must be repaired and rescanned; top alignment cannot waive it'


class LayoutVisualReview(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.p=Presentation();self.p.slide_width=Pt(960);self.p.slide_height=Pt(540)
        self.s=self.p.slides.add_slide(self.p.slide_layouts[6]);self.records=[];self.ids=[]
        self.image=self.root/'slide-01.png';im=Image.new('RGB',(960,540),'white');d=ImageDraw.Draw(im)
        for x,y,w,h,text,size in [(40,100,300,40,'Left code heading',20),(40,160,300,100,'Left body content',16),
                                  (500,45,300,40,'Right explanatory heading',20),(500,110,300,100,'Right independent explanation',16)]:
            s=self.s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));s.text=text;s.text_frame.paragraphs[0].runs[0].font.size=Pt(size)
            self.ids.append(s.shape_id);font=ImageFont.load_default(size=size);bb=d.textbbox((x,y),text,font=font)
            d.text((x,y),text,font=font,fill=(25,45,75))
            self.records.append({'slide':1,'kind':'text','shape_id':s.shape_id,'audit_path':str(s.shape_id),'shape':s.name,'text':text,'font_pt':size,
              'left_pt':x,'top_pt':y,'width_pt':w,'height_pt':h,'bound_left_pt':bb[0],'bound_top_pt':bb[1],'bound_width_pt':bb[2]-bb[0],'bound_height_pt':bb[3]-bb[1]})
        im.save(self.image);self.pptx=self.root/'source.pptx';self.p.save(self.pptx)
        self.render={'schema':'render-evidence/v2','sha256':digest(self.pptx),'snapshot':str(self.pptx),'slide_count':1,
                     'status':'rendered_pending_visual_review','issues':[],'text_records':self.records,
                     'raster_hashes':{'1':digest(self.image)},'render_finished_utc':'2026-10-18T00:00:00+00:00'}
        self.coverage=audit(self.pptx,self.render)
        self.f=next(f for f in self.coverage['findings'] if f['reason']=='rendered_row_text_misaligned' and f['shape_ids']==[self.ids[0],self.ids[3]])
        self.errors=[ROW_ERROR];self.make_evidence()
    def tearDown(self):self.tmp.cleanup()
    def make_evidence(self):
        self.proof=layout_context(self.f,self.s,self.render,self.image,self.errors)
        bb=self.proof['required_context_box_px'];crop=self.root/'local.png'
        with Image.open(self.image) as im:im.crop(bb).save(crop)
        details=[{'path':str(crop),'sha256':digest(crop),'crop_box_px':bb}]
        self.record={'code':CODE,'pptx_sha256':self.render['sha256'],'image_sha256':digest(self.image),'machine_context':copy.deepcopy(self.proof),
          'measurement_limitation':NOTE,'local_observation':NOTE,'global_observation':NOTE,'returned_to_full_page':True,
          'normal_size_viewed':True,'no_visible_layout_defect':True,'local_result':'pass','global_result':'pass','unresolved':[],
          'target_observations':{str(i):NOTE for i in self.proof['context_object_ids']},'details':details,'reviewed_at':'2026-10-18T00:03:00+00:00'}
        self.peer={'finding_id':self.f['id'],'slide':1,'code':CODE,'reviewer_id':'synthetic-independent-reviewer',
          'pptx_sha256':self.render['sha256'],'image_sha256':digest(self.image),'machine_context':copy.deepcopy(self.proof),
          'local_views':[{'sha256':details[0]['sha256'],'crop_box_px':bb}],'decision':'visually_resolved','local_viewed':True,
          'full_page_viewed':True,'normal_size_viewed':True,'no_visible_layout_defect':True,'unresolved':[],
          'local_observation':NOTE,'global_observation':NOTE,'review_source':NOTE,'reviewed_at':'2026-10-18T00:02:00+00:00',
          'context_observations':{str(i):NOTE for i in self.proof['context_object_ids']},
          'rule_assessments':[{'original_error':e,'why_rule_failed_but_visual_passed':NOTE,'no_visible_layout_defect':True} for e in self.errors],
          'relationship_assessment':{'same_alignment_group':False,'same_metric_group':False,'structural_explanation':NOTE}}
        self.peer_path=self.root/'synthetic-review.json';self.write_peer()
        self.disposition={'visual_fallback':self.record}
    def write_peer(self):
        self.peer_path.write_text(json.dumps(self.peer),encoding='utf8')
        self.record['independent_review']={'path':str(self.peer_path),'sha256':digest(self.peer_path)}
    def call(self):return validate_layout_visual_review(self.f,self.disposition,self.image,self.s,self.render,'synthetic-author',self.errors)
    def check(self,accepted):
        error,resolution=self.call();self.assertEqual(error is None,accepted,error);self.assertEqual(resolution is not None,accepted)
    def test_unfilled_cross_column_relationship_has_complete_path(self):
        self.assertTrue(all(not s._element.xpath('./p:spPr/a:solidFill') for s in self.s.shapes));self.check(True)
        self.assertEqual(self.call()[1]['original_finding'],self.f)
    def test_numeric_threshold_is_preserved_not_weakened(self):
        self.assertGreater(self.f['center_delta_pt'],2);self.check(True)
        self.assertEqual(self.call()[1]['rule_failures'],[ROW_ERROR])
    def test_same_group_optical_difference_requires_actual_review_not_automatic_veto(self):
        self.peer['relationship_assessment'].update(same_alignment_group=True,same_metric_group=True);self.write_peer();self.check(True)
    def test_actual_visible_number_title_mismatch_cannot_close(self):
        self.peer['relationship_assessment'].update(same_alignment_group=True,same_metric_group=True)
        self.peer['no_visible_layout_defect']=False;self.peer['unresolved']=['The numeral and associated title visibly disagree.'];self.write_peer();self.check(False)
    def test_author_only_role_label_cannot_close(self):self.record.pop('independent_review');self.disposition['object_relation']='different roles';self.check(False)
    def test_no_rule_output_cannot_close(self):self.errors=[];self.check(False)
    def test_rewritten_original_delta_fails(self):self.record['machine_context']['original_finding']['center_delta_pt']=0;self.check(False)
    def test_missing_original_rule_failure_fails(self):self.record['machine_context']['rule_failures']=[];self.check(False)
    def test_stale_input(self):self.record['pptx_sha256']='OLD';self.check(False)
    def test_stale_image(self):self.record['image_sha256']='OLD';self.check(False)
    def test_omitted_native_object(self):self.record['machine_context']['page_object_ids'].pop();self.check(False)
    def test_new_transparent_proxy_is_not_an_exemption(self):
        s=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(35),Pt(95),Pt(780),Pt(120));s.fill.background()
        self.check(False)
    def test_new_blank_shape_invalidates_old_object_inventory(self):
        self.s.shapes.add_textbox(Pt(900),Pt(480),Pt(30),Pt(20));self.check(False)
    def test_omitted_context_observation(self):self.peer['context_observations'].pop(next(iter(self.peer['context_observations'])));self.write_peer();self.check(False)
    def test_fake_self_reviewer(self):self.peer['reviewer_id']='synthetic-author';self.write_peer();self.check(False)
    def test_fake_review_provenance(self):self.peer['review_source']='pass';self.write_peer();self.check(False)
    def test_missing_render_record_is_input_error(self):self.render['text_records'].clear();self.check(False)
    def test_unknown_measurement_is_preserved_for_visual_inspection(self):
        self.render['text_records'][0].pop('bound_left_pt');self.make_evidence();self.check(True)
    def test_wrong_crop_pixels(self):
        path=Path(self.record['details'][0]['path'])
        with Image.open(path) as original:size=original.size
        Image.new('RGB',size,'white').save(path)
        self.record['details'][0]['sha256']=digest(path);self.check(False)
    def test_full_page_is_not_a_local_view(self):self.record['details'][0]['crop_box_px']=[0,0,960,540];self.check(False)
    def test_unreviewed_rule_error_fails(self):self.peer['rule_assessments']=[];self.write_peer();self.check(False)
    def test_pre_render_review_fails(self):self.peer['reviewed_at']='2026-10-17T00:00:00+00:00';self.write_peer();self.check(False)
    def test_spacing_region_and_new_layout_reason_use_same_protocol(self):
        for reason in ('header_body_boundary_review','rendered_ink_column_misaligned','text_region_relation_requires_review','future_layout_rule'):
            self.f={**self.f,'reason':reason,'gaps_pt':[8,26],'threshold_pt':12};self.errors=['Original rule reports its current numeric mismatch.'];self.make_evidence();self.check(True)
    def gate(self):
        import verify_coverage_review as gate
        cp=self.root/'coverage.json';cp.write_text(json.dumps(self.coverage),encoding='utf8')
        (self.root/'audit.json').write_text(json.dumps(self.render),encoding='utf8')
        lock=self.root/'lock.json';lock.write_text('{}',encoding='utf8')
        items=[{'finding_id':f['id'],'decision':'pending'} for f in self.coverage['findings']]
        item=next(i for i in items if i['finding_id']==self.f['id'])
        item.update(self.disposition,decision='verified_in_context',image_sha256=digest(self.image),criterion=NOTE,observation=NOTE,
                    object_relation=NOTE,measurements={'original_finding':self.f})
        review={'skill_lock':{'path':str(lock),'sha256':digest(lock)},'sha256':digest(self.pptx),'coverage_sha256':digest(cp),'author_id':'synthetic-author','items':items}
        stubs={'check_skill_state':{'changed_files':[]},'validate_render_evidence':None,'verify_connections':[],
               'validate_gap_audit':([],[]),'validate_manual_dispositions':[],'review_fonts':[],'check_visual_passes':[]}
        with ExitStack() as stack:
            for name,result in stubs.items():stack.enter_context(patch.object(gate,name,return_value=result))
            return gate.verify(self.pptx,cp,review,self.root)
    def test_real_coverage_dispatch_resolves_target_and_preserves_other_failures(self):
        result=self.gate();self.assertFalse([i for i in result['issues'] if i.get('id')==self.f['id']],result)
        self.assertTrue(result['issues']);self.assertEqual(result['status'],'failed')
        self.assertEqual(result['visual_resolutions'][0]['original_finding'],self.f)
    def test_real_dispatch_without_independent_review_keeps_hard_failure(self):
        self.peer_path.unlink();self.assertTrue([i for i in self.gate()['issues'] if i.get('id')==self.f['id']])
    def test_actual_42pt_numeral_14pt_title_remains_flagged_and_needs_review(self):
        p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540);s=p.slides.add_slide(p.slide_layouts[6]);records=[];ids=[]
        for x,size,label in [(40,42,'01'),(300,14,'Corresponding heading')]:
            sh=s.shapes.add_textbox(Pt(x),Pt(40),Pt(180),Pt(48));sh.text=label;sh.text_frame.paragraphs[0].runs[0].font.size=Pt(size);ids.append(sh.shape_id)
            records.append({'slide':1,'kind':'text','shape_id':sh.shape_id,'shape':sh.name,'text':label,'font_pt':size,
              'left_pt':x,'top_pt':40,'width_pt':180,'height_pt':48,'bound_left_pt':x,'bound_top_pt':40,'bound_width_pt':100,'bound_height_pt':size})
        path=self.root/'number-title.pptx';p.save(path);render={**self.render,'sha256':digest(path),'snapshot':str(path),'text_records':records}
        finding=next(f for f in audit(path,render)['findings'] if f['reason']=='rendered_row_text_misaligned' and f['shape_ids']==ids)
        self.assertEqual(finding['center_delta_pt'],14)
        error,resolution=validate_layout_visual_review(finding,{},self.image,s,render,'synthetic-author',[ROW_ERROR])
        self.assertIsNotNone(error);self.assertIsNone(resolution)
    def test_completed_preflight_failure_can_still_be_diagnosed_without_rewriting_it(self):
        before={'reason':'text_overflow','slide':1,'shape_id':self.ids[0]};self.render.update(status='failed_preflight',issues=[before])
        self.assertEqual(audit(self.pptx,self.render),self.coverage);self.assertEqual(self.render['issues'],[before])
    def test_non_chart_raw_spacing_failure_can_use_identical_layout_protocol(self):
        from audit_layout_gaps import audit as gap_audit
        from verify_coverage_review import validate_gap_audit
        spec={'pages':[{'slide':1,'manual_axes':{'vertical':NOTE},'comparisons':[{'name':'cross-columns','axis':'horizontal','target_gap_pt':240,
               'blocks':[{'shape_ids':[self.ids[0]]},{'shape_ids':[self.ids[3]]}]}]}]}
        raw=gap_audit(self.pptx,spec);self.assertEqual(len(raw['issues']),1)
        self.f={'id':'gap-raw:1','slide':1,'reason':'original_gap_layout_rule','original_issue':raw['issues'][0],'original_comparison':raw['comparisons'][0]}
        self.errors=[json.dumps(raw['issues'][0],ensure_ascii=False,sort_keys=True)];self.make_evidence()
        layout=self.root/'layout.json';layout.write_text(json.dumps(spec),encoding='utf8');report=self.root/'gap.json';report.write_text(json.dumps(raw),encoding='utf8')
        review={'author_id':'synthetic-author','gap_audit':{'layout_path':str(layout),'layout_sha256':digest(layout),'report_path':str(report),'report_sha256':digest(report),
                    'visible_reviews':[],'layout_reviews':[{'finding_id':self.f['id'],**self.disposition}]}}
        errors,resolutions=validate_gap_audit(self.pptx,review,self.render,self.root)
        self.assertEqual(errors,[]);self.assertEqual(resolutions[0]['original_finding']['original_issue'],raw['issues'][0])
        review['gap_audit']['layout_reviews']=[];self.assertTrue(validate_gap_audit(self.pptx,review,self.render,self.root)[0])




if __name__=='__main__':unittest.main()
