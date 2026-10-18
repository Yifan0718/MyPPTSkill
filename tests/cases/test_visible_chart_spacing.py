"""Synthetic adapter contracts, not PowerPoint renders or visual signoffs."""
import sys
sys.dont_write_bytecode=True
import copy,json,shutil,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement
from audit_layout_gaps import audit
import visible_chart_spacing as vs


class VisibleChartSpacingTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='visible-spacing-test-');self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.pptx=self.root/'synthetic.pptx';self.snapshot=self.root/'snapshot.pptx';self.image=self.root/'slide-01.png'
        self.p=Presentation();self.p.slide_width=Pt(500);self.p.slide_height=Pt(540)
        self.s=self.p.slides.add_slide(self.p.slide_layouts[6]);data=ChartData();data.categories=['A'];data.add_series('S',[1])
        self.chart=self.s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Pt(50),Pt(50),Pt(300),Pt(200),data)
        self.caption=self.s.shapes.add_textbox(Pt(50),Pt(255),Pt(300),Pt(30));self.caption.text='Synthetic caption'
        self.after=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(50),Pt(300),Pt(400),Pt(50))
        self.after.fill.solid();self.after.fill.fore_color.rgb=RGBColor(30,60,90);self.after.line.fill.background()
        self.im=Image.new('RGB',(500,540),'white');d=ImageDraw.Draw(self.im)
        d.rectangle((80,75,320,199),fill=(30,60,90));d.rectangle((130,220,240,229),fill=(30,60,90))
        d.rectangle((70,260,190,269),fill=(80,100,125));d.rectangle((50,300,449,349),fill=(30,60,90))
        self.refresh()

    def refresh(self):
        self.p.save(self.pptx);shutil.copyfile(self.pptx,self.snapshot);self.im.save(self.image)
        self.spec={'pages':[{'slide':1,'comparisons':[{'name':'chart-caption-container','axis':'vertical','target_gap_pt':12,'minimum_gap_pt':12,'blocks':[{'name':str(s.shape_id),'shape_ids':[s.shape_id]} for s in (self.chart,self.caption,self.after)]}],'manual_axes':{'horizontal':'Synthetic fixture has one vertical relationship; no full deck signoff is claimed.'}}]}
        self.raw=audit(self.pptx,self.spec);self.row=copy.deepcopy(self.raw['comparisons'][0])
        self.render={'schema':'render-evidence/v2','sha256':vs._sha(self.pptx),'snapshot':str(self.snapshot),'slide_count':1,
            'status':'rendered_pending_visual_review','issues':[],'raster_hashes':{'1':vs._sha(self.image)},
            'text_records':[{'slide':1,'audit_path':str(self.caption.shape_id),'shape_id':self.caption.shape_id,'text':self.caption.text}]}

    def measure(self,row=None):
        return vs.measure_relation(self.pptx,self.render,self.root,1,self.row if row is None else row)

    def test_valid_visible_gaps_retain_original_failed_frame_row(self):
        r=self.measure()
        self.assertEqual(r['measurement_status'],'measured_valid_pending_visual_review')
        self.assertEqual(r['visible_edges_px'],[230,260,270,300])
        self.assertEqual(r['visible_gaps_pt'],[30,30])
        self.assertEqual(r['frame_gaps_pt'],[5,15])
        self.assertEqual(r['original_frame_comparison'],self.row)
        self.assertTrue(self.raw['issues'])

    def test_whole_chart_includes_content_below_legend(self):
        ImageDraw.Draw(self.im).rectangle((150,240,160,245),fill='black');self.refresh()
        r=self.measure();self.assertEqual(r['chart_bottom_px'],246)
        self.assertEqual(r['measurement_status'],'measured_defect')

    def test_unbalanced_old_caption_position_is_defect(self):
        d=ImageDraw.Draw(self.im);d.rectangle((50,255,349,284),fill='white');d.rectangle((70,273,190,282),fill=(80,100,125))
        self.caption.top=Pt(268);self.caption.height=Pt(20);self.refresh()
        r=self.measure();self.assertEqual(r['measurement_status'],'measured_defect')
        self.assertIn('visible_peer_gap_imbalance',r['defects'])

    def test_true_eight_pt_gaps_do_not_lower_minimum(self):
        self.chart.height=Pt(182);self.caption.top=Pt(233);self.caption.height=Pt(17);self.after.top=Pt(256)
        d=ImageDraw.Draw(self.im);d.rectangle((50,233,449,349),fill='white');d.rectangle((70,238,190,247),fill=(80,100,125));d.rectangle((50,256,449,305),fill=(30,60,90))
        self.refresh();r=self.measure()
        self.assertEqual(r['visible_gaps_pt'],[8,8]);self.assertEqual(r['minimum_gap_pt'],12)
        self.assertEqual(r['measurement_status'],'measured_defect')

    def test_actual_overlap_cannot_pass(self):
        self.caption.top=Pt(220)
        d=ImageDraw.Draw(self.im);d.rectangle((50,255,349,284),fill='white');d.rectangle((130,225,240,235),fill=(80,100,125))
        self.refresh();r=self.measure()
        self.assertNotEqual(r['measurement_status'],'measured_valid_pending_visual_review')

    def test_picture_is_not_a_native_chart(self):
        picture=self.root/'source.png';Image.new('RGB',(300,200),'navy').save(picture)
        self.s.shapes._spTree.remove(self.chart._element)
        self.chart=self.s.shapes.add_picture(str(picture),Pt(50),Pt(50),Pt(300),Pt(200));self.refresh()
        self.assertEqual(self.measure()['measurement_status'],'invalid_relation')

    def test_missing_object_is_invalid(self):
        row=copy.deepcopy(self.row);row['blocks'][0]['shape_ids']=[99999]
        self.assertEqual(self.measure(row)['measurement_status'],'invalid_relation')

    def test_renamed_local_background_proxy_is_not_exempt(self):
        extra=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(100),Pt(251),Pt(20),Pt(2));extra.name='background:fake'
        self.refresh();r=self.measure()
        self.assertEqual(r['measurement_status'],'invalid_relation');self.assertIn(extra.shape_id,r['unlisted_shape_ids'])

    def test_omitted_nearer_container_is_invalid(self):
        extra=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(50),Pt(290),Pt(300),Pt(3));self.refresh()
        r=self.measure();self.assertEqual(r['measurement_status'],'invalid_relation');self.assertIn(extra.shape_id,r['unlisted_shape_ids'])

    def test_frame_measurement_rewrite_is_invalid(self):
        row=copy.deepcopy(self.row);row['gaps_pt']=[30,30]
        self.assertEqual(self.measure(row)['measurement_status'],'invalid_relation')

    def test_transparent_after_proxy_is_invalid(self):
        alpha=OxmlElement('a:alpha');alpha.set('val','0');self.after._element.xpath('./p:spPr/a:solidFill/*')[0].append(alpha);self.refresh()
        self.assertEqual(self.measure()['measurement_status'],'invalid_relation')

    def test_texture_is_manual_and_locks_known_caption_after_edges(self):
        d=ImageDraw.Draw(self.im)
        for y in range(50,250,2):d.line((47,y,49,y),fill=(120,130,140))
        self.refresh();r=self.measure()
        self.assertEqual(r['measurement_status'],'manual_review_required')
        self.assertEqual(r['known_edge_indices'],[1,2,3]);self.assertEqual(r['visible_edges_px'],[None,260,270,300])

    def test_invisible_white_container_does_not_pass(self):
        self.after.fill.fore_color.rgb=RGBColor(255,255,255);ImageDraw.Draw(self.im).rectangle((50,300,449,349),fill='white');self.refresh()
        r=self.measure();self.assertEqual(r['measurement_status'],'manual_review_required');self.assertIsNone(r['visible_edges_px'][3])

    def test_unknown_chart_cannot_hide_known_eight_pt_lower_gap(self):
        self.caption.height=Pt(18);self.after.top=Pt(278)
        d=ImageDraw.Draw(self.im);d.rectangle((50,278,449,349),fill='white');d.rectangle((50,278,449,327),fill=(30,60,90))
        for y in range(50,250,2):d.line((47,y,49,y),fill=(120,130,140))
        self.refresh();r=self.measure()
        self.assertEqual(r['measurement_status'],'measured_defect')
        self.assertEqual(r['visible_gaps_pt'],[None,8])

    def test_rounded_opaque_after_is_supported(self):
        self.after._element.xpath('./p:spPr/a:prstGeom')[0].set('prst','roundRect')
        d=ImageDraw.Draw(self.im);d.rectangle((50,300,449,349),fill='white');d.rounded_rectangle((50,300,449,349),radius=10,fill=(30,60,90));self.refresh()
        self.assertEqual(self.measure()['measurement_status'],'measured_valid_pending_visual_review')

    def test_lowered_original_minimum_is_not_accepted(self):
        row=copy.deepcopy(self.row);row['minimum_gap_pt']=8
        self.assertEqual(self.measure(row)['measurement_status'],'invalid_relation')

    def test_stale_png_is_hard_failure(self):
        Image.new('RGB',(500,540),'white').save(self.image)
        self.assertEqual(self.measure()['measurement_status'],'invalid_relation')

    def test_optional_dependency_missing_is_manual(self):
        with patch.object(vs,'np',None):r=self.measure()
        self.assertEqual(r['measurement_status'],'manual_review_required')

    def enable_discovery(self):
        self.after.text='Synthetic decision';self.refresh()

    def discover(self):
        return vs.discover_relations(self.pptx,self.render,self.root)

    def test_discovery_is_independent_of_layout_and_names(self):
        self.after.text='Synthetic decision';self.caption.name='background:still-a-caption';self.chart.name='unclassified';self.refresh()
        rows=self.discover();self.assertEqual(len(rows),1)
        self.assertEqual([b['shape_ids'][0] for b in rows[0]['blocks']],[self.chart.shape_id,self.caption.shape_id,self.after.shape_id])
        self.assertTrue(rows[0]['name'].startswith('visible-chart-caption-p1-'))
        self.assertEqual(self.measure(rows[0])['measurement_status'],'measured_valid_pending_visual_review')

    def test_discovery_detects_defect_even_when_original_frame_report_passes(self):
        self.caption.top=Pt(262);self.caption.height=Pt(20);self.after.top=Pt(294);self.after.text='Synthetic decision'
        d=ImageDraw.Draw(self.im);d.rectangle((50,255,449,349),fill='white');d.rectangle((70,267,190,276),fill=(80,100,125));d.rectangle((50,294,449,343),fill=(30,60,90));self.refresh()
        self.assertEqual(self.raw['issues'],[])
        rows=self.discover();self.assertEqual(len(rows),1)
        self.assertEqual(self.measure(rows[0])['measurement_status'],'measured_defect')

    def test_discovery_does_not_jump_over_formula(self):
        self.after.text='Synthetic decision'
        formula=self.s.shapes.add_textbox(Pt(50),Pt(287),Pt(300),Pt(8));formula.text='P = lambda / mu'
        self.refresh();self.assertEqual(self.discover(),[])

    def test_discovery_does_not_jump_over_another_module(self):
        self.after.text='Synthetic decision'
        extra=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(50),Pt(287),Pt(300),Pt(8));extra.fill.solid();extra.fill.fore_color.rgb=RGBColor(90,110,130);extra.line.fill.background()
        self.refresh();self.assertEqual(self.discover(),[])

    def test_discovery_does_not_use_label_inside_another_container(self):
        self.after.text='Synthetic decision'
        panel=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(45),Pt(240),Pt(310),Pt(50));panel.fill.solid();panel.fill.fore_color.rgb=RGBColor(225,230,235);panel.line.fill.background()
        self.refresh();self.assertEqual(self.discover(),[])

    def test_discovery_does_not_call_empty_decoration_a_conclusion(self):
        self.assertEqual(self.discover(),[])

    def test_stricter_original_minimum_is_preserved(self):
        row=copy.deepcopy(self.row);row['minimum_gap_pt']=35
        r=self.measure(row)
        self.assertEqual(r['minimum_gap_pt'],35)
        self.assertEqual(r['measurement_status'],'measured_defect')

    def test_discovery_stale_pixels_raise_instead_of_silent_no_matches(self):
        self.enable_discovery();Image.new('RGB',(500,540),'white').save(self.image)
        with self.assertRaisesRegex(ValueError,'stale'):self.discover()

    def test_cli_execution_failure_is_nonzero_and_structured(self):
        import subprocess
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);report=root/'failed.json'
            run=subprocess.run([sys.executable,str(Path(vs.__file__)),str(root/'missing.pptx'),str(root/'missing-render'),str(report)],capture_output=True)
            self.assertEqual(run.returncode,3)
            data=json.loads(report.read_text(encoding='utf8'))
            self.assertEqual(data['status'],'failed');self.assertIn('error',data)


if __name__=='__main__':unittest.main(verbosity=2)
