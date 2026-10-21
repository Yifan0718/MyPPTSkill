"""Generate 4 negative/4 corrected pages for REAL PowerPoint visual inspection.
Usage: python generate_usage_visual_fixture.py <fresh-output-directory>
This is a diagnostic fixture, never a signed-off customer deck.
"""
from pathlib import Path
import sys,json
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_DATA_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from PIL import Image,ImageDraw

def create(out):
 out=Path(out);out.mkdir(parents=True,exist_ok=False)
 p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540)
 im=Image.new('RGB',(600,240),'white');d=ImageDraw.Draw(im);d.rectangle((100,75,400,160),fill='#98AABB');d.ellipse((400,95,460,155),fill='#2781EA');im.save(out/'scene.png')
 def text(sl,x,y,w,h,txt,size=20,color='26334D',bold=False):
  sh=sl.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));sh.text=txt;tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=Pt(0);tf.margin_top=tf.margin_bottom=Pt(0)
  for pa in tf.paragraphs:
   pa.alignment=PP_ALIGN.LEFT
   for r in pa.runs:r.font.name='Microsoft YaHei';r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(color);r.font.bold=bold
  return sh
 def table(sl):
  sh=sl.shapes.add_table(3,3,Pt(48),Pt(130),Pt(864),Pt(168));sh.name='table:explanation'
  vals=[['模型预测什么','真实答案来源','如何纠正'],['未来视觉表示','真实画面经编码器转换','比较预测与真实表示'],['物体后果差','同起点两种顺序的结果','比较预测差与真实差']]
  for i in range(3):
   for j in range(3):
    c=sh.table.cell(i,j);c.text=vals[i][j];c.margin_left=c.margin_right=Pt(10);c.margin_top=c.margin_bottom=Pt(6);c.vertical_anchor=MSO_ANCHOR.MIDDLE;c.fill.solid();c.fill.fore_color.rgb=RGBColor.from_string('263B65' if i==0 else 'EDF1F5')
    for pa in c.text_frame.paragraphs:
     pa.alignment=PP_ALIGN.LEFT
     for r in pa.runs:r.font.name='Microsoft YaHei';r.font.size=Pt(18);r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '26334D');r.font.bold=i==0
  return sh
 for good in (False,True):
  for kind in range(4):
   sl=p.slides.add_slide(p.slide_layouts[6]);n=len(p.slides)
   text(sl,48,30,864,48,('修复对照' if good else '故意反例')+'｜'+['说明表对齐','表格—说明—流程间距','原生图例与数字','色调与主体完整性'][kind],28,bold=True)
   text(sl,48,87,864,28,'合成测试内容，用于规则与实际视觉交叉核验。',14,color='66728B')
   if kind in (0,1):
    t=table(sl)
    if kind==0:
     if not good:
      for row in t.table.rows:
       for j in (1,2):row.cells[j].text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER
     text(sl,48,350,864,80,'三列均为说明文字。检查表头、正文与单元格留白。',20)
    else:
     y=322 if good else 365;h=48
     text(sl,48,y,864,h,'每段动作包含5条控制指令，三段合计15步。\n真实执行提供答案，训练用误差更新模型。',18,color='66728B')
     sy=y+h+24 if good else y+h+5
     for k,lab in enumerate(['分别预测','对照真实答案','按误差更新']):
      sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(48+300*k),Pt(sy),Pt(264),Pt(50));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string('263B65');sh.line.fill.background()
      text(sl,62+300*k,sy+12,235,30,lab,18,'FFFFFF',True)
   elif kind==2:
    data=CategoryChartData();data.categories=['场景一','场景二','场景三'];data.add_series('基线',[54.4,48,42]);data.add_series('新方法',[52.8,44,37])
    sh=sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Pt(48),Pt(130),Pt(864),Pt(300),data);sh.name='chart:legend-fixture';c=sh.chart;c.has_legend=True;c.legend.position=XL_LEGEND_POSITION.TOP;c.legend.include_in_layout=False;c.legend.font.size=Pt(16)
    plot=c.plots[0];plot.has_data_labels=True;plot.data_labels.position=XL_DATA_LABEL_POSITION.OUTSIDE_END;plot.data_labels.font.size=Pt(16)
    c.value_axis.major_gridlines.format.line.color.rgb=RGBColor.from_string('DDE4EE');c.value_axis.major_gridlines.format.line.width=Pt(.6);c.value_axis.tick_labels.font.size=Pt(14);c.category_axis.tick_labels.font.size=Pt(16)
    text(sl,48,455,864,30,'模拟位置误差：预测位置与真实位置的距离，单位mm，越低越好。',18)
   else:
    colors=['E8A92F','41483D','51443E'] if not good else ['F6BD45','24405E','EFF3F5']
    for k,col in enumerate(colors):
     x=48+k*300;sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(145),Pt(264),Pt(90));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(col);sh.line.fill.background();sh.text=['参与共创','清晰层次','统一色调'][k]
     for pa in sh.text_frame.paragraphs:
      pa.alignment=PP_ALIGN.CENTER
      for r in pa.runs:r.font.name='Microsoft YaHei';r.font.size=Pt(24);r.font.bold=True;r.font.color.rgb=RGBColor.from_string((['181818','FFFFFF','181818'][k]) if good else '15509B')
     sh.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    pic=sl.shapes.add_picture(str(out/'scene.png'),Pt(48),Pt(260),width=Pt(600),height=Pt(220))
    if not good:pic.crop_right=.29
    else:
     pic.crop_left=80/600;pic.crop_right=120/600;pic.crop_top=55/240;pic.crop_bottom=60/240;pic.width=Pt(500);pic.height=Pt(156.25);pic.top=Pt(280)
    text(sl,580 if good else 690,310 if good else 300,332 if good else 220,100,'检查底色协调、\n字色与主体裁切。',20)
   text(sl,866,490,46,16,f'{n}/8',10)
 p.save(out/'fixture.pptx')
 (out/'expected.json').write_text(json.dumps({'negative_pages':[1,2,3,4],'corrected_pages':[5,6,7,8],'note':'Page 3 legend position is set by the companion COM step before rendering; all pages require actual visual review.'},indent=2),encoding='utf8')
 print(out/'fixture.pptx')
if __name__=='__main__':create(sys.argv[1])
