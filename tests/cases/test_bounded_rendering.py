import contextlib,io,json,subprocess,sys,tempfile,unittest,time
from pathlib import Path
from unittest.mock import patch,Mock
from pptx import Presentation
from PIL import Image
from run_bounded import run
from render_evidence import validate_bounded_run
from render_slides import render,_run_converter

class BoundedTests(unittest.TestCase):
 def test_success(self):self.assertEqual(run([sys.executable,'-c','pass'],5),0)
 def test_nonzero_child_exit_propagates(self):self.assertEqual(run([sys.executable,'-c','raise SystemExit(7)'],5),7)
 def test_timeout_returns_recovery(self):
  err=io.StringIO()
  with contextlib.redirect_stderr(err):code=run([sys.executable,'-c','import time;time.sleep(5)'],.1)
  self.assertEqual(code,3);self.assertEqual(json.loads(err.getvalue())['reason'],'task_tool_timeout')
 def test_invalid_timeout_does_not_launch(self):
  with patch('run_bounded.subprocess.Popen') as launch,contextlib.redirect_stderr(io.StringIO()):
   for value in (0,-1,float('nan'),float('inf')):self.assertEqual(run(['tool'],value),3)
   launch.assert_not_called()
 def test_missing_executable_is_structured(self):
  with patch('run_bounded.subprocess.Popen',side_effect=FileNotFoundError('missing')),contextlib.redirect_stderr(io.StringIO()):self.assertEqual(run(['tool'],5),3)
 def test_cancellation_stops_only_owned_command(self):
  child=Mock();child.wait.side_effect=[KeyboardInterrupt(),0]
  with patch('run_bounded.subprocess.Popen',return_value=child),patch('run_bounded.stop_owned_tree') as cleanup,contextlib.redirect_stderr(io.StringIO()):self.assertEqual(run(['tool'],5),130)
  cleanup.assert_called_once_with(child)

 def test_timeout_stops_task_descendants(self):
  with tempfile.TemporaryDirectory() as td:
   flag=Path(td)/'late.txt';worker=Path(td)/'worker.py';parent=Path(td)/'parent.py'
   worker.write_text('import time\nfrom pathlib import Path\ntime.sleep(1.5)\nPath('+repr(str(flag))+').write_text("late")\n','utf8')
   parent.write_text('import subprocess,sys,time\nsubprocess.Popen([sys.executable,'+repr(str(worker))+'])\ntime.sleep(8)\n','utf8')
   with contextlib.redirect_stderr(io.StringIO()):self.assertEqual(run([sys.executable,str(parent)],.5),3)
   time.sleep(1.6);self.assertFalse(flag.exists())

class RunEvidenceTests(unittest.TestCase):
 def setUp(self):
  self.t=tempfile.TemporaryDirectory();self.addCleanup(self.t.cleanup);self.w=Path(self.t.name);self.path=self.w/'tool-run.json'
 def record(self,status='finished',code=0):
  self.path.write_text(json.dumps({'schema':'bounded-run/v1','run_id':'run1','status':status,'exit_code':code}),'utf8')
  return {'bounded_run_id':'run1','status':'rendered_pending_visual_review'}
 def test_success_record(self):
  self.assertEqual(run([sys.executable,'-c','pass'],5,self.path),0);r=json.loads(self.path.read_text());self.assertEqual(r['status'],'finished');self.assertEqual(r['exit_code'],0)
 def test_existing_status_is_preserved(self):
  self.path.write_text('old','utf8')
  with contextlib.redirect_stderr(io.StringIO()):self.assertEqual(run([sys.executable,'-c','pass'],5,self.path),3)
  self.assertEqual(self.path.read_text(),'old')
 def test_success_validation(self):validate_bounded_run(self.record(),self.w)
 def test_timeout_never_reuses_completed_audit(self):
  with self.assertRaisesRegex(ValueError,'timed out'):validate_bounded_run(self.record('timed_out',3),self.w)
 def test_missing_expected_record_rejected(self):
  with self.assertRaisesRegex(ValueError,'Missing'):validate_bounded_run({'bounded_run_id':'run1'},self.w)
 def test_failed_process_rejected(self):
  with self.assertRaisesRegex(ValueError,'invocation failed'):validate_bounded_run(self.record('finished',7),self.w)
 def test_preflight_review_exit_allowed(self):
  r=self.record('finished',2);r['status']='failed_preflight';validate_bounded_run(r,self.w)
 def test_running_record_rejected(self):
  with self.assertRaisesRegex(ValueError,'incomplete'):validate_bounded_run(self.record('running',0),self.w)
 def test_mismatched_run_id_rejected(self):
  r=self.record();r['bounded_run_id']='old'
  with self.assertRaisesRegex(ValueError,'identity'):validate_bounded_run(r,self.w)
 def test_legacy_without_wrapper_evidence_supported(self):validate_bounded_run({},self.w)

class FallbackRenderTests(unittest.TestCase):
 def setUp(self):
  self.t=tempfile.TemporaryDirectory();self.addCleanup(self.t.cleanup);self.w=Path(self.t.name);self.p=self.w/'input.pptx'
  p=Presentation();p.slides.add_slide(p.slide_layouts[6]);p.save(self.p);self.calls=[]
 def fake(self,command,**kwargs):
  self.calls.append((command,kwargs))
  if command[0]=='soffice':(Path(command[command.index('--outdir')+1])/'input.pdf').write_bytes(b'fake-pdf')
  else:Image.new('RGB',(80,45),'white').save(Path(command[-1]).parent/'slide-1.png')
  return Mock(returncode=0)
 def test_both_converters_bounded_and_profiles_isolated(self):
  with patch('render_slides._run_converter',side_effect=self.fake):
   render(self.p,self.w/'a',144,3);render(self.p,self.w/'b',144,3)
  self.assertTrue(all(k['timeout']==3 for _,k in self.calls));profiles=[next(x for x in c if x.startswith('-env:UserInstallation=')) for c,k in self.calls if c[0]=='soffice'];self.assertEqual(len(set(profiles)),2)
 def test_timeout_propagates_without_second_stage(self):
  with patch('render_slides._run_converter',side_effect=subprocess.TimeoutExpired('soffice',1)) as run:
   with self.assertRaises(subprocess.TimeoutExpired):render(self.p,self.w/'a',144,1)
  self.assertEqual(run.call_count,1)
 def test_partial_invalid_png_does_not_pass(self):
  def broken(command,**kwargs):
   self.fake(command,**kwargs)
   if command[0]=='pdftoppm':(Path(command[-1]).parent/'slide-1.png').write_bytes(b'not an image')
  with patch('render_slides._run_converter',side_effect=broken):
   with self.assertRaises(OSError):render(self.p,self.w/'a',144,1)
 def test_converter_timeout_kills_descendants(self):
  flag=self.w/'converter-late.txt';worker=self.w/'worker.py';parent=self.w/'parent.py'
  worker.write_text('import time\nfrom pathlib import Path\ntime.sleep(1.5)\nPath('+repr(str(flag))+').write_text("late")\n','utf8')
  parent.write_text('import subprocess,sys,time\nsubprocess.Popen([sys.executable,'+repr(str(worker))+'])\ntime.sleep(8)\n','utf8')
  with self.assertRaises(subprocess.TimeoutExpired):_run_converter([sys.executable,str(parent)],timeout=.5)
  time.sleep(1.6);self.assertFalse(flag.exists())
 def test_wrong_page_number_does_not_pass(self):
  def wrong(command,**kwargs):
   self.fake(command,**kwargs)
   if command[0]=='pdftoppm':(Path(command[-1]).parent/'slide-1.png').rename(Path(command[-1]).parent/'slide-3.png')
  with patch('render_slides._run_converter',side_effect=wrong):
   with self.assertRaisesRegex(RuntimeError,'numbering'):render(self.p,self.w/'a',144,1)
if __name__=='__main__':unittest.main()
