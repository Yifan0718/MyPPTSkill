"""Conservative interstitial spacing contract draft, not visual signoff.
Opaque solid rectangles and actual rendered text glyph bounds form visible rows.
Unsupported contours remain pending for evidence-bound local/full-page review.
"""
import math
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

def _box(s):
    return [s.left/12700,s.top/12700,(s.left+s.width)/12700,(s.top+s.height)/12700]

def _opaque_rectangle(s, allow_white=False):
    try:
        if s.auto_shape_type != MSO_AUTO_SHAPE_TYPE.RECTANGLE or abs(s.rotation)>1e-6:return False
        if not s._element.xpath('./p:spPr/a:ln/a:noFill'):return False
        if s._element.xpath('./p:nvSpPr/p:cNvPr[@hidden="1" or @hidden="true"]'):return False
        fills=s._element.xpath('./p:spPr/a:solidFill')
        if len(fills)!=1:return False
        colors=list(fills[0])
        if len(colors)!=1 or colors[0].tag.rsplit('}',1)[-1]!='srgbClr':return False
        rgb=colors[0].get('val','').upper()
        if len(rgb)!=6 or any(c not in '0123456789ABCDEF' for c in rgb):return False
        if rgb=='FFFFFF' and not allow_white:return False
        # White-on-unknown background and inherited theme paints are not proof
        # of a visible boundary; caller must use a separate verified contour path.
        # Reject inherited/partial/relative alpha rather than guessing visibility.
        for node in colors[0]:
            tag=node.tag.rsplit('}',1)[-1]
            if tag.startswith('alpha') and (tag!='alpha' or int(node.get('val','0'))!=100000):return False
        return True
    except (AttributeError,ValueError,TypeError):return False

def _full_slide_background_ids(slide):
    """Prove background paint from native geometry and back-to-front order.

    Only a bottom-most prefix of full-slide, unrotated, explicitly opaque RGB
    rectangles with no text or line qualifies. Each is behind every local
    content object, including row labels discovered later. A foreground cover
    or a shape merely named "background" never receives this exemption.
    """
    try:
        presentation=slide.part.package.presentation_part.presentation
        width,height=presentation.slide_width,presentation.slide_height
        if width is None or height is None or min(width,height)<=0:return set()
        backgrounds=set()
        for shape in slide.shapes:
            if not _opaque_rectangle(shape,allow_white=True):break
            if shape.has_text_frame and shape.text.strip():break
            # Use native EMUs, without measurement tolerance: even a small
            # uncovered strip means this is not proven full-slide paint.
            if shape.left>0 or shape.top>0 or shape.left+shape.width<width or shape.top+shape.height<height:break
            backgrounds.add(shape.shape_id)
        return backgrounds
    except (AttributeError,ValueError,TypeError):return set()

def validate_interstitial_spacing(finding, disposition, slide, render, scale):
    """Return None only for a measured, complete, nearest local relationship.
    Caller validates current render provenance and actual detail/full-page evidence.
    """
    if not isinstance(scale,(float,int)) or not math.isfinite(scale) or scale<=0:return 'Invalid layout scale'
    if finding.get('reason')!='header_body_boundary_review':return 'Interstitial contract only applies to header/body review'
    c=disposition.get('spacing_contract',{})
    if c.get('kind')!='interstitial_text':return 'Missing interstitial_text contract'
    if not isinstance(c.get('hierarchy_observation'),str) or len(c['hierarchy_observation'].strip())<20:return 'Explain the actual local relationship and why global gaps belong to different hierarchy levels'
    mid=c.get('middle_shape_id');before=c.get('before_shape_ids');after=c.get('after_shape_id')
    if type(mid) is not int or type(after) is not int or not isinstance(before,list) or not before or any(type(i) is not int for i in before):return 'Invalid adjacent object IDs'
    ids=[mid,after]+before
    if len(set(ids))!=len(ids):return 'Duplicate or overlapping contract object IDs'
    shapes={s.shape_id:s for s in slide.shapes}
    if any(i not in shapes for i in ids):return 'Contract references missing objects'
    middle=shapes[mid]
    if not middle.has_text_frame or not middle.text.strip():return 'Middle object must be actual text'
    records=[r for r in render.get('text_records',[]) if r.get('slide')==finding.get('slide') and r.get('shape_id')==mid and r.get('kind')=='text']
    if len(records)!=1:return 'Need exactly one current rendered middle text record'
    r=records[0];keys=('bound_left_pt','bound_top_pt','bound_width_pt','bound_height_pt')
    if any(type(r.get(k)) not in (int,float) or not math.isfinite(r[k]) for k in keys):return 'Missing finite rendered glyph bounds'
    x,y,w,h=[r[k] for k in keys]
    if w<=0 or h<=0:return 'Empty rendered glyph bounds'
    mb=[x,y,x+w,y+h];frame=_box(middle);eps=.25*scale
    if mb[0]<frame[0]-eps or mb[1]<frame[1]-eps or mb[2]>frame[2]+eps or mb[3]>frame[3]+eps:return 'Rendered middle text exceeds its frame'
    if not _opaque_rectangle(shapes[after]):return 'After boundary must be a native opaque solid rectangle'
    boxes={i:_box(s) for i,s in shapes.items()}
    backgrounds=_full_slide_background_ids(slide)
    # Frame span identifies all modules in the explanation's row; it never
    # substitutes for glyph edges in coverage or vertical gap measurements.
    def overlaps_x(b):return min(b[2],frame[2])-max(b[0],frame[0])>eps
    containers={i:b for i,b in boxes.items() if i!=mid and i not in backgrounds and _opaque_rectangle(shapes[i]) and overlaps_x(b)}
    # Native nonempty text is visible only through current measured glyph bounds,
    # never through its unfilled frame. Opaque cards retain their visible fill box.
    text_bounds={}
    for sid,shape in shapes.items():
        if sid==mid or sid in containers or not shape.has_text_frame or not shape.text.strip():continue
        rr=[a for a in render.get('text_records',[]) if a.get('slide')==finding.get('slide') and a.get('shape_id')==sid and a.get('kind')=='text']
        if len(rr)!=1 or any(type(rr[0].get(k)) not in (int,float) or not math.isfinite(rr[0][k]) for k in keys):continue
        tx,ty,tw,th=[rr[0][k] for k in keys]
        tb=[tx,ty,tx+tw,ty+th];fb=boxes[sid]
        if tw<=0 or th<=0 or abs(shape.rotation)>1e-6:continue
        if tb[0]<fb[0]-eps or tb[1]<fb[1]-eps or tb[2]>fb[2]+eps or tb[3]>fb[3]+eps:continue
        if overlaps_x(tb):text_bounds[sid]=tb
    preceding={i:b for i,b in containers.items() if b[3]<=mb[1]+eps}
    following={i:b for i,b in containers.items() if b[1]>=mb[3]-eps}
    if not preceding or not following:return 'Missing real adjacent containers'
    baseline=max(b[3] for b in preceding.values())
    anchor_ids={i for i,b in preceding.items() if abs(b[3]-baseline)<=eps}
    anchor_top=max(preceding[i][1] for i in anchor_ids)
    if anchor_top>=baseline-eps:return 'Nearest containers have no common row'
    visible={**containers,**text_bounds}
    def contained(b,outer):
        return b[0]>=outer[0]-eps and b[1]>=outer[1]-eps and b[2]<=outer[2]+eps and b[3]<=outer[3]+eps
    row_ids=set(anchor_ids)
    for sid,b in visible.items():
        if sid in anchor_ids or sid==after:continue
        if b[3]>mb[1]+eps:continue
        if min(b[3],baseline)-max(b[1],anchor_top)<=eps:continue
        # Text painted inside a real row card is represented by that card; a
        # freestanding left row label remains an independent required member.
        if sid in text_bounds and any(contained(b,containers[i]) for i in anchor_ids):continue
        row_ids.add(sid)
    if set(before)!=row_ids:return 'Before IDs must exactly cover the nearest visible row including independent glyph labels; do not omit or choose distant objects'
    row=[visible[i] for i in before]
    nearest_bottom=max(b[3] for b in row)
    if min(b[0] for b in row)>mb[0]+eps or max(b[2] for b in row)<mb[2]-eps:return 'Before row does not cover the rendered text horizontal range'
    # Any other unmeasured module intersecting this row cannot be silently
    # treated as invisible merely because it has no qualifying paint/record.
    for sid,b in boxes.items():
        if sid in ids or sid in visible or sid in backgrounds:continue
        if overlaps_x(b) and min(b[3],baseline)-max(b[1],anchor_top)>eps:
            if any(contained(b,containers[i]) for i in anchor_ids):continue
            return 'Unverified object intersects the preceding row: '+str(sid)
    nearest_top=min(b[1] for b in following.values())
    following_ids={i for i,b in following.items() if abs(b[1]-nearest_top)<=eps}
    if following_ids!={after}:return 'After object must be the unique nearest visible container'
    ab=boxes[after]
    if ab[0]>mb[0]+eps or ab[2]<mb[2]-eps:return 'After container does not cover the rendered text horizontal range'
    # Reject every unlisted local object in the corridor, regardless of its
    # author-written role. Only independently proven full-slide background paint
    # is excluded; transparent proxy frames and group shapes still block.
    corridor=[frame[0],nearest_bottom,frame[2],nearest_top]
    for sid,frame_bounds in boxes.items():
        b=text_bounds.get(sid,frame_bounds)
        if sid in ids or sid in backgrounds:continue
        if min(b[2],corridor[2])-max(b[0],corridor[0])>eps and min(b[3],corridor[3])-max(b[1],corridor[1])>eps:
            return 'Unlisted object intersects the local adjacency corridor: '+str(sid)
    above=mb[1]-nearest_bottom;below=nearest_top-mb[3]
    if min(above,below)<12*scale:return 'Actual local gap is below 12pt scaled minimum'
    if abs(above-below)>max(3*scale,.2*(above+below)/2):return 'Actual local whitespace is unbalanced; repair and rerender'
    return None
