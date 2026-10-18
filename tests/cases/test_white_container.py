"""Current raster contrast may prove a real white card; never a blank proxy."""
import tempfile,unittest
from pathlib import Path
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from verify_coverage_review import validate_region_contract,white_container_visible_in_pixels

class WhiteContainer(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.image=Path(self.tmp.name)/'page.png'
        self.p=Presentation();self.p.slide_width=Pt(960);self.p.slide_height=Pt(540);self.s=self.p.slides.add_slide(self.p.slide_layouts[6])
        self.cards=[];self.texts=[]
        for x in (60,500):
            c=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,Pt(x),Pt(100),Pt(350),Pt(300));c.fill.solid();c.fill.fore_color.rgb=RGBColor(255,255,255);c.line.fill.background();self.cards.append(c)
            t=self.s.shapes.add_textbox(Pt(x+20),Pt(130),Pt(280),Pt(60));t.text='Card contents';self.texts.append(t)
        self.f={'shape_ids':[s.shape_id for s in self.texts]};self.d={'relation_contract':{'kind':'separate_visible_containers','shape_ids':self.f['shape_ids'],'containers':{str(t.shape_id):c.shape_id for t,c in zip(self.texts,self.cards)},'visible_relation_explanation':'Two actual white cards are separated by the gray page background.'}}
        self.paint()
    def tearDown(self):self.tmp.cleanup()
    def paint(self,bg=(230,230,230),fill='white'):
        im=Image.new('RGB',(960,540),bg);d=ImageDraw.Draw(im)
        for x in (60,500):d.rounded_rectangle([x,100,x+349,399],radius=12,fill=fill);d.text((x+20,130),'Card contents',fill='black')
        im.save(self.image)
    def check(self):return validate_region_contract(self.f,self.d,self.s,1,self.image,960)
    def test_actual_white_cards_on_gray_pass(self):self.assertIsNone(self.check())
    def test_white_on_white_is_not_a_proven_boundary(self):self.paint(bg='white');self.assertIsNotNone(self.check())
    def test_transparent_proxy_cannot_borrow_visible_white_pixels(self):
        a=OxmlElement('a:alpha');a.set('val','0');self.cards[0]._element.xpath('./p:spPr/a:solidFill/*')[0].append(a);self.assertIsNotNone(self.check())
    def test_wrong_pixels_fail(self):self.paint(fill=(230,230,230));self.assertIsNotNone(self.check())
    def test_missing_current_pixels_fail(self):self.image.unlink();self.assertIsNotNone(self.check())
    def test_containment_still_enforced(self):self.texts[0].left=Pt(20);self.assertIsNotNone(self.check())
    def test_rotated_contour_needs_other_evidence(self):self.cards[0].rotation=15;self.assertIsNotNone(self.check())
    def test_full_slide_white_paint_is_not_a_local_card(self):
        c=self.cards[0];c.left=c.top=0;c.width=self.p.slide_width;c.height=self.p.slide_height
        self.assertFalse(white_container_visible_in_pixels(c,self.image,960))

if __name__=='__main__':unittest.main()
