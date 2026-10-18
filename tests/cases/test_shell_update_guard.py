"""Real temporary packages; no network, live-template edits, or fabricated signoff."""
import unittest,tempfile,contextlib,io,json,subprocess,sys,shutil
from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
from lxml import etree as E
import local_update_guard as guard
import build_manifest

class ShellUpdates(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.base=Path(self.tmp.name)
  self.root=self.base/'root';self.root.mkdir()
  for name,data in [('SKILL.md','local'),('VERSION','1.0.0')]: (self.root/name).write_text(data)
  with contextlib.redirect_stdout(io.StringIO()):build_manifest.build(self.root)
  self.candidate=self.base/'candidate';shutil.copytree(self.root,self.candidate)
  (self.candidate/'VERSION').write_text('1.1.0')
  t=self.candidate/'templates/new';t.mkdir(parents=True)
  p=Presentation();p.slides.add_slide(p.slide_layouts[6]);p.save(t/'template.pptx')
  (t/'preview.png').write_bytes(b'upstream promotional artwork must never be shown')
 def stage(self):
  self.st=self.base/'stage'
  with contextlib.redirect_stdout(io.StringIO()):guard.prepare(self.root,self.st,self.candidate)
  (self.st/'review.json').write_text(json.dumps({'status':'approved_no_conflict','reviewed_by':'test','conflicts':[]}))
 def test_new_template_is_clean_and_preview_deferred_without_blocking(self):
  self.stage()
  with contextlib.redirect_stdout(io.StringIO()):guard.apply(self.root,self.st)
  with ZipFile(self.root/'templates/new/template.pptx') as z:
   self.assertFalse(any(n.startswith('docProps/thumbnail.') for n in z.namelist()))
   self.assertEqual(E.fromstring(z.read('ppt/viewProps.xml')).get('lastView'),'sldView')
  self.assertFalse((self.root/'templates/new/preview.png').exists())
  result=json.loads((self.st/'result.json').read_text())
  self.assertEqual(result['status'],'applied');self.assertEqual(result['first_slide_preview_refresh_required'],['templates/new/preview.png'])
 def test_malformed_new_template_rolls_back_every_file(self):
  (self.candidate/'templates/new/template.pptx').write_bytes(b'bad zip')
  self.stage()
  with self.assertRaises(RuntimeError),contextlib.redirect_stdout(io.StringIO()):guard.apply(self.root,self.st)
  self.assertEqual((self.root/'VERSION').read_text(),'1.0.0')
  self.assertFalse((self.root/'templates/new/template.pptx').exists())
  self.assertEqual(json.loads((self.st/'result.json').read_text())['status'],'failed_rolled_back')
 def test_clean_local_template_and_preview_survive_update_byte_identical(self):
  t=self.root/'templates/new';t.mkdir(parents=True)
  (t/'template.pptx').write_bytes(b'locally normalized protected bytes')
  (t/'preview.png').write_bytes(b'locally generated first slide')
  self.stage()
  with contextlib.redirect_stdout(io.StringIO()):guard.apply(self.root,self.st)
  self.assertEqual((t/'template.pptx').read_bytes(),b'locally normalized protected bytes')
  self.assertEqual((t/'preview.png').read_bytes(),b'locally generated first slide')
 def test_public_update_wrapper_propagates_failure_exit(self):
  script=Path(guard.__file__).with_name('apply_update.py')
  r=subprocess.run([sys.executable,'-X','utf8',str(script),'apply','--root',str(self.root),'--stage',str(self.base/'missing')],capture_output=True,encoding='utf8')
  self.assertEqual(r.returncode,4,r.stdout+r.stderr)
  self.assertIn('Update refused or failed',r.stderr)
if __name__=='__main__':unittest.main()
