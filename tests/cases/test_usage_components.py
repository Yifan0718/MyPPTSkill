import unittest,tempfile,hashlib,json,copy
from pathlib import Path
from unittest.mock import patch
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from PIL import Image
from audit_visual_components import scan,contrast
from audit_layout_coverage import audit as coverage_audit
from verify_readability_review import verify,sha
from verify_delivery import verify as final_verify

class ComponentTests(unittest.TestCase):
 def setUp(self):
  self.p=Presentation();self.p.slide_width=Pt(960);self.p.slide_height=Pt(540);self.s=self.p.slides.add_slide(self.p.slide_layouts[6])
 def table(self,mixed=False,numeric=False):
  sh=self.s.shapes.add_table(3,3,Pt(40),Pt(120),Pt(800),Pt(160));t=sh.table
  for i in range(3):
   for j in range(3):
    c=t.cell(i,j);c.text=('项目' if j==0 else str(i+1)) if numeric and i else '说明内容'
    for pa in c.text_frame.paragraphs:pa.alignment=PP_ALIGN.RIGHT if numeric and j and i else PP_ALIGN.LEFT
  if mixed:t.cell(1,1).text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER
  return sh
 def reasons(self,render=None):return [x['reason'] for x in scan(self.s,1,render or {})]
 def test_narrative_columns_mixed_alignment_caught(self):
  self.table(True);self.assertIn('table_alignment_consistency_review',self.reasons())
 def test_uniform_narrative_columns(self):
  self.table();self.assertNotIn('table_alignment_consistency_review',self.reasons())
 def test_numeric_right_text_left_is_valid(self):
  self.table(numeric=True);self.assertNotIn('table_alignment_consistency_review',self.reasons())
 def test_unknown_inherited_alignment_requires_measurement(self):
  t=self.table();t.table.cell(1,1).text_frame.paragraphs[0].alignment=None
  self.assertIn('table_alignment_measurement_required',self.reasons())
 def test_inherited_resolved_by_real_record(self):
  t=self.table();t.table.cell(1,1).text_frame.paragraphs[0].alignment=None
  r={'text_records':[{'slide':1,'audit_path':f'{t.shape_id}/table/2/2','paragraph_alignment':1}]}
  self.assertNotIn('table_alignment_measurement_required',self.reasons(r))
 def test_merged_cell_does_not_crash(self):
  t=self.table();t.table.cell(0,0).merge(t.table.cell(0,2));self.reasons()
 def test_cell_margin_mismatch_caught(self):
  t=self.table();t.table.cell(1,1).margin_left=Pt(30);self.assertIn('table_alignment_consistency_review',self.reasons())
 def test_bad_table_note_flow_gaps(self):
  self.table();m=self.s.shapes.add_textbox(Pt(40),Pt(340),Pt(800),Pt(40));m.text='必要的训练说明'
  self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(40),Pt(385),Pt(800),Pt(60))
  self.assertIn('interstitial_module_spacing_review',self.reasons())
 def test_balanced_table_note_flow_gaps(self):
  self.table();m=self.s.shapes.add_textbox(Pt(40),Pt(300),Pt(800),Pt(40));m.text='必要的训练说明'
  self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(40),Pt(360),Pt(800),Pt(60))
  self.assertNotIn('interstitial_module_spacing_review',self.reasons())
 def color(self,fg,bg):
  s=self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(40),Pt(40),Pt(200),Pt(80));s.fill.solid();s.fill.fore_color.rgb=RGBColor(*bg)
  s.text='参与共创';s.text_frame.paragraphs[0].runs[0].font.color.rgb=RGBColor(*fg)
 def test_yellow_blue_aesthetic_candidate(self):
  self.color((20,65,140),(255,185,35));self.assertIn('filled_text_color_review',self.reasons())
 def test_yellow_black_clear(self):
  self.color((20,20,20),(255,185,35));self.assertNotIn('filled_text_color_review',self.reasons())
 def test_light_yellow_white_is_not_forced(self):
  self.color((255,255,255),(255,235,180));self.assertIn('filled_text_color_review',self.reasons())
 def test_dark_white_clear(self):
  self.color((255,255,255),(25,40,70));self.assertNotIn('filled_text_color_review',self.reasons())
 def test_dark_black_candidate(self):
  self.color((0,0,0),(25,40,70));self.assertIn('filled_text_color_review',self.reasons())
 def test_color_candidates_survive_json_roundtrip_and_full_rescan(self):
  # Exercise the exact full-report equality required by coverage verification.
  # The blank raster is a synthetic scanner fixture, never visual signoff.
  self.color((20,65,140),(255,185,35))
  self.color((255,255,255),(255,235,180))
  self.color((0,0,0),(25,40,70))
  with tempfile.TemporaryDirectory() as td:
   root=Path(td);candidate=root/'candidate.pptx';self.p.save(candidate)
   Image.new('RGB',(960,540),'white').save(root/'slide-01.png')
   records=[{'slide':1,'audit_path':str(s.shape_id),'shape_id':s.shape_id,'shape':s.name,'kind':'text','text':s.text,'font_pt':24,
             'left_pt':40,'top_pt':40,'width_pt':200,'height_pt':80,
             'bound_left_pt':48,'bound_top_pt':48,'bound_width_pt':160,'bound_height_pt':28} for s in self.s.shapes]
   render={'sha256':sha(candidate),'slide_count':1,'snapshot':str(candidate),'status':'rendered_pending_visual_review','text_records':records}
   original=coverage_audit(candidate,render)
   report=root/'coverage.json';report.write_text(json.dumps(original),'utf8')
   persisted=json.loads(report.read_text('utf8'))
   recomputed=coverage_audit(candidate,render)
   self.assertEqual(persisted,recomputed)
   colors=[f for f in persisted['findings'] if f['reason']=='filled_text_color_review']
   self.assertEqual(len(colors),3)
   self.assertEqual([f['id'] for f in original['findings']],[f['id'] for f in persisted['findings']])
   self.assertEqual(colors[0]['color_pairs'][0]['foreground'],[20,65,140])
   self.assertEqual(colors[0]['color_pairs'][0]['background'],[255,185,35])
   self.assertEqual(colors[0]['color_pairs'][0]['contrast'],round(contrast((20,65,140),(255,185,35)),3))
 def test_inherited_color_candidate_survives_json_roundtrip(self):
  s=self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(40),Pt(40),Pt(200),Pt(80))
  s.fill.solid();s.fill.fore_color.rgb=RGBColor(255,185,35);s.text='继承字色'
  findings=scan(self.s,1,{})
  self.assertEqual(json.loads(json.dumps(findings)),findings)
  self.assertEqual(findings[0]['color_pairs'],[{'foreground':'inherited_or_theme','background':[255,185,35]}])
 def chart(self):
  d=CategoryChartData();d.categories=['A','B'];d.add_series('收益',[4,8]);s=self.s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Pt(40),Pt(150),Pt(800),Pt(300),d);s.chart.has_legend=True;return s
 def test_missing_legend_measurement_routes_review(self):
  self.chart();self.assertIn('chart_legend_geometry_required',self.reasons())
 def test_legend_near_label(self):
  s=self.chart();r={'chart_legend_records':[{'slide':1,'audit_path':str(s.shape_id),'bounds_pt':[100,200,200,220]}],'chart_label_records':[{'slide':1,'audit_path':str(s.shape_id),'bounds_pt':[110,221,150,240]}]}
  self.assertIn('chart_legend_clearance_review',self.reasons(r))
 def test_declared_labels_missing_are_not_silent(self):
  sh=self.chart();sh.chart.plots[0].has_data_labels=True
  r={'chart_legend_records':[{'slide':1,'audit_path':str(sh.shape_id),'bounds_pt':[100,200,200,220]}]}
  self.assertIn('chart_label_geometry_required',self.reasons(r))
 def test_invalid_label_bounds_are_reviewed(self):
  sh=self.chart();sh.chart.plots[0].has_data_labels=True
  r={'chart_label_records':[{'slide':1,'audit_path':str(sh.shape_id),'bounds_pt':[1,2,float('nan'),4]}]}
  self.assertIn('chart_label_geometry_required',self.reasons(r))
 def test_no_legend_no_labels_is_valid(self):
  sh=self.chart();sh.chart.has_legend=False
  self.assertNotIn('chart_label_geometry_required',self.reasons())
 def test_other_chart_label_not_misassociated(self):
  s=self.chart();r={'chart_legend_records':[{'slide':1,'audit_path':str(s.shape_id),'bounds_pt':[100,200,200,220]}],'chart_label_records':[{'slide':1,'audit_path':'999','bounds_pt':[110,221,150,240]}]}
  self.assertNotIn('chart_legend_clearance_review',self.reasons(r))

class ReadabilityTests(unittest.TestCase):
 def setUp(self):
  self.t=tempfile.TemporaryDirectory();self.addCleanup(self.t.cleanup);self.w=Path(self.t.name);self.r=self.w/'render';self.r.mkdir()
  p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);t=s.shapes.add_textbox(Pt(30),Pt(50),Pt(400),Pt(100));t.text='三段动作，每段五步，共十五步。相对基线误差变化 −8%。'
  self.path=self.w/'candidate.pptx';p.save(self.path);Image.new('RGB',(960,540),'white').save(self.r/'slide-01.png')
  self.audit={'sha256':sha(self.path),'slide_count':1,'text_records':[{'slide':1,'audit_path':str(t.shape_id),'text':t.text}]};(self.r/'audit.json').write_text(json.dumps(self.audit),'utf8')
  self.ev={'slide':1,'audit_path':str(t.shape_id),'quote':'每段五步'}
  self.review={'schema':'audience-review/v1','sha256':sha(self.path),'audience':'首次接触该研究的技术同行','expression_mode':'advantage_led','chapter_navigation':{'mode':'none','basis':'user','description':'用户明确不要单独的章节标识'},'unresolved':[], 'definitions':[], 'examples':[], 'feedback':[], 'pages':[{'slide':1,'image_sha256':sha(self.r/'slide-01.png'),'kind':'general','result':'pass','unresolved':[],'purpose':'解释控制步和动作段的数量关系','plain_language':'三段动作一共执行十五个控制步','operation_or_reading':'先理解每段五步再理解整体十五步','metrics_or_symbols':'负百分数表示相对基线误差下降','evidence_to_conclusion':'页面写明各段步数及带符号变化量','evidence':[self.ev],'visual_checks':{k:'实际查看页面中的对应对象和文字关系' for k in ('content_reading','layout','local_detail','full_page_return','cross_page','palette')},'visual_targets_reviewed':[],'local_views':[]}]}
 def check(self):return verify(self.path,self.r,self.review)['issues']
 def test_valid_evidence_is_pending_not_signed_off(self):
  self.assertEqual(self.check(),[]);self.assertEqual(verify(self.path,self.r,self.review)['status'],'audience_evidence_complete_pending_semantic_signoff')
 def test_missing_page(self):self.review['pages']=[];self.assertTrue(self.check())
 def test_duplicate_page(self):self.review['pages']*=2;self.assertTrue(self.check())
 def test_stale_image(self):self.review['pages'][0]['image_sha256']='OLD';self.assertTrue(self.check())
 def test_chat_only_explanation(self):self.ev['quote']='这是仅在聊天解释过的东西';self.assertTrue(self.check())
 def test_future_definition(self):self.review['definitions']=[{'introduced_on':2,'used_on':[1],'evidence':self.ev}];self.assertTrue(self.check())
 def test_unresolved_feedback(self):self.review['feedback']=[{'id':'a','category':'terminology','issue':'图中缩写没有具体解释','status':'open'}];self.assertTrue(self.check())
 def test_feedback_needs_artifact_change(self):
  self.review['feedback']=[{'id':'a','category':'terminology','issue':'图中缩写没有具体解释','status':'resolved','affected_slides':[1],'similar_checked_slides':[1],'observation':'已经在聊天里解释给用户听了','changes':[]}];self.assertTrue(self.check())
 def test_feedback_with_current_evidence(self):
  self.review['feedback']=[{'id':'a','category':'terminology','issue':'动作段与步数关系没有解释','status':'resolved','affected_slides':[1],'similar_checked_slides':[1],'observation':'页面现已显示每段五步及总共十五步','changes':[self.ev]}];self.assertEqual(self.check(),[])
 def test_example_needs_selection(self):self.review['pages'][0]['kind']='example';self.assertTrue(self.check())
 def test_missing_visual_stage_cannot_be_code_only(self):del self.review['pages'][0]['visual_checks']['palette'];self.assertTrue(self.check())
 def test_blank_reconstruction(self):self.review['pages'][0]['plain_language']='pass';self.assertTrue(self.check())
 def test_mixed_text_image_fields_cannot_bypass_missing_file(self):
  e={**self.ev,'crop_path':str(self.w/'missing.png'),'crop_box_px':[0,0,1,1],'crop_sha256':'OLD','image_sha256':'OLD','observation':'虚假的局部记录不能覆盖真实查看'}
  self.review['pages'][0]['local_views']=[e];self.assertTrue(self.check())
 def test_null_visual_checks_preserve_earlier_errors(self):
  self.review['sha256']='OLD';self.review['pages'][0]['visual_checks']=None
  errors=self.check();self.assertTrue(any('hash mismatch' in e for e in errors));self.assertTrue(any('must be an object' in e for e in errors))
 def add_table_target(self):
  prs=Presentation(self.path);t=prs.slides[0].shapes.add_table(2,2,Pt(80),Pt(200),Pt(300),Pt(100));prs.save(self.path)
  self.review['sha256']=sha(self.path);self.audit['sha256']=sha(self.path);(self.r/'audit.json').write_text(json.dumps(self.audit),'utf8');self.review['pages'][0]['visual_targets_reviewed']=[str(t.shape_id)]
  pp=960/(prs.slide_width/12700);return [int(v*pp) for v in (80,200,380,300)]
 def crop(self,box,name):
  cp=self.w/name;Image.open(self.r/'slide-01.png').crop(box).save(cp)
  return {'slide':1,'image_sha256':sha(self.r/'slide-01.png'),'crop_path':str(cp),'crop_sha256':sha(cp),'crop_box_px':box,'observation':'实际局部包含表格边界和对应单元格内容'}
 def test_unrelated_background_crop_rejected(self):
  self.add_table_target();self.review['pages'][0]['local_views']=[self.crop([0,0,1,1],'tiny.png')];self.assertTrue(self.check())
 def test_split_local_views_cover_complete_component(self):
  l,t,r,b=self.add_table_target();m=(l+r)//2
  self.review['pages'][0]['local_views']=[self.crop([l-2,t-2,m+2,b+2],'left.png'),self.crop([m-2,t-2,r+2,b+2],'right.png')];self.assertEqual(self.check(),[])
 def test_stale_crop_hash_rejected(self):
  l,t,r,b=self.add_table_target();e=self.crop([l-2,t-2,r+2,b+2],'whole-table.png');e['crop_sha256']='OLD';self.review['pages'][0]['local_views']=[e];self.assertTrue(self.check())
 def test_valid_selection_is_bound_to_source(self):
  self.review['pages'][0]['kind']='example';source=self.w/'data.txt';source.write_text('case=2, baseline=100, method=92','utf8')
  x={'slide':1,'purpose':'展示实际改善的方向与幅度','selection_scope':'比较两份已有的评测样本','source_path':str(source),'source_sha256':sha(source),'source_locator':'数据文件中的第二个样本记录','why_representative':'本例展示大幅改善，总体用完整统计说明','what_visible':'图中数值与标签完整展示所选样本','claim_check':'本例仅承担直观解释的具体职责','alternatives_considered':['第一例方向变化小，第二例方向和幅度更清楚']}
  self.review['examples']=[x];self.assertEqual(self.check(),[]);source.write_text('changed','utf8');self.assertTrue(self.check())
 def test_successful_subchecks_remain_pending_signoff(self):
  with patch('verify_delivery.verify_layout',return_value={'status':'evidence_complete_pending_visual_signoff','issues':[]}),patch('verify_delivery.verify_audience',return_value={'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}),patch('verify_delivery.load',return_value={}):
   self.assertEqual(final_verify(self.path,None,None,self.r,None)['status'],'evidence_complete_pending_visual_signoff')
 def test_invalid_page_list_retains_hash_error(self):
  self.review['sha256']='OLD';self.review['pages']=None;errors=self.check()
  self.assertTrue(any('hash mismatch' in e for e in errors));self.assertIn('Invalid audience page list',errors)
 def test_mixed_target_types_report_field_error(self):
  self.review['pages'][0]['visual_targets_reviewed']=['2',{}];self.assertTrue(any('inventory incomplete' in e for e in self.check()))
 def test_noninteger_example_page_rejected(self):
  source=self.w/'source.txt';source.write_text('measured values','utf8')
  x={key:'具体的来源选择和图上观察记录' for key in ('purpose','selection_scope','source_locator','why_representative','what_visible','claim_check')}
  x.update(slide='1',source_path=str(source),source_sha256=sha(source),alternatives_considered=['已比较两份可用的原始样本记录'])
  self.review['examples']=[x];self.assertIn('Example slide must be a valid integer page number',self.check())
 def test_definition_names_required(self):
  self.review['definitions']=[{'introduced_on':1,'used_on':[1],'evidence':self.ev}];self.assertIn('Definition needs term and display_name',self.check())
 def test_valid_named_definition(self):
  self.review['definitions']=[{'term':'segment','display_name':'动作段','introduced_on':1,'used_on':[1],'evidence':self.ev}];self.assertEqual(self.check(),[])
 def test_bad_example_path_preserves_hash_error(self):
  self.review['sha256']='OLD';self.review['examples']=[{'slide':1,'source_path':None}];errors=self.check()
  self.assertIn('Audience review candidate hash mismatch',errors);self.assertIn('Example source_path must be a nonempty string',errors)
 def test_unhashable_feedback_pages_report_invalid(self):
  self.review['feedback']=[{'id':'a','category':'terminology','issue':'这是必须解释的术语问题','status':'resolved','observation':'当前页面已写明对应的术语定义','affected_slides':[{}],'similar_checked_slides':[1],'changes':[self.ev]}];self.assertIn('Invalid affected pages',self.check())
 def test_feedback_requires_category(self):
  self.review['feedback']=[{'id':'a','issue':'这是必须解释的术语问题','status':'resolved','observation':'当前页面已写明对应的术语定义','affected_slides':[1],'similar_checked_slides':[1],'changes':[self.ev]}];self.assertIn('Feedback category is required',self.check())
 def test_huge_crop_rejected_before_allocation(self):
  e=self.crop([0,0,20,20],'small.png');e['crop_box_px']=[0,0,10**12,10**12];self.review['pages'][0]['local_views']=[e]
  with patch('PIL.Image.Image.crop',side_effect=AssertionError('invalid crop must not allocate')):
   self.assertIn('Visual evidence crop outside page',self.check())
 def test_missing_page_preserves_earlier_error(self):
  self.review['sha256']='OLD';(self.r/'slide-01.png').unlink();errors=self.check()
  self.assertIn('Audience review candidate hash mismatch',errors);self.assertTrue(any('image unavailable' in e for e in errors))
 def test_malformed_text_records_are_located(self):
  self.audit['text_records']=[None,{'slide':1,'audit_path':[],'text':'wrong'}];(self.r/'audit.json').write_text(json.dumps(self.audit),'utf8')
  self.assertTrue(any('text_records[0]' in e for e in self.check()))
 def test_duplicate_text_records_rejected(self):
  self.audit['text_records']*=2;(self.r/'audit.json').write_text(json.dumps(self.audit),'utf8');self.assertTrue(any('Duplicate render text record' in e for e in self.check()))
 def test_invalid_subcheck_result_still_runs_next(self):
  with patch('verify_delivery.verify_layout',return_value=None),patch('verify_delivery.verify_audience',return_value={'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}) as reader,patch('verify_delivery.load',return_value={}):
   result=final_verify(self.path,None,None,self.r,None);self.assertEqual(result['status'],'failed');reader.assert_called_once();self.assertIn('audience',result['checks'])
 def test_input_disappears_after_checks_keeps_results(self):
  def remove(*args):
   self.path.unlink();return {'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}
  with patch('verify_delivery.verify_layout',return_value={'status':'failed','issues':['prior layout issue']}),patch('verify_delivery.verify_audience',side_effect=remove),patch('verify_delivery.load',return_value={}):
   result=final_verify(self.path,None,None,self.r,None);self.assertEqual(result['status'],'failed');self.assertEqual(result['checks']['layout']['issues'],['prior layout issue']);self.assertTrue(any(x['stage']=='integrity' for x in result['issues']))
 def test_evidence_file_changes_during_gate_rejected(self):
  source=self.w/'review.json';source.write_text('{}','utf8')
  def mutate(*args):
   source.write_text('{"different":true}','utf8');return {'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}
  with patch('verify_delivery.verify_layout',return_value={'status':'evidence_complete_pending_visual_signoff','issues':[]}),patch('verify_delivery.verify_audience',side_effect=mutate),patch('verify_delivery.load',return_value={}):
   result=final_verify(self.path,None,source,self.r,None);self.assertTrue(any(x['stage']=='integrity' for x in result['issues']))
 def test_page_pixels_change_during_final_gate(self):
  def mutate(*args):
   Image.new('RGB',(960,540),'red').save(self.r/'slide-01.png');return {'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}
  with patch('verify_delivery.verify_layout',return_value={'status':'evidence_complete_pending_visual_signoff','issues':[]}),patch('verify_delivery.verify_audience',side_effect=mutate),patch('verify_delivery.load',return_value={}):
   result=final_verify(self.path,None,None,self.r,None);self.assertTrue(any(x['stage']=='integrity' for x in result['issues']))
 def test_indirect_crop_change_during_final_gate(self):
  crop=self.w/'crop.png';Image.new('RGB',(20,20),'white').save(crop);meta=self.w/'review.json';meta.write_text(json.dumps({'local_views':[{'crop_path':str(crop)}]}),'utf8')
  def mutate(*args):
   Image.new('RGB',(20,20),'red').save(crop);return {'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}
  with patch('verify_delivery.verify_layout',return_value={'status':'evidence_complete_pending_visual_signoff','issues':[]}),patch('verify_delivery.verify_audience',side_effect=mutate),patch('verify_delivery.load',return_value={}):
   result=final_verify(self.path,None,meta,self.r,None);self.assertTrue(any(x['stage']=='integrity' and x.get('path')==str(crop.resolve()) for x in result['issues']))
 def test_new_render_evidence_during_gate_is_rejected(self):
  def mutate(*args):
   (self.r/'tool-run.json').write_text('{"status":"running"}','utf8');return {'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}
  with patch('verify_delivery.verify_layout',return_value={'status':'evidence_complete_pending_visual_signoff','issues':[]}),patch('verify_delivery.verify_audience',side_effect=mutate),patch('verify_delivery.load',return_value={}):
   result=final_verify(self.path,None,None,self.r,None);self.assertTrue(any('dependency paths' in x.get('reason','') for x in result['issues']))
 def test_rules_rechecked_after_audience_stage(self):
  meta=self.w/'review.json';meta.write_text('{"skill_lock":{"path":"lock.json"}}','utf8')
  with patch('verify_delivery.verify_layout',return_value={'status':'evidence_complete_pending_visual_signoff','issues':[]}),patch('verify_delivery.verify_audience',return_value={'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}),patch('verify_delivery.load',return_value={'skill_lock':{'path':'lock.json'}}),patch('verify_delivery.check_skill_state',return_value={'changed_files':['references/rule.md']}):
   result=final_verify(self.path,None,meta,self.r,None);self.assertTrue(any(x['stage']=='rules' for x in result['issues']))
 def test_gate_propagates_audience_failure(self):
  with patch('verify_delivery.verify_layout',return_value={'status':'evidence_complete_pending_visual_signoff','issues':[]}),patch('verify_delivery.verify_audience',return_value={'status':'failed','issues':['missing definition']}),patch('verify_delivery.load',return_value={}):
   self.assertEqual(final_verify(self.path,None,None,self.r,None)['status'],'failed')
 def test_gate_continues_audience_after_layout_exception(self):
  with patch('verify_delivery.verify_layout',side_effect=ValueError('layout input invalid')),patch('verify_delivery.verify_audience',return_value={'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}) as reader,patch('verify_delivery.load',return_value={}):
   self.assertEqual(final_verify(self.path,None,None,self.r,None)['status'],'failed');reader.assert_called_once()
 def test_unknown_child_status_fails(self):
  with patch('verify_delivery.verify_layout',return_value={'status':'unknown','issues':[]}),patch('verify_delivery.verify_audience',return_value={'status':'audience_evidence_complete_pending_semantic_signoff','issues':[]}),patch('verify_delivery.load',return_value={}):
   self.assertEqual(final_verify(self.path,None,None,self.r,None)['status'],'failed')
class UpdatePreservationTests(unittest.TestCase):
 def test_new_rules_and_gate_survive_upstream_replacement(self):
  import local_update_guard as guard
  with tempfile.TemporaryDirectory() as td:
   base=Path(td);live=base/'live';incoming=base/'incoming';live.mkdir();incoming.mkdir()
   (live/'VERSION').write_text('1.0.20','utf8');(live/'SKILL.md').write_text('local entry','utf8');(live/'manifest.json').write_text('{"files":{}}','utf8')
   rels=['references/audience-and-evidence.md','references/revision-control.md','scripts/audit_visual_components.py','scripts/verify_readability_review.py','scripts/verify_delivery.py','scripts/report_io.py','scripts/run_bounded.py']
   expected={}
   source_root=Path(__file__).resolve().parent.parent
   for rel in rels:
    dst=live/rel;dst.parent.mkdir(parents=True,exist_ok=True);dst.write_bytes((source_root/rel).read_bytes());expected[rel]=dst.read_bytes()
    up=incoming/rel;up.parent.mkdir(parents=True,exist_ok=True);up.write_text('upstream replacement','utf8')
   (incoming/'VERSION').write_text('1.0.21','utf8');(incoming/'SKILL.md').write_text('upstream entry','utf8')
   stage=base/'stage';guard.prepare(live,stage,incoming)
   (stage/'review.json').write_text(json.dumps({'status':'approved_no_conflict','reviewed_by':'synthetic regression fixture','conflicts':[],'user_confirmation':''}),'utf8')
   guard.apply(live,stage)
   for rel,data in expected.items():self.assertEqual((live/rel).read_bytes(),data)
   self.assertEqual((live/'VERSION').read_text(),'1.0.21')

if __name__=='__main__':unittest.main()
