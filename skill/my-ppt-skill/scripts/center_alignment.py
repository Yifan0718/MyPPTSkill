"""Current-pixel centering candidates for native doughnut charts.

Frames/COM plot bounds locate the search only. They never stand in for visible
host, hole or text ink centers. Every result needs local + full-page review;
reliable misalignment is a repair finding, uncertainty is never a silent pass.
"""
import math
from pptx.enum.shapes import MSO_SHAPE_TYPE

try:
    import cv2
    import numpy as np
except (ImportError,OSError):
    cv2 = np = None

MISALIGNED = 'centered_text_visible_host_misaligned'
REVIEW = 'centered_text_visible_host_review'


def _bbox(mask):
    yy, xx = np.nonzero(mask)
    return None if not len(xx) else [int(xx.min()), int(yy.min()), int(xx.max()+1), int(yy.max()+1)]


def _center(box):
    return [(box[0]+box[2])/2, (box[1]+box[3])/2]


def _unknown(reason, crop=None):
    return {'measurement_status': 'manual_review_required', 'measurement_limitation': reason,
            'crop_box_px': crop, 'local_and_full_page_review_required': True,
            'internal_review_required': True}


def _segment(rgb, roi, regions, threshold):
    x0, y0, x1, y1 = roi
    tile = rgb[y0:y1, x0:x1]
    border = np.concatenate([tile[0], tile[-1], tile[:,0], tile[:,-1]])
    background = np.median(border, axis=0)
    if np.quantile(np.max(np.abs(border.astype(float)-background), axis=1), .95) > 5:
        return None, 'Nonuniform or occluded search boundary; visual segmentation required'
    foreground = np.max(np.abs(tile.astype(float)-background), axis=2) > threshold
    count, labels, stats, _ = cv2.connectedComponentsWithStats(foreground.astype(np.uint8), 8)
    candidates = []
    for label in range(1, count):
        if stats[label, cv2.CC_STAT_AREA] < .025*tile.shape[0]*tile.shape[1]:
            continue
        host = (labels == label).astype(np.uint8)
        contours, hierarchy = cv2.findContours(host, cv2.RETR_CCOMP, cv2.CHAIN_APPROX_SIMPLE)
        if hierarchy is None:
            continue
        holes = [i for i,v in enumerate(hierarchy[0]) if v[3] >= 0 and cv2.contourArea(contours[i]) > .03*tile.shape[0]*tile.shape[1]]
        if len(holes) != 1:
            continue
        hole = np.zeros_like(host)
        cv2.drawContours(hole, contours, holes[0], 1, -1)
        hole = cv2.erode(hole, np.ones((5,5), np.uint8))
        hb, ib = _bbox(host), _bbox(hole)
        if hb is None or ib is None:
            continue
        if max(abs(a-b) for a,b in zip(_center(hb), _center(ib))) > 2:
            return None, 'Visible host/hole centers disagree; shadow or asymmetric contour requires review'
        inner = foreground & (hole > 0)
        if np.count_nonzero(inner) < 8:
            continue
        allowed = np.zeros_like(host)
        masks = []
        for region in regions:
            l,t,r,b = region
            l,r = max(0,l-x0), min(tile.shape[1],r-x0)
            t,b = max(0,t-y0), min(tile.shape[0],b-y0)
            mask = np.zeros_like(host)
            if r > l and b > t:
                mask[t:b,l:r] = 1
                allowed[t:b,l:r] = 1
            masks.append(mask)
        def global_box(box):
            return [box[0]+x0,box[1]+y0,box[2]+x0,box[3]+y0]
        partial = {'host_bbox_px':global_box(hb),'hole_bbox_px':global_box(ib)}
        if np.any(foreground & (allowed > 0) & (hole == 0) & (labels != label)):
            return partial, 'Discovered text region includes ink outside the hole; do not measure a truncated text group'
        if np.any(inner & (allowed == 0)):
            return partial, 'Unaccounted current foreground inside hole; inspect every internal object'
        if any(not np.any(inner & (mask > 0)) for mask in masks):
            return partial, 'At least one discovered text member has no reliably isolated ink in the hole'
        rows = np.flatnonzero(np.any(inner, axis=1))
        bands = []
        for y in rows:
            if bands and y-bands[-1][1] <= 2:
                bands[-1][1] = int(y)+1
            else:
                bands.append([int(y),int(y)+1])
        line_boxes = []
        for top,bottom in bands:
            line = np.zeros_like(inner);line[top:bottom] = inner[top:bottom]
            line_boxes.append(global_box(_bbox(line)))
        candidates.append({'host_bbox_px':global_box(hb), 'hole_bbox_px':global_box(ib),
            'text_group_bbox_px':global_box(_bbox(inner)),
            'member_ink_bounds_px':[global_box(_bbox(inner & (mask > 0))) for mask in masks],
            'line_ink_bounds_px':line_boxes})
    if len(candidates) != 1:
        return None, 'No unique connected visible host and closed hole; local visual measurement required'
    return candidates[0], None


def measure_visible_hole(image, roi, text_regions, pixels_per_pt, scale=1):
    """Measure the union of all discovered text members, including multiple lines.

    Each member/line is retained for internal review, not forced individually to
    the hole center. This pure helper does not accept author-written centers.
    """
    if cv2 is None or np is None:
        return _unknown('Pixel component dependency unavailable; retain local + full-page visual review', roi)
    try:
        if not text_regions or not math.isfinite(pixels_per_pt) or pixels_per_pt <= 0:
            return _unknown('Missing text members or valid pixel scale', roi)
        rgb = np.asarray(image.convert('RGB'))
        if len(roi) != 4 or any(type(v) is not int for v in roi) or not 0 <= roi[0] < roi[2] <= image.width or not 0 <= roi[1] < roi[3] <= image.height:
            return _unknown('Current host search rectangle is invalid', roi)
        runs,errors = [],[]
        for threshold in (8,12,18):
            result, error = _segment(rgb, roi, text_regions, threshold)
            runs.append(result)
            errors.append(error)
        if any(errors):
            partial = {}
            for field in ('host_bbox_px','hole_bbox_px','text_group_bbox_px'):
                if all(r is not None and field in r for r in runs) and np.max(np.ptp(np.array([r[field] for r in runs]),axis=0)) <= 1:
                    partial[field] = runs[1][field]
            return {**_unknown(next(e for e in errors if e),roi),**partial,'measurement_known_fields':sorted(partial)}
        for field in ('host_bbox_px','hole_bbox_px','text_group_bbox_px'):
            if np.max(np.ptp(np.array([r[field] for r in runs]), axis=0)) > 1:
                return _unknown('Visible bounds change across contrast thresholds; measurement is uncertain', roi)
        result = runs[1]
        hc, tc = _center(result['hole_bbox_px']), _center(result['text_group_bbox_px'])
        delta = [tc[i]-hc[i] for i in (0,1)]
        tolerance = max(2.0, .75*scale*pixels_per_pt)
        overlaps = []
        for i,a in enumerate(result['member_ink_bounds_px']):
            for j,b in enumerate(result['member_ink_bounds_px'][i+1:], i+1):
                if min(a[2],b[2])-max(a[0],b[0]) > 1 and min(a[3],b[3])-max(a[1],b[1]) > 1:
                    overlaps.append([i,j])
        hb = result['host_bbox_px']
        crop = [max(0,hb[0]-8),max(0,hb[1]-8),min(image.width,hb[2]+8),min(image.height,hb[3]+8)]
        return {**result, 'measurement_status':'measured_misalignment' if max(abs(v) for v in delta) >= tolerance else 'measured_aligned',
            'host_center_px':_center(result['host_bbox_px']), 'hole_center_px':hc, 'text_group_center_px':tc,
            'delta_px':delta, 'delta_pt':[v/pixels_per_pt for v in delta], 'tolerance_px':tolerance,
            'pixels_per_pt':pixels_per_pt, 'segmentation_thresholds':[8,12,18], 'crop_box_px':crop,
            'internal_overlap_member_indices':overlaps, 'internal_review_required':True,
            'measurement_known_fields':['host_bbox_px','hole_bbox_px','text_group_bbox_px','member_ink_bounds_px','line_ink_bounds_px'],
            'local_and_full_page_review_required':True}
    except (ValueError,TypeError,IndexError,OverflowError,cv2.error) as exc:
        return _unknown('Pixel measurement unavailable: '+str(exc), roi)


def _walk(shapes, prefix=''):
    for shape in shapes:
        path = prefix+str(shape.shape_id)
        yield path, shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes, path+'/group/')


def scan_doughnut_centers(prs, slide_no, render, raster, inventory=None):
    """Discover native doughnuts and overlapping central text independently of names.

    Recursive audit paths preserve grouped members. Uncertain group transforms,
    native effects, missing pixels or unreadable contours remain explicit reviews.
    """
    from render_evidence import native_inventory
    inv = inventory if inventory is not None else native_inventory(prs)
    slide = prs.slides[slide_no-1]
    entries = list(_walk(slide.shapes))
    rows = {(r.get('slide'),r.get('audit_path')):r for r in render.get('text_records',[])}
    width,height = prs.slide_width/12700,prs.slide_height/12700
    sx,sy = (raster.width/width,raster.height/height) if raster is not None else (1,1)
    scale = height/540
    findings = []
    for path,chart in entries:
        # Inspect the native plot tag without requiring python-pptx to classify
        # every other chart type (some valid mixed/extension charts are unknown).
        if not chart.has_chart or not chart.chart._element.xpath('.//c:doughnutChart'):
            continue
        cb = inv['bounds'].get((slide_no,path))
        uncertain_chart = cb is None
        if cb is None:
            cb = [chart.left/12700,chart.top/12700,(chart.left+chart.width)/12700,(chart.top+chart.height)/12700]
        l,t,r,b = cb
        central = [l+.15*(r-l),t+.15*(b-t),r-.15*(r-l),b-.15*(b-t)]
        members = []
        for member_path,member in entries:
            if not member.has_text_frame or not member.text.strip():
                continue
            record = rows.get((slide_no,member_path))
            mb = inv['bounds'].get((slide_no,member_path))
            if mb is None and record is not None:
                vals = [record.get(k) for k in ('left_pt','top_pt','width_pt','height_pt')]
                if all(type(v) in (int,float) and math.isfinite(v) for v in vals):
                    mb = [vals[0],vals[1],vals[0]+vals[2],vals[1]+vals[3]]
            if mb is None:
                continue
            mc = _center(mb)
            overlaps = min(mb[2],central[2]) > max(mb[0],central[0]) and min(mb[3],central[3]) > max(mb[1],central[1])
            if overlaps and l <= mc[0] <= r and t <= mc[1] <= b:
                members.append((member_path,member,mb,record))
        if not members:
            continue
        chart_owner=int(path.split('/')[0])
        text_owners=[int(p.split('/')[0]) for p,_,_,_ in members]
        ids = [chart_owner]+text_owners
        roi = [max(0,math.floor(l*sx)),max(0,math.floor(t*sy)),min(raster.width,math.ceil(r*sx)) if raster else math.ceil(r),min(raster.height,math.ceil(b*sy)) if raster else math.ceil(b)]
        detail = {'shape_ids':sorted(set(ids)), 'chart_id':chart_owner, 'chart_leaf_id':chart.shape_id, 'chart_audit_path':path,
            'text_ids':sorted(set(text_owners)), 'text_leaf_ids':[member.shape_id for _,member,_,_ in members], 'text_audit_paths':[p for p,_,_,_ in members],
            'texts':[member.text for _,member,_,_ in members], 'relation':'text_group_in_native_doughnut_hole',
            'pptx_sha256':render.get('sha256'), 'image_sha256':render.get('raster_hashes',{}).get(str(slide_no)),
            'host_search_box_px':roi}
        if raster is None:
            measured = _unknown('Current page raster unavailable; retain explicit local/full-page review', None)
        elif '/group/' in path or any('/group/' in p for p,_,_,_ in members):
            measured = _unknown('Nested chart/text group retains complete audit paths and requires visible contour review',roi)
        elif uncertain_chart or any(inv['bounds'].get((slide_no,p)) is None for p,_,_,_ in members):
            measured = _unknown('Rotated/unknown group geometry requires visible contour review',roi)
        elif chart._element.xpath('.//a:outerShdw | .//a:innerShdw | .//a:glow | .//a:softEdge') or chart.chart._element.xpath('.//a:outerShdw | .//a:innerShdw | .//a:glow | .//a:softEdge'):
            measured = _unknown('Native chart effects require visible contour review',roi)
        else:
            regions = [[max(0,math.floor(mb[0]*sx)-2),max(0,math.floor(mb[1]*sy)-2),min(raster.width,math.ceil(mb[2]*sx)+2),min(raster.height,math.ceil(mb[3]*sy)+2)] for _,_,mb,_ in members]
            measured = measure_visible_hole(raster,roi,regions,sx,scale)
        if measured['measurement_status']=='manual_review_required' and raster is not None:
            # Include complete discovered member frames, never only an arbitrary
            # center patch. Frame bounds are context, not claimed visible centers.
            boxes=[cb]+[mb for _,_,mb,_ in members]
            measured['crop_box_px']=[max(0,math.floor(min(b[0] for b in boxes)*sx)-2),max(0,math.floor(min(b[1] for b in boxes)*sy)-2),min(raster.width,math.ceil(max(b[2] for b in boxes)*sx)+2),min(raster.height,math.ceil(max(b[3] for b in boxes)*sy)+2)]
        measured.setdefault('measurement_known_fields',[])
        reason = MISALIGNED if measured['measurement_status']=='measured_misalignment' else REVIEW
        findings.append({'reason':reason, **detail, **measured})
    return findings
