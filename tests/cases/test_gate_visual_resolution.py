"""SYNTHETIC contract fixtures only. No record here is real visual signoff."""
import copy
import hashlib
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Pt

from connector_integrity import (connector_finding, endpoint_findings, inventory,
                                 verify_connections)
from layout_visual_review import CODE, layout_context
from render_evidence import (digest, native_inventory, preflight_finding,
                             validate_render_evidence)


OBS='SYNTHETIC TEST fixture validates the evidence contract only; no visual signoff occurred.'


class VisualGateTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='synthetic-visual-gate-')
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name).resolve()
        self.prs=Presentation();self.prs.slide_width=Pt(960);self.prs.slide_height=Pt(540)
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.left=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(50),Pt(110),Pt(40),Pt(40))
        self.right=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(200),Pt(100),Pt(100),Pt(60))
        for shape in (self.left,self.right):
            shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(30,70,110)
        self.line=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(90),Pt(130),Pt(197),Pt(130))
        self.image=self.root/'slide-01.png'
        self.pptx=self.root/'synthetic-input.pptx'
        self.snapshot=self.root/'synthetic-snapshot.pptx'
        self.prs.save(self.pptx);shutil.copyfile(self.pptx,self.snapshot)
        im=Image.new('RGB',(960,540),'white');draw=ImageDraw.Draw(im)
        draw.rectangle((50,110,90,150),fill='#334466')
        draw.rectangle((200,100,300,160),fill='#334466')
        draw.line((90,130,197,130),fill='black',width=3)
        draw.text((10,10),'SYNTHETIC TEST PIXELS',fill='black');im.save(self.image)
        self.render={'schema':'render-evidence/v2','sha256':digest(self.pptx),'snapshot':str(self.snapshot),
                     'slide_count':1,'status':'rendered_pending_visual_review','issues':[],
                     'raster_hashes':{'1':digest(self.image)},'text_records':[],'chart_geometry_records':[],
                     'manual_review':[{'review_id':'m1','slide':1,'audit_path':'master-layout','reason':'master_layout_visual_review'}],
                     'render_finished_utc':'2026-10-18T09:00:00+00:00'}
        self.record={'slide':1,'shape_id':self.line.shape_id,'image_sha256':digest(self.image),
                     'returned_to_full_page':True,'path_observation':OBS,'semantic_relation':OBS,'endpoints':[]}
        for k,x,target in ((0,90,self.left),(1,197,self.right)):
            box=[int(x)-12,118,int(x)+12,142];path=self.root/f'endpoint-{k}.png'
            im.crop(box).save(path)
            self.record['endpoints'].append({'kind':'shape_boundary','target_shape_id':target.shape_id,
                'crop_box_px':box,'crop_path':str(path),'visible_tip_and_boundary_observation':OBS})
        self.review={'author_id':'SYNTHETIC_AUTHOR_ONLY','connector_reviews':[self.record]}
        self.counter=0

    def proof(self,finding,failures):
        self.counter+=1
        context=layout_context(finding,self.slide,self.render,self.image,failures)
        # Two real local regions cover even a full-page context, without ever
        # claiming that a full page is itself an independent local view.
        boxes=[[0,0,500,540],[460,0,960,540]];details=[]
        with Image.open(self.image) as im:
            for k,box in enumerate(boxes):
                path=self.root/f'synthetic-detail-{self.counter}-{k}.png';im.crop(box).save(path)
                details.append({'path':str(path),'sha256':digest(path),'crop_box_px':box})
        peer={'finding_id':finding['id'],'slide':1,'code':CODE,'pptx_sha256':self.render['sha256'],
              'image_sha256':digest(self.image),'machine_context':context,
              'local_views':[{'sha256':d['sha256'],'crop_box_px':d['crop_box_px']} for d in details],
              'decision':'visually_resolved','local_viewed':True,'full_page_viewed':True,
              'normal_size_viewed':True,'no_visible_layout_defect':True,'unresolved':[],
              'reviewer_id':'SYNTHETIC_PEER_ONLY','reviewed_at':'2026-10-18T09:01:00+00:00',
              'local_observation':OBS,'global_observation':OBS,'review_source':OBS,
              'context_observations':{str(i):OBS for i in context['context_object_ids']},
              'rule_assessments':[{'original_error':f,'why_rule_failed_but_visual_passed':OBS,'no_visible_layout_defect':True} for f in failures],
              'relationship_assessment':{'same_alignment_group':False,'same_metric_group':False,'structural_explanation':OBS}}
        path=self.root/f'synthetic-peer-{self.counter}.json';path.write_text(json.dumps(peer),encoding='utf8')
        return {'visual_fallback':{'code':CODE,'pptx_sha256':self.render['sha256'],'image_sha256':digest(self.image),
                'machine_context':context,'measurement_limitation':OBS,'local_observation':OBS,'global_observation':OBS,
                'returned_to_full_page':True,'normal_size_viewed':True,'no_visible_layout_defect':True,
                'unresolved':[],'local_result':'pass','global_result':'pass',
                'target_observations':{str(i):OBS for i in context['context_object_ids']},'details':details,
                'independent_review':{'path':str(path),'sha256':digest(path)},'reviewed_at':'2026-10-18T09:02:00+00:00'}}

    def preflight(self,reason='outside_safe_margin'):
        issue={'slide':1,'audit_path':str(self.right.shape_id),'shape':self.right.name,'reason':reason,
               'minimum_pt':25.2,'actual_pt':24.9}
        self.render.update(status='failed_preflight',issues=[issue])
        finding=preflight_finding(issue,1,self.prs,native_inventory(self.prs))
        failures=[json.dumps(issue,ensure_ascii=False,sort_keys=True)]
        self.review['render_preflight_reviews']=[{'finding_id':finding['id'],**self.proof(finding,failures)}]
        return issue

    def connection_proof(self,k=1):
        row=inventory(self.slide)[0];endpoint=self.record['endpoints'][k]
        failures=[str(e) for e in endpoint_findings(row,k,endpoint,self.slide,inventory(self.slide),1) if e.category=='layout']
        if row['endpoints_pt'] is None:failures.append('Endpoint local evidence invalid: Endpoint geometry unavailable')
        finding=connector_finding(1,row,k,endpoint,self.record)
        endpoint['layout_visual_review']=self.proof(finding,failures)
        return failures

    def verify(self):
        resolutions=[]
        issues=verify_connections(self.prs,self.review,self.root,render=self.render,resolutions=resolutions)
        return issues,resolutions

    def test_preflight_strict_default_and_unchanged_originals(self):
        self.preflight();original=copy.deepcopy(self.render)
        with self.assertRaises(ValueError):validate_render_evidence(self.pptx,self.render,self.root)
        resolutions=[]
        validate_render_evidence(self.pptx,self.render,self.root,review=self.review,resolutions=resolutions)
        self.assertEqual(self.render,original)
        self.assertEqual(resolutions[0]['original_finding']['original_issue'],original['issues'][0])

    def test_preflight_all_current_layout_categories(self):
        for reason in ('text_height_overflow','text_width_overflow','short_label_wrapped','punctuation_only_line',
                       'outside_safe_margin','native_chart_data_labels_overlap','page_number_zone_overlap'):
            with self.subTest(reason=reason):
                self.preflight(reason)
                validate_render_evidence(self.pptx,self.render,self.root,review=self.review,resolutions=[])

    def test_preflight_nonlayout_constraints_never_close(self):
        self.preflight()
        for reason in ('font_too_small','run_font_too_small','page_number_missing_or_duplicate',
                       'page_number_sequence_or_total_wrong','non_page_number_footer_forbidden',
                       'input_changed_during_render','render_exception','new_unknown_check'):
            with self.subTest(reason=reason):
                self.render['issues'][0]['reason']=reason
                with self.assertRaisesRegex(ValueError,'not a visual layout judgment'):
                    validate_render_evidence(self.pptx,self.render,self.root,review=self.review,resolutions=[])

    def test_preflight_execution_failure_never_enters_visual_adapter(self):
        self.preflight();self.render['status']='failed_render'
        with self.assertRaisesRegex(ValueError,'not complete'):
            validate_render_evidence(self.pptx,self.render,self.root,review=self.review,resolutions=[])

    def test_preflight_complete_inventory_and_snapshot_still_required(self):
        self.preflight();original=copy.deepcopy(self.render)
        for change in (lambda r:r.update(slide_count=2),lambda r:r.update(snapshot=str(self.image)),
                       lambda r:r.update(raster_hashes={}),lambda r:r.update(manual_review=[])):
            self.render=copy.deepcopy(original);change(self.render)
            with self.assertRaises(ValueError):
                validate_render_evidence(self.pptx,self.render,self.root,review=self.review,resolutions=[])

    def test_preflight_exact_issue_coverage_is_required(self):
        self.preflight();self.render['issues'].append(copy.deepcopy(self.render['issues'][0]))
        with self.assertRaisesRegex(ValueError,'every original issue exactly'):
            validate_render_evidence(self.pptx,self.render,self.root,review=self.review,resolutions=[])

    def test_preflight_author_self_review_is_rejected(self):
        self.preflight();self.review['author_id']='SYNTHETIC_PEER_ONLY'
        with self.assertRaisesRegex(ValueError,'different actual reviewer'):
            validate_render_evidence(self.pptx,self.render,self.root,review=self.review,resolutions=[])

    def test_preflight_rule_values_cannot_change_after_review(self):
        self.preflight();self.render['issues'][0]['actual_pt']=23
        with self.assertRaisesRegex(ValueError,'changes/omits original'):
            validate_render_evidence(self.pptx,self.render,self.root,review=self.review,resolutions=[])

    def test_connection_strict_default_and_complete_resolution(self):
        failures=self.connection_proof();original=copy.deepcopy(self.review)
        self.assertTrue(verify_connections(self.prs,self.review,self.root))
        issues,resolutions=self.verify();self.assertEqual(issues,[])
        self.assertEqual(resolutions[0]['rule_failures'],failures);self.assertEqual(self.review,original)

    def test_connection_stale_local_pixels_still_fail(self):
        self.connection_proof();Image.new('RGB',(24,24),'red').save(self.record['endpoints'][1]['crop_path'])
        self.assertTrue(any('exact current' in i['reason'] for i in self.verify()[0]))

    def test_connection_missing_endpoint_still_fails(self):
        self.connection_proof();self.record['endpoints'].pop(0)
        self.assertTrue(any('Both endpoints' in i['reason'] for i in self.verify()[0]))

    def test_connection_missing_native_inventory_still_fails(self):
        self.connection_proof();self.review['connector_reviews']=[]
        self.assertTrue(any('independent native inventory' in i['reason'] for i in self.verify()[0]))

    def test_connection_author_self_review_is_rejected(self):
        self.connection_proof();self.review['author_id']='SYNTHETIC_PEER_ONLY'
        self.assertTrue(any('different actual reviewer' in i['reason'] for i in self.verify()[0]))

    def test_connection_missing_peer_file_is_rejected(self):
        self.connection_proof()
        proof=self.record['endpoints'][1]['layout_visual_review']['visual_fallback']
        proof['independent_review']['path']=str(self.root/'missing-peer.json')
        self.assertTrue(any('Missing/stale independent' in i['reason'] for i in self.verify()[0]))

    def test_connection_visible_defect_cannot_be_closed(self):
        self.connection_proof()
        self.record['endpoints'][1]['layout_visual_review']['visual_fallback']['no_visible_layout_defect']=False
        self.assertTrue(any('visible layout defect' in i['reason'] for i in self.verify()[0]))

    def test_connection_full_page_alone_is_not_local_evidence(self):
        self.connection_proof()
        proof=self.record['endpoints'][1]['layout_visual_review']['visual_fallback']
        proof['details']=[{'path':str(self.image),'sha256':digest(self.image),'crop_box_px':[0,0,960,540]}]
        self.assertTrue(any('genuine nonempty crop' in i['reason'] for i in self.verify()[0]))

    def test_connection_missing_target_remains_nonlayout(self):
        endpoint=self.record['endpoints'][1];endpoint['target_shape_id']=9999
        errors=endpoint_findings(inventory(self.slide)[0],1,endpoint,self.slide,inventory(self.slide),1)
        self.assertTrue(any(e.category=='integrity' and 'Missing real target' in e for e in errors))

    def test_endpoint_layout_failure_does_not_hide_data_mapping_failure(self):
        marker=self.slide.shapes.add_shape(MSO_SHAPE.OVAL,Pt(196),Pt(126),Pt(8),Pt(8))
        marker.fill.solid();marker.fill.fore_color.rgb=RGBColor(10,20,30)
        source=self.root/'synthetic-data.csv';source.write_text('v\n10\n',encoding='utf8')
        proof={'kind':'data_marker_center','target_shape_id':marker.shape_id,'data_mapping':{
            'source_path':str(source),'source_sha256':digest(source),'row_index':0,'column':'v','axis':'x','offset_pt':102,'scale_pt':10}}
        errors=endpoint_findings(inventory(self.slide)[0],1,proof,self.slide,inventory(self.slide),1)
        self.assertTrue(any(e.category=='layout' and 'center' in e for e in errors))
        self.assertTrue(any(e.category=='integrity' and 'source data mapping' in e for e in errors))

    def test_layout_failure_does_not_hide_omitted_adjacent_stroke(self):
        marks=[]
        for y in (220,230,245,260):
            marks.append(self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(30),Pt(y),Pt(70),Pt(y)))
        rows=inventory(self.slide);row=next(r for r in rows if r['shape_id']==marks[0].shape_id)
        proof={'kind':'icon_stroke','cluster_shape_ids':[m.shape_id for m in marks[:3]]}
        errors=endpoint_findings(row,0,proof,self.slide,rows,1)
        self.assertTrue(any(e.category=='integrity' and 'omits a matching adjacent stroke' in e for e in errors))

    def test_unsupported_geometry_requires_both_current_crops_and_full_context(self):
        # A native block arrow has no implemented endpoint measurement, but its
        # complete current geometry can be reviewed without changing the shape.
        self.slide.shapes._spTree.remove(self.line._element)
        block=self.slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Pt(90),Pt(120),Pt(107),Pt(20))
        block.fill.solid();block.fill.fore_color.rgb=RGBColor(10,20,30)
        self.line=block;self.record['shape_id']=block.shape_id
        self.prs.save(self.pptx);shutil.copyfile(self.pptx,self.snapshot);self.render['sha256']=digest(self.pptx)
        validate_render_evidence(self.pptx,self.render,self.root)
        for k in (0,1):self.connection_proof(k)
        issues,resolutions=self.verify()
        self.assertEqual(issues,[]);self.assertEqual(len(resolutions),2)
        self.assertTrue(all(any('geometry unavailable' in e.lower() for e in r['rule_failures']) for r in resolutions))

    def test_block_arrow_cannot_be_relabelled_as_open_decoration(self):
        block=self.slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Pt(90),Pt(120),Pt(107),Pt(20))
        rows=inventory(self.slide);row=next(r for r in rows if r['shape_id']==block.shape_id)
        errors=endpoint_findings(row,0,{'kind':'open_nonrelational','purpose':'separator'},self.slide,rows,1)
        self.assertTrue(any(e.category=='integrity' and 'self-exempted' in e for e in errors))


if __name__=='__main__':unittest.main(verbosity=2)
