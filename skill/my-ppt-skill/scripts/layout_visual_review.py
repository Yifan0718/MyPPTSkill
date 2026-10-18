"""Uniform visual reconciliation of layout-rule results, never input failures.

Callers first run their original rule and provenance checks. Only their layout
results enter this adapter; no diagram/error-string whitelist changes the
thresholds or erases findings. A second actual reviewer must explain each
machine failure against exact local pixels and the complete current page.
This validates the evidence chain, not whether a person/agent really looked.
"""
import copy
import math
import json
from pathlib import Path
from PIL import Image, ImageChops
from pptx.enum.shapes import MSO_SHAPE_TYPE
from measurement_fallback import digest, _specific, _time
from interstitial_spacing import _full_slide_background_ids

CODE='layout_rule_visual_resolution'


def _target_ids(value, known):
    result=set()
    if isinstance(value,dict):
        for key,v in value.items():
            if key.endswith('_id') and type(v) is int:
                if v not in known:raise ValueError('Layout finding references an unknown native object')
                result.add(v)
            elif key.endswith('_ids') and isinstance(v,list):
                if any(type(i) is not int or i not in known for i in v):raise ValueError('Layout finding object list is invalid')
                result.update(v)
            elif isinstance(v,(list,dict)):result.update(_target_ids(v,known))
    elif isinstance(value,list):
        for v in value:result.update(_target_ids(v,known))
    return result


def layout_context(finding, slide, render, image_path, rule_failures):
    """Recompute evidence to bind. The caller supplies its ORIGINAL rule output."""
    if not isinstance(finding,dict) or not isinstance(finding.get('id'),str) or type(finding.get('slide')) is not int:
        raise ValueError('Layout review needs the intact identified original finding')
    if not isinstance(rule_failures,list) or not rule_failures or any(not isinstance(s,str) or not s.strip() for s in rule_failures):
        raise ValueError('Layout review needs every original rule failure, not a replacement pass')
    shapes={};paths={}
    def visit(items,prefix=''):
        for s in items:
            if s.shape_id in shapes:raise ValueError('Ambiguous duplicate native object ID')
            shapes[s.shape_id]=s;paths[s.shape_id]=prefix+str(s.shape_id)
            if s.shape_type==MSO_SHAPE_TYPE.GROUP:visit(s.shapes,paths[s.shape_id]+'/group/')
    visit(slide.shapes)
    backgrounds=_full_slide_background_ids(slide)
    targets=_target_ids(finding,set(shapes)) or (set(shapes)-backgrounds)
    if not targets:raise ValueError('No native layout context was discovered')
    prs=slide.part.package.presentation_part.presentation
    from render_evidence import native_inventory
    inv=native_inventory(prs);page_number=finding['slide']
    width,height=prs.slide_width/12700,prs.slide_height/12700;scale=height/540
    def bounds(s):return inv['bounds'].get((page_number,paths[s.shape_id])) or [0,0,width,height]
    bb=[bounds(shapes[i]) for i in targets]
    region=[max(0,min(b[0] for b in bb)-12*scale),max(0,min(b[1] for b in bb)-12*scale),
            min(width,max(b[2] for b in bb)+12*scale),min(height,max(b[3] for b in bb)+12*scale)]
    context=set(targets)
    for sid,s in shapes.items():
        b=bounds(s)
        if sid not in backgrounds and min(b[2],region[2])>=max(b[0],region[0]) and min(b[3],region[3])>=max(b[1],region[1]):context.add(sid)
    with Image.open(image_path) as source:im=source.convert('RGB')
    sx,sy=im.width/width,im.height/height
    native={};measured={}
    from audit_layout_coverage import ink_bounds
    for sid in sorted(context):
        s=shapes[sid]
        native[str(sid)]={'audit_path':paths[sid],'bounds_pt':inv['bounds'].get((page_number,paths[sid])),
                          'rotation':float(s.rotation),'shape_type':int(s.shape_type),
                          'text':s.text if s.has_text_frame else None}
        records=[r for r in render.get('text_records',[]) if r.get('slide')==finding['slide'] and r.get('kind')=='text' and r.get('shape_id')==sid]
        if s.has_text_frame and s.text.strip():
            if len(records)!=1:raise ValueError('Current layout context has missing/ambiguous rendered text evidence')
            record=records[0]
            keys=('bound_left_pt','bound_top_pt','bound_width_pt','bound_height_pt')
            finite=all(type(record.get(k)) in (int,float) and math.isfinite(record[k]) for k in keys)
            v=ink_bounds(record,im,sx) if finite else None
            measured[str(sid)]={'render_record':copy.deepcopy(record),'rendered_bounds_pt':[record.get(k) for k in keys],
                                'ink_bounds_px':v['bounds_px'] if v else None}
    # Rotation is retained in native evidence. Use the whole page for context
    # when a rotated object's envelope cannot be proven from its frame; tiled
    # local crops permit full coverage without pretending a full page is local.
    if any(inv['bounds'].get((page_number,paths[i])) is None for i in context):crop=[0,0,im.width,im.height]
    else:
        boxes=[bounds(shapes[i]) for i in context]
        crop=[max(0,math.floor(min(b[0] for b in boxes)*sx)-2),max(0,math.floor(min(b[1] for b in boxes)*sy)-2),
              min(im.width,math.ceil(max(b[2] for b in boxes)*sx)+2),min(im.height,math.ceil(max(b[3] for b in boxes)*sy)+2)]
    return {'original_finding':copy.deepcopy(finding),'rule_failures':list(rule_failures),
            'page_object_ids':sorted(shapes),'page_object_paths':sorted(p for n,p in inv['objects'] if n==page_number),
            'target_ids':sorted(targets),'context_object_ids':sorted(context),
            'native_objects':native,'current_text_measurements':measured,'required_context_box_px':crop}


def _covered(required, boxes):
    """Exact rectangle-union coverage; permits several genuine local views."""
    l,t,r,b=required
    xs=sorted({l,r}|{max(l,min(r,v)) for box in boxes for v in (box[0],box[2])})
    for a,z in zip(xs,xs[1:]):
        if a==z:continue
        intervals=sorted((max(t,box[1]),min(b,box[3])) for box in boxes if box[0]<=a and box[2]>=z)
        end=t
        for start,stop in intervals:
            if start>end:break
            end=max(end,stop)
        if end<b:return False
    return True


def validate_layout_visual_review(finding, disposition, image_path, slide, render, author_id, rule_failures):
    """Return (error, resolution), preserving original offsets/errors verbatim."""
    try:
        proof=layout_context(finding,slide,render,image_path,rule_failures)
        record=disposition.get('visual_fallback')
        if not isinstance(record,dict) or record.get('code')!=CODE:
            raise ValueError('Layout rule unresolved: supply current local + full-page independent visual reconciliation')
        image_sha=digest(image_path);pptx_sha=render['sha256']
        if record.get('pptx_sha256')!=pptx_sha or record.get('image_sha256')!=image_sha or render['raster_hashes'].get(str(finding['slide']))!=image_sha:
            raise ValueError('Layout review input/image is stale')
        if record.get('machine_context')!=proof:
            raise ValueError('Layout review changes/omits original failures, numeric values, native objects, or current pixel evidence')
        for key in ('measurement_limitation','local_observation','global_observation'):
            if not _specific(record.get(key)):raise ValueError('Layout review needs specific rule/local/full-page observations')
        if any(record.get(k) is not True for k in ('returned_to_full_page','normal_size_viewed','no_visible_layout_defect')) or record.get('unresolved')!=[] or record.get('local_result')!='pass' or record.get('global_result')!='pass':
            raise ValueError('A visible layout defect or unresolved local/full-page concern remains; continue repair')
        observations=record.get('target_observations',{})
        keys=set(map(str,proof['context_object_ids']))
        if set(observations)!=keys or not all(_specific(v) for v in observations.values()):raise ValueError('Every discovered local context object needs an actual observation')
        details=record.get('details');boxes=[];views=[]
        if not isinstance(details,list) or not details:raise ValueError('Provide at least one current local view, separate from the full page')
        with Image.open(image_path) as opened:
            page=opened.convert('RGB')
            for detail in details:
                box=detail.get('crop_box_px');path=Path(detail.get('path',''))
                if not isinstance(box,list) or len(box)!=4 or any(type(v) is not int for v in box) or not (0<=box[0]<box[2]<=page.width and 0<=box[1]<box[3]<=page.height) or box==[0,0,page.width,page.height]:
                    raise ValueError('Local view must be a genuine nonempty crop; full-page inspection is separate')
                if not path.is_file() or digest(path)!=detail.get('sha256'):raise ValueError('Missing/stale local view')
                with Image.open(path) as image:
                    local=image.convert('RGB');expected=page.crop(box)
                    if local.size!=expected.size or ImageChops.difference(local,expected).getbbox():raise ValueError('Local view is not exact current rendered pixels')
                boxes.append(box);views.append({'sha256':detail['sha256'],'crop_box_px':box})
        if not _covered(proof['required_context_box_px'],boxes):raise ValueError('Local views omit discovered relationship members or boundary context')
        ref=record.get('independent_review',{});path=Path(ref.get('path',''))
        if not path.is_file() or digest(path)!=ref.get('sha256'):raise ValueError('Missing/stale independent actual visual review')
        peer=json.loads(path.read_text(encoding='utf-8-sig'));reviewer=peer.get('reviewer_id')
        if not isinstance(author_id,str) or not author_id.strip() or not isinstance(reviewer,str) or not reviewer.strip() or reviewer==author_id:
            raise ValueError('A different actual reviewer is required; author self-exemption cannot close layout rules')
        expected={'finding_id':finding['id'],'slide':finding['slide'],'code':CODE,'pptx_sha256':pptx_sha,'image_sha256':image_sha,
                  'machine_context':proof,'local_views':views,'decision':'visually_resolved','local_viewed':True,
                  'full_page_viewed':True,'normal_size_viewed':True,'no_visible_layout_defect':True,'unresolved':[]}
        if any(peer.get(k)!=v for k,v in expected.items()):raise ValueError('Independent review has not closed this exact original finding and all current context')
        for key in ('local_observation','global_observation','review_source'):
            if not _specific(peer.get(key)):raise ValueError('Independent review needs specific local/global observations and actual provenance')
        peer_observations=peer.get('context_observations',{})
        if set(peer_observations)!=keys or not all(_specific(v) for v in peer_observations.values()):raise ValueError('Independent review omitted native context objects')
        assessments=peer.get('rule_assessments',[])
        if not isinstance(assessments,list) or len(assessments)!=len(rule_failures):raise ValueError('Independent review must explain every preserved rule failure')
        for failure,assessment in zip(rule_failures,assessments):
            if assessment.get('original_error')!=failure or not _specific(assessment.get('why_rule_failed_but_visual_passed')) or assessment.get('no_visible_layout_defect') is not True:
                raise ValueError('Every layout-rule failure needs a specific visual reconciliation, not a role/skip label')
        relation=peer.get('relationship_assessment',{})
        if any(type(relation.get(k)) is not bool for k in ('same_alignment_group','same_metric_group')) or not _specific(relation.get('structural_explanation')):
            raise ValueError('Independent review must explicitly assess actual alignment/metric grouping and visible structure')
        if not _time(render['render_finished_utc'])<=_time(peer['reviewed_at'])<=_time(record['reviewed_at']):raise ValueError('Layout visual review predates rendering or independent inspection')
        return None,{'finding_id':finding['id'],'code':CODE,'resolution':'visually_resolved_pending_signoff',
                     'original_finding':copy.deepcopy(finding),'rule_failures':list(rule_failures),'reviewer_id':reviewer,
                     'independent_review_sha256':ref['sha256'],'image_sha256':image_sha,'local_views':views}
    except (ValueError,TypeError,KeyError,AttributeError,IndexError,OSError) as exc:
        return str(exc),None
