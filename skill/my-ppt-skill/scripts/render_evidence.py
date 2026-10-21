"""Draft render-evidence/v2 gate. Synthetic tests verify contracts, not visual quality.

Integration: call validate_render_evidence before coverage scans; propagate
validate_manual_dispositions issues into final evidence status. Never interpret
successful structural validation as visual signoff.
"""
from pathlib import Path
from io import BytesIO
import copy
import json
import hashlib
import math
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SCHEMA = 'render-evidence/v2'
CHART_ERROR = 'native_plot_bounds_unavailable'

# Categories belong to the renderer's checks, not reviewer-provided labels.
# This exhaustively describes the current renderer's preflight rules. Unknown
# new checks fail closed until their actual contract is classified here.
PREFLIGHT_RULE_CATEGORY = {
    'text_height_overflow': 'layout', 'text_width_overflow': 'layout',
    'short_label_wrapped': 'layout', 'punctuation_only_line': 'layout',
    'outside_safe_margin': 'layout', 'native_chart_data_labels_overlap': 'layout',
    'page_number_zone_overlap': 'layout',
    'font_too_small': 'font_constraint', 'run_font_too_small': 'font_constraint',
    'non_page_number_footer_forbidden': 'content_constraint',
    'page_number_sequence_or_total_wrong': 'content_constraint',
    'page_number_missing_or_duplicate': 'inventory_constraint',
    'input_changed_during_render': 'provenance_constraint',
    'render_exception': 'execution_failure',
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest().upper()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def nonempty(value):
    return isinstance(value, str) and bool(value.strip())


def records(value, label):
    require(isinstance(value, list) and all(isinstance(r, dict) for r in value), label + ' must be an array of objects')
    return value


def key(row):
    require(type(row.get('slide')) is int and nonempty(row.get('audit_path')), 'Invalid slide/audit_path')
    return row['slide'], row['audit_path']


def keyed(rows, label):
    result = {}
    for row in records(rows, label):
        k = key(row)
        require(k not in result, label + ' duplicate: ' + str(k))
        result[k] = row
    return result


def native_inventory(prs):
    """Paths are independent of COM kind/name and include nested group children."""
    objects, texts, charts, pictures, groups = set(), set(), set(), set(), set()
    bounds = {};text_values={}
    def visit(shapes, slide, prefix='', transform=(1, 1, 0, 0), uncertain=False):
        for shape in shapes:
            path = prefix + str(shape.shape_id)
            k = (slide, path)
            require(k not in objects, 'Duplicate native object path: ' + str(k))
            objects.add(k)
            sx, sy, ox, oy = transform
            left, top, width, height = [v / 12700 for v in (shape.left, shape.top, shape.width, shape.height)]
            unmeasured = uncertain or shape.rotation != 0
            bounds[k] = None if unmeasured else [ox + sx * left, oy + sy * top, ox + sx * (left + width), oy + sy * (top + height)]
            if shape.has_text_frame and shape.text.strip():
                texts.add(k)
                text_values[k]=shape.text
            if shape.has_table:
                for ri, row in enumerate(shape.table.rows, 1):
                    for ci, cell in enumerate(row.cells, 1):
                        cell_key = (slide, path + '/table/' + str(ri) + '/' + str(ci))
                        objects.add(cell_key)
                        cl = left + sum(c.width for c in list(shape.table.columns)[:ci-1]) / 12700
                        ct = top + sum(r.height for r in list(shape.table.rows)[:ri-1]) / 12700
                        cw, ch = shape.table.columns[ci-1].width/12700, shape.table.rows[ri-1].height/12700
                        bounds[cell_key] = None if unmeasured else [ox+sx*cl, oy+sy*ct, ox+sx*(cl+cw), oy+sy*(ct+ch)]
                        if cell.text.strip():
                            texts.add(cell_key)
                            text_values[cell_key]=cell.text
            if shape.has_chart:
                charts.add(k)
            if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
                pictures.add(k)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                groups.add(k)
                try:
                    xf = shape._element.xpath('./p:grpSpPr/a:xfrm')[0]
                    off, ext = xf.xpath('./a:chOff')[0], xf.xpath('./a:chExt')[0]
                    cx, cy = int(off.get('x'))/12700, int(off.get('y'))/12700
                    cw, ch = int(ext.get('cx'))/12700, int(ext.get('cy'))/12700
                    fx, fy = width/cw, height/ch
                    child_transform = (sx*fx, sy*fy, ox+sx*(left-cx*fx), oy+sy*(top-cy*fy))
                    child_uncertain = unmeasured or xf.get('flipH') in ('1','true') or xf.get('flipV') in ('1','true')
                except (IndexError, TypeError, ValueError, ZeroDivisionError):
                    child_transform, child_uncertain = transform, True
                visit(shape.shapes, slide, path + '/group/', child_transform, child_uncertain)
    for n, slide in enumerate(prs.slides, 1):
        visit(slide.shapes, n)
        objects.add((n, 'master-layout'))
        bounds[n, 'master-layout'] = [0, 0, prs.slide_width/12700, prs.slide_height/12700]
    return {'objects': objects, 'texts': texts, 'charts': charts, 'pictures': pictures, 'groups': groups, 'bounds': bounds,'text_values':text_values}


def finite_vector(value, length, label):
    require(isinstance(value, (list, tuple)) and len(value) == length and
            all(type(v) in (int, float) and math.isfinite(v) for v in value), 'Invalid finite ' + label)


def opaque_uniform_picture_color(prs, object_key):
    """Read the embedded source, never a reviewer-supplied uniformity claim."""
    slide_number, path = object_key
    shapes = prs.slides[slide_number - 1].shapes
    parts = path.split('/group/')
    for index, part in enumerate(parts):
        shape = next((s for s in shapes if str(s.shape_id) == part), None)
        if shape is None:
            return None
        if index < len(parts) - 1:
            if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
                return None
            shapes = shape.shapes
    if shape.shape_type not in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
        return None
    try:
        with Image.open(BytesIO(shape.image.blob)) as embedded:
            if getattr(embedded, 'n_frames', 1) != 1:
                return None
            extrema = embedded.convert('RGBA').getextrema()
            if extrema[3] != (255, 255) or any(lo != hi for lo, hi in extrema[:3]):
                return None
            return tuple(lo for lo, hi in extrema[:3])
    except (AttributeError, KeyError, OSError, ValueError, Image.DecompressionBombError):
        return None


def validate_uniform_picture_context(prs, original, item, source, target_px, target_pixels):
    """Narrow exception for visible constant-color native picture interiors.

    A current local crop must surround every object edge. Contrast is checked
    immediately outside each edge, so unrelated marks elsewhere cannot rescue
    an invisible white-on-white image. This remains evidence validation only.
    """
    color = opaque_uniform_picture_color(prs, key(original))
    require(color is not None, 'Blank/constant pixels require an opaque uniform native picture source')
    require(target_pixels.getextrema() == tuple((v, v) for v in color),
            'Uniform picture pixels do not match the native source color')
    detail = item.get('detail_crop_box_px')
    require(isinstance(detail, list) and len(detail) == 4 and all(type(v) is int for v in detail),
            'Uniform picture requires explicit local detail_crop_box_px')
    require(nonempty(item.get('detail_crop_path')), 'Uniform picture requires explicit detail_crop_path')
    require(0 <= detail[0] < detail[2] <= source.width and 0 <= detail[1] < detail[3] <= source.height,
            'Uniform picture local detail outside image')
    require(detail != [0, 0, source.width, source.height],
            'Uniform picture needs a local detail crop, not the full page')
    left, top, right, bottom = target_px
    require(detail[0] < left and detail[1] < top and detail[2] > right and detail[3] > bottom,
            'Uniform picture local detail must include context outside all object boundaries')
    page = source.convert('RGB')
    with Image.open(item['detail_crop_path']) as detail_file:
        expected, actual = page.crop(detail), detail_file.convert('RGB')
        require(expected.size == actual.size and ImageChops.difference(expected, actual).getbbox() is None,
                'Uniform picture local detail does not match current page')
    adjacent_strips = [(left - 1, top, left, bottom), (right, top, right + 1, bottom),
                       (left, top - 1, right, top), (left, bottom, right, bottom + 1)]
    uniform_extrema = tuple((v, v) for v in color)
    require(all(page.crop(strip).getextrema() != uniform_extrema for strip in adjacent_strips),
            'Uniform picture lacks visible contrast at its object boundaries')


def preflight_finding(issue, index, prs, inv):
    """Bind an original renderer row to independently discovered native targets."""
    require(PREFLIGHT_RULE_CATEGORY.get(issue.get('reason')) == 'layout',
            'Preflight issue is not a visual layout judgment: ' + str(issue.get('reason')))
    n = issue.get('slide')
    require(type(n) is int and 1 <= n <= len(prs.slides), 'Preflight issue has no current slide')
    finding = {'id': f'render-preflight:{index}', 'slide': n,
               'type': 'render_preflight_layout', 'original_issue': copy.deepcopy(issue)}
    path = issue.get('audit_path')
    if path is not None:
        require((n, path) in inv['objects'], 'Preflight issue refers to a missing native object')
        if path != 'master-layout':
            finding['shape_id'] = int(path.split('/')[0])
    else:
        # Older page-number overlap rows identify shapes by their actual names.
        names = [issue[k] for k in ('shape', 'overlapping_shape') if k in issue]
        shapes = list(prs.slides[n - 1].shapes)
        ids = []
        for name in names:
            matches = [s.shape_id for s in shapes if s.name == name]
            require(len(matches) == 1, 'Preflight object name is missing or ambiguous')
            ids.extend(matches)
        if ids:
            finding['shape_ids'] = sorted(set(ids))
    return finding


def validate_preflight_dispositions(render, review, render_dir, prs, inv):
    """Close layout judgments only; preserve the original report and every row."""
    issues = render['issues']
    require(isinstance(review, dict), 'Rendering has unresolved issues; structured visual review is required')
    done = records(review.get('render_preflight_reviews'), 'render_preflight_reviews')
    ids = [r.get('finding_id') for r in done]
    require(all(isinstance(i, str) for i in ids) and len(ids) == len(set(ids)) and
            set(ids) == {f'render-preflight:{i}' for i in range(1, len(issues) + 1)},
            'Preflight dispositions must cover every original issue exactly')
    by_id = {r['finding_id']: r for r in done}
    author_id = review.get('author_id')
    results, errors = [], []
    from layout_visual_review import validate_layout_visual_review
    for i, issue in enumerate(issues, 1):
        try:
            finding = preflight_finding(issue, i, prs, inv)
            # Deterministic complete original row, including measured numbers.
            failures = [json.dumps(issue, ensure_ascii=False, sort_keys=True)]
            error, resolution = validate_layout_visual_review(
                finding, by_id[finding['id']], Path(render_dir) / f'slide-{finding["slide"]:02}.png',
                prs.slides[finding['slide'] - 1], render, author_id, failures)
            if error:
                errors.append(finding['id'] + ': ' + error)
            else:
                results.append(resolution)
        except (ValueError, TypeError, KeyError, OSError) as exc:
            errors.append(f'render-preflight:{i}: {exc}')
    require(not errors, 'Unresolved rendering preflight: ' + '; '.join(errors))
    return results


def validate_bounded_run(render,render_dir):
    path=Path(render_dir)/'tool-run.json';expected=render.get('bounded_run_id')
    if not expected and not path.exists():return
    require(path.is_file(),'Missing bounded tool-run evidence; use a fresh render run')
    record=json.loads(path.read_text(encoding='utf-8-sig'))
    require(isinstance(record,dict) and record.get('schema')=='bounded-run/v1','Invalid bounded tool-run evidence')
    require(isinstance(expected,str) and bool(expected) and record.get('run_id')==expected,'Bounded run identity mismatch')
    require(record.get('status')=='finished','Render invocation timed out, was cancelled, or is incomplete; rerender in a fresh directory')
    code=record.get('exit_code')
    require(type(code)==int and (code==0 or (code==2 and render.get('status')=='failed_preflight')),'Render invocation failed; inspect its exit status before retrying')


def validate_render_evidence(pptx, render, render_dir, *, diagnostic_attempt=False, review=None, resolutions=None):
    """Raises ValueError on stale, incomplete, duplicate, or unsupported evidence.

    raster_hashes is a dictionary mapping 1-based slide number strings to SHA256.
    Returns native inventory for callers; unresolved manual reviews remain pending.
    """
    pptx, render_dir = Path(pptx), Path(render_dir).resolve()
    require(isinstance(render, dict) and render.get('schema') == SCHEMA, 'Require render-evidence/v2; rerender with the current renderer')
    validate_bounded_run(render,render_dir)
    fallback_requested = review is not None and not diagnostic_attempt
    require(render.get('status') in (('rendered_pending_visual_review','failed_preflight') if diagnostic_attempt or fallback_requested else ('rendered_pending_visual_review',)), 'Rendering is not complete')
    require(isinstance(render.get('issues'),list) and all(isinstance(i,dict) for i in render['issues']) and
            (diagnostic_attempt or fallback_requested or render['issues']==[]), 'Rendering has unresolved or missing issues state')
    require(render.get('status') != 'failed_preflight' or bool(render['issues']), 'Failed preflight has no original issues')
    expected_hash = digest(pptx)
    require(render.get('sha256') == expected_hash, 'PPTX hash mismatch')
    require(nonempty(render.get('snapshot')), 'Missing snapshot path')
    snapshot = Path(render['snapshot'])
    require(snapshot.is_absolute() and snapshot.is_file(), 'Snapshot must be an existing absolute file')
    require(snapshot.resolve().parent == render_dir, 'Snapshot must belong to the checked render directory')
    require(digest(snapshot) == expected_hash, 'Actual snapshot hash mismatch')
    prs = Presentation(pptx)
    count = len(prs.slides)
    require(type(render.get('slide_count')) is int and render['slide_count'] == count, 'Slide count mismatch')
    hashes = render.get('raster_hashes')
    require(isinstance(hashes, dict) and set(hashes) == {str(n) for n in range(1, count + 1)}, 'Raster inventory missing/extra pages')
    for n in range(1, count + 1):
        image = render_dir / f'slide-{n:02}.png'
        require(image.is_file() and hashes[str(n)] == digest(image), 'Raster missing/stale: ' + str(n))
        with Image.open(image) as opened:
            opened.verify()
    inv = native_inventory(prs)
    text = keyed(render.get('text_records'), 'text_records')
    require(set(text) == inv['texts'], 'Text inventory mismatch; missing=' + str(sorted(inv['texts'] - set(text))) + '; extra=' + str(sorted(set(text) - inv['texts'])))
    for k, row in text.items():
        require(nonempty(row.get('text')), 'Empty text record: ' + str(k))
        require(''.join(row['text'].split())==''.join(inv['text_values'][k].split()),'Text content differs from actual PPTX: '+str(k))
        fonts=row.get('run_font_sizes_pt')
        require(isinstance(fonts,list) and fonts and all(type(v) in (int,float) and math.isfinite(v) and v>0 for v in fonts),'Unresolved effective run font sizes: '+str(k))
        dimensions = [row.get(name) for name in ('left_pt', 'top_pt', 'width_pt', 'height_pt', 'bound_left_pt', 'bound_top_pt', 'bound_width_pt', 'bound_height_pt')]
        finite_vector(dimensions, 8, 'text geometry: ' + str(k))
        require(all(dimensions[i] > 0 for i in (2, 3, 6, 7)), 'Nonpositive text dimensions: ' + str(k))
    geometry = keyed(render.get('chart_geometry_records'), 'chart_geometry_records')
    require(set(geometry) <= inv['charts'], 'Unknown chart geometry object')
    for k, row in geometry.items():
        box = row.get('plot_bounds_pt')
        finite_vector(box, 4, 'plot bounds: ' + str(k))
        require(box[2] > box[0] and box[3] > box[1], 'Empty/reversed plot bounds: ' + str(k))
    manual = records(render.get('manual_review'), 'manual_review')
    ids = [r.get('review_id') for r in manual]
    require(all(isinstance(i, str) for i in ids) and len(set(ids)) == len(ids) and set(ids) == {f'm{i}' for i in range(1, len(ids) + 1)}, 'Manual review IDs must be unique m1..mn')
    manual_keys = set()
    for row in manual:
        k = key(row)
        require(k in inv['objects'] and nonempty(row.get('reason')), 'Unknown/invalid manual review target')
        mk = (k, row['reason'])
        require(mk not in manual_keys, 'Duplicate manual object/reason')
        manual_keys.add(mk)
    for n in range(1, count + 1):
        require(((n, 'master-layout'), 'master_layout_visual_review') in manual_keys, 'Missing master/layout review: ' + str(n))
    chart_errors = {key(r) for r in manual if r['reason'] == CHART_ERROR}
    require(chart_errors <= inv['charts'], 'Chart-error targets a non-chart')
    require(not (set(geometry) & chart_errors), 'Chart cannot simultaneously have successful and unavailable geometry')
    require(set(geometry) | chart_errors == inv['charts'], 'Missing native chart geometry or explicit chart-error review')
    for category, reason in [('charts', 'chart_axes_curves_annotations_require_visual_review'), ('pictures', 'picture_internal_visual_review'), ('groups', 'group_geometry_requires_visual_review')]:
        for k in inv[category]:
            require((k, reason) in manual_keys, 'Missing internal visual review for ' + str(k) + ': ' + reason)
    require(digest(pptx) == expected_hash and digest(snapshot) == expected_hash, 'Input changed during evidence validation')
    # Run every original snapshot/inventory/content check before attempting a
    # layout reconciliation. A visual observation cannot cure missing evidence.
    if fallback_requested and render['issues']:
        resolved = validate_preflight_dispositions(render, review, render_dir, prs, inv)
        if resolutions is not None:
            resolutions.extend(resolved)
    return inv


def validate_manual_dispositions(render, review, render_dir):
    """Every renderer manual/error row requires its own current pixel evidence.

    Caller must first validate_render_evidence. Returns issues for propagation;
    never treats merely supplying an error reason as closing a manual review.
    """
    issues = []
    try:
        prs = Presentation(render['snapshot'])
        inventory = native_inventory(prs)
        pending = {r['review_id']: r for r in records(render.get('manual_review'), 'manual_review')}
        done = records(review.get('render_manual_reviews'), 'render_manual_reviews')
        ids = [r.get('review_id') for r in done]
        require(all(isinstance(i, str) for i in ids) and len(set(ids)) == len(ids) and set(ids) == set(pending), 'Render manual dispositions must cover pending IDs exactly')
        for item in done:
            try:
                original = pending[item['review_id']]
                require(key(item) == key(original), 'Manual disposition target mismatch')
                require(item.get('decision') == 'verified_in_context' and item.get('returned_to_full_page') is True, 'Manual review remains unresolved')
                require(nonempty(item.get('observation')) and item['observation'].strip().lower() not in ('ok', 'pass', 'passed', 'n/a', '通过', '不适用'), 'Missing specific manual observation')
                path = Path(render_dir) / f"slide-{original['slide']:02}.png"
                require(item.get('image_sha256') == digest(path) == render['raster_hashes'][str(original['slide'])], 'Manual evidence image stale')
                box = item.get('crop_box_px')
                require(isinstance(box, list) and len(box) == 4 and all(type(v) is int for v in box), 'Invalid manual crop box')
                require(nonempty(item.get('crop_path')), 'Missing manual crop file')
                with Image.open(path) as source, Image.open(item['crop_path']) as local:
                    require(0 <= box[0] < box[2] <= source.width and 0 <= box[1] < box[3] <= source.height, 'Manual crop outside image')
                    a, b = source.convert('RGB').crop(box), local.convert('RGB')
                    require(a.size == b.size and ImageChops.difference(a, b).getbbox() is None, 'Manual crop not current rendered pixels')
                    target = inventory['bounds'].get(key(original))
                    if target is None:
                        require(box == [0, 0, source.width, source.height], 'Unmeasured geometry requires current full-page crop')
                        detail = item.get('detail_crop_box_px')
                        require(isinstance(detail, list) and len(detail) == 4 and all(type(v) is int for v in detail), 'Unmeasured geometry also needs explicit local detail')
                        require(0 <= detail[0] < detail[2] <= source.width and 0 <= detail[1] < detail[3] <= source.height, 'Invalid unmeasured local detail box')
                        with Image.open(item['detail_crop_path']) as detail_file:
                            a, b = source.convert('RGB').crop(detail), detail_file.convert('RGB')
                            require(a.size == b.size and ImageChops.difference(a,b).getbbox() is None, 'Local detail does not match current page')
                    else:
                        px, py = source.width/(prs.slide_width/12700), source.height/(prs.slide_height/12700)
                        target_px = [max(0, math.floor(target[0]*px)), max(0, math.floor(target[1]*py)), min(source.width, math.ceil(target[2]*px)), min(source.height, math.ceil(target[3]*py))]
                        require(box[0] <= target_px[0] and box[1] <= target_px[1] and box[2] >= target_px[2] and box[3] >= target_px[3], 'Manual crop omits its native object bounds')
                        # Check target pixels, not unrelated content elsewhere in a large crop.
                        a = source.convert('RGB').crop(target_px)
                    require(a.width > 0 and a.height > 0, 'Blank/constant pixels cannot close a visual review')
                    if not any(lo != hi for lo, hi in a.getextrema()):
                        require(target is not None and key(original) in inventory['pictures'],
                                'Blank/constant pixels cannot close a visual review')
                        validate_uniform_picture_context(prs, original, item, source, target_px, a)
            except (ValueError, TypeError, KeyError, OSError) as exc:
                issues.append({'review_id': item.get('review_id'), 'reason': str(exc)})
    except (ValueError, TypeError, KeyError, OSError) as exc:
        issues.append({'reason': str(exc)})
    return issues
