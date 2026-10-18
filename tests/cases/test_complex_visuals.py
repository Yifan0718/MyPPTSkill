"""Build an editable two-slide visual regression fixture (all data simulated).
Usage: python test_complex_visuals.py output.pptx
Render with render_powerpoint_audit.ps1, then inspect both slides visually.
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData, BubbleChartData
from pptx.oxml.xmlchemy import OxmlElement

def main():
    if len(sys.argv)!=2:
        raise SystemExit('Usage: test_complex_visuals.py output.pptx')
    P=Presentation();P.slide_width=Inches(13.333333);P.slide_height=Inches(7.5)
    BLUE='426DB2';NAVY='2C3E70';INK='24334B';GRAY='59677C';LIGHT='EFF3FA';RED='AC433D';WHITE='FFFFFF'
    FONT='Microsoft YaHei'
    def rgb(c):return RGBColor.from_string(c)
    def tx(s,t,x,y,w,h,sz=14,bold=False,color=INK,role='body',align=PP_ALIGN.LEFT):
     sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.name=role+':'+t[:22]
     tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
     tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
     for i,line in enumerate(t.split('\n')):
      p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_after=Pt(3);p.line_spacing=1.1
      r=p.add_run();r.text=line;r.font.name=FONT;r.font.size=Pt(sz);r.font.bold=bold;r.font.color.rgb=rgb(color)
      ea=OxmlElement('a:ea');ea.set('typeface',FONT);r._r.get_or_add_rPr().append(ea)
     return sh
    def title(s,n,t,sub):
     tx(s,t,.55,.43,12.23,.5,28,True,NAVY,'title');tx(s,sub,.55,1.12,12.23,.42,14,False,GRAY,'subtitle')
     tx(s,f'{n:02d} / 02',11.8,6.94,.98,.2,10,False,GRAY,'page-number',PP_ALIGN.RIGHT)
    def style(c):
     c.font.name=FONT;c.font.size=Pt(12);c.has_title=False;c.has_legend=False
     try:
      for ax in [c.category_axis,c.value_axis]:
       ax.tick_labels.font.size=Pt(12);ax.tick_labels.font.name=FONT;ax.tick_label_position=XL_TICK_LABEL_POSITION.LOW
     except (ValueError,AttributeError):pass
    def series_color(sr,col):
     sr.format.fill.solid();sr.format.fill.fore_color.rgb=rgb(col);sr.format.line.fill.background()

    s=P.slides.add_slide(P.slide_layouts[6]);title(s,1,'学习任务诊断：难度、提示与结果','模拟样本：5 类任务、100 次尝试；图表共享同一份数据，用于验证复杂图表能力。')
    cats=['配方','顶点','零点','参数','综合'];counts=[20,25,15,20,20]
    independent=[15,16,8,9,6];prompted=[3,6,4,6,7];unfinished=[2,3,3,5,7]
    rates=[(b+c)/n for b,c,n in zip(prompted,unfinished,counts)]
    tx(s,'任务难度与非独立完成率',.55,1.67,6.8,.35,18,True,NAVY,'short')
    b=BubbleChartData()
    for name,x,y,size in zip(cats,[1.3,2.1,3.0,3.8,4.7],rates,counts):b.add_series(name).add_data_point(x,y,size)
    c=s.shapes.add_chart(XL_CHART_TYPE.BUBBLE, Inches(.55), Inches(2.15), Inches(6.5), Inches(3.96),b).chart;style(c)
    c.category_axis.minimum_scale=0;c.category_axis.maximum_scale=6;c.category_axis.major_unit=1
    c.value_axis.minimum_scale=0;c.value_axis.maximum_scale=1;c.value_axis.major_unit=.2
    c.value_axis.tick_labels.number_format='0%';c.value_axis.tick_labels.number_format_is_linked=False
    for ax,label in [(c.category_axis,'任务难度（预设 1–5）'),(c.value_axis,'非独立完成率')]:
     ax.has_title=True;ax.axis_title.text_frame.text=label
     for p in ax.axis_title.text_frame.paragraphs:
      for r in p.runs:r.font.name=FONT;r.font.size=Pt(12)
    for sr in c.series:series_color(sr,BLUE)
    c.plots[0].bubble_scale=55;c.plots[0].has_data_labels=True
    dl=c.plots[0].data_labels;dl.show_series_name=True;dl.show_value=False;dl.show_bubble_size=False;dl.font.size=Pt(12);dl.font.name=FONT;dl.position=XL_LABEL_POSITION.ABOVE
    tx(s,'气泡大小代表尝试次数；非独立完成率 =（提示后完成 + 未完成）÷ 尝试次数。',.55,6.22,6.5,.55,12,False,GRAY)

    tx(s,'各任务完成结构（次）',7.45,1.67,5.33,.35,18,True,NAVY,'short')
    d=CategoryChartData();d.categories=cats
    for name,vals in [('独立完成',independent),('提示后完成',prompted),('未完成',unfinished)]:d.add_series(name,vals)
    c=s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED,Inches(7.45),Inches(2.13),Inches(5.33),Inches(2.28),d).chart;style(c)
    for sr,col in zip(c.series,[BLUE,'94ADD4',RED]):series_color(sr,col)
    c.value_axis.minimum_scale=0;c.value_axis.maximum_scale=25;c.value_axis.major_unit=5
    c.plots[0].has_data_labels=True;dl=c.plots[0].data_labels;dl.show_value=True;dl.font.size=Pt(12);dl.font.name=FONT;dl.position=XL_LABEL_POSITION.CENTER
    tx(s,'整体完成结构',7.45,4.59,5.33,.35,18,True,NAVY,'short')
    d=CategoryChartData();d.categories=['独立完成','提示后完成','未完成'];d.add_series('尝试次数',[54,26,20])
    c=s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT,Inches(7.45),Inches(5.02),Inches(2.25),Inches(1.64),d).chart;style(c);c.plots[0].hole_size=62
    for pt,col in zip(c.series[0].points,[BLUE,'94ADD4',RED]):pt.format.fill.solid();pt.format.fill.fore_color.rgb=rgb(col);pt.format.line.fill.background()
    for i,(t,col) in enumerate([('独立完成   54 次 / 54%',BLUE),('提示后完成 26 次 / 26%','6481AE'),('未完成     20 次 / 20%',RED)]):tx(s,t,9.9,5.12+i*.48,2.88,.34,14,True,col,'short')
    s.notes_slide.notes_text_frame.text='全部数据为构造样本。任务总数20、25、15、20、20；独立完成15、16、8、9、6；提示后完成3、6、4、6、7；未完成2、3、3、5、7。合计100，三类总计54/26/20。难度值人为预设，不代表实测难度。原生气泡图、堆叠条形图和环形图均可编辑。'

    s=P.slides.add_slide(P.slide_layouts[6]);title(s,2,'错题复习流程：条件分支与反馈回路','8 个可编辑节点；“否”分支返回首错定位，“是”分支进入迁移与记录。')
    nodes=[]
    def node(t,body,x,y,diamond=False):
     sh=s.shapes.add_shape(MSO_SHAPE.DIAMOND if diamond else MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.7), Inches(1.1))
     sh.name='node:'+t;sh.fill.solid();sh.fill.fore_color.rgb=rgb(LIGHT);sh.line.color.rgb=rgb(BLUE);sh.line.width=Pt(1)
     for st in list(sh._element.xpath('./p:style')):sh._element.remove(st)
     if diamond:tx(s,t,x+.36,y+.33,1.98,.4,14,True,NAVY,'short',PP_ALIGN.CENTER)
     else:
      tx(s,t,x+.18,y+.16,2.34,.3,16,True,NAVY,'short',PP_ALIGN.CENTER)
      tx(s,body,x+.18,y+.62,2.34,.32,12,False,INK,'body',PP_ALIGN.CENTER)
     nodes.append(sh);return sh
    def arrow(points,color=BLUE):
     for i,((x1,y1),(x2,y2)) in enumerate(zip(points,points[1:])):
      sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2));sh.name='flow:route'
      sh.line.color.rgb=rgb(color);sh.line.width=Pt(1.5)
      if i==len(points)-2:
       el=OxmlElement('a:tailEnd');el.set('type','triangle');sh._element.spPr.get_or_add_ln().append(el)
    xs=[.65,3.55,6.45,9.35]
    for x,t,body in zip(xs,['保留原解','定位首错','分级提示','独立变式'],['题目、步骤与疑问','记录第一处错误','先方向，再中间步骤','收起提示后作答新题']):node(t,body,x,2.05)
    node('三项均正确？','',9.35,4.45,True)
    node('新题迁移','改变系数或开口方向',6.45,4.45)
    node('记录归档','保存过程与提示等级',3.55,4.45)
    node('阶段总结','保留仍需复测的错因',.65,4.45)
    for a,b in zip(xs,xs[1:]):arrow([(a+2.7,2.6),(b,2.6)])
    arrow([(10.7,3.15),(10.7,4.45)])
    arrow([(9.35,5.0),(9.03,5.0),(9.03,4.2),(7.8,4.2),(7.8,4.45)])
    tx(s,'是',8.1,3.79,.5,.3,12,True,BLUE,'short')
    arrow([(6.45,5.0),(6.25,5.0)]);arrow([(3.55,5.0),(3.35,5.0)])
    arrow([(12.05,5),(12.55,5),(12.55,1.9),(4.9,1.9),(4.9,2.05)],RED)
    tx(s,'否：返回首错定位',6.8,1.58,2.7,.3,12,True,RED,'short')
    tx(s,'判断依据',.65,6.12,1.8,.35,16,True,NAVY,'short')
    tx(s,'配方等价、顶点与最值正确、零点代回成立；同时确认该轮没有使用提示。',2.5,6.12,10.2,.45,14)
    s.notes_slide.notes_text_frame.text='教学流程设计示例，不代表验证过的教育干预。流程节点和连线均为原生对象。是分支进入新题迁移、记录归档、阶段总结；否分支回到定位首错。'

    target=Path(sys.argv[1]);target.parent.mkdir(parents=True,exist_ok=True);P.save(target)
    q=Presentation(target)
    assert len(q.slides)==2
    assert sum(sh.has_chart for sh in q.slides[0].shapes)==3
    assert len([sh for sh in q.slides[1].shapes if sh.name.startswith('node:')])==8
    assert all(a+b+c==n for a,b,c,n in zip(independent,prompted,unfinished,counts))
    print(f'Saved {target}; 3 native charts; 8 native nodes; shared data totals verified.')


if __name__ == '__main__':
    main()
