import tempfile
import unittest
import hashlib
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
from connector_integrity import inventory, endpoint_error, verify_connections, structural_role_error, data_role_error


class ConnectionTests(unittest.TestCase):
    def make(self, gap=0, length=20):
        p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540)
        s=p.slides.add_slide(p.slide_layouts[6])
        target=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(200),Pt(100),Pt(100),Pt(60))
        target.fill.solid();target.fill.fore_color.rgb=RGBColor(255,180,40)
        c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(200-gap-length),Pt(130),Pt(200-gap),Pt(130))
        arrow=OxmlElement('a:tailEnd');arrow.set('type','triangle');c._element.spPr.get_or_add_ln().append(arrow)
        return p,s,target,c

    def test_connected_arrow_passes_geometry(self):
        p,s,t,c=self.make();rows=inventory(s)
        self.assertIsNone(endpoint_error(rows[0],1,{'kind':'shape_boundary','target_shape_id':t.shape_id},s,rows,1))

    def test_detached_arrow_fails(self):
        p,s,t,c=self.make(gap=3);rows=inventory(s)
        self.assertIn('detached',endpoint_error(rows[0],1,{'kind':'shape_boundary','target_shape_id':t.shape_id},s,rows,1))

    def test_short_terminal_fails_even_when_coordinates_touch(self):
        p,s,t,c=self.make(length=3.6);rows=inventory(s)
        self.assertIn('too short',endpoint_error(rows[0],1,{'kind':'shape_boundary','target_shape_id':t.shape_id},s,rows,1))

    def test_name_and_decoration_cannot_exempt_arrow(self):
        p,s,t,c=self.make();c.name='background:approved';rows=inventory(s)
        self.assertEqual(len(rows),1)
        self.assertIn('self-exempted',endpoint_error(rows[0],1,{'kind':'open_nonrelational','purpose':'axis'},s,rows,1))

    def test_missing_inventory_blocks(self):
        p,s,t,c=self.make()
        self.assertTrue(verify_connections(p,{},Path('.')))

    def test_invisible_target_rejected(self):
        p,s,t,c=self.make();t.fill.background();rows=inventory(s)
        self.assertIsNotNone(endpoint_error(rows[0],1,{'kind':'shape_boundary','target_shape_id':t.shape_id},s,rows,1))

    def test_group_is_not_silently_omitted(self):
        p,s,t,c=self.make();s.shapes.add_group_shape([c]);rows=inventory(s)
        self.assertEqual(len(rows),1);self.assertFalse(rows[0]['geometry_supported'])

    def test_crop_pixels_and_complete_positive_gate(self):
        p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540)
        s=p.slides.add_slide(p.slide_layouts[6])
        c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(100),Pt(100),Pt(200),Pt(100))
        with tempfile.TemporaryDirectory() as td:
            d=Path(td);im=Image.new('RGB',(960,540),'white');im.save(d/'slide-01.png')
            record={'slide':1,'shape_id':c.shape_id,'image_sha256':hashlib.sha256((d/'slide-01.png').read_bytes()).hexdigest().upper(),
                    'returned_to_full_page':True,'path_observation':'Separator across two columns','semantic_relation':'Section separation','endpoints':[]}
            for i,x in enumerate((100,200)):
                box=[x-10,90,x+10,110];path=d/f'{i}.png';im.crop(box).save(path)
                record['endpoints'].append({'kind':'open_nonrelational','purpose':'separator','crop_box_px':box,'crop_path':str(path),'visible_tip_and_boundary_observation':'Open separator end with clearance'})
            self.assertEqual(verify_connections(p,{'connector_reviews':[record]},d),[])
            Image.new('RGB',(20,20),'red').save(d/'1.png')
            self.assertTrue(any('exact current' in x['reason'] for x in verify_connections(p,{'connector_reviews':[record]},d)))

    def test_block_arrow_is_not_omitted(self):
        p,s,t,c=self.make()
        a=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Pt(30),Pt(50),Pt(30),Pt(20))
        self.assertTrue(next(r for r in inventory(s) if r['shape_id']==a.shape_id)['block_arrow'])

    def test_real_icon_cluster_and_false_role(self):
        p,s,t,c=self.make();marks=[]
        for y in (200,210,220):marks.append(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(40),Pt(y),Pt(70),Pt(y)))
        rows=inventory(s);r=next(r for r in rows if r['shape_id']==marks[0].shape_id)
        proof={'kind':'icon_stroke','cluster_shape_ids':[m.shape_id for m in marks]}
        self.assertIsNone(structural_role_error(r,proof,s,rows,1))
        self.assertIsNotNone(structural_role_error(rows[0],proof,s,rows,1))
        proof['cluster_shape_ids'].pop()
        self.assertIsNotNone(structural_role_error(r,proof,s,rows,1))

    def test_axis_needs_real_ticks_and_labels(self):
        p,s,t,c=self.make();c.begin_x=Pt(50);c.begin_y=Pt(300);c.end_x=Pt(300);c.end_y=Pt(300)
        ticks=[];labels=[]
        for x in (80,180):
            ticks.append(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x),Pt(294),Pt(x),Pt(306)))
            q=s.shapes.add_textbox(Pt(x-10),Pt(310),Pt(20),Pt(20));q.text=str(x);labels.append(q)
        rows=inventory(s);r=rows[0]
        proof={'kind':'arrow_axis','axis_shape_id':c.shape_id,'tick_shape_ids':[t.shape_id for t in ticks],'label_shape_ids':[t.shape_id for t in labels]}
        self.assertIsNone(structural_role_error(r,proof,s,rows,1))
        proof['label_shape_ids']=[]
        self.assertIsNotNone(structural_role_error(r,proof,s,rows,1))
        proof['label_shape_ids']=[t.shape_id for t in labels]
        ticks[0].begin_y=Pt(200);ticks[0].end_y=Pt(210)
        self.assertIsNotNone(structural_role_error(r,proof,s,inventory(s),1))

    def test_data_marker_requires_real_source_value_and_actual_center(self):
        p,s,t,c=self.make();c._element.spPr.get_or_add_ln().remove(c._element.xpath('./p:spPr/a:ln/a:tailEnd')[0])
        marker=s.shapes.add_shape(MSO_SHAPE.OVAL,Pt(196),Pt(126),Pt(8),Pt(8))
        marker.fill.solid();marker.fill.fore_color.rgb=RGBColor(20,40,60)
        with tempfile.TemporaryDirectory() as td:
            src=Path(td)/'data.csv';src.write_text('v\n10\n',encoding='utf8')
            proof={'kind':'data_marker_center','target_shape_id':marker.shape_id,'data_mapping':{'source_path':str(src),'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest().upper(),'row_index':0,'column':'v','axis':'x','offset_pt':100,'scale_pt':10}}
            row=inventory(s)[0]
            self.assertIsNone(data_role_error(row,1,proof,s,1))
            proof['data_mapping']['offset_pt']=102
            self.assertIsNotNone(data_role_error(row,1,proof,s,1))
            proof['data_mapping']['offset_pt']=100
            row['arrows']={'tailEnd':'triangle'}
            self.assertIsNotNone(data_role_error(row,1,proof,s,1))

    def test_reference_line_mapping_and_malformed_source(self):
        p,s,t,c=self.make();c._element.spPr.get_or_add_ln().remove(c._element.xpath('./p:spPr/a:ln/a:tailEnd')[0])
        c.begin_x=Pt(100);c.begin_y=Pt(300);c.end_x=Pt(200);c.end_y=Pt(200)
        with tempfile.TemporaryDirectory() as td:
            data=Path(td)/'data.csv';data.write_text('v\n1\n')
            spec={'equation':'y=x','x_axis':{'offset_pt':100,'scale_pt':10},'y_axis':{'offset_pt':300,'scale_pt':-10},'endpoints_data':[[0,0],[10,10]],'source_data_path':str(data),'source_data_sha256':hashlib.sha256(data.read_bytes()).hexdigest().upper()}
            source=Path(td)/'ref.json'
            def check(spec):
                source.write_text(json.dumps(spec));proof={'kind':'data_reference_line','data_mapping':{'source_path':str(source),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest().upper()}}
                return data_role_error(inventory(s)[0],1,proof,s,1)
            self.assertIsNone(check(spec))
            self.assertIsNotNone(check({**spec,'equation':'y=2*x'}))
            self.assertIsNotNone(check({'extra':3}))
            spec['endpoints_data'][1]=[11,11]
            self.assertIsNotNone(check(spec))


if __name__=='__main__':unittest.main()
