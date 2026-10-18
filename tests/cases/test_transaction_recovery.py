"""Regression tests for update transaction recovery and metadata safety."""
from pathlib import Path
import sys,tempfile,json,io,contextlib,subprocess,unittest
from types import SimpleNamespace
from unittest.mock import patch
sys.dont_write_bytecode=True
sys.path.insert(0,str(Path(__file__).resolve().parent))
import local_update_guard as g
import check_update as cu
import strip_pptx_metadata as metadata
from pptx import Presentation

class ReleaseFailurePaths(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
  self.base=Path(self.temp.name);self.live=self.base/'live';self.candidate=self.base/'candidate';self.stage=self.base/'stage'
  for root,version,value in ((self.live,'1.0.0','before'),(self.candidate,'1.0.1','after')):
   root.mkdir();self.write(root,'SKILL.md','rules');self.write(root,'VERSION',version);self.write(root,'a.txt',value);self.write(root,'b.txt',value)
  self.write(self.live,'manifest.json',{'files':{k:{'sha256':g.digest(p)} for k,p in g.files(self.live).items()}})
  self.write(self.live,'local-overrides/registry.json',{'protected_paths':['deleted-user-rule.md']})
 def write(self,root,name,value):
  p=root/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(value if isinstance(value,str) else json.dumps(value),encoding='utf8')
 def prepare(self):
  with contextlib.redirect_stdout(io.StringIO()):g.prepare(self.live,self.stage,self.candidate)
  self.write(self.stage,'review.json',{'status':'approved_no_conflict','conflicts':[],'reviewed_by':'isolated fixture'})
 def invoke(self):
  with patch.object(sys,'argv',['update','apply','--root',str(self.live),'--stage',str(self.stage)]),contextlib.redirect_stdout(io.StringIO()),contextlib.redirect_stderr(io.StringIO()):return g.main()
 def test_normal_update_preserves_extra_requirement(self):
  self.write(self.live,'extra.md','user requirement');self.write(self.candidate,'extra.md','upstream replacement')
  self.prepare();self.assertEqual(self.invoke(),0);self.assertEqual((self.live/'extra.md').read_text(),'user requirement')
 def test_pending_review_refused_without_file_mutation(self):
  self.prepare();self.write(self.stage,'review.json',{'status':'pending'})
  self.assertEqual(self.invoke(),4);self.assertEqual((self.live/'a.txt').read_text(),'before')
 def test_conflict_without_confirmation_refused(self):
  self.prepare();self.write(self.stage,'review.json',{'status':'approved_no_conflict','conflicts':['fixture'],'reviewed_by':'test'})
  self.assertEqual(self.invoke(),4);self.assertEqual((self.live/'VERSION').read_text(),'1.0.0')
 def test_local_change_after_staging_refused(self):
  self.prepare();self.write(self.live,'new-user-rule.md','new requirement')
  self.assertEqual(self.invoke(),4);self.assertEqual((self.live/'a.txt').read_text(),'before')
 def test_candidate_change_after_staging_refused(self):
  self.prepare();self.write(self.stage/'incoming','a.txt','modified after review')
  self.assertEqual(self.invoke(),4);self.assertEqual((self.live/'a.txt').read_text(),'before')
 def test_ordinary_copy_failure_rolls_back(self):
  self.prepare();original=g.shutil.copy2
  def fail(src,dst,*a,**kw):
   if Path(dst)==self.live/'b.txt':raise OSError('fixture copy failure')
   return original(src,dst,*a,**kw)
  with patch.object(g.shutil,'copy2',fail):code=self.invoke()
  self.assertEqual(code,4);self.assertEqual((self.live/'a.txt').read_text(),'before')
 def test_rollback_failure_is_not_success(self):
  self.prepare();original=g.shutil.copy2
  def fail(src,dst,*a,**kw):
   if Path(dst)==self.live/'b.txt' or (Path(dst)==self.live/'a.txt' and 'rollback' in Path(src).parts):raise OSError('fixture failure')
   return original(src,dst,*a,**kw)
  with patch.object(g.shutil,'copy2',fail):code=self.invoke()
  self.assertEqual(code,4);self.assertEqual(json.loads((self.stage/'result.json').read_text())['status'],'rollback_failed')
 def test_final_result_failure_rolls_back_committed_files(self):
  self.prepare();original=g.os.replace
  def fail(src,dst,*a,**kw):
   if Path(dst)==self.stage/'result.json' and json.loads(Path(src).read_text()).get('status')=='applied':raise OSError('fixture receipt failure')
   return original(src,dst,*a,**kw)
  with patch.object(g.os,'replace',fail):code=self.invoke()
  self.assertEqual(code,4);self.assertEqual((self.live/'a.txt').read_text(),'before','final receipt failure left live files changed')
 def test_registry_partial_write_keeps_original_valid_registry(self):
  registry=self.live/'local-overrides/registry.json';before=registry.read_bytes()
  with patch.object(g.os,'fsync',side_effect=OSError('fixture interrupted temp write')),self.assertRaises(OSError):self.prepare()
  self.assertEqual(registry.read_bytes(),before,'live protection registry was truncated')
 def test_ctrl_c_during_apply_restores_preupdate_files(self):
  self.prepare();original=g.shutil.copy2
  def interrupt(src,dst,*a,**kw):
   if Path(dst)==self.live/'b.txt':raise KeyboardInterrupt('fixture Ctrl+C')
   return original(src,dst,*a,**kw)
  with patch.object(g.shutil,'copy2',interrupt):
   try:self.invoke()
   except KeyboardInterrupt:pass
  self.assertEqual((self.live/'a.txt').read_text(),'before','Ctrl+C left a partial update')
 def test_update_check_timeout_has_controlled_failure(self):
  self.write(self.live,'updates.json',{'update_source':'git+https://example.invalid/repo#main'})
  exception=None;code=None
  with patch.object(sys,'argv',['check','--skill-root',str(self.live)]),patch.object(cu,'fetch_remote_updates',side_effect=subprocess.TimeoutExpired('git clone',30)),contextlib.redirect_stdout(io.StringIO()),contextlib.redirect_stderr(io.StringIO()):
   try:code=cu.main()
   except Exception as exc:exception=type(exc).__name__
  self.assertIsNone(exception,'timeout escaped the controlled update-check path');self.assertNotEqual(code,0)
 def test_update_check_network_error_is_nonzero(self):
  self.write(self.live,'updates.json',{'update_source':'git+https://example.invalid/repo#main'})
  with patch.object(sys,'argv',['check','--skill-root',str(self.live)]),patch.object(cu,'fetch_remote_updates',side_effect=OSError('offline')),contextlib.redirect_stderr(io.StringIO()):code=cu.main()
  self.assertEqual(code,3)
 def test_metadata_cleanup_does_not_overwrite_external_edit(self):
  p=self.base/'candidate.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.core_properties.title='before';prs.save(p)
  original=metadata.transform
  def interleave(data,author=None):
   external=Presentation(p);external.core_properties.title='external edit';external.save(p)
   return original(data,author)
  with patch.object(metadata,'transform',interleave):
   try:metadata.strip_metadata(p)
   except Exception:pass
  self.assertEqual(Presentation(p).core_properties.title,'external edit','concurrent edit was overwritten')
  self.assertFalse(list(self.base.glob('*.tmp')))
 def test_metadata_normal_cleanup_preserves_other_properties(self):
  p=self.base/'candidate.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.core_properties.title='retained title';prs.save(p)
  metadata.strip_metadata(p,'Known Author')
  self.assertEqual(metadata.findings(p.read_bytes(),'Known Author'),[])
  self.assertEqual(Presentation(p).core_properties.title,'retained title')
  self.assertFalse(list(self.base.glob('*.tmp')))
 def test_git_check_is_bounded_and_noninteractive(self):
  with patch.object(cu.subprocess,'run',return_value=SimpleNamespace(stdout=b'{}')) as run:
   cu.fetch_remote_updates('git+https://example.invalid/repo#main')
  self.assertEqual(run.call_count,2)
  for call in run.call_args_list:
   self.assertGreater(call.kwargs['timeout'],0)
   self.assertEqual(call.kwargs['env']['GIT_TERMINAL_PROMPT'],'0')
 def test_stage_clone_timeout_controlled_and_preserves_live_files(self):
  self.write(self.live,'updates.json',{'update_source':'git+https://example.invalid/repo#main'})
  before={k:g.digest(p) for k,p in g.files(self.live).items()}
  with patch.object(sys,'argv',['update','prepare','--root',str(self.live),'--stage',str(self.stage)]),patch.object(g.subprocess,'run',side_effect=subprocess.TimeoutExpired('git clone',180)) as run,contextlib.redirect_stderr(io.StringIO()):code=g.main()
  self.assertEqual(code,4);self.assertEqual(before,{k:g.digest(p) for k,p in g.files(self.live).items()})
  self.assertGreater(run.call_args.kwargs['timeout'],0)
  self.assertEqual(run.call_args.kwargs['env']['GIT_TERMINAL_PROMPT'],'0')
 def test_registry_replace_failure_preserves_original_and_cleans_temp(self):
  registry=self.live/'local-overrides/registry.json';before=registry.read_bytes();original=g.os.replace
  def fail(src,dst):
   if Path(dst)==registry:raise OSError('fixture replace failure')
   return original(src,dst)
  with patch.object(g.os,'replace',fail),self.assertRaises(OSError):self.prepare()
  self.assertEqual(registry.read_bytes(),before);self.assertFalse(list(registry.parent.glob('*.tmp')))
 def test_unwritable_failure_receipt_still_reports_rollback(self):
  self.prepare();original=g.atomic_json
  def fail(path,value):
   if Path(path)==self.stage/'result.json' and value.get('status')!='applying_not_complete':raise OSError('fixture status disk failure')
   return original(path,value)
  output=io.StringIO();error=io.StringIO()
  with patch.object(g,'atomic_json',fail),patch.object(sys,'argv',['update','apply','--root',str(self.live),'--stage',str(self.stage)]),contextlib.redirect_stdout(output),contextlib.redirect_stderr(error):code=g.main()
  self.assertEqual(code,4);self.assertEqual((self.live/'a.txt').read_text(),'before')
  self.assertNotIn('APPLIED.',output.getvalue());self.assertIn('failed_rolled_back',error.getvalue())

if __name__=='__main__':unittest.main()
