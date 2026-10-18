import unittest
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from interstitial_spacing import validate_interstitial_spacing
results=[]
def make(above=25.2,below=25.2):
 p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
 def rect(x,y,w,h):
  a=s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));a.fill.solid();a.fill.fore_color.rgb=RGBColor(30,60,90);a.line.fill.background();return a
 a=rect(100,100,95,60);b=rect(205,100,95,60)
 y=160+above;m=s.shapes.add_textbox(Pt(100),Pt(y),Pt(200),Pt(20));m.text='Actual explanation'
 z=rect(100,y+20+below,200,40)
 f={'slide':1,'reason':'header_body_boundary_review'};d={'spacing_contract':{'kind':'interstitial_text','middle_shape_id':m.shape_id,'before_shape_ids':[a.shape_id,b.shape_id],'after_shape_id':z.shape_id,'hierarchy_observation':'The explanation belongs between the two panel containers and conclusion; header whitespace belongs to a separate page hierarchy.'}}
 r={'text_records':[dict(slide=1,shape_id=m.shape_id,kind='text',bound_left_pt=100,bound_top_pt=y,bound_width_pt=200,bound_height_pt=20)]}
 return p,s,rect,a,b,m,z,f,d,r
class Contracts(unittest.TestCase):
 def check(self,name,expected,mutate=lambda v:None,above=25.2,below=25.2):
  v=make(above,below);mutate(v);p,s,_,a,b,m,z,f,d,r=v;error=validate_interstitial_spacing(f,d,s,r,1);results.append(dict(case=name,expected_accept=expected,error=error));self.assertEqual(error is None,expected,error)
 def test_balanced(self):self.check('balanced_25_2',True)
 def test_unbalanced(self):self.check('unbalanced_14_2_36',False,above=14.2,below=36)
 def test_omit_row(self):self.check('omit_before',False,lambda v:v[8]['spacing_contract'].update(before_shape_ids=[v[3].shape_id]))
 def test_far_before(self):self.check('far_before',False,lambda v:v[8]['spacing_contract'].update(before_shape_ids=[v[2](100,30,200,40).shape_id]))
 def test_far_after(self):self.check('far_after',False,lambda v:v[8]['spacing_contract'].update(after_shape_id=v[2](100,400,200,40).shape_id))
 def test_transparent(self):
  def change(v):
   color=v[3]._element.xpath('./p:spPr/a:solidFill/*')[0];a=OxmlElement('a:alpha');a.set('val','0');color.append(a)
  self.check('transparent_before',False,change)
 def test_white_proxy(self):
  def change(v):v[3].fill.fore_color.rgb=RGBColor(255,255,255)
  self.check('white_proxy',False,change)
 def test_rotated(self):
  def change(v):v[3].rotation=20
  self.check('rotated_container',False,change)
 def test_nofill(self):self.check('nofill_before',False,lambda v:v[3].fill.background())
 def test_extra_body(self):self.check('omitted_body',False,lambda v:v[2](130,175,20,5))
 def test_fake_role(self):
  def change(v):v[2](130,175,20,5).name='background:proxy'
  self.check('wrong_role_body',False,change)
 def test_self_reported(self):
  def change(v):v[8]['measurements']={'above_visual_gap_pt':25.2,'below_visual_gap_pt':25.2}
  self.check('lying_numbers',False,change,above=14.2,below=36)
 def test_small(self):self.check('small_balanced',False,above=8,below=8)
 def test_missing_bounds(self):self.check('missing_bounds',False,lambda v:v[9]['text_records'][0].pop('bound_top_pt'))
 def test_no_observation(self):self.check('no_hierarchy_observation',False,lambda v:v[8]['spacing_contract'].pop('hierarchy_observation'))
 def test_duplicate(self):self.check('duplicate_before',False,lambda v:v[8]['spacing_contract']['before_shape_ids'].append(v[3].shape_id))
 def test_table(self):
  def change(v):
   tb=v[1].shapes.add_table(2,2,Pt(100),Pt(100),Pt(200),Pt(60));v[8]['spacing_contract']['before_shape_ids']=[tb.shape_id]
  self.check('unverified_table_contour',False,change)
class FullRowContracts(unittest.TestCase):
 def check_row(self,name,expected,variant=None):
  p,sl,rect,a,b,m,z,f,d,r=make()
  a.left=Pt(190);a.width=Pt(80);b.left=Pt(280);b.width=Pt(80)
  third=rect(370,100,79,60)
  label=sl.shapes.add_textbox(Pt(49),Pt(125),Pt(115),Pt(30));label.text='Row label'
  m.left=Pt(49);m.width=Pt(400);z.left=Pt(49);z.width=Pt(400)
  r['text_records'][0].update(bound_left_pt=49,bound_width_pt=400)
  lr=dict(slide=1,shape_id=label.shape_id,kind='text',bound_left_pt=49,bound_top_pt=130,bound_width_pt=100,bound_height_pt=20)
  r['text_records'].append(lr)
  d['spacing_contract']['before_shape_ids']=[a.shape_id,b.shape_id,third.shape_id,label.shape_id]
  if variant=='omit_label':d['spacing_contract']['before_shape_ids'].remove(label.shape_id)
  if variant=='omit_card':d['spacing_contract']['before_shape_ids'].remove(b.shape_id)
  if variant=='blank_proxy':label.text='';r['text_records'].remove(lr)
  if variant=='missing_record':r['text_records'].remove(lr)
  if variant=='far_label':label.top=Pt(30);lr['bound_top_pt']=35
  if variant=='wide_empty_frame':lr['bound_left_pt']=80;lr['bound_width_pt']=69
  if variant=='inner_text':
   inner=sl.shapes.add_textbox(Pt(200),Pt(110),Pt(50),Pt(30));inner.text='Card text'
   r['text_records'].append(dict(slide=1,shape_id=inner.shape_id,kind='text',bound_left_pt=200,bound_top_pt=115,bound_width_pt=45,bound_height_pt=20))
  err=validate_interstitial_spacing(f,d,sl,r,1);results.append(dict(case=name,expected_accept=expected,error=err));self.assertEqual(err is None,expected,err)
 def test_full_row(self):self.check_row('full_row_label_three_cards',True)
 def test_omit_label(self):self.check_row('full_row_omit_label',False,'omit_label')
 def test_omit_card(self):self.check_row('full_row_omit_card',False,'omit_card')
 def test_blank_proxy(self):self.check_row('full_row_blank_proxy',False,'blank_proxy')
 def test_missing_record(self):self.check_row('full_row_missing_glyph_record',False,'missing_record')
 def test_far_label(self):self.check_row('full_row_far_label',False,'far_label')
 def test_frame_padding(self):self.check_row('full_row_empty_frame_padding',False,'wide_empty_frame')
 def test_internal_card_text(self):self.check_row('full_row_internal_card_text',True,'inner_text')

class BackgroundContracts(unittest.TestCase):
 def fixture(self,color=(240,244,248),bottom=True,above=25.2,below=25.2):
  v=make(above,below);p,sl=v[:2]
  bg=sl.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,0,0,p.slide_width,p.slide_height)
  bg.fill.solid();bg.fill.fore_color.rgb=RGBColor(*color);bg.line.fill.background()
  # Deliberately unrelated to any reserved role or name.
  bg.name='Plain rectangle'
  if bottom:
   sl.shapes._spTree.remove(bg._element);sl.shapes._spTree.insert(2,bg._element)
  return v,bg
 def validate(self,v):
  return validate_interstitial_spacing(v[7],v[8],v[1],v[9],1)
 def set_alpha(self,bg,value,tag='alpha'):
  alpha=OxmlElement('a:'+tag);alpha.set('val',str(value));bg._element.xpath('./p:spPr/a:solidFill/*')[0].append(alpha)
 def test_white_full_slide_background(self):
  v,bg=self.fixture((255,255,255));self.assertIsNone(self.validate(v))
 def test_tinted_full_slide_background(self):
  v,bg=self.fixture();self.assertIsNone(self.validate(v))
 def test_background_on_custom_slide_extent(self):
  v,bg=self.fixture();p=v[0];p.slide_width=Pt(960);p.slide_height=Pt(540);bg.width=p.slide_width;bg.height=p.slide_height
  self.assertIsNone(self.validate(v))
 def test_full_bleed_background(self):
  v,bg=self.fixture();bg.left=Pt(-10);bg.top=Pt(-10);bg.width=v[0].slide_width+Pt(20);bg.height=v[0].slide_height+Pt(20)
  self.assertIsNone(self.validate(v))
 def test_explicitly_opaque_background(self):
  v,bg=self.fixture();self.set_alpha(bg,100000);self.assertIsNone(self.validate(v))
 def test_stacked_backgrounds_behind_all_content(self):
  v,bg=self.fixture();sl=v[1]
  second=sl.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,0,0,v[0].slide_width,v[0].slide_height)
  second.fill.solid();second.fill.fore_color.rgb=RGBColor(255,255,255);second.line.fill.background()
  sl.shapes._spTree.remove(second._element);sl.shapes._spTree.insert(3,second._element)
  self.assertIsNone(self.validate(v))
 def test_foreground_covers_still_block(self):
  for color in ((255,255,255),(240,244,248)):
   with self.subTest(color=color):
    v,bg=self.fixture(color,bottom=False);bg.name='background:genuine-full-slide'
    self.assertIsNotNone(self.validate(v))
 def test_cover_between_relation_members_still_blocks(self):
  v,bg=self.fixture();tree=v[1].shapes._spTree;tree.remove(bg._element);v[3]._element.addnext(bg._element)
  self.assertIsNotNone(self.validate(v))
 def test_cover_above_unlisted_label_still_blocks(self):
  v,bg=self.fixture();sl=v[1]
  label=sl.shapes.add_textbox(Pt(100),Pt(175),Pt(100),Pt(8));label.text='Intervening label'
  sl.shapes._spTree.remove(label._element);sl.shapes._spTree.insert(2,label._element)
  self.assertIsNotNone(self.validate(v))
 def test_local_background_named_rectangles_still_block(self):
  for color in ((255,255,255),(240,244,248)):
   with self.subTest(color=color):
    v,bg=self.fixture(color);bg.left=Pt(130);bg.top=Pt(175);bg.width=Pt(20);bg.height=Pt(5);bg.name='background:proxy'
    self.assertIsNotNone(self.validate(v))
 def test_almost_full_slide_rectangle_still_blocks(self):
  v,bg=self.fixture();bg.width=v[0].slide_width-1
  self.assertIsNotNone(self.validate(v))
 def test_transparent_partial_or_relative_alpha_still_blocks(self):
  for tag,value in (('alpha',0),('alpha',50000),('alphaMod',100000),('alphaOff',0)):
   with self.subTest(tag=tag,value=value):
    v,bg=self.fixture();self.set_alpha(bg,value,tag);self.assertIsNotNone(self.validate(v))
 def test_missing_fill_still_blocks(self):
  v,bg=self.fixture();bg.fill.background();self.assertIsNotNone(self.validate(v))
 def test_theme_fill_still_blocks(self):
  v,bg=self.fixture();fill=bg._element.xpath('./p:spPr/a:solidFill')[0];fill.remove(fill[0]);color=OxmlElement('a:schemeClr');color.set('val','accent1');fill.append(color)
  self.assertIsNotNone(self.validate(v))
 def test_line_still_blocks(self):
  v,bg=self.fixture();bg.line.color.rgb=RGBColor(0,0,0);self.assertIsNotNone(self.validate(v))
 def test_rotation_still_blocks(self):
  v,bg=self.fixture();bg.rotation=5;self.assertIsNotNone(self.validate(v))
 def test_background_text_still_blocks(self):
  v,bg=self.fixture();bg.text='This is actual content';self.assertIsNotNone(self.validate(v))
 def test_oval_still_blocks(self):
  v,bg=self.fixture();bg._element.xpath('./p:spPr/a:prstGeom')[0].set('prst','ellipse');self.assertIsNotNone(self.validate(v))
 def test_hidden_rectangle_is_not_proven_background(self):
  for hidden in ('1','true'):
   with self.subTest(hidden=hidden):
    v,bg=self.fixture();bg._element.xpath('./p:nvSpPr/p:cNvPr')[0].set('hidden',hidden);self.assertIsNotNone(self.validate(v))
 def test_background_does_not_bypass_gap_validation(self):
  v,bg=self.fixture(above=14.2,below=36);self.assertIn('unbalanced',self.validate(v))

if __name__=='__main__':
 unittest.main(verbosity=2)
