"""Conservative component candidates; actual local/full-page review closes them."""
import math,re,statistics
from pptx.enum.dml import MSO_COLOR_TYPE,MSO_FILL_TYPE
from pptx.enum.text import PP_ALIGN

REASONS={'table_alignment_consistency_review','table_alignment_measurement_required',
         'chart_legend_clearance_review','chart_legend_geometry_required','chart_label_geometry_required',
         'interstitial_module_spacing_review','filled_text_color_review'}

def numeric(text):
    return bool(re.fullmatch(r'[+−–-]?\d[\d,.]*(?:\s*(?:%|‰|mm|cm|m|s|ms|px|kWh|W|°))?',text.strip()))

def box(shape):return [shape.left/12700,shape.top/12700,(shape.left+shape.width)/12700,(shape.top+shape.height)/12700]
def finite_box(b):return isinstance(b,(list,tuple)) and len(b)==4 and all(type(x) in (int,float) and math.isfinite(x) for x in b) and b[2]>b[0] and b[3]>b[1]
def rgb(color):
    # Evidence must compare identically before and after JSON serialization.
    try:return list(color.rgb) if color.type==MSO_COLOR_TYPE.RGB else None
    except (ValueError,AttributeError):return None
def contrast(a,b):
    def lum(c):
        v=[x/255 for x in c];v=[x/12.92 if x<=.04045 else ((x+.055)/1.055)**2.4 for x in v]
        return sum(x*w for x,w in zip(v,(.2126,.7152,.0722)))
    x,y=sorted((lum(a),lum(b)));return (y+.05)/(x+.05)

def scan(slide,n,render,scale=1):
    out=[]
    def add(reason,**kw):out.append({'reason':reason,**kw})
    records={r.get('audit_path'):r for r in render.get('text_records',[]) if r.get('slide')==n}
    def visit(shapes,parent='',owner=None):
        for s in shapes:
            path=(parent+'/group/' if parent else '')+str(s.shape_id);sid=owner or s.shape_id
            if hasattr(s,'shapes'):visit(s.shapes,path,sid)
            if s.has_table:
                rows=[];unknown=[]
                for ri,row in enumerate(s.table.rows,1):
                    cells=[]
                    for ci,cell in enumerate(row.cells,1):
                        if cell.is_spanned or not cell.text.strip():continue
                        cp=f'{path}/table/{ri}/{ci}';rec=records.get(cp,{})
                        aligns={p.alignment for p in cell.text_frame.paragraphs if p.text.strip()}
                        if None in aligns:
                            effective=rec.get('paragraph_alignment')
                            aligns.discard(None)
                            if effective in (1,2,3):aligns.add({1:PP_ALIGN.LEFT,2:PP_ALIGN.CENTER,3:PP_ALIGN.RIGHT}[effective])
                            else:unknown.append(cp)
                        cells.append({'audit_path':cp,'row':ri,'column':ci,'role':'number' if numeric(cell.text) else 'text',
                                      'alignment':sorted(int(a) for a in aligns),'vertical_anchor':str(cell.vertical_anchor),
                                      'margins_pt':[cell.margin_left/12700,cell.margin_right/12700,cell.margin_top/12700,cell.margin_bottom/12700]})
                    rows.append(cells)
                mismatches=[]
                # Compare equivalent roles within a row, and each body column.
                groups=[g for cells in rows for role in ('text','number') if len(g:=[c for c in cells if c['role']==role])>1]
                groups += [[c for row in rows[1:] for c in row if c['column']==ci] for ci in range(1,len(s.table.columns)+1)]
                for group in groups:
                    if len(group)>1 and (len({tuple(c['alignment']) for c in group})>1 or len({c['vertical_anchor'] for c in group})>1 or any(max(c['margins_pt'][k] for c in group)-min(c['margins_pt'][k] for c in group)>2*scale for k in range(4))):
                        mismatches.append(group)
                if mismatches:add('table_alignment_consistency_review',shape_ids=[sid],audit_path=path,cell_groups=mismatches)
                if unknown:add('table_alignment_measurement_required',shape_ids=[sid],audit_path=path,audit_paths=unknown)
            if s.has_chart:
                try:has_legend=s.chart.has_legend
                except (ValueError,AttributeError):has_legend=True
                try:
                    label_expected=any(plot.has_data_labels for plot in s.chart.plots)
                    expected_count=sum(len(series.points) for series in s.chart.series) if label_expected else 0
                except (ValueError,AttributeError,TypeError):label_expected=True;expected_count=None
                label_records=[x for x in render.get('chart_label_records',[]) if x.get('slide')==n and x.get('audit_path')==path]
                if label_expected and (not label_records or any(not finite_box(x.get('bounds_pt')) for x in label_records) or (expected_count is not None and len(label_records)!=expected_count)):
                    add('chart_label_geometry_required',shape_ids=[sid],audit_path=path,expected_count=expected_count,measured_count=len(label_records))
                if has_legend:
                    geom=next((x for x in render.get('chart_legend_records',[]) if x.get('slide')==n and x.get('audit_path')==path),None)
                    if not geom or not finite_box(geom.get('bounds_pt')):
                        add('chart_legend_geometry_required',shape_ids=[sid],audit_path=path)
                    else:
                        lb=geom['bounds_pt'];near=[]
                        for label in render.get('chart_label_records',[]):
                            if label.get('slide')!=n or label.get('audit_path')!=path:continue
                            b=label.get('bounds_pt')
                            if not finite_box(b):continue
                            dx=max(lb[0]-b[2],b[0]-lb[2],0);dy=max(lb[1]-b[3],b[1]-lb[3],0)
                            if math.hypot(dx,dy)<4*scale:near.append({'label':label,'distance_pt':math.hypot(dx,dy)})
                        if near:add('chart_legend_clearance_review',shape_ids=[sid],audit_path=path,legend=geom,near_labels=near,threshold_pt=4*scale)
            if s.has_text_frame and s.text.strip():
                try:bg=rgb(s.fill.fore_color) if s.fill.type==MSO_FILL_TYPE.SOLID else None
                except (AttributeError,ValueError):bg=None
                if bg:
                    runs=[r for p in s.text_frame.paragraphs for r in p.runs if r.text.strip()]
                    colors=[rgb(r.font.color) for r in runs];review=[]
                    for fg in colors:
                        if fg is None:review.append({'foreground':'inherited_or_theme','background':bg})
                        elif contrast(fg,bg)<4.5 or max(fg)-min(fg)>32:
                            review.append({'foreground':fg,'background':bg,'contrast':round(contrast(fg,bg),3),'colored_foreground':max(fg)-min(fg)>32})
                    if review:add('filled_text_color_review',shape_ids=[sid],audit_path=path,color_pairs=review)
    visit(slide.shapes)
    # Independent table -> unboxed explanation -> lower module candidates.
    shapes=list(slide.shapes)
    for table in [s for s in shapes if s.has_table]:
        tb=box(table)
        for mid in shapes:
            if not mid.has_text_frame or not mid.text.strip() or mid.shape_id==table.shape_id:continue
            rec=records.get(str(mid.shape_id),{})
            if all(type(rec.get(k)) in (int,float) and math.isfinite(rec[k]) for k in ('bound_left_pt','bound_top_pt','bound_width_pt','bound_height_pt')):
                mb=[rec['bound_left_pt'],rec['bound_top_pt'],rec['bound_left_pt']+rec['bound_width_pt'],rec['bound_top_pt']+rec['bound_height_pt']]
            else:mb=box(mid)
            overlap=min(tb[2],mb[2])-max(tb[0],mb[0])
            if overlap<.5*min(tb[2]-tb[0],mb[2]-mb[0]) or not 0<=mb[1]-tb[3]<=100*scale:continue
            after=[]
            for s in shapes:
                if s.shape_id in (table.shape_id,mid.shape_id) or s.name.startswith('page-number:'):continue
                b=box(s)
                if b[2]<=b[0] or b[3]<=b[1]:continue
                if 0<=b[1]-mb[3]<=100*scale and min(b[2],mb[2])-max(b[0],mb[0])>0:after.append((b[1],s))
            if not after:continue
            y=min(x[0] for x in after);lower=[s for yy,s in after if abs(yy-y)<2*scale]
            a,b=mb[1]-tb[3],y-mb[3]
            if min(a,b)<12*scale or abs(a-b)>max(3*scale,(a+b)/2*.2):
                add('interstitial_module_spacing_review',shape_ids=[table.shape_id,mid.shape_id]+[s.shape_id for s in lower],
                    gaps_pt=[a,b],text_bounds_pt=mb,minimum_pt=12*scale,balance_tolerance_pt=max(3*scale,(a+b)/2*.2))
    return out
