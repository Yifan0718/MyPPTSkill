"""Isolated control-flow checks for release test entrypoints; no Office use."""
from pathlib import Path
import unittest,tempfile,json,sys,io,contextlib,types
from unittest.mock import patch
from pptx import Presentation
import run_integration as integration
import run_tests as runner

class Entrypoints(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.base=Path(self.temp.name)
  self.root=self.base/'skill';(self.root/'scripts').mkdir(parents=True)
  for name in ('render_powerpoint_audit.ps1','verify_coverage_review.py'):(self.root/'scripts'/name).write_text('fixture')
  (self.root/'SKILL.md').write_text('fixture')
  self.src=self.base/'input.pptx';p=Presentation();p.slides.add_slide(p.slide_layouts[6]);p.save(self.src)
 def integration_case(self,exit_code,status,issues):
  output=self.base/'result'
  def fake(command,**kwargs):
   if any(str(x).endswith('render_powerpoint_audit.ps1') for x in command):
    directory=Path(command[command.index('-OutputDir')+1]);directory.mkdir()
    (directory/'audit.json').write_text(json.dumps({'status':status,'slide_count':1,'issues':issues,'manual_review':[]}))
    return types.SimpleNamespace(returncode=exit_code,stdout='',stderr='fixture error' if exit_code else '')
   return types.SimpleNamespace(returncode=0,stdout='',stderr='')
  shell=types.ModuleType('pptx_delivery_shell');shell.verify_shell=lambda *a:[]
  evidence=types.ModuleType('render_evidence');evidence.validate_render_evidence=lambda *a,**kw:None
  argv=['integration','--skill-root',str(self.root),'--input',str(self.src),'--output',str(output)]
  with patch.object(sys,'argv',argv),patch.object(integration.shutil,'which',return_value='powershell'),patch.object(integration.subprocess,'run',side_effect=fake),patch.dict(sys.modules,{'pptx_delivery_shell':shell,'render_evidence':evidence}),contextlib.redirect_stdout(io.StringIO()):code=integration.main()
  return code,json.loads((output/'integration-summary.json').read_text())
 def test_nonzero_execution_cannot_be_hidden_by_complete_report(self):
  code,r=self.integration_case(1,'rendered_pending_visual_review',[]);self.assertEqual(code,1);self.assertEqual(r['status'],'failed')
 def test_preflight_issues_remain_pending_review(self):
  code,r=self.integration_case(1,'failed_preflight',[{'kind':'fixture layout issue'}]);self.assertEqual(code,2);self.assertEqual(r['status'],'render_complete_rule_review_required')
 def test_pipeline_success_does_not_claim_visual_signoff(self):
  code,r=self.integration_case(0,'rendered_pending_visual_review',[]);self.assertEqual(code,0);self.assertFalse(r['visual_signoff']);self.assertIn('visual_review_pending',r['status'])
 def test_missing_integration_input_rejected(self):
  with patch.object(sys,'argv',['integration','--input',str(self.base/'missing.pptx'),'--output',str(self.base/'out')]),contextlib.redirect_stderr(io.StringIO()):self.assertEqual(integration.main(),2)
 def test_missing_skill_rejected(self):
  with patch.object(sys,'argv',['runner','--skill-root',str(self.base/'missing'),'--output',str(self.base/'out')]),contextlib.redirect_stderr(io.StringIO()):self.assertEqual(runner.main(),2)
 def test_outputs_inside_installation_rejected(self):
  with patch.object(sys,'argv',['runner','--skill-root',str(self.root),'--output',str(self.root/'out')]),contextlib.redirect_stderr(io.StringIO()):self.assertEqual(runner.main(),2)

if __name__=='__main__':unittest.main()
