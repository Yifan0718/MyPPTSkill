"""Measure peer gaps declared in a layout plan; visual signoff remains required.
Usage: python audit_layout_gaps.py file.pptx layout.json report.json [--slides 2]
Spec: {pages:[{slide:2, comparisons:[{name:'rows',axis:'vertical',
blocks:[{name:'top',shape_ids:[6,12]},{name:'middle',shape_ids:[39]},
{name:'conclusion',shape_ids:[67]}],target_gap_pt:18}]}]}
All pages need both axes checked, or an explicit manual_axes reason.
"""
from pathlib import Path
import argparse,hashlib,json,math
from pptx import Presentation
def bounds(block,shapes):
    points=[]
    for sid in block['shape_ids']:
        if sid not in shapes:raise ValueError(f'Missing top-level shape id {sid}')
        s=shapes[sid];points.append((s.left/12700,s.top/12700,(s.left+s.width)/12700,(s.top+s.height)/12700))
    if not points:raise ValueError('Empty block')
    return [min(x[0] for x in points),min(x[1] for x in points),max(x[2] for x in points),max(x[3] for x in points)]
def audit(pptx,spec,only=None):
    digest=lambda:hashlib.sha256(pptx.read_bytes()).hexdigest().upper()
    before=digest();p=Presentation(pptx);scale=p.slide_height/12700/540
    page_numbers=[x['slide'] for x in spec['pages']]
    selected=list(range(1,len(p.slides)+1)) if only is None else only
    for numbers in (page_numbers,selected):
        if not numbers or len(set(numbers))!=len(numbers) or any(type(n) is not int or not 1<=n<=len(p.slides) for n in numbers):
            raise ValueError('Page lists must contain unique valid positive slide numbers')
    pages={x['slide']:x for x in spec['pages']};issues=[];rows=[];manual=[]
    for n in selected:
        page=pages.get(n)
        if not page:issues.append({'slide':n,'reason':'missing_layout_plan'});continue
        shapes={s.shape_id:s for s in p.slides[n-1].shapes};covered=set()
        for c in page.get('comparisons',[]):
            axis=c['axis'];covered.add(axis)
            if axis not in ('horizontal','vertical'):raise ValueError('Invalid axis')
            for block in c['blocks']:
                ids=block['shape_ids']
                if len(ids)!=len(set(ids)):
                    raise ValueError('Duplicate shape ids inside a block')
                if len(ids)>1:
                    parts=[bounds({'shape_ids':[sid]},shapes) for sid in ids]
                    # A union box alone does not establish a component. A real
                    # containing object is provisional evidence, still visually reviewed.
                    containers=[sid for sid,box in zip(ids,parts) if all(box[0]<=b[0]+.5 and box[1]<=b[1]+.5 and box[2]>=b[2]-.5 and box[3]>=b[3]-.5 for b in parts)]
                    if not containers:
                        issues.append({'slide':n,'name':c['name'],'block':block.get('name'),'reason':'unverified_multi_shape_block','shape_ids':ids})
                    else:
                        manual.append({'slide':n,'block':block.get('name'),'reason':'verify_real_shared_container_and_internal_spacing','container_ids':containers,'shape_ids':ids})
            boxes=[bounds(b,shapes) for b in c['blocks']]
            if len(boxes)<2:raise ValueError('Need at least 2 blocks')
            lo,hi=(0,2) if axis=='horizontal' else (1,3)
            gaps=[boxes[i+1][lo]-boxes[i][hi] for i in range(len(boxes)-1)]
            target=float(c.get('target_gap_pt',sum(gaps)/len(gaps)/scale))*scale
            minimum=float(c.get('minimum_gap_pt',12))*scale;tolerance=max(3*scale,abs(target)*.2)
            if not math.isfinite(target) or not math.isfinite(minimum) or target<=0 or minimum<=0:
                raise ValueError('Gap target and minimum must be finite and positive')
            if minimum<12*scale and not c.get('compact_group_reason'):raise ValueError('A compact group below 12 pt needs compact_group_reason')
            rows.append({'slide':n,'name':c['name'],'axis':axis,'blocks':c['blocks'],'bounds_pt':boxes,'gaps_pt':gaps,'target_gap_pt':target,'tolerance_pt':tolerance,'minimum_gap_pt':minimum})
            if any(g<minimum-.25*scale for g in gaps):issues.append({'slide':n,'name':c['name'],'reason':'gap_below_minimum_or_overlap','gaps_pt':gaps})
            if any(abs(g-target)>tolerance for g in gaps):issues.append({'slide':n,'name':c['name'],'reason':'peer_gap_imbalance','gaps_pt':gaps,'target_gap_pt':target})
        for axis in ('horizontal','vertical'):
            if axis not in covered:
                reason=page.get('manual_axes',{}).get(axis)
                if isinstance(reason,str) and reason.strip() and reason.strip().lower() not in ('pass','ok','n/a','通过','无'):
                    manual.append({'slide':n,'axis':axis,'reason':reason})
                else:issues.append({'slide':n,'axis':axis,'reason':'uncovered_axis'})
    after=digest()
    if before!=after:issues.append({'reason':'input_changed_during_audit'})
    return {'sha256':before,'source_hash_after':after,'checked_slides':selected,'slide_count':len(p.slides),'full_deck':set(selected)==set(range(1,len(p.slides)+1)),'issues':issues,'comparisons':rows,'manual_review':manual,'status':'failed' if issues else 'geometry_pass_pending_visual_review'}
def main():
    ap=argparse.ArgumentParser();ap.add_argument('pptx',type=Path);ap.add_argument('spec',type=Path);ap.add_argument('report',type=Path);ap.add_argument('--slides',type=int,nargs='+');a=ap.parse_args()
    try:
        result=audit(a.pptx,json.loads(a.spec.read_text(encoding='utf8')),a.slides)
    except Exception as exc:
        result={'status':'failed','issues':[{'reason':'audit_exception','detail':str(exc)}]}
    a.report.parent.mkdir(parents=True,exist_ok=True);a.report.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(result['status'],len(result['issues']));return 2 if result['issues'] else 0
if __name__=='__main__':raise SystemExit(main())
