import contextlib,io,json,os,subprocess,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from pptx import Presentation
from pptx.util import Pt
from report_io import emit_report,write_report

class ReportIOTests(unittest.TestCase):
 def setUp(self):
  self.t=tempfile.TemporaryDirectory();self.addCleanup(self.t.cleanup);self.w=Path(self.t.name)
  self.source=self.w/'candidate.pptx';p=Presentation();p.slides.add_slide(p.slide_layouts[6]);p.save(self.source)
  self.before=self.source.read_bytes();self.result={'status':'failed','issues':['original layout issue']}
 def capture(self,path,**kwargs):
  err=io.StringIO()
  with contextlib.redirect_stderr(err):code=emit_report(path,self.result,**kwargs)
  return code,json.loads(err.getvalue()) if err.getvalue() else None
 def test_input_alias_is_untouched(self):
  code,out=self.capture(self.source,inputs=[self.source]);self.assertEqual(code,3);self.assertEqual(self.source.read_bytes(),self.before);self.assertIn('new writable',out['recovery'])
 def test_hardlink_input_alias_is_untouched(self):
  alias=self.w/'alias.json';os.link(self.source,alias);code,_=self.capture(alias,inputs=[self.source]);self.assertEqual(code,3);self.assertEqual(self.source.read_bytes(),self.before)
 def test_frozen_render_file_is_untouched(self):
  d=self.w/'render';d.mkdir();p=d/'audit.json';p.write_text('original','utf8');code,_=self.capture(p,frozen_dirs=[d]);self.assertEqual(code,3);self.assertEqual(p.read_text(),'original')
 def test_new_report_beside_render_works(self):
  d=self.w/'render';d.mkdir();p=self.w/'report.json';self.assertEqual(emit_report(p,self.result,inputs=[self.source],frozen_dirs=[d]),0);self.assertEqual(json.loads(p.read_text()),self.result)
 def test_replace_failure_preserves_previous_report(self):
  p=self.w/'report.json';p.write_text('previous','utf8')
  with patch('report_io.os.replace',side_effect=PermissionError('locked')):code,out=self.capture(p)
  self.assertEqual(code,3);self.assertEqual(p.read_text(),'previous');self.assertEqual(out['check_result'],self.result);self.assertEqual(list(self.w.glob('*.tmp')),[])
 def test_parent_creation_failure_is_structured(self):
  p=self.w/'new'/'report.json'
  with patch('report_io.Path.mkdir',side_effect=PermissionError('denied')):code,out=self.capture(p)
  self.assertEqual(code,3);self.assertEqual(out['stage'],'report_write')
 def test_nonfinite_result_retains_valid_json_diagnostic(self):
  self.result={'status':'failed','measurement':float('nan')};code,out=self.capture(self.w/'report.json')
  self.assertEqual(code,3);self.assertEqual(out['check_result']['measurement'],'NaN');self.assertFalse((self.w/'report.json').exists())
 def test_bad_destination_directory(self):
  code,out=self.capture(self.w);self.assertEqual(code,3);self.assertIn('not a file',out['error'])
 def test_existing_report_can_be_replaced_atomically(self):
  p=self.w/'report.json';p.write_text('previous','utf8');write_report(p,self.result);self.assertEqual(json.loads(p.read_text()),self.result)
 def test_all_eight_cli_entries_protect_pptx_output(self):
  scripts=Path(__file__).resolve().parent;meta=self.w/'meta.json';meta.write_text('{}','utf8');render=self.w/'render';render.mkdir();(render/'audit.json').write_text('{}','utf8')
  commands={
   'audit_presentation_text.py':[self.source,self.source],
   'audit_layout_coverage.py':[self.source,meta,self.source],
   'audit_layout_gaps.py':[self.source,meta,self.source],
   'check_signoff_scope.py':[self.source,meta,self.source],
   'visible_chart_spacing.py':[self.source,render,self.source],
   'verify_coverage_review.py':[self.source,meta,meta,render,self.source],
   'verify_readability_review.py':[self.source,render,meta,self.source],
   'verify_delivery.py':[self.source,meta,meta,render,meta,self.source]}
  for name,args in commands.items():
   with self.subTest(entry=name):
    p=subprocess.run([sys.executable,'-B','-X','utf8',str(scripts/name),*[str(x) for x in args]],capture_output=True,encoding='utf8',timeout=25)
    self.assertEqual(p.returncode,3,p.stdout+p.stderr);self.assertEqual(json.loads(p.stderr)['stage'],'report_write');self.assertEqual(self.source.read_bytes(),self.before)
 def test_indirect_json_source_is_protected(self):
  data=self.w/'source.json';data.write_text('{"values":[1,2]}','utf8');review=self.w/'review.json';review.write_text(json.dumps({'examples':[{'source_path':str(data)}]}),'utf8')
  before=data.read_bytes();code,_=self.capture(data,inputs=[review]);self.assertEqual(code,3);self.assertEqual(data.read_bytes(),before)
 def test_cyclic_metadata_references_terminate_and_protect(self):
  a=self.w/'a.json';b=self.w/'b.json';a.write_text(json.dumps({'path':str(b)}),'utf8');b.write_text(json.dumps({'path':str(a)}),'utf8')
  before=b.read_bytes();code,_=self.capture(b,inputs=[a]);self.assertEqual(code,3);self.assertEqual(b.read_bytes(),before)
 def test_bom_layout_spec_runs(self):
  scripts=Path(__file__).resolve().parent;spec=self.w/'layout.json'
  spec.write_text(json.dumps({'pages':[{'slide':1,'comparisons':[],'manual_axes':{'horizontal':'单页合成测试没有横向并列模块','vertical':'单页合成测试没有纵向并列模块'}}]}),'utf-8-sig')
  p=subprocess.run([sys.executable,'-B','-X','utf8',str(scripts/'audit_layout_gaps.py'),str(self.source),str(spec),str(self.w/'gap.json')],capture_output=True,encoding='utf8',timeout=25)
  self.assertEqual(p.returncode,0,p.stdout+p.stderr)
if __name__=='__main__':unittest.main()
