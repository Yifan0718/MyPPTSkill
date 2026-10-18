"""Validate full coverage or the eligibility/impact of a post-signoff user amendment.
This validates scope only, never replaces quality/visual signoff.
"""
import argparse,hashlib,json,posixpath
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation

def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest().upper()
def parts(path):
    with ZipFile(path) as z:
        if len(z.namelist())!=len(set(z.namelist())):raise ValueError('Duplicate package parts')
        return {n:z.read(n) for n in z.namelist()}
def impacted_pages(before,after):
    a,b=parts(before),parts(after)
    old=[str(s.part.partname).lstrip('/') for s in Presentation(before).slides]
    new=[str(s.part.partname).lstrip('/') for s in Presentation(after).slides]
    changed={n for n in set(a)|set(b) if a.get(n)!=b.get(n)}
    if old!=new or any(n.startswith('ppt/presentation') for n in changed):return list(range(1,len(new)+1)),sorted(changed)
    def dependencies(package,start):
        seen=set();todo=[start]
        while todo:
            part=todo.pop()
            if part in seen:continue
            seen.add(part)
            rel=posixpath.join(posixpath.dirname(part),'_rels',posixpath.basename(part)+'.rels')
            seen.add(rel)
            if rel not in package:continue
            for entry in etree.fromstring(package[rel]):
                if entry.get('TargetMode')=='External':
                    if entry.get('Type','').rsplit('/',1)[-1] not in ('hyperlink',):
                        raise ValueError('External visual/data dependencies prevent proving a local-only impact scope')
                    continue
                target=entry.get('Target','')
                target=posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(posixpath.dirname(part),target))
                if target.startswith('../'):raise ValueError('Unsafe package relationship')
                todo.append(target)
        return seen
    scope=[];covered=set()
    for i,part in enumerate(new,1):
        deps=dependencies(a,part)|dependencies(b,part);covered|=deps
        if changed&deps:scope.append(i)
    # Unknown changed global parts are conservative: they may affect every page.
    ignored={n for n in changed if n.startswith('docProps/')}
    if changed-covered-ignored:scope=list(range(1,len(new)+1))
    return scope,sorted(changed)

def check(pptx,spec):
    count=len(Presentation(pptx).slides);all_pages=list(range(1,count+1))
    if spec.get('sha256')!=sha(pptx):raise ValueError('Current PPTX hash mismatch')
    pages=spec.get('checked_slides')
    if not isinstance(pages,list) or any(type(n) is not int or n not in all_pages for n in pages) or len(pages)!=len(set(pages)):raise ValueError('Invalid/duplicate checked pages')
    mode=spec.get('mode')
    if mode=='full':
        if sorted(pages)!=all_pages:raise ValueError('Full signoff requires every page, including cover')
        return {'status':'scope_valid_pending_quality_checks','mode':'full','required_slides':all_pages,'sha256':sha(pptx)}
    if mode!='amendment':raise ValueError('Unknown signoff scope mode')
    baseline=Path(spec.get('baseline_signoff',''));before=Path(spec.get('baseline_pptx',''))
    if not baseline.is_file() or spec.get('baseline_signoff_sha256')!=sha(baseline):raise ValueError('Missing/stale baseline signoff')
    record=json.loads(baseline.read_text(encoding='utf-8-sig'))
    if record.get('status')!='signed_off' or record.get('mode')!='full' or record.get('unresolved')!=[]:raise ValueError('Local amendment requires an already completed full signoff, not pending/partial review')
    if not before.is_file() or record.get('sha256')!=sha(before):raise ValueError('Baseline PPTX missing or modified')
    if record.get('checked_slides')!=list(range(1,len(Presentation(before).slides)+1)):raise ValueError('Baseline does not cover every page')
    user=spec.get('user_reported_issue',{})
    if user.get('source')!='user' or not isinstance(user.get('reference'),str) or not user['reference'].strip() or not isinstance(user.get('quote'),str) or not user['quote'].strip():raise ValueError('Need actual post-signoff user feedback, not an agent-selected local scope')
    # Timestamps establish order, not proof of human authorship; verify the
    # referenced conversation as part of signoff, never fabricate it.
    from datetime import datetime
    signed=datetime.fromisoformat(record['signed_at']);reported=datetime.fromisoformat(user['reported_at'])
    if signed.tzinfo is None or reported.tzinfo is None or reported<=signed:raise ValueError('User issue must occur after full signoff')
    needed,changed=impacted_pages(before,pptx)
    if not set(needed).issubset(pages):raise ValueError('Unreviewed impacted pages: '+str(sorted(set(needed)-set(pages))))
    return {'status':'scope_valid_pending_quality_checks','mode':mode,'required_slides':needed,'changed_parts':changed,'sha256':sha(pptx),'baseline_sha256':sha(before)}

def main():
    ap=argparse.ArgumentParser()
    for n in ('pptx','scope','report'):ap.add_argument(n,type=Path)
    a=ap.parse_args()
    try:r=check(a.pptx,json.loads(a.scope.read_text(encoding='utf-8-sig')));code=0
    except Exception as e:r={'status':'failed','error':str(e)};code=2
    a.report.parent.mkdir(parents=True,exist_ok=True);a.report.write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf8')
    print(r['status']);return code
if __name__=='__main__':raise SystemExit(main())
