"""Synthetic contract regressions only; no PowerPoint rendering or signoff."""
import sys
sys.dont_write_bytecode = True
import shutil
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Pt

from render_evidence import (
    digest, native_inventory, validate_manual_dispositions, validate_render_evidence,
)


NAVY = (30, 60, 90)
OBJECT_BOX = [100, 100, 180, 160]
DETAIL_BOX = [90, 90, 190, 170]
PAGE_BOX = [0, 0, 960, 540]


class UniformPictureTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='uniform-picture-test-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()

    def make(self, source_image=None, target_color=NAVY, background='white', *, nested=False, picture=True):
        self.source_path = self.root / 'source.png'
        if source_image is None:
            source_image = Image.new('RGB', (80, 60), NAVY)
        source_image.save(self.source_path)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Pt(960), Pt(540)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shapes, prefix = slide.shapes, ''
        if nested:
            group = shapes.add_group_shape()
            inner = group.shapes.add_group_shape()
            shapes = inner.shapes
            prefix = f'{group.shape_id}/group/{inner.shape_id}/group/'
        if picture:
            target = shapes.add_picture(str(self.source_path), Pt(100), Pt(100), Pt(80), Pt(60))
        else:
            target = shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Pt(100), Pt(100), Pt(80), Pt(60))
        if nested:
            # Give group reviews nonconstant pixels without changing the picture.
            shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Pt(200), Pt(100), Pt(10), Pt(10))
        candidate = self.root / 'fixture.pptx'
        prs.save(candidate)
        snapshot = self.root / 'snapshot.pptx'
        shutil.copyfile(candidate, snapshot)
        self.page_path = self.root / 'slide-01.png'
        page = Image.new('RGB', (960, 540), background)
        page.paste(target_color, tuple(OBJECT_BOX))
        # This unrelated mark keeps master/layout nonconstant even for white-on-white.
        page.paste((180, 40, 40), (800, 400, 820, 420))
        page.save(self.page_path)
        manual = [{'slide': 1, 'audit_path': prefix + str(target.shape_id),
                   'reason': 'picture_internal_visual_review' if picture else 'shape_visual_review'}]
        for slide_number, path in sorted(native_inventory(prs)['groups']):
            manual.append({'slide': slide_number, 'audit_path': path,
                           'reason': 'group_geometry_requires_visual_review'})
        manual.append({'slide': 1, 'audit_path': 'master-layout', 'reason': 'master_layout_visual_review'})
        for number, row in enumerate(manual, 1):
            row['review_id'] = f'm{number}'
        self.render = {
            'schema': 'render-evidence/v2', 'sha256': digest(candidate),
            'snapshot': str(snapshot), 'slide_count': 1,
            'status': 'rendered_pending_visual_review', 'issues': [],
            'raster_hashes': {'1': digest(self.page_path)},
            'text_records': [], 'chart_geometry_records': [], 'manual_review': manual,
        }
        self.review = {'render_manual_reviews': [dict(
            row, decision='verified_in_context', returned_to_full_page=True,
            image_sha256=digest(self.page_path), crop_box_px=PAGE_BOX.copy(),
            crop_path=str(self.page_path),
            observation='Synthetic contract probe: navy swatch interior matches its source; all four edges have white context. No visual signoff.',
        ) for row in manual]}
        self.set_detail(DETAIL_BOX)
        validate_render_evidence(candidate, self.render, self.root)

    @property
    def item(self):
        return self.review['render_manual_reviews'][0]

    def set_detail(self, box):
        detail_path = self.root / 'detail.png'
        with Image.open(self.page_path) as page:
            page.crop(box).save(detail_path)
        self.item.update(detail_crop_box_px=list(box), detail_crop_path=str(detail_path))

    def refresh_page_hashes(self):
        self.render['raster_hashes']['1'] = digest(self.page_path)
        for row in self.review['render_manual_reviews']:
            row['image_sha256'] = digest(self.page_path)

    def issues(self):
        return validate_manual_dispositions(self.render, self.review, self.root)

    def assert_picture_rejected(self, message):
        errors = self.issues()
        self.assertTrue(any(row.get('review_id') == 'm1' and message in row['reason'] for row in errors), errors)

    def test_opaque_uniform_rgb_and_rgba_sources_with_current_local_context_pass(self):
        for mode, color in [('RGB', NAVY), ('RGBA', NAVY + (255,))]:
            with self.subTest(mode=mode):
                self.make(Image.new(mode, (80, 60), color))
                self.assertEqual(self.issues(), [])

    def test_nested_picture_source_resolves_by_native_path(self):
        self.make(nested=True)
        self.assertEqual(self.issues(), [])

    def test_embedded_blob_is_used_instead_of_original_external_asset(self):
        self.make()
        Image.new('RGBA', (80, 60), (255, 255, 255, 0)).save(self.source_path)
        self.assertEqual(self.issues(), [])

    def test_matching_visible_white_picture_on_navy_context_passes(self):
        self.make(Image.new('RGB', (80, 60), 'white'), target_color='white', background=NAVY)
        self.item['observation'] = 'Synthetic contract probe: opaque white source matches interior; each edge meets navy background. No visual signoff.'
        self.assertEqual(self.issues(), [])

    def test_white_on_white_fails_despite_distant_content_and_specific_observation(self):
        self.make(Image.new('RGB', (80, 60), 'white'), target_color='white')
        self.assert_picture_rejected('lacks visible contrast')

    def test_uniform_target_must_match_embedded_source(self):
        self.make(target_color=(31, 60, 90))
        self.assert_picture_rejected('do not match the native source color')

    def test_nonuniform_native_source_cannot_use_constant_picture_exception(self):
        source = Image.new('RGB', (80, 60), NAVY)
        source.putpixel((40, 30), (31, 60, 90))
        self.make(source)
        self.item['uniform_picture'] = True  # A reviewer claim grants no exemption.
        self.assert_picture_rejected('opaque uniform native picture source')

    def test_transparent_and_translucent_source_pixels_fail(self):
        for alpha in (0, 128, 254):
            with self.subTest(alpha=alpha):
                source = Image.new('RGBA', (80, 60), NAVY + (255,))
                source.putpixel((40, 30), NAVY + (alpha,))
                self.make(source)
                self.assert_picture_rejected('opaque uniform native picture source')

    def test_palette_transparency_is_not_lost_by_rgb_conversion(self):
        source = Image.new('P', (80, 60), 0)
        source.putpalette(list(NAVY) + [0] * (768 - 3))
        source.info['transparency'] = 0
        self.make(source)
        self.assert_picture_rejected('opaque uniform native picture source')

    def test_constant_native_shape_does_not_gain_picture_exception(self):
        self.make(picture=False)
        self.assert_picture_rejected('Blank/constant pixels')

    def test_both_explicit_local_detail_fields_are_required(self):
        for field in ('detail_crop_box_px', 'detail_crop_path'):
            with self.subTest(field=field):
                self.make()
                self.item.pop(field)
                self.assert_picture_rejected('requires explicit')

    def test_full_page_is_not_a_local_detail_substitute(self):
        self.make()
        self.set_detail(PAGE_BOX)
        self.assert_picture_rejected('not the full page')

    def test_detail_must_include_context_outside_every_object_edge(self):
        for box in (OBJECT_BOX, [100, 90, 190, 170], [90, 100, 190, 170],
                    [90, 90, 180, 170], [90, 90, 190, 160]):
            with self.subTest(box=box):
                self.make()
                self.set_detail(box)
                self.assert_picture_rejected('context outside all object boundaries')

    def test_each_boundary_needs_adjacent_visible_contrast(self):
        strips = [(99, 100, 100, 160), (180, 100, 181, 160),
                  (100, 99, 180, 100), (100, 160, 180, 161)]
        for strip in strips:
            with self.subTest(strip=strip):
                self.make()
                with Image.open(self.page_path) as opened:
                    page = opened.copy()
                page.paste(NAVY, strip)
                page.save(self.page_path)
                self.refresh_page_hashes()
                self.set_detail(DETAIL_BOX)
                self.assert_picture_rejected('lacks visible contrast')

    def test_unrelated_current_blank_detail_fails(self):
        self.make()
        self.set_detail([500, 300, 600, 380])
        self.assert_picture_rejected('context outside all object boundaries')

    def test_unrelated_primary_crop_cannot_be_rescued_by_valid_detail(self):
        self.make()
        blank_path = self.root / 'blank-primary.png'
        with Image.open(self.page_path) as page:
            page.crop([500, 300, 600, 380]).save(blank_path)
        self.item.update(crop_box_px=[500, 300, 600, 380], crop_path=str(blank_path))
        self.assert_picture_rejected('omits its native object bounds')

    def test_detail_pixels_must_match_the_current_render_exactly(self):
        self.make()
        detail_path = Path(self.item['detail_crop_path'])
        with Image.open(detail_path) as opened:
            detail = opened.copy()
        detail.putpixel((0, 0), (254, 255, 255))
        detail.save(detail_path)
        self.assert_picture_rejected('local detail does not match current page')

    def test_full_page_path_cannot_replace_local_detail_file(self):
        self.make()
        self.item['detail_crop_path'] = str(self.page_path)
        self.assert_picture_rejected('local detail does not match current page')

    def test_existing_hash_observation_and_full_page_requirements_still_apply(self):
        changes = [('image_sha256', '0' * 64, 'image stale'),
                   ('observation', 'OK', 'specific manual observation'),
                   ('observation', ' ', 'specific manual observation'),
                   ('returned_to_full_page', False, 'remains unresolved')]
        for field, value, message in changes:
            with self.subTest(field=field, value=value):
                self.make()
                self.item[field] = value
                self.assert_picture_rejected(message)

    def test_current_primary_crop_pixels_are_still_checked(self):
        self.make()
        fake_path = self.root / 'fake-primary.png'
        with Image.open(self.page_path) as opened:
            fake = opened.copy()
        fake.putpixel((0, 0), (0, 0, 0))
        fake.save(fake_path)
        self.item['crop_path'] = str(fake_path)
        self.assert_picture_rejected('not current rendered pixels')


if __name__ == '__main__':
    unittest.main(verbosity=2)
