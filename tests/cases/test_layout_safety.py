"""Portable logic regressions; synthetic rendered bounds do not replace real rendering.
Run: python test_layout_safety.py
"""
import copy
import hashlib
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_CONNECTOR
from audit_layout_coverage import audit, segment_hits_box, ink_bounds
from verify_coverage_review import check_visual_passes


class LayoutSafetyTests(unittest.TestCase):
    def test_same_coordinate_different_ink_start(self):
        from PIL import Image,ImageDraw
        image=Image.new('RGB',(300,120),'white');draw=ImageDraw.Draw(image)
        draw.rectangle((102,23,145,35),fill='black');draw.rectangle((100,63,145,75),fill='black')
        a={'bound_left_pt':100,'bound_top_pt':20,'bound_width_pt':60,'bound_height_pt':20}
        b={**a,'bound_top_pt':60}
        aa=ink_bounds(a,image,1);bb=ink_bounds(b,image,1)
        self.assertEqual(a['bound_left_pt'],b['bound_left_pt'])
        self.assertEqual(aa['bounds_px'][0]-bb['bounds_px'][0],2)

    def test_segment_clearance(self):
        box=[10,10,20,20]
        self.assertTrue(segment_hits_box(0,15,30,15,box))
        self.assertTrue(segment_hits_box(15,0,15,30,box))
        self.assertTrue(segment_hits_box(0,0,30,30,box))
        self.assertFalse(segment_hits_box(0,5,30,5,box))
        self.assertFalse(segment_hits_box(0,0,5,5,box))

    def test_overlap_connector_and_multiline_not_skipped(self):
        with tempfile.TemporaryDirectory() as td:
            p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540)
            slide=p.slides.add_slide(p.slide_layouts[6]);records=[]
            for x,y,w,h,inkh,label in [(40,100,120,30,20,'first'),(100,110,120,30,20,'overlap'),(300,100,150,50,36,'two lines')]:
                s=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));s.text=label;s.name='body:'+label
                records.append({'slide':1,'shape_id':s.shape_id,'shape':s.name,'kind':'text','text':label,
                                'font_pt':20 if label=='overlap' else 16,'left_pt':x,'top_pt':y,'width_pt':w,'height_pt':h,
                                'bound_left_pt':x,'bound_top_pt':y,'bound_width_pt':w-10,'bound_height_pt':inkh})
            slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(25),Pt(115),Pt(240),Pt(115))
            slide.shapes.add_connector(MSO_CONNECTOR.ELBOW,Pt(25),Pt(180),Pt(240),Pt(200))
            path=Path(td)/'case.pptx';p.save(path)
            render={'sha256':hashlib.sha256(path.read_bytes()).hexdigest().upper(),'slide_count':1,
                    'status':'rendered_pending_visual_review','issues':[],'text_records':records}
            reasons={f['reason'] for f in audit(path,render)['findings']}
            self.assertIn('rendered_text_rectangles_overlap',reasons)
            self.assertIn('connector_crosses_or_touches_text',reasons)
            self.assertIn('nonstraight_connector_path_requires_review',reasons)
            self.assertIn('multiline_or_mixed_font_row_alignment_requires_review',reasons)
            self.assertIn('mixed_typography_optical_review',reasons)

    def test_two_final_passes_and_current_images(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);image=root/'slide-01.png';image.write_bytes(b'fixture image')
            image_sha=hashlib.sha256(image.read_bytes()).hexdigest().upper()
            review={'visual_passes':[]}
            for mode,keys in [('content',('readability','chart_semantics','language','quantity_unit_pairing')),('layout',('row_alignment','column_alignment','whitespace','containment_and_connectors','page_number','rules_crosscheck','optical_edges'))]:
                review['visual_passes'].append({'mode':mode,'sha256':'PPTX', 'pages':[{
                    'slide':1,'image_sha256':image_sha,'normal_size_viewed':True,'result':'pass','rule_findings_reviewed':[],'optical_targets_reviewed':[],
                    'checks':{k:'Specific fixture observation for '+k for k in keys}}]})
            self.assertEqual(check_visual_passes(review,root,'PPTX',1),[])
            self.assertTrue(check_visual_passes(review,root,'PPTX',1,[{'slide':1,'id':'unreviewed'}]))
            self.assertTrue(check_visual_passes(review,root,'PPTX',1,[],[{'slide':1,'optical_review_targets':[{'shape_id':5}]}]))
            for mutate in (
                lambda r:r['visual_passes'].pop(),
                lambda r:r['visual_passes'][1]['pages'].clear(),
                lambda r:r['visual_passes'][0].update(sha256='OLD'),
                lambda r:r['visual_passes'][0]['pages'][0]['checks'].pop('quantity_unit_pairing'),
                lambda r:r['visual_passes'][1]['pages'][0]['checks'].update(row_alignment='pass'),
                lambda r:r['visual_passes'][1]['pages'][0].update(image_sha256='OLD')):
                bad=copy.deepcopy(review);mutate(bad)
                self.assertTrue(check_visual_passes(bad,root,'PPTX',1))


if __name__=='__main__':unittest.main()
