"""Gap report control-flow tests; fixture PNGs are not render/signoff evidence."""
import json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from audit_layout_gaps import audit
from verify_coverage_review import validate_gap_audit,digest

class GapAuditGate(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.pptx=self.root/'source.pptx'
        p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540);s=p.slides.add_slide(p.slide_layouts[6])
        a=s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(100),Pt(100),Pt(100),Pt(50));b=s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(212),Pt(100),Pt(100),Pt(50));p.save(self.pptx)
        self.spec={'pages':[{'slide':1,'manual_axes':{'vertical':'A single actual content row is checked in the full page visual pass.'},'comparisons':[{'name':'row','axis':'horizontal','target_gap_pt':12,'blocks':[{'shape_ids':[a.shape_id]},{'shape_ids':[b.shape_id]}]}]}]}
        self.layout=self.root/'layout.json';self.report=self.root/'gap.json'
        self.layout.write_text(json.dumps(self.spec),encoding='utf8');self.report.write_text(json.dumps(audit(self.pptx,self.spec)),encoding='utf8')
        self.image=self.root/'slide-01.png';im=Image.new('RGB',(960,540),'white');d=ImageDraw.Draw(im);d.rectangle((100,100,199,149),fill='navy');d.rectangle((212,100,311,149),fill='navy');im.save(self.image)
        self.render={'schema':'render-evidence/v2','sha256':digest(self.pptx),'snapshot':str(self.pptx),'slide_count':1,'status':'rendered_pending_visual_review','issues':[],'raster_hashes':{'1':digest(self.image)}}
        self.review={'gap_audit':{'layout_path':str(self.layout),'layout_sha256':digest(self.layout),'report_path':str(self.report),'report_sha256':digest(self.report),'visible_reviews':[]}}
    def tearDown(self):self.tmp.cleanup()
    def call(self):return validate_gap_audit(self.pptx,self.review,self.render,self.root)
    def test_complete_raw_zero_issue_report_passes(self):self.assertEqual(self.call(),([],[]))
    def test_missing_gate_is_not_a_pass(self):self.review={};self.assertTrue(self.call()[0])
    def test_stale_spec_is_rejected(self):self.review['gap_audit']['layout_sha256']='old';self.assertTrue(self.call()[0])
    def test_raw_report_cannot_be_rewritten(self):
        r=json.loads(self.report.read_text());r['comparisons']=[];self.report.write_text(json.dumps(r));self.review['gap_audit']['report_sha256']=digest(self.report);self.assertTrue(self.call()[0])
    def test_structural_error_cannot_be_visually_waived(self):
        self.spec['pages'][0].pop('manual_axes');self.layout.write_text(json.dumps(self.spec));self.report.write_text(json.dumps(audit(self.pptx,self.spec)));self.review['gap_audit'].update(layout_sha256=digest(self.layout),report_sha256=digest(self.report));self.assertTrue(self.call()[0])
    def test_unused_visual_exemption_rejected(self):self.review['gap_audit']['visible_reviews']=[{'slide':1,'comparison_name':'fake'}];self.assertTrue(self.call()[0])
    def test_discovery_is_mandatory_even_when_raw_frame_report_passes(self):
        row={'slide':1,'name':'independent','axis':'vertical','blocks':[{'shape_ids':[2]},{'shape_ids':[3]},{'shape_ids':[4]}],'minimum_gap_pt':12}
        with patch('visible_chart_spacing.discover_relations',return_value=[row]):self.assertTrue(self.call()[0])
    def mocked_resolution(self,state='measured_valid_pending_visual_review',bad_pixels=False):
        # Native/pixel adapter correctness is tested separately. This exercises
        # final binding and original-issue retention with a deterministic adapter.
        row={'slide':1,'name':'independent','axis':'vertical','blocks':[{'shape_ids':[2]},{'shape_ids':[3]},{'shape_ids':[4]}],'minimum_gap_pt':12}
        m={'slide':1,'comparison_name':'independent','shape_ids':[2,3,4],'measurement_status':state,'image_sha256':digest(self.image),'crop_box_px':[80,80,340,180],'visible_gaps_pt':[20,20]}
        mf=self.root/'measured.json';mf.write_text(json.dumps(m));crop=self.root/'crop.png'
        with Image.open(self.image) as im:im.crop(m['crop_box_px']).save(crop)
        if bad_pixels:Image.new('RGB',(260,100),'white').save(crop)
        item={'slide':1,'comparison_name':'independent','measurement_path':str(mf),'measurement_sha256':digest(mf),'decision':'verified_in_context','returned_to_full_page':True,'observation':'Synthetic protocol fixture only; local and full-page provenance checked separately.','image_sha256':digest(self.image),'detail_evidence':{'path':str(crop),'sha256':digest(crop),'source_image_sha256':digest(self.image),'crop_box_px':m['crop_box_px']}}
        self.review['gap_audit']['visible_reviews']=[item]
        with patch('visible_chart_spacing.discover_relations',return_value=[row]),patch('visible_chart_spacing.measure_relation',return_value=m):return self.call()
    def test_complete_discovered_review_is_retained(self):
        errors,closed=self.mocked_resolution();self.assertEqual(errors,[]);self.assertEqual(len(closed),1)
    def test_actual_measured_defect_cannot_close(self):self.assertTrue(self.mocked_resolution(state='measured_defect')[0])
    def test_current_hash_of_unrelated_crop_is_not_enough(self):self.assertTrue(self.mocked_resolution(bad_pixels=True)[0])

if __name__=='__main__':unittest.main()
