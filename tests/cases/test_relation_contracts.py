"""Regression tests for null COM names, column scope and explicit hierarchy."""
import hashlib,tempfile,unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from audit_layout_coverage import audit
from verify_coverage_review import validate_region_contract,verify_plot_table_pixels,verified_optical_anchor
from skill_state import snapshot,check

class RelationTests(unittest.TestCase):
    def test_actual_ink_can_resolve_com_side_bearings_but_not_visible_offset(self):
        f={'shape_ids':[1,2,3],'ink_bounds_px':[[100,10,200,40],[100,50,150,70],[100,80,180,110]],'threshold_px':2}
        r={'alignment_contract':{'anchor':'left','shape_ids':[1,2,3]}}
        self.assertTrue(verified_optical_anchor(f,r))
        f['ink_bounds_px'][2][0]=102
        self.assertFalse(verified_optical_anchor(f,r))
        r['alignment_contract']['shape_ids']=[1,2]
        self.assertFalse(verified_optical_anchor(f,r))
        f['ink_bounds_px']=None
        self.assertFalse(verified_optical_anchor(f,r))
    def test_distinct_native_filled_text_containers(self):
        p,s=self.make();ids=[]
        for y in (100,180):
            c=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(40),Pt(y),Pt(600),Pt(40));c.text='actual panel'
            c.fill.solid();c.fill.fore_color.rgb=RGBColor(30,60,90);ids.append(c.shape_id)
        f={'shape_ids':ids,'native_anchors':['center','left']}
        r={'relation_contract':{'kind':'separate_visible_containers','shape_ids':ids,'containers':{str(i):i for i in ids},'visible_relation_explanation':'two separate filled panels'}}
        self.assertIsNone(validate_region_contract(f,r,s,1))
        s.shapes[1].top=Pt(120)
        self.assertIsNotNone(validate_region_contract(f,r,s,1))
        s.shapes[1].top=Pt(180);s.shapes[1].fill.background()
        self.assertIsNotNone(validate_region_contract(f,r,s,1))

    def test_mixed_alignment_modes_do_not_waive_visible_stack_mismatch(self):
        p,s=self.make()
        a,ra=self.text(s,100,100,200,'heading',PP_ALIGN.CENTER)
        b,rb=self.text(s,100,140,200,'849.1',PP_ALIGN.CENTER)
        c,rc=self.text(s,100,200,200,'explanation',PP_ALIGN.LEFT)
        f={'shape_ids':[a.shape_id,b.shape_id,c.shape_id],'native_anchors':['center','center','left'],
           'ink_bounds_px':[[120,100,180,120],[100,140,200,170],[105,200,215,240]],'threshold_px':2}
        r={'relation_contract':{'kind':'different_alignment_roles','shape_ids':f['shape_ids'],'visible_relation_explanation':'same metric group'}}
        self.assertIsNotNone(validate_region_contract(f,r,s,1))
        r['relation_contract']['common_visual_anchor']='center'
        self.assertIsNotNone(validate_region_contract(f,r,s,1))
        f['ink_bounds_px']=[[100,100,180,120],[100,140,200,170],[100,200,215,240]]
        r['relation_contract']['common_visual_anchor']='left'
        self.assertIsNone(validate_region_contract(f,r,s,1))
        f['ink_bounds_px']=None
        self.assertIsNotNone(validate_region_contract(f,r,s,1))
        f['native_anchors']=['left','left','left']
        f['ink_bounds_px']=[[100,100,180,120],[100,140,200,170],[100,200,215,240]]
        r['relation_contract']['kind']='common_visual_anchor'
        self.assertIsNone(validate_region_contract(f,r,s,1))
        f['ink_bounds_px'][2][0]=103
        self.assertIsNotNone(validate_region_contract(f,r,s,1))
        scan=self.scan(p,[ra,rb,rc])
        self.assertTrue(any(set(x.get('shape_ids',[]))==set(f['shape_ids']) for x in scan['findings'] if x['reason']=='text_region_relation_requires_review'))

    def make(self):
        p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540)
        return p,p.slides.add_slide(p.slide_layouts[6])
    def text(self,s,x,y,w,label,align=PP_ALIGN.LEFT):
        sh=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(30));sh.text=label;sh.text_frame.paragraphs[0].alignment=align
        return sh,{'slide':1,'shape_id':sh.shape_id,'shape':sh.name,'kind':'text','text':label,
          'font_pt':16,'left_pt':x,'top_pt':y,'width_pt':w,'height_pt':30,
          'bound_left_pt':x,'bound_top_pt':y,'bound_width_pt':w,'bound_height_pt':20}
    def scan(self,p,records):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'test.pptx';p.save(path)
            return audit(path,{'sha256':hashlib.sha256(path.read_bytes()).hexdigest().upper(),'slide_count':1,
              'status':'rendered_pending_visual_review','issues':[],'text_records':records})
    def test_spanning_summary_does_not_bridge_columns_and_null_cell_kept(self):
        p,s=self.make();a,ra=self.text(s,40,100,300,'left');b,rb=self.text(s,520,100,300,'right');c,rc=self.text(s,40,150,880,'summary')
        null={**ra,'kind':'table','shape':None,'shape_id':None,'text':'unnamed cell'}
        report=self.scan(p,[ra,rb,rc,null])
        self.assertFalse(any(f['reason']=='rendered_column_text_anchor_mismatch' for f in report['findings']))
        f=next(f for f in report['findings'] if f['reason']=='text_region_relation_requires_review')
        r={'relation_contract':{'kind':'spanning_region','shape_ids':f['shape_ids'],'visible_relation_explanation':'Whole-page summary below separate panels'}}
        self.assertIsNone(validate_region_contract(f,r,s,1))
        self.assertTrue(any(x['text']=='unnamed cell' for x in report['pages'][0]['optical_text_inventory']))
    def test_real_small_offset_still_fails(self):
        p,s=self.make();a,ra=self.text(s,100,100,200,'label');b,rb=self.text(s,108,145,200,'number')
        self.assertTrue(any(f['reason']=='rendered_column_text_anchor_mismatch' for f in self.scan(p,[ra,rb])['findings']))
    def test_centered_nodes_use_center_not_left_edge(self):
        p,s=self.make();a,ra=self.text(s,100,100,200,'a',PP_ALIGN.CENTER);b,rb=self.text(s,150,150,100,'b',PP_ALIGN.CENTER)
        self.assertFalse(any(f['reason']=='rendered_column_text_anchor_mismatch' for f in self.scan(p,[ra,rb])['findings']))
    def test_hierarchy_is_measured_not_just_named(self):
        p,s=self.make();a,ra=self.text(s,100,100,300,'system');b,rb=self.text(s,118,150,300,'strategy')
        f=next(f for f in self.scan(p,[ra,rb])['findings'] if f['reason']=='text_region_relation_requires_review')
        contract={'kind':'hierarchy_indent','shape_ids':f['shape_ids'],'indent_pt':18,
          'levels':{str(a.shape_id):0,str(b.shape_id):1},'visible_relation_explanation':'Visible system to strategy level'}
        self.assertIsNone(validate_region_contract(f,{'relation_contract':contract},s,1))
        self.assertIsNotNone(validate_region_contract(f,{'relation_contract':{**contract,'indent_pt':30}},s,1))
        self.assertIsNotNone(validate_region_contract(f,{'relation_contract':{**contract,'shape_ids':[a.shape_id]}},s,1))
        self.assertIsNotNone(validate_region_contract(f,{'relation_contract':{'kind':'ignore','shape_ids':f['shape_ids']}},s,1))
    def test_plot_table_edge_is_not_outer_frame_edge(self):
        p,s=self.make();data=CategoryChartData();data.categories=['a','b'];data.add_series('values',[1,2])
        chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Pt(100),Pt(100),Pt(300),Pt(150),data)
        table=s.shapes.add_table(2,2,Pt(100),Pt(280),Pt(300),Pt(80))
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'test.pptx'
            def check():
                p.save(path)
                render={'sha256':hashlib.sha256(path.read_bytes()).hexdigest().upper(),'slide_count':1,'status':'rendered_pending_visual_review','issues':[],'text_records':[],
                        'chart_geometry_records':[{'slide':1,'shape_id':chart.shape_id,'plot_bounds_pt':[130,110,380,240]}]}
                return audit(path,render)
            a=check();f=next(f for f in a['findings'] if f['reason']=='chart_table_visible_edge_requires_review')
            self.assertEqual(f['delta_pt'],20)
            table.width=Pt(280)
            self.assertFalse(any(f['reason']=='chart_table_visible_edge_requires_review' for f in check()['findings']))
    def test_pixel_edge_proof_rejects_real_mismatch_and_allows_com_rounding(self):
        from PIL import Image,ImageDraw
        p,s=self.make();data=CategoryChartData();data.categories=['a'];data.add_series('values',[1])
        chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Pt(100),Pt(100),Pt(300),Pt(150),data)
        table=s.shapes.add_table(2,2,Pt(100),Pt(280),Pt(300),Pt(80))
        f={'shape_ids':[chart.shape_id,table.shape_id],'plot_right_pt':376,'table_right_pt':400}
        r={'plot_table_edge_evidence':{'pixels_per_pt':1,'slide_width_pt':960,'chart_rows_px':[150,160,170],'table_rows_px':[300,310,320]}}
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'edge.png'
            def draw(right):
                im=Image.new('RGB',(960,540),'white');d=ImageDraw.Draw(im);d.rectangle((100,120,380,220),fill='black');d.rectangle((100,280,right,359),fill='#CBD5E1');im.save(path)
            draw(400)
            self.assertIsNotNone(verify_plot_table_pixels(f,r,path,s,1))
            table.width=Pt(280);f['table_right_pt']=380;draw(380)
            self.assertIsNone(verify_plot_table_pixels(f,r,path,s,1))
    def test_skill_freeze_detects_edits_and_new_rules(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)/'skill';root.mkdir();(root/'SKILL.md').write_text('original')
            lock=Path(td)/'lock.json';snapshot(lock,root)
            self.assertEqual(check(lock,root)['status'],'unchanged')
            (root/'SKILL.md').write_text('changed')
            self.assertEqual(check(lock,root)['status'],'changed')
            (root/'SKILL.md').write_text('original');(root/'new-rule.md').write_text('new')
            self.assertEqual(check(lock,root)['changed_files'],['new-rule.md'])

if __name__=='__main__':unittest.main()
