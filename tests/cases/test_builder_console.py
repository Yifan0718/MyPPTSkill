"""Legacy console logging must not alter data or interrupt the builder."""
import json,os,subprocess,sys,tempfile,unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt

class BuilderConsole(unittest.TestCase):
    def run_case(self,encoding,valid=True):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)/'console_⭐';root.mkdir()
            p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);shape=s.shapes.add_textbox(Pt(40),Pt(50),Pt(300),Pt(40));shape.text='Original ⭐ 科学'
            template=root/'template.pptx';p.save(template)
            new='科学 ∑ 🧪 完整内容'
            spec={'selected_slides':[1],'edits':[{'slide':1,'address':{'shape_id':shape.shape_id,'paragraph':0,'run':0},'expected_text':shape.text if valid else 'Mismatch ⭐','new_text':new}]}
            edits=root/'edits.json';edits.write_text(json.dumps(spec,ensure_ascii=False),encoding='utf8')
            # Exercise pre-run auto-detection logging with an emoji in the path.
            (root/'detail.json').write_text('{"pages":[]}',encoding='utf8')
            out=root/'out.pptx';env=os.environ.copy();env['PYTHONIOENCODING']=encoding+':strict';env['PYTHONUTF8']='0'
            result=subprocess.run([sys.executable,str(Path(__file__).with_name('build_pptx.py')),str(template),str(edits),str(out),'--strict'],capture_output=True,env=env)
            log=(result.stdout+result.stderr).decode(encoding)
            self.assertNotIn('UnicodeEncodeError',log)
            if valid:
                self.assertEqual(result.returncode,0,log);self.assertEqual(Presentation(out).slides[0].shapes[0].text,new)
            else:
                self.assertNotEqual(result.returncode,0);self.assertFalse(out.exists());self.assertIn('expected_text mismatch',log)
            return log
    def test_gbk_console_build_keeps_unicode_in_pptx(self):self.assertIn('\\u2b50',self.run_case('gbk'))
    def test_ascii_console_build_keeps_unicode_in_pptx(self):self.assertIn('\\u2b50',self.run_case('ascii'))
    def test_utf8_console_preserves_literal_log(self):self.assertIn('⭐',self.run_case('utf-8'))
    def test_strict_failure_remains_nonzero_on_legacy_console(self):self.run_case('gbk',False)

if __name__=='__main__':unittest.main()
