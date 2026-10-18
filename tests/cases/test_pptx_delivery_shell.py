import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0,str(Path(__file__).resolve().parent))
import pptx_delivery_shell as shell
from PIL import Image

def png(color='blue'):
    out=BytesIO(); Image.new('RGB',(1280,720),color).save(out,format='PNG'); return out.getvalue()

def package(extra=None):
    parts={
        '[Content_Types].xml':f'<Types xmlns="{shell.CT}"><Default Extension="jpeg" ContentType="image/jpeg"/><Override PartName="/docProps/thumbnail.jpeg" ContentType="image/jpeg"/></Types>'.encode(),
        '_rels/.rels':f'<Relationships xmlns="{shell.R}"><Relationship Id="rId1" Type="{shell.THUMB_TYPE}" Target="docProps/thumbnail.jpeg"/><Relationship Id="rId2" Type="{shell.THUMB_TYPE}" Target="ppt/media/ad.png"/></Relationships>'.encode(),
        'docProps/thumbnail.jpeg':b'old advertisement',
        'ppt/media/ad.png':png('red'),
        'docProps/core.xml':b'<core><creator>Actual author</creator><title>Keep title</title></core>',
        'ppt/slides/slide1.xml':b'body unchanged',
        'ppt/notesSlides/notesSlide1.xml':b'notes unchanged',
        'ppt/charts/chart1.xml':b'chart unchanged',
    }
    parts.update(extra or {})
    out=BytesIO()
    with ZipFile(out,'w') as z:
        for name,data in parts.items(): z.writestr(name,data)
    return out.getvalue()

class DeliveryShellTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name); self.pptx=self.root/'candidate.pptx'; self.first=self.root/'slide-01.png'
        self.pptx.write_bytes(package()); self.first.write_bytes(png())

    def finalize(self):
        return shell.finalize(self.pptx,shell.sha(self.pptx.read_bytes()),self.first,self.root/'proof.json')

    def test_stale_ads_multiple_relationships_replaced(self):
        proof=self.finalize(); self.assertEqual([],shell.verify_shell(self.pptx,self.first))
        parts,_,_=shell.unpack(self.pptx.read_bytes())
        self.assertNotEqual(b'old advertisement',parts[shell.THUMB]); self.assertNotIn('ppt/media/ad.png',parts)
        self.assertEqual(proof['content_fingerprint_before'],proof['content_fingerprint_after'])
        with Image.open(BytesIO(parts[shell.THUMB])) as im: self.assertEqual((320,180),im.size)

    def test_only_shell_changes_author_and_body_unchanged(self):
        before,_,_=shell.unpack(self.pptx.read_bytes()); self.finalize(); after,_,_=shell.unpack(self.pptx.read_bytes())
        for name in ('docProps/core.xml','ppt/slides/slide1.xml','ppt/notesSlides/notesSlide1.xml','ppt/charts/chart1.xml'):
            self.assertEqual(before[name],after[name])

    def test_missing_view_and_relationships_completed(self):
        self.finalize(); parts,_,_=shell.unpack(self.pptx.read_bytes())
        self.assertIn(shell.VIEW,parts); self.assertIn('ppt/_rels/presentation.xml.rels',parts)
        self.assertEqual([],shell.verify_shell(self.pptx,self.first))

    def test_normal_view_preserves_guide_grid(self):
        self.pptx.write_bytes(package({shell.VIEW:f'<p:viewPr xmlns:p="{shell.P}" lastView="outlineView" gridSpacing="keep"><p:normalViewPr showOutlineIcons="1"/><p:slideViewPr custom="keep"/></p:viewPr>'.encode()}))
        self.finalize(); parts,_,_=shell.unpack(self.pptx.read_bytes()); root=shell.ET.fromstring(parts[shell.VIEW])
        self.assertEqual('keep',root.get('gridSpacing')); self.assertEqual('keep',root.find('{'+shell.P+'}slideViewPr').get('custom'))
        self.assertEqual('1',root.find('{'+shell.P+'}normalViewPr').get('showOutlineIcons'))
        self.assertEqual([],shell.verify_shell(self.pptx,self.first))

    def test_reused_media_not_deleted(self):
        rel=f'<Relationships xmlns="{shell.R}"><Relationship Id="img" Type="image" Target="../media/ad.png"/></Relationships>'.encode()
        self.pptx.write_bytes(package({'ppt/slides/_rels/slide1.xml.rels':rel})); self.finalize()
        parts,_,_=shell.unpack(self.pptx.read_bytes())
        self.assertEqual(png('red'),parts['ppt/media/ad.png']); self.assertEqual(rel,parts['ppt/slides/_rels/slide1.xml.rels'])

    def test_encoded_reused_media_not_deleted(self):
        rel=f'<Relationships xmlns="{shell.R}"><Relationship Id="img" Type="image" Target="../media/ad%20image.png"/></Relationships>'.encode()
        thumb=f'<Relationships xmlns="{shell.R}"><Relationship Id="t" Type="{shell.THUMB_TYPE}" Target="ppt/media/ad image.png"/></Relationships>'.encode()
        self.pptx.write_bytes(package({'ppt/slides/_rels/slide1.xml.rels':rel,'_rels/.rels':thumb,'ppt/media/ad image.png':png('red')})); self.finalize()
        parts,_,_=shell.unpack(self.pptx.read_bytes()); self.assertEqual(png('red'),parts['ppt/media/ad image.png'])

    def test_reused_canonical_preserved_with_separate_thumbnail(self):
        rel=f'<Relationships xmlns="{shell.R}"><Relationship Id="img" Type="image" Target="../../{shell.THUMB}"/></Relationships>'.encode()
        original=package({'ppt/slides/_rels/slide1.xml.rels':rel,shell.THUMB:png('red')}); self.pptx.write_bytes(original)
        self.finalize()
        parts,_,_=shell.unpack(self.pptx.read_bytes())
        self.assertEqual(png('red'),parts[shell.THUMB])
        self.assertEqual(rel,parts['ppt/slides/_rels/slide1.xml.rels'])
        self.assertIn('docProps/thumbnail.1.jpeg',parts)
        self.assertEqual([],shell.verify_shell(self.pptx,self.first))

    def test_clean_intermediate_no_thumbnail_author_preserved(self):
        before,_,_=shell.unpack(self.pptx.read_bytes()); shell.clean_intermediate(self.pptx)
        parts,_,_=shell.unpack(self.pptx.read_bytes()); self.assertNotIn(shell.THUMB,parts)
        self.assertEqual(before['docProps/core.xml'],parts['docProps/core.xml'])
        self.assertEqual('sldView',shell.ET.fromstring(parts[shell.VIEW]).get('lastView'))

    def test_stale_source_hash_refused(self):
        original=self.pptx.read_bytes()
        with self.assertRaisesRegex(ValueError,'source hash'): shell.finalize(self.pptx,'0'*64,self.first,self.root/'proof.json')
        self.assertEqual(original,self.pptx.read_bytes())

    def test_source_race_before_commit_refused(self):
        original=shell.transform
        def race(*args):
            result=original(*args); self.pptx.write_bytes(b'concurrent edit'); return result
        with patch.object(shell,'transform',side_effect=race):
            with self.assertRaisesRegex(ValueError,'changed before commit'): self.finalize()
        self.assertEqual(b'concurrent edit',self.pptx.read_bytes()); self.assertFalse((self.root/'proof.json').exists())
        self.assertFalse(list(self.root.glob('*.tmp')))

    def test_wrong_first_image_rejected_by_final_gate(self):
        self.first.write_bytes(png('red')); self.finalize(); self.first.write_bytes(png('blue'))
        self.assertTrue(any('differs' in e for e in shell.verify_shell(self.pptx,self.first)))

    def test_missing_thumbnail_fails_final_gate(self):
        shell.clean_intermediate(self.pptx); self.assertTrue(shell.verify_shell(self.pptx,self.first))

    def test_wrong_view_fails_final_gate(self):
        self.finalize(); parts,_,_=shell.unpack(self.pptx.read_bytes())
        parts[shell.VIEW]=parts[shell.VIEW].replace(b'sldView',b'outlineView')
        with ZipFile(self.pptx,'w') as z:
            for name,data in parts.items(): z.writestr(name,data)
        self.assertTrue(any('view' in e for e in shell.verify_shell(self.pptx,self.first)))

    def test_malformed_thumbnail_relation_cannot_delete_body(self):
        rel=f'<Relationships xmlns="{shell.R}"><Relationship Id="bad" Type="{shell.THUMB_TYPE}" Target="ppt/slides/slide1.xml"/></Relationships>'.encode()
        self.pptx.write_bytes(package({'_rels/.rels':rel})); self.finalize()
        parts,_,_=shell.unpack(self.pptx.read_bytes()); self.assertEqual(b'body unchanged',parts['ppt/slides/slide1.xml'])

    def test_thumbnail_content_types_cleanup_only_unused(self):
        self.finalize(); parts,_,_=shell.unpack(self.pptx.read_bytes())
        self.assertFalse(any(e.get('Extension')=='jpeg' for e in shell.ET.fromstring(parts['[Content_Types].xml'])))
        self.pptx.write_bytes(package({'ppt/media/keep.jpeg':b'image bytes'})); self.finalize()
        parts,_,_=shell.unpack(self.pptx.read_bytes()); self.assertIn(b'jpeg',parts['[Content_Types].xml'])
        self.assertEqual(b'image bytes',parts['ppt/media/keep.jpeg'])

    def test_proof_binds_actual_input_and_output(self):
        before=shell.sha(self.pptx.read_bytes()); proof=self.finalize()
        self.assertEqual(before,proof['before_sha256']); self.assertEqual(shell.sha(self.pptx.read_bytes()),proof['after_sha256'])
        self.assertEqual(shell.sha(self.first.read_bytes()),proof['first_slide_sha256'])
        self.assertIn('docProps/core.xml',proof['preserved_parts_sha256'])

    def test_orphan_stale_thumbnail_rejected(self):
        self.finalize()
        with ZipFile(self.pptx,'a') as z: z.writestr('docProps/thumbnail.png',b'ad')
        self.assertTrue(any('stale thumbnail' in e for e in shell.verify_shell(self.pptx,self.first)))

    def test_non_widescreen_aspect_ratio(self):
        out=BytesIO(); Image.new('RGB',(1200,900),'blue').save(out,format='PNG'); self.first.write_bytes(out.getvalue())
        self.finalize(); parts,_,_=shell.unpack(self.pptx.read_bytes())
        with Image.open(BytesIO(parts[shell.THUMB])) as im: self.assertEqual((320,240),im.size)

    def test_fit_to_window_preserves_existing_pan_and_scale(self):
        view=f'<p:viewPr xmlns:p="{shell.P}" xmlns:a="{shell.A}"><p:slideViewPr><p:cSldViewPr showGuides="1"><p:cViewPr varScale="0"><p:scale><a:sx n="400" d="100"/><a:sy n="400" d="100"/></p:scale><p:origin x="99999" y="99999"/></p:cViewPr><p:guideLst/></p:cSldViewPr></p:slideViewPr></p:viewPr>'.encode()
        self.pptx.write_bytes(package({shell.VIEW:view})); self.finalize()
        parts,_,_=shell.unpack(self.pptx.read_bytes()); root=shell.ET.fromstring(parts[shell.VIEW])
        camera=root.find('{'+shell.P+'}slideViewPr/{'+shell.P+'}cSldViewPr/{'+shell.P+'}cViewPr')
        self.assertEqual('1',camera.get('varScale')); self.assertEqual('99999',camera.find('{'+shell.P+'}origin').get('x'))
        self.assertEqual('400',camera.find('{'+shell.P+'}scale/{'+shell.A+'}sx').get('n'))
        self.assertEqual([],shell.verify_shell(self.pptx,self.first))
        parts[shell.VIEW]=parts[shell.VIEW].replace(b'varScale="1"',b'varScale="0"')
        with ZipFile(self.pptx,'w') as z:
            for name,data in parts.items(): z.writestr(name,data)
        self.assertTrue(any('fit' in error for error in shell.verify_shell(self.pptx,self.first)))

    def test_missing_or_invalid_png_never_changes_source(self):
        original=self.pptx.read_bytes(); self.first.unlink()
        with self.assertRaises(OSError): self.finalize()
        self.assertEqual(original,self.pptx.read_bytes())
        self.first.write_bytes(b'not a PNG')
        with self.assertRaises(OSError): self.finalize()
        self.assertEqual(original,self.pptx.read_bytes())

    def test_duplicate_zip_rejected_without_mutation(self):
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            with ZipFile(self.pptx,'a') as z: z.writestr('ppt/slides/slide1.xml',b'duplicate')
        original=self.pptx.read_bytes()
        with self.assertRaisesRegex(ValueError,'Duplicate ZIP'): self.finalize()
        self.assertEqual(original,self.pptx.read_bytes())

    def test_final_coverage_gate_cannot_waive_shell(self):
        # All other evidence components are isolated here; this verifies only
        # the final gate's control flow, not actual visual/render approval.
        import json
        import contextlib
        from pptx import Presentation
        import verify_coverage_review as gate
        prs=Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); prs.save(self.pptx)
        self.finalize(); digest=shell.sha(self.pptx.read_bytes())
        lock=self.root/'lock.json'; lock.write_text('{}')
        coverage={'sha256':digest,'status':'coverage_pass_pending_visual_review','slide_count':1,'pages':[],'findings':[]}
        coverage_path=self.root/'coverage.json'; coverage_path.write_text(json.dumps(coverage))
        (self.root/'audit.json').write_text('{}')
        review={'sha256':digest,'coverage_sha256':shell.sha(coverage_path.read_bytes()),
                'skill_lock':{'path':str(lock),'sha256':shell.sha(lock.read_bytes())},'items':[],
                'visual_fallback':{'decision':'pretend_pass'}}
        with contextlib.ExitStack() as stack:
            mocks={'check_skill_state':{'changed_files':[]},'rescan':coverage,'validate_render_evidence':None,
                   'verify_connections':[],'validate_gap_audit':([],[]),'validate_manual_dispositions':[],
                   'review_fonts':[],'check_visual_passes':[]}
            for name,value in mocks.items(): stack.enter_context(patch.object(gate,name,return_value=value))
            self.assertEqual([],gate.verify(self.pptx,coverage_path,review,self.root)['issues'])
            self.first.write_bytes(png('red'))
            result=gate.verify(self.pptx,coverage_path,review,self.root)
            self.assertEqual('failed',result['status'])
            self.assertTrue(any(row.get('stage')=='delivery_shell' for row in result['issues']))
            self.assertTrue(all(' -Width 1280' in row['recovery'] for row in result['issues'] if row.get('stage')=='delivery_shell'))

if __name__=='__main__': unittest.main()
