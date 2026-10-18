"""Measure chart/caption/container spacing from current pixels, not chart padding.

The caller must bind the original layout and freshly recomputed gap report and
run validate_render_evidence first. This adapter independently rechecks the exact
three-single-object comparison, current input/image provenance, types and local
completeness. It never deletes raw frame-model issues or accepts supplied edges.
"""
from pathlib import Path
import copy
import hashlib
import json
import math
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from interstitial_spacing import _full_slide_background_ids

try:
    import numpy as np
except (ImportError,OSError):
    np = None

EDGE_NAMES = ['chart_bottom_px','caption_top_px','caption_bottom_px','after_top_px']


def _sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest().upper()


def _box(shape):
    return [shape.left/12700,shape.top/12700,(shape.left+shape.width)/12700,(shape.top+shape.height)/12700]


def _inside(inner,outer,epsilon=.25):
    return inner[0]>=outer[0]-epsilon and inner[1]>=outer[1]-epsilon and inner[2]<=outer[2]+epsilon and inner[3]<=outer[3]+epsilon


def _opaque_container(shape):
    if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or shape.auto_shape_type not in (MSO_AUTO_SHAPE_TYPE.RECTANGLE,MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE):
        return 'invalid','Following object is not a native rectangle/rounded rectangle'
    if shape._element.xpath('./p:nvSpPr/p:cNvPr[@hidden="1" or @hidden="true"]'):
        return 'invalid','Following container is hidden'
    fills=shape._element.xpath('./p:spPr/a:solidFill')
    if not fills:
        return 'invalid','Following object has no actual opaque solid fill'
    colors=list(fills[0])
    if len(colors)!=1 or colors[0].tag.rsplit('}',1)[-1]!='srgbClr':
        return 'manual','Theme/inherited container paint needs visual verification'
    value=colors[0].get('val','')
    if len(value)!=6 or any(c not in '0123456789ABCDEFabcdef' for c in value):
        return 'invalid','Invalid native container paint'
    for node in colors[0]:
        tag=node.tag.rsplit('}',1)[-1]
        if tag.startswith('alpha') and (tag!='alpha' or node.get('val')!='100000'):
            return 'invalid','Transparent/partially opaque containers cannot act as boundary proxies'
        if not tag.startswith('alpha'):
            return 'manual','Transformed native container color needs visual verification'
    if not shape._element.xpath('./p:spPr/a:ln/a:noFill'):
        return 'manual','Container stroke is not independently resolved'
    return 'opaque',tuple(int(value[i:i+2],16) for i in (0,2,4))


def _pixel_box(box,sx,sy,width,height,pad=0):
    return [max(0,math.floor(box[0]*sx)-pad),max(0,math.floor(box[1]*sy)-pad),
            min(width,math.ceil(box[2]*sx)+pad),min(height,math.ceil(box[3]*sy)+pad)]


def _ink(rgb,box,*,container_rgb=None):
    """Whole-object pixel bounds; no scanline/ROI can be chosen by the writer."""
    l,t,r,b=box
    if r-l<3 or b-t<3:return None,'Object pixel rectangle is empty or too small'
    if l<3 or t<3 or r>rgb.shape[1]-3 or b>rgb.shape[0]-3:
        return None,'Object lacks surrounding background context'
    outside=np.concatenate([rgb[t-3:t,l:r].reshape(-1,3),rgb[b:b+3,l:r].reshape(-1,3),
                            rgb[t:b,l-3:l].reshape(-1,3),rgb[t:b,r:r+3].reshape(-1,3)])
    background=np.median(outside,axis=0)
    if np.quantile(np.max(np.abs(outside.astype(float)-background),axis=1),.98)>6:
        return None,'Surrounding background is textured/occluded; current boundary is uncertain'
    tile=rgb[t:b,l:r]
    if container_rgb is None:
        inner_border=np.concatenate([tile[0],tile[-1],tile[:,0],tile[:,-1]])
        if np.quantile(np.max(np.abs(inner_border.astype(float)-background),axis=1),.95)>6:
            return None,'Object border has visible paint/overlap; cannot ignore its frame as blank padding'
    else:
        # Confirm actual native paint exists in the upper central interior, before
        # any internal caption. White-on-white invisible proxies remain unknown.
        strip=tile[max(1,min(5,tile.shape[0]//4)):max(2,min(9,tile.shape[0]//2)),tile.shape[1]//4:3*tile.shape[1]//4]
        if not strip.size or np.mean(np.max(np.abs(strip.astype(int)-np.array(container_rgb)),axis=2)<=8)<.8:
            return None,'Current container pixels do not reliably confirm its opaque native paint'
    contrast=np.max(np.abs(tile.astype(float)-background),axis=2)
    boxes=[]
    for threshold in (18,24,30):
        mask=contrast>threshold
        yy,xx=np.nonzero(mask)
        if not len(xx):return None,'No reliably visible object pixels against current background'
        boxes.append([l+int(xx.min()),t+int(yy.min()),l+int(xx.max())+1,t+int(yy.max())+1])
    if np.max(np.ptp(np.array(boxes),axis=0))>1:
        return None,'Object bounds are unstable across contrast thresholds'
    # Conservative extrema preserve the tightest actual visible gap.
    return [min(v[0] for v in boxes),min(v[1] for v in boxes),max(v[2] for v in boxes),max(v[3] for v in boxes)],None


def measure_relation(pptx,render,render_dir,slide,comparison):
    """Return an immutable-comparison measurement; malformed/proxy input is hard.

    comparison is an original audit_layout_gaps row or an independently generated
    discover_relations row, including bounds_pt/gaps_pt. The caller binds original
    rows to their layout/raw report and discovered rows to a fresh native scan.
    Pixel uncertainty preserves every independently known
    edge using visible_edges_px=[chart_bottom,caption_top,caption_bottom,after_top]
    and known_edge_indices. A successful measure is still pending visual review.
    """
    result={'slide':slide,'comparison_name':comparison.get('name') if isinstance(comparison,dict) else None,
            'shape_ids':[],'measurement_status':'invalid_relation','visible_edges_px':[None]*4,
            'known_edge_indices':[],'measurement_known_fields':[],
            'local_and_full_page_review_required':True,'raw_issues_must_be_retained':True}
    def invalid(message):
        return {**result,'measurement_status':'invalid_relation','error':message}
    try:
        pptx=Path(pptx);render_dir=Path(render_dir).resolve()
        sha=_sha(pptx);prs=Presentation(pptx)
        if type(slide) is not int or not 1<=slide<=len(prs.slides):return invalid('Invalid actual slide number')
        if not isinstance(render,dict) or render.get('schema')!='render-evidence/v2' or render.get('sha256')!=sha or render.get('slide_count')!=len(prs.slides) or render.get('status') not in ('rendered_pending_visual_review','failed_preflight'):
            return invalid('Current completed v2 render provenance is required')
        snapshot=Path(render.get('snapshot',''))
        if not snapshot.is_file() or snapshot.resolve().parent!=render_dir or _sha(snapshot)!=sha:
            return invalid('Current snapshot is missing, stale or belongs to another render directory')
        image_path=render_dir/f'slide-{slide:02}.png';image_sha=_sha(image_path)
        if render.get('raster_hashes',{}).get(str(slide))!=image_sha:return invalid('Current PNG hash differs from render evidence')
        result.update(pptx_sha256=sha,image_sha256=image_sha)
        if not isinstance(comparison,dict) or comparison.get('slide')!=slide or comparison.get('axis')!='vertical' or not isinstance(comparison.get('name'),str) or not comparison['name'].strip():
            return invalid('Need the exact recomputed named vertical comparison row')
        blocks=comparison.get('blocks')
        if not isinstance(blocks,list) or len(blocks)!=3:return invalid('This adapter requires exactly three original blocks')
        if any(not isinstance(block,dict) or not isinstance(block.get('shape_ids'),list) or len(block['shape_ids'])!=1 or type(block['shape_ids'][0]) is not int for block in blocks):
            return invalid('Each original block must contain exactly one native object')
        ids=[block['shape_ids'][0] for block in blocks];result['shape_ids']=ids
        shapes={s.shape_id:s for s in prs.slides[slide-1].shapes}
        if len(set(ids))!=3 or any(sid not in shapes for sid in ids):return invalid('Missing, duplicate or non-top-level relationship object')
        chart,caption,after=[shapes[sid] for sid in ids]
        if not chart.has_chart:return invalid('First object must be a real native chart, never a picture/proxy')
        if caption.shape_type!=MSO_SHAPE_TYPE.TEXT_BOX or not caption.text.strip():return invalid('Middle object must be an independent native text box')
        state,paint=_opaque_container(after)
        if state=='invalid':return invalid(paint)
        boxes=[_box(s) for s in (chart,caption,after)];c,m,a=boxes
        scale=prs.slide_height/12700/540
        frame_gaps=[m[1]-c[3],a[1]-m[3]]
        result.update(frame_bounds_pt=boxes,frame_gaps_pt=frame_gaps,original_frame_comparison=copy.deepcopy(comparison))
        expected_bounds=comparison.get('bounds_pt');expected_gaps=comparison.get('gaps_pt')
        if not isinstance(expected_bounds,list) or len(expected_bounds)!=3 or not isinstance(expected_gaps,list) or len(expected_gaps)!=2:
            return invalid('Comparison must come from the recomputed original gap audit, including frame measurements')
        if any(not isinstance(row,list) or len(row)!=4 for row in expected_bounds):return invalid('Invalid original comparison bounds')
        values=[*expected_gaps,*[v for row in expected_bounds for v in row]]
        if any(type(v) not in (int,float) or not math.isfinite(v) for v in values) or any(abs(v-w)>1e-6 for row,box in zip(expected_bounds,boxes) for v,w in zip(row,box)) or any(abs(v-w)>1e-6 for v,w in zip(expected_gaps,frame_gaps)):
            return invalid('Original frame comparison differs from independently recomputed native geometry')
        minimum=comparison.get('minimum_gap_pt')
        if type(minimum) not in (int,float) or not math.isfinite(minimum) or minimum<12*scale:
            return invalid('This visible-spacing adapter cannot reduce the 12pt scaled minimum')
        result.update(minimum_gap_pt=max(12*scale,minimum),minimum12scaled=12*scale,
            comparison_sha256=hashlib.sha256(json.dumps(comparison,sort_keys=True,ensure_ascii=False,separators=(',',':')).encode()).hexdigest().upper())
        if not c[1]<m[1]<a[1] or min(c[2],m[2])-max(c[0],m[0])<=0 or not a[0]<=m[0] or not a[2]>=m[2]:
            return invalid('Native order/horizontal relationship does not support chart-caption-container spacing')
        # Known full-slide paint is safe; a renamed local proxy is never excluded.
        backgrounds=_full_slide_background_ids(prs.slides[slide-1])
        corridor=[min(c[0],m[0]),c[1],max(c[2],m[2]),a[1]]
        omitted=[]
        for sid,other in shapes.items():
            if sid in ids or sid in backgrounds:continue
            box=_box(other)
            if _inside(box,a):continue  # actual content within the opaque final container
            if min(box[2],corridor[2])-max(box[0],corridor[0])>.25*scale and min(box[3],corridor[3])-max(box[1],corridor[1])>.25*scale:
                omitted.append(sid)
        if omitted:
            result['unlisted_shape_ids']=omitted
            return invalid('Unlisted native objects intersect the complete relationship; do not omit neighbours')
        rows=[r for r in render.get('text_records',[]) if r.get('slide')==slide and r.get('audit_path')==str(caption.shape_id)]
        if len(rows)!=1 or ''.join(rows[0].get('text','').split())!=''.join(caption.text.split()):
            return invalid('Current caption text record does not exactly identify the native text')
        with Image.open(image_path) as image:
            image=image.convert('RGB');sx=image.width/(prs.slide_width/12700);sy=image.height/(prs.slide_height/12700)
            union=[min(b[0] for b in boxes),min(b[1] for b in boxes),max(b[2] for b in boxes),max(b[3] for b in boxes)]
            result['crop_box_px']=_pixel_box(union,sx,sy,image.width,image.height,4)
            result.update(pixels_per_pt_x=sx,pixels_per_pt_y=sy,segmentation_thresholds=[18,24,30])
            if np is None:return {**result,'measurement_status':'manual_review_required','measurement_limitation':['NumPy unavailable; retain exact local + full-page visual review']}
            rgb=np.asarray(image);pixel_boxes=[_pixel_box(b,sx,sy,image.width,image.height) for b in boxes]
            errors=[]
            for i,(shape,pixel_box) in enumerate(zip((chart,caption,after),pixel_boxes)):
                if abs(shape.rotation)>1e-6:
                    errors.append(f'Object {ids[i]} is rotated; current contour requires visual measurement');continue
                if shape._element.xpath('.//a:outerShdw | .//a:innerShdw | .//a:glow | .//a:softEdge') or (i==0 and chart.chart._element.xpath('.//a:outerShdw | .//a:innerShdw | .//a:glow | .//a:softEdge')):
                    errors.append(f'Object {ids[i]} has visual effects; current contour requires review');continue
                if i==2 and state!='opaque':errors.append(paint);continue
                box,error=_ink(rgb,pixel_box,container_rgb=paint if i==2 else None)
                if error:errors.append(f'Object {ids[i]}: '+error);continue
                field=['chart_ink_bbox_px','caption_ink_bbox_px','after_visible_bbox_px'][i]
                result[field]=box;result['measurement_known_fields'].append(field)
                indices=[0] if i==0 else ([1,2] if i==1 else [3])
                edge_values=[box[3]] if i==0 else ([box[1],box[3]] if i==1 else [box[1]])
                for index,value in zip(indices,edge_values):
                    result['visible_edges_px'][index]=value;result['known_edge_indices'].append(index)
                    result[EDGE_NAMES[index]]=value;result['measurement_known_fields'].append(EDGE_NAMES[index])
            if _sha(pptx)!=sha or _sha(image_path)!=image_sha:return invalid('Inputs changed during visible spacing measurement')
            edges=result['visible_edges_px']
            partial_gaps=[(edges[1]-edges[0])/sy if edges[0] is not None and edges[1] is not None else None,
                          (edges[3]-edges[2])/sy if edges[2] is not None and edges[3] is not None else None]
            partial_target=sum(partial_gaps)/2 if all(g is not None for g in partial_gaps) else None
            result.update(visible_gaps_pt=partial_gaps,visible_target_gap_pt=partial_target,
                          balance_tolerance_pt=max(3*scale,abs(partial_target)*.2) if partial_target is not None else None)
            if any(g is not None and g<result['minimum_gap_pt'] for g in partial_gaps):
                return {**result,'measurement_status':'measured_defect','visible_gaps_pt':partial_gaps,
                        'defects':['visible_gap_below_minimum_or_overlap'],'measurement_limitation':errors}
            if errors:
                return {**result,'measurement_status':'manual_review_required','measurement_limitation':errors}
            bottom,top,end,next_top=result['visible_edges_px']
            gaps=[(top-bottom)/sy,(next_top-end)/sy]
            target=sum(gaps)/2;tolerance=max(3*scale,abs(target)*.2)
            defects=[]
            if min(gaps)<result['minimum_gap_pt']:defects.append('visible_gap_below_minimum_or_overlap')
            if abs(gaps[0]-gaps[1])>tolerance:defects.append('visible_peer_gap_imbalance')
            result.update(visible_gaps_pt=gaps,visible_target_gap_pt=target,balance_tolerance_pt=tolerance,
                defects=defects,measurement_status='measured_defect' if defects else 'measured_valid_pending_visual_review')
            return result
    except (OSError,ValueError,TypeError,KeyError,IndexError,AttributeError,OverflowError) as exc:
        return invalid(str(exc))


def discover_relations(pptx,render,render_dir):
    """Find only high-confidence native chart -> short caption -> colored band.

    No layout/spec/name is read. Nearest native neighbours must form the complete
    three-object relation; formulas, other modules and text in another container
    cannot be skipped. Pixel defects/uncertainty remain candidates, never filters
    for choosing only favourable measured groups. Other layouts keep existing
    whole-page/chart visual review rather than being forced into this adapter.
    Raises ValueError for stale evidence, which must not masquerade as no matches.
    """
    pptx=Path(pptx);render_dir=Path(render_dir).resolve();sha=_sha(pptx)
    prs=Presentation(pptx);scale=prs.slide_height/12700/540
    if not isinstance(render,dict) or render.get('schema')!='render-evidence/v2' or render.get('sha256')!=sha or render.get('slide_count')!=len(prs.slides) or render.get('status') not in ('rendered_pending_visual_review','failed_preflight'):
        raise ValueError('Discovery requires current completed v2 render evidence')
    snapshot=Path(render.get('snapshot',''))
    if not snapshot.is_file() or snapshot.resolve().parent!=render_dir or _sha(snapshot)!=sha:
        raise ValueError('Discovery snapshot is missing, stale or in another directory')
    images=[]
    for n in range(1,len(prs.slides)+1):
        path=render_dir/f'slide-{n:02}.png';current=_sha(path)
        if render.get('raster_hashes',{}).get(str(n))!=current:raise ValueError('Discovery page pixels are stale: '+str(n))
        images.append((path,current))
    found=[]
    for n,slide in enumerate(prs.slides,1):
        shapes={s.shape_id:s for s in slide.shapes};boxes={sid:_box(s) for sid,s in shapes.items()}
        backgrounds=_full_slide_background_ids(slide)
        def overlap(a,b):return min(a[2],b[2])-max(a[0],b[0])
        def next_objects(box,skip,max_gap):
            choices=[]
            for sid,b in boxes.items():
                if sid in skip or sid in backgrounds or min(b[2]-b[0],b[3]-b[1])<=0:continue
                if -.25*scale<=b[1]-box[3]<=max_gap and overlap(box,b)>=.5*min(box[2]-box[0],b[2]-b[0]):
                    choices.append((b[1],sid))
            if not choices:return []
            nearest=min(top for top,_ in choices)
            return [sid for top,sid in choices if abs(top-nearest)<=.25*scale]
        for chart in slide.shapes:
            if not chart.has_chart or abs(chart.rotation)>1e-6:continue
            cb=boxes[chart.shape_id];cw=cb[2]-cb[0]
            if cw<.3*(prs.slide_width/12700) or cb[3]-cb[1]<.15*(prs.slide_height/12700):continue
            first=next_objects(cb,{chart.shape_id},36*scale)
            if len(first)!=1:continue
            caption=shapes[first[0]];mb=boxes[caption.shape_id]
            if caption.shape_type!=MSO_SHAPE_TYPE.TEXT_BOX or not caption.text.strip() or abs(caption.rotation)>1e-6:continue
            text=caption.text.strip()
            if len(text)>180 or len(text.replace('\v','\n').splitlines())>2 or mb[3]-mb[1]>32*scale:continue
            if not .35*cw<=mb[2]-mb[0]<=1.25*cw or overlap(cb,mb)<.75*min(cw,mb[2]-mb[0]):continue
            # A label within a real separate container belongs to that module.
            contained=False
            for sid,s in shapes.items():
                if sid in (chart.shape_id,caption.shape_id) or sid in backgrounds:continue
                state,_=_opaque_container(s)
                if state=='opaque' and _inside(mb,boxes[sid]):contained=True;break
            if contained:continue
            second=next_objects(mb,{chart.shape_id,caption.shape_id},60*scale)
            if len(second)!=1:continue
            after=shapes[second[0]];ab=boxes[after.shape_id]
            state,color=_opaque_container(after)
            if state!='opaque' or color==(255,255,255) or abs(after.rotation)>1e-6:continue
            if not 20*scale<=ab[3]-ab[1]<=90*scale or ab[0]>mb[0] or ab[2]<mb[2]:continue
            has_label=after.has_text_frame and bool(after.text.strip())
            has_label=has_label or any(sid!=after.shape_id and s.has_text_frame and s.text.strip() and _inside(boxes[sid],ab) for sid,s in shapes.items())
            if not has_label:continue
            ids=[chart.shape_id,caption.shape_id,after.shape_id]
            row={'slide':n,'name':f'visible-chart-caption-p{n}-'+ '-'.join(map(str,ids)),
                'axis':'vertical','blocks':[{'name':str(sid),'shape_ids':[sid]} for sid in ids],
                'bounds_pt':[cb,mb,ab],'gaps_pt':[mb[1]-cb[3],ab[1]-mb[3]],
                'target_gap_pt':12*scale,'tolerance_pt':3*scale,'minimum_gap_pt':12*scale,
                'discovery_source':'native_chart_caption_container/v1'}
            # Current pixels and the complete intervening-object inventory are
            # checked by the same adapter used for declared groups. A real defect
            # or an uncertain contour stays mandatory; invalid cherry-picked
            # object combinations are not promoted to this specialised relation.
            measured=measure_relation(pptx,render,render_dir,n,row)
            if measured['measurement_status']=='invalid_relation':
                if 'caption text record' in measured.get('error',''):raise ValueError(measured['error'])
                continue
            found.append(row)
    if _sha(pptx)!=sha or _sha(snapshot)!=sha or any(_sha(path)!=value for path,value in images):
        raise ValueError('Inputs changed during independent chart-caption discovery')
    return sorted(found,key=lambda row:(row['slide'],row['name']))


def main():
    import argparse
    from render_evidence import validate_render_evidence
    parser=argparse.ArgumentParser(description='Independently discover and measure native chart/caption/container spacing')
    parser.add_argument('pptx',type=Path);parser.add_argument('render_dir',type=Path);parser.add_argument('report',type=Path)
    args=parser.parse_args()
    try:
        render=json.loads((args.render_dir/'audit.json').read_text(encoding='utf-8-sig'))
        validate_render_evidence(args.pptx,render,args.render_dir)
        rows=discover_relations(args.pptx,render,args.render_dir)
        records=[{'comparison':row,'measurement':measure_relation(args.pptx,render,args.render_dir,row['slide'],row)} for row in rows]
        result={'schema':'visible-chart-spacing/v1','sha256':_sha(args.pptx),'relations':records,
                'status':'review_required' if records else 'no_specialized_relations_pending_full_visual_review'}
        code=2 if records else 0
    except Exception as exc:
        result={'schema':'visible-chart-spacing/v1','status':'failed','error':str(exc),'relations':[]};code=3
    args.report.parent.mkdir(parents=True,exist_ok=True)
    args.report.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(result['status'],len(result['relations']))
    return code


if __name__=='__main__':raise SystemExit(main())
