"""Bind every reduced-size text region to actual render evidence, never names alone."""
import json,math,re
from pathlib import Path
from datetime import datetime
from PIL import Image,ImageChops
from pptx import Presentation
from render_evidence import validate_render_evidence,digest,native_inventory

def review_fonts(pptx,render,review):
    issues=[];prs=Presentation(pptx);scale=prs.slide_height/12700/540
    needed={}
    for r in render['text_records']:
        size=min(r['run_font_sizes_pt'])/scale
        if size>=12:continue
        if '/table/' not in r['audit_path'] and re.fullmatch(r'[1-9][0-9]*/[1-9][0-9]*',r['text'].strip()) and ((r.get('shape') or '').startswith('page-number:') or r['top_pt']>=prs.slide_height/12700*.9) and size>=10:continue
        needed[(r['slide'],r['audit_path'])]=r
    rows=review.get('font_exceptions',[])
    keys=[(r.get('slide'),r.get('audit_path')) for r in rows]
    ids=[r.get('exception_id') for r in rows]
    if len(set(keys))!=len(keys) or set(keys)!=set(needed) or any(not isinstance(i,str) or not i.strip() for i in ids) or len(ids)!=len(set(ids)):
        return [{'reason':'Font exceptions must exactly cover every reduced-size region with unique exception IDs'}]
    for item in rows:
        r=needed[item['slide'],item['audit_path']];size=min(r['run_font_sizes_pt'])/scale
        def fail(message):issues.append({'exception_id':item['exception_id'],'reason':message})
        for field in ('justification','final_visual_observation','content_role'):
            if not isinstance(item.get(field),str) or not item[field].strip():fail('Missing '+field)
        if size<8:fail('Below the absolute 8pt auxiliary floor');continue
        if '/table/' in r['audit_path']:
            if size<10 or item.get('content_role')!='ordinary_data_table':fail('Tables cannot use 8pt or disguise body prose as a chart label')
        elif item.get('content_role')!='complex_graph_auxiliary':fail('Reduced font is restricted to necessary auxiliary graphic text')
        if size>=10:continue
        for field in ('observed_10pt_conflict','content_must_remain','layout_and_split_analysis'):
            if not isinstance(item.get(field),str) or not item[field].strip():fail('Missing actual 8pt evidence: '+field)
        attempts=item.get('attempts',[])
        if len(attempts)<2:fail('8pt needs initial and layout-adjusted actual 10pt renders');continue
        hashes=[];layout_hashes=[];last_time=None
        for attempt in attempts:
            try:
                candidate=Path(attempt['pptx']);folder=Path(attempt['render_dir'])
                report=json.loads((folder/'audit.json').read_text(encoding='utf-8-sig'))
                # Failed layout preflight is expected evidence here, never a final pass.
                inv=validate_render_evidence(candidate,report,folder,diagnostic_attempt=True)
                match=[x for x in report['text_records'] if x['slide']==attempt['slide'] and x['audit_path']==attempt['audit_path']]
                if len(match)!=1:raise ValueError('10pt target missing')
                target=match[0];attempt_scale=Presentation(candidate).slide_height/12700/540
                if abs(min(target['run_font_sizes_pt'])/attempt_scale-10)>.01:raise ValueError('Attempt target was not actually 10pt')
                if ''.join(target['text'].split())!=''.join(r['text'].split()):raise ValueError('10pt attempt did not preserve the same required text')
                if not attempt.get('observed_conflict') or not attempt.get('layout_action'):raise ValueError('Missing actual conflict/layout attempt observations')
                end=datetime.fromisoformat(report['render_finished_utc']);start=datetime.fromisoformat(render['render_started_utc'])
                if end>=start or (last_time and end<=last_time):raise ValueError('10pt render evidence is not ordered before final render')
                last_time=end;hashes.append(digest(candidate))
                import hashlib
                layout_hashes.append(hashlib.sha256(Presentation(candidate).slides[attempt['slide']-1]._element.xml.encode()).hexdigest())
                image=folder/f"slide-{attempt['slide']:02}.png";box=attempt['crop_box_px']
                if len(box)!=4 or any(type(v) is not int for v in box):raise ValueError('Invalid 10pt crop coordinates')
                with Image.open(image) as src,Image.open(attempt['crop_path']) as crop:
                    if not 0<=box[0]<box[2]<=src.width or not 0<=box[1]<box[3]<=src.height:raise ValueError('10pt crop outside page')
                    actual=src.convert('RGB').crop(box);stored=crop.convert('RGB')
                    if actual.size!=stored.size or ImageChops.difference(actual,stored).getbbox():raise ValueError('10pt crop differs from actual render')
                    bounds=inv['bounds'][attempt['slide'],attempt['audit_path']]
                    if bounds is None:raise ValueError('10pt target geometry must be resolved for exception evidence')
                    ap=Presentation(candidate);px=src.width/(ap.slide_width/12700);py=src.height/(ap.slide_height/12700)
                    target_box=[math.floor(bounds[0]*px),math.floor(bounds[1]*py),math.ceil(bounds[2]*px),math.ceil(bounds[3]*py)]
                    if box[0]>target_box[0] or box[1]>target_box[1] or box[2]<target_box[2] or box[3]<target_box[3]:raise ValueError('10pt crop omits target region')
            except (KeyError,ValueError,TypeError,OSError) as exc:fail(str(exc))
        if len(set(hashes))<2:fail('Two different actual 10pt layout attempts are required, not copied evidence')
        if len(set(layout_hashes))<2:fail('10pt slide content/layout never changed; metadata-only changes are not layout attempts')
    return issues
