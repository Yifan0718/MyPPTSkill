"""Synthetic protocol regressions; these fixtures are not visual signoffs."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Pt
from pptx.dml.color import RGBColor
from measurement_fallback import digest, eligible_code, validate_visual_fallback


class VisualFallback(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.root=Path(self.temp.name)
        self.p=Presentation();self.p.slide_width=Pt(960);self.p.slide_height=Pt(540)
        self.s=self.p.slides.add_slide(self.p.slide_layouts[6])
        self.a=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(100),Pt(100),Pt(260),Pt(60))
        self.a.fill.solid();self.a.fill.fore_color.rgb=RGBColor(30,60,90);self.a.line.fill.background()
        self.m=self.s.shapes.add_textbox(Pt(100),Pt(185),Pt(260),Pt(20));self.m.text='Measured explanation'
        self.z=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,Pt(100),Pt(230),Pt(260),Pt(40))
        self.f={'id':'s1-f1','slide':1,'reason':'header_body_boundary_review','shape_ids':[self.a.shape_id,self.m.shape_id,self.z.shape_id]}
        self.image=self.root/'slide-01.png';im=Image.new('RGB',(960,540),'white');d=ImageDraw.Draw(im)
        d.rectangle([100,100,359,159],fill=(30,60,90));d.text((100,185),'Measured explanation',fill=(30,60,90));d.rounded_rectangle([100,230,359,269],radius=8,fill=(30,60,90));im.save(self.image)
        self.crop=self.root/'local.png';im.crop((90,90,370,280)).save(self.crop)
        self.error='After boundary must be a native opaque solid rectangle'
        self.render={'sha256':'A'*64,'raster_hashes':{'1':digest(self.image)},'render_finished_utc':'2026-10-18T00:00:00+00:00'}
        self.render['text_records']=[{'slide':1,'shape_id':self.m.shape_id,'kind':'text','bound_left_pt':100,'bound_top_pt':185,'bound_width_pt':240,'bound_height_pt':20}]
        self.record={'code':'unsupported_adjacent_contour','original_error':self.error,'pptx_sha256':'A'*64,'image_sha256':digest(self.image),
            'page_object_ids':[s.shape_id for s in self.s.shapes],
            'measurement_limitation':'The actual rounded rectangle cannot be resolved by the rectangle-only contour classifier.',
            'local_observation':'The full relationship shows two clear equal gaps around the intermediate explanation.',
            'global_observation':'The full-page composition remains balanced and the related modules remain visibly separated.',
            'local_result':'pass','global_result':'pass','unresolved':[],'returned_to_full_page':True,'normal_size_viewed':True,
            'target_observations':{str(i):'The complete native member and its visible boundary were compared in the local context.' for i in self.f['shape_ids']},
            'detail':{'path':str(self.crop),'sha256':digest(self.crop),'crop_box_px':[90,90,370,280]},
            'visible_measurements':{'vertical_edges_px':[160,185,205,230]},'reviewed_at':'2026-10-18T00:03:00+00:00'}
        self.peer={'reviewer_id':'independent-test-reviewer','finding_id':self.f['id'],'slide':1,'code':self.record['code'],
            'original_error':self.error,'pptx_sha256':'A'*64,'image_sha256':digest(self.image),'detail_sha256':digest(self.crop),
            'decision':'visually_resolved','unresolved':[],'local_viewed':True,'full_page_viewed':True,
            'visible_measurements':copy.deepcopy(self.record['visible_measurements']),
            'local_observation':'Synthetic test record of the local two-gap contour relationship, not a real visual approval.',
            'global_observation':'Synthetic test record for full-page review completeness, not a real visual approval.',
            'review_source':'Synthetic unittest fixture; deliberately not real review provenance or a deliverable signoff.',
            'reviewed_at':'2026-10-18T00:02:00+00:00'}
        self.peer_path=self.root/'independent.json';self.write_peer()
        self.disposition={'spacing_contract':{'before_shape_ids':[self.a.shape_id],'middle_shape_id':self.m.shape_id,'after_shape_id':self.z.shape_id},'visual_fallback':self.record}
    def tearDown(self):self.temp.cleanup()
    def write_peer(self):
        self.peer_path.write_text(json.dumps(self.peer),encoding='utf8');self.record['independent_review']={'path':str(self.peer_path),'sha256':digest(self.peer_path)}
    def check(self,accepted):
        error,resolution=validate_visual_fallback(self.f,self.disposition,'interstitial',self.error,self.image,self.s,self.render,'test-author')
        self.assertEqual(error is None,accepted,error)
        self.assertEqual(resolution is not None,accepted)
    def test_complete_supported_uncertainty(self):self.check(True)
    def test_stale_input(self):self.record['pptx_sha256']='B'*64;self.check(False)
    def test_stale_pixels(self):self.record['image_sha256']='B'*64;self.check(False)
    def test_full_page_is_not_local(self):self.record['detail']['crop_box_px']=[0,0,960,540];self.check(False)
    def test_wrong_detail_pixels(self):Image.new('RGB',(280,190),'white').save(self.crop);self.record['detail']['sha256']=digest(self.crop);self.check(False)
    def test_omitted_member(self):self.record['page_object_ids'].pop();self.check(False)
    def test_omitted_observation(self):self.record['target_observations'].pop(str(self.m.shape_id));self.check(False)
    def test_wrong_error(self):self.record['original_error']='different';self.check(False)
    def test_no_global_return(self):self.record['returned_to_full_page']=False;self.check(False)
    def test_failed_local(self):self.record['local_result']='fail';self.check(False)
    def test_unresolved_peer(self):self.peer['unresolved']=['gap uncertain'];self.write_peer();self.check(False)
    def test_self_review(self):self.peer['reviewer_id']='test-author';self.write_peer();self.check(False)
    def test_missing_review_file(self):self.peer_path.unlink();self.check(False)
    def test_pre_render_review(self):self.peer['reviewed_at']='2026-10-17T23:00:00+00:00';self.write_peer();self.check(False)
    def test_missing_time_zone(self):self.peer['reviewed_at']='2026-10-18T00:02:00';self.write_peer();self.check(False)
    def test_wrong_peer_measurement(self):self.peer['visible_measurements']['vertical_edges_px'][0]=150;self.write_peer();self.check(False)
    def test_measured_unbalanced_cannot_close(self):
        self.record['visible_measurements']['vertical_edges_px']=[150,185,205,230]
        self.peer['visible_measurements']=copy.deepcopy(self.record['visible_measurements']);self.write_peer();self.check(False)
    def test_missing_actual_pixel_measurement(self):self.record.pop('visible_measurements');self.check(False)
    def test_missing_peer_context(self):self.peer['global_observation']='pass';self.write_peer();self.check(False)
    def test_cannot_shrink_known_container_to_fake_gap(self):self.a.height=Pt(80);self.check(False)
    def test_cannot_move_known_glyph_to_fake_gap(self):self.render['text_records'][0]['bound_top_pt']=165;self.check(False)
    def test_omitted_nearer_object_rejected_even_with_full_inventory(self):
        other=self.s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Pt(100),Pt(170),Pt(260),Pt(10))
        other.fill.solid();other.fill.fore_color.rgb=RGBColor(30,60,90);other.line.fill.background()
        self.record['page_object_ids'].append(other.shape_id);self.check(False)
    def plot_case(self,claimed_chart_edge):
        from pptx.chart.data import ChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from verify_coverage_review import verify_plot_table_pixels
        p=Presentation();p.slide_width=Pt(960);p.slide_height=Pt(540);s=p.slides.add_slide(p.slide_layouts[6])
        data=ChartData();data.categories=['A'];data.add_series('S',[1])
        chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Pt(100),Pt(100),Pt(260),Pt(70),data)
        table=s.shapes.add_table(2,2,Pt(100),Pt(250),Pt(240),Pt(50))
        im=Image.new('RGB',(960,540),'white');d=ImageDraw.Draw(im)
        d.rectangle([100,100,299,169],fill=(30,60,90));d.rectangle([100,250,299,274],fill=(30,60,90))
        im.save(self.image);im.crop((90,90,370,310)).save(self.crop)
        f={'id':'plot-test','slide':1,'reason':'chart_table_visible_edge_requires_review','shape_ids':[chart.shape_id,table.shape_id],'plot_right_pt':300,'table_right_pt':340,'delta_pt':40}
        disp={'plot_table_edge_evidence':{'pixels_per_pt':1,'slide_width_pt':960,'chart_rows_px':[110,120,130],'table_rows_px':[260,270,280]}}
        error=verify_plot_table_pixels(f,disp,self.image,s,1);self.assertEqual(error,'No measurable visible edge on this scanline')
        self.record.update(code='unmeasurable_scanline',original_error=error,image_sha256=digest(self.image),page_object_ids=f['shape_ids'],
            target_observations={str(i):'Synthetic protocol test; not an actual independent review or visual quality signoff.' for i in f['shape_ids']},
            detail={'path':str(self.crop),'sha256':digest(self.crop),'crop_box_px':[90,90,370,310]},
            visible_measurements={'anchor':'right','anchors_px':{str(chart.shape_id):claimed_chart_edge,str(table.shape_id):claimed_chart_edge}})
        self.peer.update(finding_id=f['id'],code=self.record['code'],original_error=error,image_sha256=digest(self.image),detail_sha256=digest(self.crop),visible_measurements=self.record['visible_measurements']);self.write_peer()
        disp['visual_fallback']=self.record;self.render['raster_hashes']['1']=digest(self.image)
        return validate_visual_fallback(f,disp,'plot_table',error,self.image,s,self.render,'test-author')
    def test_known_chart_scanline_cannot_be_changed(self):self.assertIsNotNone(self.plot_case(320)[0])
    def test_unknown_table_scanline_can_be_visually_resolved(self):self.assertIsNone(self.plot_case(299)[0])
    def center_case(self,change_known=False,known_misalignment=False):
        f={'id':'center-protocol','slide':1,'reason':'centered_text_visible_host_review','measurement_status':'manual_review_required',
            'shape_ids':self.f['shape_ids'],'crop_box_px':[90,90,370,280],'measurement_known_fields':['text_group_bbox_px'],
            'text_group_bbox_px':[190,175,270,195]}
        if known_misalignment:f['measurement_status']='measured_misalignment'
        error='Visible host contour or hosted text ink is unmeasurable'
        self.record.update(code='uncertain_visible_center',original_error=error,
            visible_measurements={'reference':'visible_hole','host_bbox_px':[100,100,360,270],
                                  'hole_bbox_px':[110,120,350,250],'text_group_bbox_px':[190,175,270,195]})
        if change_known:self.record['visible_measurements']['text_group_bbox_px']=[191,175,271,195]
        self.peer.update(finding_id=f['id'],code=self.record['code'],original_error=error,visible_measurements=self.record['visible_measurements']);self.write_peer()
        return validate_visual_fallback(f,self.disposition,'center',error,self.image,self.s,self.render,'test-author')
    def test_unknown_center_has_complete_evidence_path(self):self.assertIsNone(self.center_case()[0])
    def test_unknown_center_cannot_rewrite_known_text_ink(self):self.assertIsNotNone(self.center_case(change_known=True)[0])
    def test_known_center_mismatch_cannot_enter_fallback(self):self.assertIsNotNone(self.center_case(known_misalignment=True)[0])
    def chart_gap_case(self,change_known=False,minimum=12):
        f={'id':'gap-protocol','slide':1,'reason':'native_chart_spacing_unmeasurable','measurement_status':'manual_review_required',
           'shape_ids':self.f['shape_ids'],'crop_box_px':[90,90,370,280],'known_edge_indices':[1,2,3],
           'visible_edges_px':[None,185,205,230],'minimum_gap_pt':minimum}
        error='Visible chart spacing boundaries are unmeasurable'
        self.record.update(code='uncertain_chart_spacing',original_error=error,visible_measurements={'vertical_edges_px':[160,185,205,230]})
        if change_known:self.record['visible_measurements']['vertical_edges_px'][1]=186
        self.peer.update(finding_id=f['id'],code=self.record['code'],original_error=error,visible_measurements=self.record['visible_measurements']);self.write_peer()
        return validate_visual_fallback(f,self.disposition,'chart_gap',error,self.image,self.s,self.render,'test-author')
    def test_unknown_chart_gap_has_complete_evidence_path(self):self.assertIsNone(self.chart_gap_case()[0])
    def test_unknown_chart_gap_preserves_known_caption_edge(self):self.assertIsNotNone(self.chart_gap_case(change_known=True)[0])
    def test_unknown_chart_gap_preserves_stricter_minimum(self):self.assertIsNotNone(self.chart_gap_case(minimum=30)[0])
    def test_no_generic_bypass(self):
        for stage in ('interstitial','region','plot_table'):
            for error in ('Actual local gap is below 12pt scaled minimum','Actual local whitespace is unbalanced; repair and rerender','Actual visible chart/table edges remain misaligned','PPTX hash mismatch','Missing finite rendered glyph bounds','Before IDs must exactly cover the nearest visible row including independent glyph labels; do not omit or choose distant objects'):
                self.assertIsNone(eligible_code(stage,error))


if __name__=='__main__':unittest.main()
